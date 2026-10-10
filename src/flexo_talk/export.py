"""Writing a deck: one PowerPoint file, one PDF, and an SVG and a PNG per slide."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from flexo.drawing import Drawing, Group, Image, Shape, Text, ink_bounds, read_drawing
from flexo.export import rasterise
from flexo.pdf import ARTIFACT, Tag, Tagger, write_pdf
from flexo.portable import portable_svg
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Pt

from flexo_talk.deck import Deck, DeckBuild, RenderedSlide, Slide, _Figure, _Image
from flexo_talk.pptx import (
    Placement,
    add_drawing,
    add_heading,
    add_list,
    add_reveals,
    add_table,
    linking,
    plain_names,
    slide_pictures,
)
from flexo_talk.pptx import editable_maths as editable_maths_of

FORMATS = ("pptx", "pdf", "svg", "png")


def file_stem(identifier: object) -> str:
    """A deck's id as the name of its files: its own words, with nothing that would take
    them out of the folder they are written to (a slash, a leading dot) -- or "deck"."""

    import re

    words = re.sub(r"[\\/:\x00-\x1f]+", "-", str(identifier or "")).strip(" .-")
    return words or "deck"


def build_deck(
    deck: Deck,
    directory: Path,
    formats: tuple[str, ...],
    *,
    handout: bool = False,
    editable_maths: bool = True,
    images: Path | None = None,
    png_width: int | None = None,
) -> DeckBuild:
    """Write the deck into ``directory``: its PowerPoint and PDF, named after the deck,
    and an SVG and a PNG of each slide, beside them as talk-01.png, or in a folder of
    their own (``images``) as slide-01.png. The PDF has a page per step of each list a
    slide reveals, or, as a ``handout``, one page per slide, showing it whole. A PNG is
    ``png_width`` pixels wide, as tall as the slide's proportions make it -- else drawn
    at 144 dpi (1920 pixels across a 16:9 slide)."""

    unknown = set(formats) - set(FORMATS)
    if unknown:
        raise ValueError(f"Unknown {'format' if len(unknown) == 1 else 'formats'} {', '.join(sorted(unknown))}. "
                         f"Available formats: {', '.join(FORMATS)}.")
    directory.mkdir(parents=True, exist_ok=True)
    if "svg" in formats:
        rendered = deck.render()
    else:
        # Only an SVG carries its fonts: a PDF embeds its own, a PNG draws outlines, and a
        # PowerPoint names them, so none is cut down and embedded for nothing.
        from flexo.svg_resources import fonts_linked

        with fonts_linked():
            rendered = deck.render()
    name = file_stem(deck.id)
    diagnostics = [message for item in rendered for message in item.diagnostics]
    svgs: list[Path] = []
    pngs: list[Path] = []
    # Numbered as wide as the count needs, so the files sort in order: talk-100 after talk-099.
    digits = max(2, len(str(len(rendered))))
    folder, prefix = (directory, name) if images is None else (images, "slide")
    # (The pixels across a slide's points, at 72 to the inch: 144 dpi draws a point as two.)
    dpi = 144.0 if png_width is None else png_width * 72.0 / deck.style.width
    if images is not None and {"svg", "png"} & set(formats):
        folder.mkdir(parents=True, exist_ok=True)
    # (Numbered as the slides are shown: a slide skipped leaves no gap -- compose._slide_number.)
    for number, item in enumerate(rendered, 1):
        stem = f"{prefix}-{number:0{digits}d}"
        if "svg" in formats:
            path = folder / f"{stem}.svg"
            path.write_text(item.svg, encoding="utf-8")
            svgs.append(path)
        if "png" in formats:
            # Drawn from outlines, so the preview shows exactly the glyphs measured.
            path = folder / f"{stem}.png"
            path.write_bytes(rasterise(portable_svg(item.svg), dpi=dpi))
            pngs.append(path)
    pdf = None
    if "pdf" in formats:
        # A slide that reveals its list takes a page per step, unless this is a handout.
        pages = [
            page
            for item in rendered
            for page in ([item.svg] if handout else [item.at_step(step) for step in range(1, item.steps + 1)])
        ]
        title, author = _named(deck)
        pdf = write_pdf(pages, directory / f"{name}.pdf", title=title or name, author=author,
                        tags=_tagger(rendered))
    pptx = None
    if "pptx" in formats:
        pptx = directory / f"{name}.pptx"
        write_pptx(deck, rendered, pptx, editable_maths=editable_maths)
    notes = tuple(note for item in rendered for note in item.notes)
    return DeckBuild(pptx, pdf, tuple(svgs), tuple(pngs), tuple(diagnostics), notes)


def write_pptx(
    deck: Deck, rendered: list[RenderedSlide], target: Path, *, groups: bool = True, editable_maths: bool = True
) -> Path:
    """The deck as PowerPoint: every shape native, lists as lists, notes as notes.

    With ``editable_maths``, its maths is PowerPoint's own equations -- editable in its
    equation editor, set in its maths font -- over flexo's drawing of each, which
    other slide programs show instead (see ``pptx.alternate``).
    """

    presentation = Presentation()
    presentation.slide_width = Pt(deck.style.width)
    presentation.slide_height = Pt(deck.style.height)
    _properties(presentation, deck)
    layouts = {layout.name: layout for layout in presentation.slide_layouts}
    for item in rendered:
        drawing = read_drawing(item.svg)
        formulas = _drop_lists(drawing.root)
        _dissolve_regions(drawing.root)
        under, headings = _take_headings(drawing.root, item)
        # The layout the slide's headings come from: a content slide's titled, or blank.
        kinds = {kind for _, kind in headings}
        chosen = _LAYOUTS.get(item.slide.layout, "Title Only" if "title" in kinds else "Blank") if kinds else "Blank"
        slide = presentation.slides.add_slide(layouts[chosen])
        for empty in list(slide.placeholders):
            # The layout's own, empty: the slide's are its headings, and nothing else shows.
            empty._element.getparent().remove(empty._element)
        with linking(slide):
            page = deck.background
            own = item.slide.backdrop
            if own and own.startswith("#"):
                page = own
            if page:
                # A picture page (paper) comes as the slide's first picture; under it, the theme's page.
                colour = str(page) if str(page).startswith("#") else deck.palette.get("canvas")
                slide.background.fill.solid()
                slide.background.fill.fore_color.rgb = RGBColor.from_string(colour.lstrip("#").upper())
            tree = slide.shapes._spTree
            drawn = {"background": False, "groups": groups, "pictures": slide_pictures(slide),
                     "backdrop": deck.palette.get("canvas")}
            for shape in under:
                add_drawing(tree, Drawing(drawing.width, drawing.height, Group(None, [shape])), Placement(), **drawn)
            worded = {words.id: words for words in item.worded}
            span = (deck.style.margin, deck.style.width - deck.style.margin)
            # The headings' words as written, for the breaks typed in them.
            written = {f"{item.slide.id}.title": item.slide.title_runs,
                       f"{item.slide.id}.subtitle": item.slide.subtitle_runs}
            under = any(kind in {"subTitle", "body"} for _, kind in headings)
            for heading, kind in headings:
                runs = written.get(heading.id or "")
                add_heading(
                    tree, heading, kind, words=worded.get(heading.id), formulas=_formulas(heading)
                    if isinstance(heading, Group) else None, deck=deck, palette=deck.palette, editable=editable_maths,
                    span=span, source="".join(run.text for run in runs) if runs else None,
                    above=under and kind in {"ctrTitle", "title"},
                )
            add_drawing(tree, drawing, Placement(), **drawn)
            _describe(tree, item.slide)
            if editable_maths:
                editable_maths_of(tree, drawing, item.worded, deck, deck.palette)
            clicks: list[tuple[int, int, tuple[int, int] | None]] = []
            for layout in item.lists:
                shape_id = add_list(tree, deck, layout, formulas=formulas.get(layout.id), editable=editable_maths)
                outer = [index for index, (level, _, _) in enumerate(layout.items) if level == 0]
                if layout.reveal and outer:
                    # Each outer item appears with the items under it; any before the first
                    # are there from the start, as in the PDF.
                    ends = [*outer[1:], len(layout.items)]
                    ranges = [(first, end - 1) for first, end in zip(outer, ends, strict=True)]
                    steps = layout.stepped if len(layout.stepped) == len(ranges) else range(len(ranges))
                    clicks += [(step, shape_id, one) for step, one in zip(steps, ranges, strict=True)]
            for layout in item.tables:
                add_table(tree, deck, layout, formulas=formulas.get(layout.id), editable=editable_maths)
            # A block that builds in appears whole, on its click: a table with its caption.
            # Only a shape on the slide itself: one within a group is not animated by itself.
            shapes: dict[str, int] = {}
            for child in tree:
                if (named := next(child.iter(f"{{{_PML}}}cNvPr"), None)) is not None:
                    shapes.setdefault(named.get("name", ""), int(named.get("id", "0")))
            clicks += [(step, shapes[shown], None) for step, name in item.builds
                       for shown in (name, f"{name}.caption") if shown in shapes]
            if item.slide.notes_text:
                slide.notes_slide.notes_text_frame.text = item.slide.notes_text
            steps = sorted({step for step, _, _ in clicks})
            add_reveals(slide._element, [[(shape_id, one) for at, shape_id, one in clicks if at == step]
                                         for step in steps])
            _read_in_order(tree, item.slide)
            # Last, as nothing more finds a shape by its id: each named as a person would.
            plain_names(tree, {f"{item.slide.id}.{region.name}.{index}": _BLOCK_NAMES.get(type(block).__name__, "Group")
                               for region in item.slide.regions.values() for index, block in enumerate(region.blocks)})
    buffer = BytesIO()
    presentation.save(buffer)
    target.write_bytes(buffer.getvalue())
    return target


_PML = "http://schemas.openxmlformats.org/presentationml/2006/main"


def _read_in_order(tree, slide: Slide) -> None:
    """A slide's shapes in the order they are read (a slide program's Selection Pane, a
    screen reader), as the PDF reads them: its headings and what is drawn under them, then
    its body top to bottom (a left column before a right) -- its lists and tables where they
    are drawn, not after the rest -- then its footnotes, then its number and footer."""

    import re

    regions = {name: order for order, name in enumerate(slide.regions)}
    block = re.compile(rf"{re.escape(slide.id)}\.([^.]+)\.(\d+)(?:\..*)?")

    def rank(child) -> tuple[int, int, int, bool]:
        named = next(child.iter(f"{{{_PML}}}cNvPr"), None)
        name = named.get("name", "") if named is not None else ""
        rest = name.removeprefix(f"{slide.id}.")
        if rest in {"number", "footer"}:
            return 3, 0, 0, False
        if rest.startswith("footnote"):
            return 2, 0, 0, False
        found = block.fullmatch(name)
        if found and found.group(1) in regions:
            # A table's caption after the table, as it is drawn under it.
            return 1, regions[found.group(1)], int(found.group(2)), name.endswith(".caption")
        return 0, 0, 0, False

    shapes = [child for child in tree if child.tag not in {f"{{{_PML}}}nvGrpSpPr", f"{{{_PML}}}grpSpPr"}]
    for child in sorted(shapes, key=rank):
        tree.append(child)


def _spec_of(block: _Figure):
    spec = block.figure
    return spec if hasattr(spec, "nodes") else getattr(spec, "spec", None)


def _said_words(runs: object) -> str:
    text = runs if isinstance(runs, str) else "".join(getattr(run, "text", "") for run in runs or ())
    return " ".join(text.split())


_REGULATES = {
    "inhibition": "inhibits", "catalysis": "catalyses", "stimulation": "stimulates",
    "necessary": "is necessary for", "modulation": "modulates",
}
"""What a line with each of SBGN's heads says one thing does to another: not "→", which reads
as "goes to" (or "activates") whatever the line's head means."""


def _figure_said(block: _Figure) -> str:
    """A figure's alt text, from its shapes: a figure of lines read along them, in the
    order they flow, each branch said with its words ("A flow chart: Purify CA → Mix CA
    with IP6 → Tubes formed? (yes: Cryo-EM grids; no: back to Mix CA with IP6)"), and what
    each regulating line does ("SPINK1 inhibits Trypsin"); one of shapes alone by their
    words, in order, a drawn thing by its parts (a timeline's in time order); one of
    molecules alone by what each shows."""

    spec = _spec_of(block)
    nodes = list(getattr(spec, "nodes", ()) or ())
    if not nodes:
        return ""
    if all(node.kind == "structure" for node in nodes):
        return " ".join(_structures_said(block).values())
    names = {node.id: _said_words(node.label) or node.id for node in nodes}
    for group in getattr(spec, "groups", ()) or ():
        names.setdefault(group.id, _said_words(getattr(group, "text", "")) or group.id)
    onward: dict[str, list[tuple[str, str]]] = {}
    regulated: list[str] = []
    for edge in getattr(spec, "edges", ()) or ():
        source, target, label = edge.source.node_id, edge.target.node_id, _said_words(edge.label)
        if (verb := _REGULATES.get(getattr(edge, "head", "arrow"))) is not None:
            # A line that regulates says what it does, in words, after the lines that flow.
            said = f"{names.get(source, source)} {verb} {names.get(target, target)}"
            regulated.append(f"{said} ({label})" if label else said)
            continue
        onward.setdefault(source, []).append((target, label))
    for net in getattr(spec, "nets", ()) or ():
        for source in net.sources:
            for index, target in enumerate(net.targets):
                label = _said_words(net.label) if not index else ""
                onward.setdefault(source.node_id, []).append((target.node_id, label))
    # A drawn thing in it (a timeline, a protein) is read by its parts, a sentence of its own.
    drawn = " ".join(f"{said[:1].upper()}{said[1:]}." for node in nodes if (said := _drawn_said(node)))
    named = {*(target for branches in onward.values() for target, _ in branches), *onward}
    if not onward and not regulated:
        words = [names[node.id] for node in nodes if not _drawn_said(node) and _said_words(node.label)]
        return " ".join(filter(None, [f"A figure: {', '.join(words)}." if words else "", drawn]))
    seen: set[str] = set()

    def along(node: str, path: tuple[str, ...]) -> str:
        seen.add(node)
        said = names.get(node, node)
        branches = onward.get(node, [])

        def towards(target: str) -> str:
            if target in path or target == node:
                return f"back to {names.get(target, target)}"
            return names.get(target, target) if target in seen else along(target, (*path, node))

        if len(branches) == 1:
            target, label = branches[0]
            return f"{said}{f' ({label})' if label else ''} → {towards(target)}"
        if branches:
            ways = [f"{label}: {towards(target)}" if label else towards(target) for target, label in branches]
            # A question's branches are said by their answers; a fan's, as where it leads.
            onto = " → " if not any(label for _, label in branches) else " "
            return f"{said}{onto}({'; '.join(ways)})"
        return said

    reached = {target for branches in onward.values() for target, _ in branches}
    starts = [node.id for node in nodes if node.id in onward and node.id not in reached]
    walks = [along(start, ()) for start in starts]
    walks += [along(node.id, ()) for node in nodes if node.id not in seen and node.id in onward]
    walks += regulated
    mentioned = {*seen, *(node.id for node in nodes if any(names[node.id] in said for said in regulated))}
    walks += [names[node.id] for node in nodes
              if node.id not in mentioned and node.id not in named and _said_words(node.label)]
    kind = "flow chart" if any(node.kind in {"terminal", "decision"} for node in nodes) else "figure"
    return " ".join(filter(None, [f"A {kind}: {'; '.join(walks)}", drawn]))


_TO = "\u2013"
"""The dash a range of numbers is said with (1 to 15, as 1\u201315)."""

_PART_NAMES = {
    "promoter": "promoter", "rbs": "ribosome binding site", "cds": "coding sequence", "gene": "gene",
    "terminator": "terminator", "operator": "operator", "origin": "origin", "insulator": "insulator",
}
"""A construct's or plasmid's part, by its type, as said."""


def _drawn_said(node) -> str:
    """A drawn thing (a timeline, a protein, a construct, a plasmid, a well plate) in words:
    what it is, its name, and its parts in their order -- a timeline's events and spans in
    time order, a protein's domains and sites by residue. Empty for anything else."""

    def records(name: str) -> list[dict]:
        return [dict(getattr(item, "items", item)) for item in node.property(name) or ()]

    def number(value: object) -> str:
        return f"{value:g}" if isinstance(value, int | float) else str(value)

    name = _said_words(node.label)
    if node.kind == "timeline":
        unit = str(node.property("unit") or "")
        at = [(float(item.get("at", 0)), 0, f"{number(item.get('at', 0))} {_said_words(item.get('label', ''))}")
              for item in records("events")]
        spans = [(float(item.get("start", 0)), 1, f"{number(item.get('start', 0))}{_TO}{number(item.get('end', 0))} "
                  f"{_said_words(item.get('label', ''))}") for item in records("spans")]
        said = "; ".join(text.strip() for _, _, text in sorted(at + spans, key=lambda item: item[:2]))
        head = f"a timeline{f', {name}' if name else ''}{f', in {unit}' if unit else ''}"
        return f"{head}: {said}" if said else head
    if node.kind == "protein":
        parts = []
        # From its first residue to its last, a site among the domains where it is.
        for item in sorted(records("features"), key=lambda item: float(item.get("at", item.get("start", 0)))):
            label = _said_words(item.get("label", "")) or str(item.get("type", "feature"))
            if "at" in item:
                parts.append(f"{label} at {number(item['at'])}")
            elif "start" in item:
                parts.append(f"{label} {number(item['start'])}{_TO}{number(item.get('end', item['start']))}")
            else:
                parts.append(label)
        length = node.property("length")
        head = f"a protein{f', {name}' if name else ''}{f', {number(length)} residues' if length else ''}"
        return f"{head}: {'; '.join(parts)}" if parts else head
    if node.kind in {"construct", "plasmid"}:
        parts = []
        for item in records("parts" if node.kind == "construct" else "features"):
            kind = _PART_NAMES.get(str(item.get("type", "")), str(item.get("type", "")))
            label = _said_words(item.get("label", ""))
            parts.append(f"{label} ({kind})" if label and kind and kind.lower() != label.lower() else label or kind)
        length = node.property("length") if node.kind == "plasmid" else None
        head = f"a {node.kind}{f', {name}' if name else ''}{f', {number(length)} bp' if length else ''}"
        return f"{head}: {' → '.join(parts) if node.kind == 'construct' else '; '.join(parts)}" if parts else head
    if node.kind == "wellplate":
        wells = node.property("wells") or 96
        groups = [f"{_said_words(item.get('label', ''))} in {item.get('wells', '')}".strip()
                  for item in records("groups")]
        head = f"a {number(wells)}-well plate{f', {name}' if name else ''}"
        return f"{head}: {'; '.join(groups)}" if groups else head
    return ""


def _structures_said(block: _Figure) -> dict[str, str]:
    """The alt text of each molecule a figure draws as a picture, by the picture's id in
    the figure: what it is, by its file's own title where it has one, and which entry it is
    ("Molecular structure: HIV capsid C-terminal domain (1A8O)") -- or what the figure calls
    it, where the file says nothing."""

    import re

    said = {}
    for node in getattr(_spec_of(block), "nodes", ()) or ():
        if getattr(node, "kind", "") != "structure":
            continue
        source = str(node.property("source") or "").strip()
        name = Path(source).name
        stem = re.sub(r"\.(pdb|cif|mmcif|ent)$", "", name, flags=re.IGNORECASE)
        entry = stem.upper() if re.fullmatch(r"[0-9][A-Za-z0-9]{3}", stem) else ""
        title = _structure_title(source)
        label = _said_words(node.label)
        # The figure's own name for it, unless that is only the entry's id or the title.
        own = label if label and label.upper() not in {entry, title.upper()} else ""
        named = f"{own}: {title}" if own and title else own or title
        if not named:
            text = f"Molecular structure {entry}." if entry else f"Molecular structure, from {name}."
        else:
            text = f"Molecular structure: {named}{f' ({entry})' if entry and entry not in named.upper() else ''}."
        said[f"{node.id}.molecule"] = text
    return said


_SMALL_WORDS = frozenset({
    "A", "AN", "THE", "OF", "IN", "ON", "AT", "TO", "BY", "FOR", "AND", "OR", "WITH", "FROM", "ITS", "AS",
    "IS", "ARE", "INTO", "ONTO", "VIA", "PER", "BOUND",
})
"""Words a title set in capitals has that are words, not names (HIV, DNA, E2)."""


def _structure_title(source: str) -> str:
    """The title a structure's file gives it (a PDB file's TITLE, an mmCIF file's
    ``_struct.title``, else its molecule's name), in sentence case where it is set in
    capitals, as the Protein Data Bank's files are: names (HIV, DNA, E2) stay as they are."""

    import re

    try:
        with open(source, encoding="utf-8", errors="replace") as handle:
            head = handle.read(400_000)
    except OSError:
        return ""
    title = " ".join(re.sub(r"^TITLE\s+(\d+\s)?", "", line).strip()
                     for line in head.splitlines() if line.startswith("TITLE "))
    if not title:
        found = re.search(r"^_struct\.title\s+(?:'([^']*)'|\"([^\"]*)\"|;\s*\n(.*?)\n;|(\S+))", head, re.M | re.S)
        title = next((group for group in found.groups() if group), "") if found else ""
    if not title:
        found = re.search(r"^COMPND\s+\d*\s*MOLECULE:\s*([^;\n]*)", head, re.M)
        title = found.group(1) if found else ""
    title = " ".join(title.split()).strip(" .")
    if not title or title.upper() != title:
        return title

    def word(part: str) -> str:
        if any(character.isdigit() for character in part) or (len(part) <= 3 and part not in _SMALL_WORDS):
            return part
        return part.lower()

    said = " ".join("-".join(word(part) for part in token.split("-")) for token in title.split())
    return said[:1].upper() + said[1:]


def _describe(tree: etree._Element, slide: Slide) -> None:
    """Each picture's description as its alt text, on the shape drawn for it: a figure's,
    and each molecule's in it, as well as a picture's own."""

    described = _descriptions(slide)
    for properties in tree.iter(f"{{{_PML}}}cNvPr") if described else ():
        name = properties.get("name", "")
        # (A picture with a caption is grouped with it: described as its group is.)
        text = described.get(name) or described.get(name.removesuffix(".picture"))
        if not text:
            continue
        shape = properties.getparent().getparent()
        above = [group.find(f"{{{_PML}}}nvGrpSpPr/{{{_PML}}}cNvPr")
                 for group in shape.iterancestors(f"{{{_PML}}}grpSp")]
        if any(text in (group.get("descr") or "") for group in above if group is not None):
            # Said by the group it is in (a picture with its caption, a figure of one
            # molecule): marked decorative, so a screen reader says it once.
            _decorative(properties)
        else:
            properties.set("descr", text)


_DECORATIVE = "http://schemas.microsoft.com/office/drawing/2017/decorative"


def _decorative(properties: etree._Element) -> None:
    """A shape marked decorative, as PowerPoint's "Mark as Decorative" does: passed over by
    a screen reader, and not asked for alt text."""

    a = "http://schemas.openxmlformats.org/drawingml/2006/main"
    extensions = etree.SubElement(properties, f"{{{a}}}extLst")
    extension = etree.SubElement(extensions, f"{{{a}}}ext", uri="{C183D7F6-B498-43B3-948B-1728B52AA6E4}")
    etree.SubElement(extension, f"{{{_DECORATIVE}}}decorative", nsmap={"adec": _DECORATIVE}, val="1")


def _descriptions(slide: Slide) -> dict[str, str]:
    """What each picture on a slide shows, in words, by the id it is drawn with: a
    picture's own description, a figure's words, each molecule's in a figure."""

    described = {
        # A description given (a picture's, a figure's) is said in place of one made up.
        f"{slide.id}.{region.name}.{index}": block.description or (
            _figure_said(block) if isinstance(block, _Figure) else "")
        for region in slide.regions.values()
        for index, block in enumerate(region.blocks)
        if isinstance(block, _Image | _Figure)
    }
    for region in slide.regions.values():
        for index, block in enumerate(region.blocks):
            if isinstance(block, _Figure):
                molecules = _structures_said(block)
                if block.description and len(molecules) == 1:
                    molecules = dict.fromkeys(molecules, block.description)
                described |= {f"{slide.id}.{region.name}.{index}.{key}": text for key, text in molecules.items()}
    return {key: text for key, text in described.items() if text}


_BLOCK_NAMES = {
    "_Bullets": "List", "_Words": "Text", "_Figure": "Figure", "_Image": "Picture", "_Plot": "Plot", "_Table": "Table",
    "_Gallery": "Gallery", "_Code": "Code", "_Quote": "Quote", "_Stats": "Numbers", "_Callout": "Callout",
    "_Missing": "Missing Picture", "_Math": "Equation",
}
"""What PowerPoint's Selection Pane calls each kind of block on a slide."""

_BLOCKS = {"figure": "Figure", "bullets": "L", "table": "Table", "code": "P", "quote": "BlockQuote",
           "stats": "Div", "callout": "Div", "gallery": "Div", "missing": "Div", "invalid": "Div",
           "picture": "Figure"}
"""The structure element each kind of block on a slide is, in a tagged PDF."""


def _tagger(rendered: list[RenderedSlide]) -> Tagger:
    """How a deck's PDF is tagged (``flexo.pdf.Tagger``), as a screen reader reads a
    slide: its title a heading, then its blocks in order -- paragraphs, lists of items in
    their levels, tables of rows and header cells, figures and pictures described in words,
    formulas -- and its bands, rules, marks and page numbers passed over."""

    import re

    described: dict[str, str] = {}
    levels: dict[str, list[int]] = {}
    headed: dict[str, bool] = {}
    for item in rendered:
        described |= _descriptions(item.slide)
        levels |= {layout.id: [level for level, _, _ in layout.items] for layout in item.lists}
        headed |= {layout.id: layout.header for layout in item.tables}

    def holder(within: Tag | None, kind: str) -> tuple[tuple[str, str], ...] | None:
        """The path to the innermost element of ``kind`` ``within`` is in, if any."""

        path = within.path if within is not None else ()
        found = [index for index, (named, _) in enumerate(path) if named == kind]
        return path[: found[-1] + 1] if found else None

    def listed(base: tuple[tuple[str, str], ...], ident: str) -> Tag | None:
        # An item's words, or its bullet or number: in its item, within the item it is under.
        owner = base[-1][1]
        match = re.fullmatch(rf"{re.escape(owner)}\.(\d+)(\.mark)?", ident)
        if match is None:
            return None
        index, steps = int(match.group(1)), levels.get(owner, [])
        chain = [index]
        while chain[0] < len(steps) and (above := next(
                (j for j in range(chain[0] - 1, -1, -1) if steps[j] < steps[chain[0]]), None)) is not None:
            chain.insert(0, above)
        path = list(base)
        for depth, at in enumerate(chain):
            if depth:
                path.append(("L", f"{owner}.{chain[depth - 1]}.list"))
            path.append(("LI", f"{owner}.{at}"))
        path.append(("Lbl", f"{ident}") if match.group(2) else ("LBody", f"{ident}.body"))
        return Tag(tuple(path))

    def tag(item: object, within: Tag | None) -> Tag | None:
        ident = getattr(item, "id", None) or ""
        role = ident.split(".", 1)[1] if "." in ident else ""
        inside = within.path[-1][0] if within is not None and within.path else None
        owner = holder(within, "Figure") or holder(within, "Table")
        if owner is not None and ident == f"{owner[-1][1]}.caption":
            # A picture's or table's caption: its own, read with it.
            return Tag((*owner, ("Caption", ident)))
        if isinstance(item, Group):
            block = _BLOCKS.get(item.data.get("data-flexo-talk", ""))
            if block is not None and inside not in {"Figure", "Formula"}:
                # A figure is described by its words, and by what each molecule in it shows;
                # a picture, by what it is said to show.
                alt = described.get(ident, "A figure" if item.data.get("data-flexo-talk") != "picture" else "")
                alt = " ".join([alt, *(text for key, text in described.items()
                                       if key.startswith(f"{ident}.") and text not in alt)])
                return Tag((*(within.path if within else ()), (block, ident)), alt=alt if block == "Figure" else "")
            source = item.data.get("data-flexo-math", "")
            if source.startswith("\\displaystyle") and inside in {"P", "LBody"}:
                # A formula displayed on a line of its own in words ($$...$$) is a formula
                # within them, described by its words, as an equation on its own is.
                from flexo.texmath import linear

                key = ident or f"{within.path[-1][1]}.formula{id(item)}"
                return Tag((*within.path, ("Formula", key)), alt=linear(source))
            # Words with a formula in them (a title's, a paragraph's, an item's) are tagged as
            # words are, the formula read in its place.
            kinds = {"data-flexo-talk", "data-flexo-entity", "data-flexo-math"}

            def formula_among(group: Group) -> bool:
                # In its words, or in one of their lines (``{id}.1``) set each on its own.
                return any(isinstance(child, Group) and (
                    "data-flexo-math" in child.data or (
                        bool(group.id) and (child.id or "").startswith(f"{group.id}.")
                        and not kinds & set(child.data) and formula_among(child)))
                    for child in group.items)

            worded = not kinds & set(item.data) and formula_among(item)
            if not (worded and role):
                return None
        if isinstance(item, Shape):
            return None if inside in {"Figure", "Formula"} else ARTIFACT
        if isinstance(item, Image):
            if inside in {"Figure", "Formula"}:
                return None
            if ident in described or re.fullmatch(r"slide\d+\.[\w-]+\.\d+", ident):
                return Tag((*(within.path if within else ()), ("Figure", ident)), alt=described.get(ident, ""))
            return ARTIFACT
        # Words: a heading, a page's furniture, an item of a list, a cell of a table; the
        # words within words (a title's, with a formula in it) are theirs.
        if inside in {"Figure", "Formula", "H1", "P", "Lbl", "LBody", "TH", "TD", "Caption"}:
            return None
        if role in {"number", "footer"} or (role.endswith(".mark") and holder(within, "L") is None):
            # A page's number and footer, a quotation's mark: furniture, not words to read.
            return ARTIFACT
        if role == "title":
            return Tag((("H1", ident),))
        if (base := holder(within, "L")) is not None and (found := listed(base, ident)):
            return found
        if (base := holder(within, "Table")) is not None:
            match = re.fullmatch(rf"{re.escape(base[-1][1])}\.(\d+)\.(\d+)", ident)
            if match:
                row = int(match.group(1))
                cell = "TH" if headed.get(base[-1][1]) and row == 0 else "TD"
                return Tag((*base, ("TR", f"{base[-1][1]}.{row}"), (cell, ident)))
            return ARTIFACT
        if within is not None and within.path:
            return Tag((*within.path, ("P", ident)))
        return Tag((("P", ident),)) if isinstance(item, Group) else None

    return tag


_HEADINGS = {
    "title": {"title": "ctrTitle", "subtitle": "subTitle"},
    "section": {"title": "title", "subtitle": "body"},
    # A statement's words are its title, as on Keynote's Statement layout: what an outline
    # and a screen reader name the slide by.
    "statement": {"title": "title"},
}
"""The placeholder (``pptx.PLACEHOLDER_KINDS``) each heading a slide draws fills in
PowerPoint, by the slide's layout and the heading's id after the slide's: on any other
layout, a title and a subtitle."""

_LAYOUTS = {"title": "Title Slide", "section": "Section Header", "statement": "Title Only"}
"""The default template's layout whose placeholders a slide of each layout fills."""


def _take_headings(root: Group, item: RenderedSlide) -> tuple[list, list[tuple[Text | Group, str]]]:
    """The slide's headings (``_HEADINGS``), each with its placeholder, taken out of its
    drawing to be written as placeholders (``add_heading``) -- first on the slide, as a
    slide's title is first read -- and what is drawn under them taken out with them, to
    go before them: a band, a backdrop. Nothing changes its place in front of or behind
    anything it overlaps."""

    slide = item.slide
    kinds = _HEADINGS.get(slide.layout, {"title": "title", "subtitle": "subTitle"})
    worded = {words.id for words in item.worded}
    prefix = f"{slide.id}."

    def kind_of(child: object) -> str | None:
        name = getattr(child, "id", None) or ""
        kind = kinds.get(name.removeprefix(prefix)) if name.startswith(prefix) else None
        if isinstance(child, Text) and not child.angle and any(
            run.text.strip() for line in child.lines for run in line.runs
        ):
            return kind
        # Words with maths in them, set as words and formulas.
        return kind if isinstance(child, Group) and name in worded else None

    def holder(group: Group) -> Group | None:
        if any(kind_of(child) for child in group.items):
            return group
        return next((found for child in group.items if isinstance(child, Group) and (found := holder(child))), None)

    parent = holder(root)
    if parent is None:
        return [], []
    headings = [(child, kind) for child in parent.items if (kind := kind_of(child))]
    first = next(index for index, child in enumerate(parent.items) if child is headings[0][0])
    # What is drawn before the headings and overlaps them, or overlaps what does, goes under them.
    boxes = [_bounds(heading, 1.0) for heading, _ in headings]
    under: list[int] = []
    for index in reversed(range(first)):
        box = _bounds(parent.items[index])
        if any(_overlap(box, other) for other in boxes):
            under.insert(0, index)
            boxes.append(box)
    taken = {id(heading) for heading, _ in headings} | {id(parent.items[index]) for index in under}
    lifted = [parent.items[index] for index in under]
    parent.items = [child for child in parent.items if id(child) not in taken]
    return lifted, headings


def _bounds(item: object, margin: float = 0.0) -> tuple[float, float, float, float]:
    left, top, right, bottom = ink_bounds(Drawing(0.0, 0.0, Group(None, [item])))  # type: ignore[list-item]
    return left - margin, top - margin, right + margin, bottom + margin


def _overlap(one: tuple[float, float, float, float], other: tuple[float, float, float, float]) -> bool:
    # What draws nothing (bounds of no size) lies over nothing.
    drawn = all(box[2] > box[0] and box[3] > box[1] for box in (one, other))
    return drawn and one[0] < other[2] and other[0] < one[2] and one[1] < other[3] and other[1] < one[3]


def _named(deck: Deck) -> tuple[str, str]:
    """The talk's title and author, as its exports name them: its title slide's (else its
    first titled slide's title), and no author but the title slide's."""

    opening = next((slide for slide in deck.slides if slide.layout == "title"), None)
    titled = opening or next((slide for slide in deck.slides if slide.title_runs), None)
    title = "".join(run.text for run in titled.title_runs) if titled else ""
    return title, str((opening.source.get("author") if opening else "") or "")


def _properties(presentation: Presentation, deck: Deck) -> None:
    """The file's own properties, as PowerPoint shows them in its Properties: the talk's
    title and author, made now -- never python-pptx's template's (its author, its 2013
    dates, "generated using python-pptx")."""

    import datetime

    properties = presentation.core_properties
    properties.title, properties.author = _named(deck)
    properties.title = properties.title or str(deck.id)
    properties.last_modified_by = properties.author
    properties.comments = ""
    properties.subject = ""
    properties.keywords = ""
    properties.revision = 1
    now = datetime.datetime.now(datetime.UTC).replace(microsecond=0)
    properties.created = properties.modified = now


def _drop_lists(group: Group, found: dict[str, list] | None = None) -> dict[str, list]:
    """Bulleted lists and tables are set natively (``add_list``, ``add_table``), so
    their drawn copies go -- all but their formulas, which native words leave room for
    and cannot draw: each list's and table's, by its id, for its writer."""

    found = {} if found is None else found
    kept: list = []
    for item in group.items:
        if isinstance(item, Group) and item.data.get("data-flexo-talk") in {"bullets", "table"}:
            if item.id:
                found[item.id] = _formulas(item)
            # A table's caption stays, words after the table set natively.
            kept += [child for child in item.items if item.id and getattr(child, "id", None) == f"{item.id}.caption"]
        else:
            kept.append(item)
    group.items = kept
    for item in group.items:
        if isinstance(item, Group):
            _drop_lists(item, found)
    return found


def _formulas(group: Group) -> list[Group]:
    found: list[Group] = []
    for item in group.items:
        if isinstance(item, Group):
            if "data-flexo-math" in item.data:
                found.append(item)
            else:
                found.extend(_formulas(item))
    return found


def _dissolve_regions(group: Group) -> None:
    """A slide's regions are groups only so the layout can move each whole; in the
    PowerPoint their shapes stand on the slide, as if placed there by hand."""

    items: list = []
    for item in group.items:
        if isinstance(item, Group):
            _dissolve_regions(item)
            if item.data.get("data-flexo-talk") == "region":
                items.extend(item.items)
                continue
        items.append(item)
    group.items = items


_ = etree
