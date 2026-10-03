"""PowerPoint shapes from a Flexo drawing: every box, line, and word native and editable.

A drawing (``flexo.drawing``) is placed straight into a slide's shape tree:

- a rectangle is a rectangle or a rounded rectangle, an ellipse an ellipse, so
  they stay editable as what they are; any other path is a freeform shape with
  the same lines and curves;
- an arrowhead is drawn exactly, as a small filled shape grouped with its line
  (PowerPoint's own line ends come in three sizes and would not match);
- text is a text box per label, its lines and runs as written -- italic,
  weight, face, colour, raised scripts -- placed so each baseline lands where
  Flexo set it. A label whose runs step back over one another (stacked
  scripts, an accent's mark) is set as one box per run instead;
- groups keep Flexo's groups and ids, so a component moves as one.

Coordinates are points on the drawing; ``place`` maps them onto the slide (an
offset and a scale), so a figure can be put anywhere on a slide at any size.
"""

from __future__ import annotations

from collections.abc import Callable
from contextlib import contextmanager
from contextvars import ContextVar
from dataclasses import dataclass, replace
from functools import cache
from itertools import pairwise
from xml.sax.saxutils import escape

from flexo.drawing import Drawing, Group, Image, Paint, Segment, Shape, Text
from lxml import etree

EMU_PER_POINT = 12700
_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_NS = f'xmlns:a="{_A}" xmlns:p="{_P}" xmlns:r="{_R}"'

ASCENT = 0.8
"""Where a text box's first baseline sits below its top, as a fraction of the
line spacing, when the spacing is set exactly -- for a face whose metrics are
unknown. A known face uses its own share (``ascent``)."""

LOOSE = 0.75
"""Where PowerPoint sets the first baseline of lines spaced wider than single (lines
opened for a tall formula), as a fraction of the spacing, whatever the face -- as
measured against PowerPoint 16 for Mac, to the nearest point."""


def ascent(face) -> float:
    """Where a slide program puts the first baseline of exactly spaced lines in
    ``face``, as a share of the spacing: the face's ascent over its full height
    (0.79 for Figtree, 0.67 for Kalam), as measured against ONLYOFFICE."""

    if face is None:
        return ASCENT
    from flexo.fonts import load_face

    loaded = load_face(face)
    total = loaded.ascent + loaded.descent
    return loaded.ascent / total if total > 0 else ASCENT


def baseline_down(face, spacing: float, plain: float) -> float:
    """How far below the top of an exactly spaced line a slide program sets its baseline:
    ``spacing`` as the words alone are spaced (``plain``), or opened for a formula."""

    return LOOSE * spacing if spacing > plain + 0.01 else ascent(face) * spacing

SCRIPT_SCALE = 0.65
"""How much smaller a slide program draws a raised or lowered run than its size
says (measured in ONLYOFFICE, which follows PowerPoint): a script is written at
its own size divided by this, so it is drawn at its own size."""

DASHES = {
    "sysDot": (1.0, 1.0),
    "sysDash": (3.0, 1.0),
    "dash": (4.0, 3.0),
    "lgDash": (8.0, 3.0),
    "dot": (1.0, 3.0),
}
"""PowerPoint's preset dashes, as (dash, gap) in line widths. Presets, not
custom dashes: every slide program draws them, and some draw custom ones solid."""


@dataclass(frozen=True, slots=True)
class Placement:
    """Where a drawing goes on a slide: points on the drawing to points on the slide."""

    x: float = 0.0
    y: float = 0.0
    scale: float = 1.0

    def point(self, x: float, y: float) -> tuple[int, int]:
        return (
            round((self.x + x * self.scale) * EMU_PER_POINT),
            round((self.y + y * self.scale) * EMU_PER_POINT),
        )

    def length(self, value: float) -> int:
        return round(value * self.scale * EMU_PER_POINT)


type Pictures = Callable[[bytes, str], str]
"""Adds a picture's bytes (and media type) to the slide; returns its relationship id."""


class _Ids:
    def __init__(
        self, start: int, pictures: Pictures | None = None, backdrop: str | None = None
    ) -> None:
        self.next = start
        self.pictures = pictures
        self.backdrop = backdrop
        """The page colour, for drawing a multiply blend (which slide programs lack)."""

    def __call__(self) -> int:
        self.next += 1
        return self.next


def add_drawing(
    tree: etree._Element,
    drawing: Drawing,
    placement: Placement | None = None,
    *,
    name: str | None = None,
    background: bool = True,
    groups: bool = True,
    pictures: Pictures | None = None,
    backdrop: str | None = None,
) -> None:
    """Append ``drawing`` to a slide's shape tree (``slide.shapes._spTree``).

    ``groups=False`` lays every shape flat on the slide instead of in Flexo's
    groups -- for renderers that cannot draw a freeform inside a group (macOS
    Quick Look). ``pictures`` adds an image's bytes to the slide (see
    ``slide_pictures``); without it, pictures are left out.
    """

    placement = placement or Placement()
    existing = [int(item) for item in tree.xpath(".//@id") if str(item).isdigit()]
    ids = _Ids(max(existing, default=1), pictures, backdrop)
    items = [
        item
        for item in drawing.root.items
        if background or not (isinstance(item, Group) and item.id == "layer.background")
    ]
    group = Group(name or drawing.root.id, items)
    element = _group(group, placement, ids)
    if element is None:
        return
    if groups:
        tree.append(element)
        return
    for leaf in element.iter(f"{{{_P}}}sp"):
        tree.append(leaf)


def _shown_name(group: Group) -> str:
    """What PowerPoint's Selection Pane calls a group: a figure's shape by its words
    ("Orders API"), a line as a line; anything else by its id, which other writers here
    find it by."""

    entity = group.data.get("data-flexo-entity")
    if entity == "component":
        words: list[str] = []

        def gather(item: object) -> None:
            if isinstance(item, Text):
                words.append("".join(run.text for line in item.lines for run in line.runs))
            elif isinstance(item, Group):
                for child in item.items:
                    gather(child)

        gather(group)
        said = " ".join(" ".join(words).split())
        if said:
            return said if len(said) <= 60 else f"{said[:59]}…"
    if entity in {"connector", "net"}:
        return "Line"
    return group.id or "group"


def _group(group: Group, placement: Placement, ids: _Ids) -> etree._Element | None:
    children = [child for item in group.items if (child := _item(item, placement, ids)) is not None]
    if not children:
        return None
    if len(children) == 1 and not group.id:
        return children[0]
    boxes = [_extent(child) for child in children]
    left = min(box[0] for box in boxes)
    top = min(box[1] for box in boxes)
    right = max(box[0] + box[2] for box in boxes)
    bottom = max(box[1] + box[3] for box in boxes)
    element = etree.fromstring(
        f"<p:grpSp {_NS}><p:nvGrpSpPr><p:cNvPr id=\"{ids()}\" name=\"{escape(_shown_name(group), {'"': '&quot;'})}\"/>"
        f"<p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr><p:grpSpPr><a:xfrm>"
        f'<a:off x="{left}" y="{top}"/><a:ext cx="{right - left}" cy="{bottom - top}"/>'
        f'<a:chOff x="{left}" y="{top}"/><a:chExt cx="{right - left}" cy="{bottom - top}"/>'
        f"</a:xfrm></p:grpSpPr></p:grpSp>"
    )
    for child in children:
        element.append(child)
    return element


def _extent(element: etree._Element) -> tuple[int, int, int, int]:
    # The shape's own transform: a picture's blip also holds an <a:ext> (its extensions).
    xfrm = element.find(f".//{{{_A}}}xfrm")
    if xfrm is None:
        xfrm = element.find(f".//{{{_P}}}xfrm")
    off = xfrm.find(f"{{{_A}}}off")
    ext = xfrm.find(f"{{{_A}}}ext")
    return int(off.get("x")), int(off.get("y")), int(ext.get("cx")), int(ext.get("cy"))


def _item(item: Shape | Text | Image | Group, placement: Placement, ids: _Ids):
    if isinstance(item, Group):
        return _group(item, placement, ids)
    if isinstance(item, Shape):
        return _shape(item, placement, ids)
    if isinstance(item, Text):
        return _text(item, placement, ids)
    if isinstance(item, Image):
        return _picture(item, placement, ids)
    return None


def slide_pictures(slide) -> Pictures:
    """A ``Pictures`` for a python-pptx slide: PNG and JPEG as they are; an SVG
    as a PNG drawn at print resolution with the SVG beside it, which PowerPoint
    shows (and can turn into shapes) and other programs replace by the PNG."""

    from io import BytesIO

    from pptx.opc.constants import RELATIONSHIP_TYPE
    from pptx.opc.package import Part

    def add(data: bytes, mime: str) -> str:
        if mime == "image/svg+xml":
            raise ValueError("an SVG picture is added with its PNG: use add_svg")
        return slide.part.get_or_add_image_part(BytesIO(data))[1]

    def add_svg(data: bytes) -> str:
        partname = slide.part.package.next_image_partname("svg")
        part = Part(partname, "image/svg+xml", slide.part.package, data)
        return slide.part.relate_to(part, RELATIONSHIP_TYPE.IMAGE)

    add.svg = add_svg  # type: ignore[attr-defined]
    return add


def _jpeg_size(data: bytes) -> tuple[int, int] | None:
    import struct

    index = 2
    while data.startswith(b"\xff\xd8") and index + 9 < len(data):
        marker, length = data[index + 1], struct.unpack(">H", data[index + 2 : index + 4])[0]
        if 0xC0 <= marker <= 0xCF and marker not in (0xC4, 0xC8, 0xCC):
            height, width = struct.unpack(">HH", data[index + 5 : index + 9])
            return width, height
        index += 2 + length
    return None


def _picture(image: Image, placement: Placement, ids: _Ids) -> etree._Element | None:
    import base64
    import re
    import struct

    if ids.pictures is None:
        return None
    match = re.match(r"data:([^;,]+);base64,(.*)", image.href, re.DOTALL)
    if match is None:
        return None
    mime, data = match.group(1), base64.b64decode(match.group(2))
    extension = ""
    if mime == "image/svg+xml":
        import resvg_py

        svg = data.decode("utf-8")
        pixels = max(1, round(image.width * placement.scale / 72.0 * 300.0))
        from flexo.fonts import font_directories

        folders = [str(folder) for folder in font_directories()]
        data = bytes(resvg_py.svg_to_bytes(svg_string=svg, dpi=72, width=pixels, font_dirs=folders))
        embed = ids.pictures(data, "image/png")
        svg_id = ids.pictures.svg(svg.encode("utf-8"))  # type: ignore[attr-defined]
        extension = (
            '<a:extLst><a:ext uri="{96DAC541-7B7A-43D3-8B79-37D633B846F1}">'
            '<asvg:svgBlip xmlns:asvg="http://schemas.microsoft.com/office/drawing/2016/SVG/main" '
            f'r:embed="{svg_id}"/></a:ext></a:extLst>'
        )
        x, y, width, height = image.x, image.y, image.width, image.height
    elif mime in {"image/png", "image/jpeg"}:
        embed = ids.pictures(data, mime)
        size = struct.unpack(">II", data[16:24]) if mime == "image/png" else _jpeg_size(data)
        x, y, width, height = image.placed(*size) if size else (image.x, image.y, image.width, image.height)
        if size and "slice" in image.fit and (width > image.width + 0.01 or height > image.height + 0.01):
            # A picture filling its box, cropped rather than stretched: a native crop.
            crop = (
                (image.x - x) / width, (image.y - y) / height,
                (x + width - image.x - image.width) / width, (y + height - image.y - image.height) / height,
            )
            extension = '<a:srcRect l="{}" t="{}" r="{}" b="{}"/>'.format(*(round(v * 100000) for v in crop))
            x, y, width, height = image.x, image.y, image.width, image.height
    else:
        return None
    blip_extension = extension if extension.startswith("<a:extLst") else ""
    crop_rect = extension if extension.startswith("<a:srcRect") else ""
    left, top = placement.point(x, y)
    flips = (' flipH="1"' if image.flip_x else "") + (' flipV="1"' if image.flip_y else "")
    label = escape(image.id or "picture", {'"': "&quot;"})
    return etree.fromstring(
        f'<p:pic {_NS}><p:nvPicPr><p:cNvPr id="{ids()}" name="{label}"/>'
        f'<p:cNvPicPr><a:picLocks noChangeAspect="1"/></p:cNvPicPr><p:nvPr/></p:nvPicPr>'
        f'<p:blipFill><a:blip r:embed="{embed}">{blip_extension}</a:blip>{crop_rect}'
        f"<a:stretch><a:fillRect/></a:stretch></p:blipFill>"
        f'<p:spPr><a:xfrm{flips}><a:off x="{left}" y="{top}"/>'
        f'<a:ext cx="{placement.length(width)}" cy="{placement.length(height)}"/></a:xfrm>'
        f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr></p:pic>'
    )


# -- shapes ----------------------------------------------------------------------------


def _colour(value: str) -> str:
    value = value.strip().lstrip("#")
    if len(value) == 3:
        value = "".join(ch * 2 for ch in value)
    return value.upper()


def _blended(colour: str | None, blend: str, backdrop: str | None) -> str | None:
    """A multiply blend drawn as a plain fill: the colour it makes over the page.

    Slide programs have no blend modes; a wash multiplied onto the page is, over
    the page, the product of the two colours, so that product is drawn instead.
    """

    if colour is None or blend != "multiply" or not backdrop:
        return colour
    from flexo.colour import to_hex, to_rgb

    try:
        return to_hex([a * b for a, b in zip(to_rgb(colour), to_rgb(backdrop), strict=True)])
    except ValueError:
        return colour


def _fill(colour: str | None, alpha: float) -> str:
    if colour is None:
        return "<a:noFill/>"
    extra = f'<a:alpha val="{round(alpha * 100000)}"/>' if alpha < 0.999 else ""
    return f'<a:solidFill><a:srgbClr val="{_colour(colour)}">{extra}</a:srgbClr></a:solidFill>'


def _line(paint: Paint, placement: Placement) -> str:
    if paint.stroke is None:
        return "<a:ln><a:noFill/></a:ln>"
    width = placement.length(paint.stroke_width)
    cap = {"round": "rnd", "square": "sq"}.get(paint.linecap, "flat")
    join = {"round": "<a:round/>", "bevel": "<a:bevel/>"}.get(paint.linejoin, '<a:miter lim="800000"/>')
    dash = f'<a:prstDash val="{dash_preset(paint)}"/>' if paint.dash else ""
    return (
        f'<a:ln w="{width}" cap="{cap}">'
        f"{_fill(paint.stroke, paint.stroke_opacity * paint.opacity)}{dash}{join}</a:ln>"
    )


def dash_preset(paint: Paint) -> str:
    """The preset dash nearest a line's dash pattern, as the eye sees it.

    A preset's dash includes its caps; a round or square cap adds a line width
    to what an SVG dash draws (a zero-length dash with round caps is a dot).
    """

    import math

    width = max(paint.stroke_width, 1e-6)
    pattern = list(paint.dash) if len(paint.dash) % 2 == 0 else list(paint.dash) * 2
    capped = 1.0 if paint.linecap in {"round", "square"} else 0.0
    on = max(pattern[0] / width + capped, 0.25)
    off = max(pattern[1] / width - capped, 0.25)
    return min(
        DASHES,
        key=lambda name: abs(math.log(on / DASHES[name][0])) + abs(math.log(off / DASHES[name][1])),
    )


def _shape(shape: Shape, placement: Placement, ids: _Ids) -> etree._Element | None:
    parts = [_geometry_shape(shape, placement, ids)]
    for head in shape.arrowheads:
        xs = [x for segment in head.outline for x, _ in segment.points]
        ys = [y for segment in head.outline for _, y in segment.points]
        outline = Shape(
            None, "path", head.paint, min(xs), min(ys), max(xs) - min(xs), max(ys) - min(ys),
            segments=head.outline,
        )
        parts.append(_geometry_shape(outline, placement, ids, name="arrowhead"))
    parts = [part for part in parts if part is not None]
    if not parts:
        return None
    return parts[0] if len(parts) == 1 else _wrap(parts, shape.id, ids)


def _wrap(parts: list[etree._Element], name: str | None, ids: _Ids) -> etree._Element:
    boxes = [_extent(part) for part in parts]
    left = min(box[0] for box in boxes)
    top = min(box[1] for box in boxes)
    right = max(box[0] + box[2] for box in boxes)
    bottom = max(box[1] + box[3] for box in boxes)
    element = etree.fromstring(
        f"<p:grpSp {_NS}><p:nvGrpSpPr><p:cNvPr id=\"{ids()}\" name=\"{escape(name or 'line')}\"/>"
        f"<p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr><p:grpSpPr><a:xfrm>"
        f'<a:off x="{left}" y="{top}"/><a:ext cx="{right - left}" cy="{bottom - top}"/>'
        f'<a:chOff x="{left}" y="{top}"/><a:chExt cx="{right - left}" cy="{bottom - top}"/>'
        f"</a:xfrm></p:grpSpPr></p:grpSp>"
    )
    for part in parts:
        element.append(part)
    return element


def _geometry_shape(
    shape: Shape, placement: Placement, ids: _Ids, *, name: str | None = None
) -> etree._Element | None:
    paint = shape.paint
    if paint.fill is None and paint.stroke is None:
        return None
    x, y = placement.point(shape.x, shape.y)
    width = max(1, placement.length(shape.width))
    height = max(1, placement.length(shape.height))
    if shape.kind == "rect" and shape.radius <= 0:
        geometry = '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
    elif shape.kind == "rect":
        # A rounded rectangle's corner is a fraction of its shorter side (at most a half).
        adjust = round(min(0.5, shape.radius / max(1e-6, min(shape.width, shape.height))) * 100000)
        corner = f'<a:gd name="adj" fmla="val {adjust}"/>'
        geometry = f'<a:prstGeom prst="roundRect"><a:avLst>{corner}</a:avLst></a:prstGeom>'
    elif shape.kind == "ellipse":
        geometry = '<a:prstGeom prst="ellipse"><a:avLst/></a:prstGeom>'
    else:
        geometry = _custom(shape.segments, shape.x, shape.y, width, height, placement, filled=paint.fill is not None)
    label = escape(name or shape.id or shape.kind, {'"': "&quot;"})
    return etree.fromstring(
        f"<p:sp {_NS}><p:nvSpPr><p:cNvPr id=\"{ids()}\" name=\"{label}\"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>"
        f'<p:spPr><a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{width}" cy="{height}"/></a:xfrm>'
        f"{geometry}{_fill(_blended(paint.fill, paint.blend, ids.backdrop), paint.fill_opacity * paint.opacity)}"
        f"{_line(paint, placement)}</p:spPr>"
        f"</p:sp>"
    )


def _custom(
    segments: tuple[Segment, ...],
    left: float,
    top: float,
    width: int,
    height: int,
    placement: Placement,
    *,
    filled: bool,
) -> str:
    def point(p: tuple[float, float]) -> str:
        return (
            f'<a:pt x="{placement.length(p[0] - left)}" y="{placement.length(p[1] - top)}"/>'
        )

    body = []
    for segment in segments:
        if segment.kind == "M":
            body.append(f"<a:moveTo>{point(segment.points[0])}</a:moveTo>")
        elif segment.kind == "L":
            body.append(f"<a:lnTo>{point(segment.points[0])}</a:lnTo>")
        elif segment.kind == "C":
            body.append(f"<a:cubicBezTo>{''.join(point(p) for p in segment.points)}</a:cubicBezTo>")
        else:
            body.append("<a:close/>")
    fill = "" if filled else ' fill="none"'
    return (
        "<a:custGeom><a:avLst/><a:gdLst/><a:ahLst/><a:cxnLst/>"
        '<a:rect l="0" t="0" r="r" b="b"/><a:pathLst>'
        f'<a:path w="{width}" h="{height}"{fill}>{"".join(body)}</a:path>'
        "</a:pathLst></a:custGeom>"
    )


# -- text ------------------------------------------------------------------------------


_LINKS: ContextVar[Callable[[str], str] | None] = ContextVar("flexo_talk_links", default=None)
"""Relates a URL to the slide being written and gives its relationship id."""


@contextmanager
def linking(slide):
    """While writing ``slide``, a run's link becomes a hyperlink on that slide."""

    from pptx.opc.constants import RELATIONSHIP_TYPE

    def relate(url: str) -> str:
        return slide.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)

    token = _LINKS.set(relate)
    try:
        yield
    finally:
        _LINKS.reset(token)


def _run_xml(run, *, size: float | None = None, baseline: int | None = None, spacing: float = 0.0) -> str:
    """One run: its face, size, style, colour, any raise, and any link; ``spacing`` (points)
    is added after each of its characters."""

    size = size if size is not None else run.size
    face = escape(_family_name(run), {'"': "&quot;"})
    # A monospace face says so: where it is not installed, a slide program puts another
    # monospace face in its place, and code keeps its columns.
    pitch = ' pitchFamily="49"' if run.face is not None and _fixed_pitch(run.face.source, run.face.index) else ""
    bold = ' b="1"' if _bold(run) else ""
    italic = ' i="1"' if run.italic else ""
    raise_ = f' baseline="{baseline}"' if baseline else ""
    if spacing:
        raise_ += f' spc="{max(-400000, min(400000, round(spacing * 100)))}"'
    link = ""
    relate = _LINKS.get()
    if getattr(run, "link", "") and relate is not None:
        # The run keeps its own colour (hlinkClr "tx"), not the theme's hyperlink blue.
        link = (
            f'<a:hlinkClick r:id="{relate(run.link)}"><a:extLst>'
            '<a:ext uri="{A12FA001-AC4F-418D-AE19-62706E023703}">'
            '<ahyp:hlinkClr xmlns:ahyp="http://schemas.microsoft.com/office/drawing/2018/hyperlinkcolor" '
            'val="tx"/></a:ext></a:extLst></a:hlinkClick>'
        )
    return (
        f'<a:r><a:rPr lang="{_lang(run.text)}" sz="{round(size * 100)}"{bold}{italic}{raise_} dirty="0">'
        f"{_fill(run.fill, 1.0)}"
        f'<a:latin typeface="{face}"{pitch}/><a:ea typeface="{face}"{pitch}/><a:cs typeface="{face}"{pitch}/>{link}'
        f"</a:rPr><a:t>{escape(run.text)}</a:t></a:r>"
    )


_LANGUAGES = (
    # The first script found decides: Persian's own letters before Arabic's.
    ("fa-IR", ((0x067E, 0x067E), (0x0686, 0x0686), (0x0698, 0x0698), (0x06A9, 0x06A9), (0x06AF, 0x06AF),
               (0x06CC, 0x06CC))),
    ("ar-SA", ((0x0600, 0x06FF), (0x0750, 0x077F), (0xFB50, 0xFDFF), (0xFE70, 0xFEFF))),
    ("he-IL", ((0x0590, 0x05FF),)),
    ("ja-JP", ((0x3040, 0x30FF),)),
    ("ko-KR", ((0x1100, 0x11FF), (0xAC00, 0xD7AF))),
    ("zh-CN", ((0x4E00, 0x9FFF), (0x3400, 0x4DBF))),
    ("th-TH", ((0x0E00, 0x0E7F),)),
    ("hi-IN", ((0x0900, 0x097F),)),
    ("el-GR", ((0x0370, 0x03FF),)),
    ("ru-RU", ((0x0400, 0x04FF),)),
)


def _lang(text: str) -> str:
    """The language a run is tagged with: by its script (a slide program checks its
    spelling, and picks Chinese or Japanese forms of a character, by it); English else."""

    codes = {ord(character) for character in text}
    for tag, ranges in _LANGUAGES:
        if any(low <= code <= high for code in codes for low, high in ranges):
            return tag
    return "en-GB"


def _family_name(run) -> str:
    """The family name a slide program finds this run's face by.

    Office picks a face by family name plus bold and italic; a weight between
    (Figtree Medium, a semibold) is its own family there, under the font's
    legacy family name (name ID 1).
    """

    return _legacy_family(run.face.source, run.face.index) or run.face.family


@cache
def _fixed_pitch(source: str, index: int) -> bool:
    from fontTools.ttLib import TTCollection, TTFont

    font = (
        TTCollection(source, lazy=True).fonts[index]
        if source.lower().endswith((".ttc", ".otc"))
        else TTFont(source, lazy=True)
    )
    return bool(font["post"].isFixedPitch) or (font["OS/2"].panose.bProportion == 9 if "OS/2" in font else False)


@cache
def _legacy_family(source: str, index: int) -> str | None:
    from fontTools.ttLib import TTCollection, TTFont

    font = (
        TTCollection(source, lazy=True).fonts[index]
        if source.lower().endswith((".ttc", ".otc"))
        else TTFont(source, lazy=True)
    )
    name = font["name"]
    # Office finds a face by its Windows family name (platform 3), which may
    # differ from the Mac one ("LM Roman 10" against "Latin Modern Roman").
    # In English when the font has it (Windows', else the Mac's): Geeza Pro names itself
    # on Windows only in Persian, Hindi and Arabic, which no slide program looks for.
    legacy = (
        name.getName(1, 3, 1, 0x409) or name.getName(1, 1, 0, 0)
        or name.getName(1, 3, 1) or name.getName(1, 3, 10)
    )
    if legacy is not None:
        return legacy.toUnicode()
    fallback = name.getDebugName(1)
    return str(fallback) if fallback else None


def _bold(run) -> bool:
    """Whether the run asks for its family's bold: a static face says so in its
    Windows style name ("Bold", not "Regular" of "Figtree SemiBold"); a variable
    face is bold from 600 up."""

    if run.face.variable:
        return run.weight >= 600
    return "bold" in _style_name(run.face.source, run.face.index).lower()


@cache
def _style_name(source: str, index: int) -> str:
    from fontTools.ttLib import TTCollection, TTFont

    font = (
        TTCollection(source, lazy=True).fonts[index]
        if source.lower().endswith((".ttc", ".otc"))
        else TTFont(source, lazy=True)
    )
    name = font["name"]
    record = name.getName(2, 3, 1, 0x409) or name.getName(2, 3, 1)
    return record.toUnicode() if record is not None else str(name.getDebugName(2) or "")


def _text(text: Text, placement: Placement, ids: _Ids) -> etree._Element | None:
    element = _flat_text(replace(text, angle=0.0) if text.angle else text, placement, ids)
    if element is None or not text.angle:
        return element
    return _turned(element, text, placement)


def _turned(element: etree._Element, text: Text, placement: Placement) -> etree._Element:
    """Turn a placed text about its pivot: a slide program turns a box about its centre,
    so the box moves to where its centre goes and turns there."""

    import math

    turn = math.radians(text.angle)
    cos, sin = math.cos(turn), math.sin(turn)
    px, py = placement.point(*text.pivot)
    boxes = [box for box in element.iter(f"{{{_A}}}xfrm") if box.getparent().tag != f"{{{_P}}}grpSpPr"]
    for box in boxes:
        off, ext = box.find(f"{{{_A}}}off"), box.find(f"{{{_A}}}ext")
        cx = int(off.get("x")) + int(ext.get("cx")) / 2.0
        cy = int(off.get("y")) + int(ext.get("cy")) / 2.0
        tx = px + cos * (cx - px) - sin * (cy - py)
        ty = py + sin * (cx - px) + cos * (cy - py)
        off.set("x", str(round(int(off.get("x")) + tx - cx)))
        off.set("y", str(round(int(off.get("y")) + ty - cy)))
        box.set("rot", str(round(text.angle * 60000) % 21600000))
    if element.tag == f"{{{_P}}}grpSp":
        # The group's own extent follows its turned children.
        own = {f"{{{_P}}}nvGrpSpPr", f"{{{_P}}}grpSpPr"}
        extents = [_extent(child) for child in element if child.tag not in own]
        left, top = min(e[0] for e in extents), min(e[1] for e in extents)
        right, bottom = max(e[0] + e[2] for e in extents), max(e[1] + e[3] for e in extents)
        group = element.find(f"{{{_P}}}grpSpPr/{{{_A}}}xfrm")
        for tag, values in (("off", (left, top)), ("chOff", (left, top))):
            node = group.find(f"{{{_A}}}{tag}")
            node.set("x", str(values[0]))
            node.set("y", str(values[1]))
        for tag in ("ext", "chExt"):
            node = group.find(f"{{{_A}}}{tag}")
            node.set("cx", str(right - left))
            node.set("cy", str(bottom - top))
    return element


def _flat_text(text: Text, placement: Placement, ids: _Ids) -> etree._Element | None:
    if not text.simple:
        # Runs stepped back over each other: each run in a box of its own.
        parts = []
        for line in text.lines:
            for run in line.runs:
                # Each box sits at its run's own baseline, so no run is raised.
                flat = replace(run, shift=0.0)
                single = Text(
                    None, "start", run.x, (type(line)(run.baseline, (flat,)),), run.size,
                    run.family, run.weight, run.fill, run.fill_role, text.line_height, True,
                )
                part = _text_box(single, placement, ids, name="run")
                if part is not None:
                    parts.append(part)
        return _wrap(parts, text.id, ids) if parts else None
    return _text_box(text, placement, ids, name=text.id)


def _text_box(
    text: Text, placement: Placement, ids: _Ids, *, name: str | None, breaks: bool = False, steps: bool = False
) -> etree._Element | None:
    """Words as one text box, each line where Flexo set it: a paragraph a line, or with
    ``breaks`` one paragraph broken where Flexo broke it (a title is one paragraph). With
    ``steps``, a run set further on than the one before it ends (past a mark or a script
    drawn apart from the words) is spaced out to where it was set."""

    lines = text.lines
    spacing = text.line_height or text.size * 1.2
    left = min(line.left for line in lines)
    right = max(line.right for line in lines)
    slack = 2.0
    first = lines[0].baseline
    top = first - ascent(lines[0].runs[0].face if lines[0].runs else None) * spacing
    height = spacing * len(lines)
    align = {"start": "l", "middle": "ctr", "end": "r"}[text.anchor]
    written = []
    for line in lines:
        # A right-to-left line is written as it is read; the slide program orders it.
        rtl = bool(line.logical) and _rtl_text("".join(run.text for run in line.logical))
        order = line.logical or line.runs
        gaps = [0.0] * len(order)
        if steps and not line.logical:
            for index, (run, after) in enumerate(pairwise(order)):
                gap = after.x - (run.x + run.width)
                gaps[index] = gap * placement.scale if gap > 0.01 else 0.0
        written.append((rtl, "".join(_written_run(run, gap) for run, gap in zip(order, gaps, strict=True))))

    def paragraph(rtl: bool, runs: str) -> str:
        return (
            f'<a:p><a:pPr algn="{align}"{' rtl="1"' if rtl else ""}>'
            f'<a:lnSpc><a:spcPts val="{round(spacing * placement.scale * 100)}"/></a:lnSpc>'
            f'<a:spcBef><a:spcPts val="0"/></a:spcBef><a:spcAft><a:spcPts val="0"/></a:spcAft></a:pPr>'
            f"{runs}</a:p>"
        )

    if breaks and len({rtl for rtl, _ in written}) == 1:
        # Exactly spaced, a broken line sits where the next paragraph's would.
        br = f'<a:br><a:rPr lang="en-GB" sz="{round(text.size * 100)}" dirty="0"/></a:br>'
        paragraphs = [paragraph(written[0][0], br.join(runs for _, runs in written))]
    else:
        paragraphs = [paragraph(rtl, runs) for rtl, runs in written]
    # Room to spare on the side the words are not set against, so they start (or end)
    # where Flexo set them.
    spare = {"start": 0.0, "middle": slack, "end": 2 * slack}[text.anchor]
    x, y = placement.point(left - spare, top)
    width = placement.length(right - left + 2 * slack)
    box = placement.length(height)
    label = escape(name or "text", {'"': "&quot;"})
    body = "".join(paragraphs)
    if placement.scale != 1.0:
        body = _scaled(body, placement.scale)
    return etree.fromstring(
        f"<p:sp {_NS}><p:nvSpPr><p:cNvPr id=\"{ids()}\" name=\"{label}\"/><p:cNvSpPr txBox=\"1\"/><p:nvPr/></p:nvSpPr>"
        f'<p:spPr><a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{width}" cy="{box}"/></a:xfrm>'
        f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/></p:spPr>'
        f'<p:txBody><a:bodyPr wrap="none" lIns="0" tIns="0" rIns="0" bIns="0" anchor="t" rtlCol="0">'
        f"<a:noAutofit/></a:bodyPr><a:lstStyle/>{body}</p:txBody></p:sp>"
    )


def _written_run(run, spacing: float = 0.0) -> str:
    """A run of a text box, raised or lowered as set, and ``spacing`` (points) more after
    its last character than its own advance."""

    options = {}
    if run.shift:
        # A raised or lowered run is drawn smaller than its size: write
        # it larger so it is drawn at its own, raised as a percentage.
        written = run.size / SCRIPT_SCALE
        options = {"size": written, "baseline": round(run.shift / written * 100000)}
    if not spacing:
        return _run_xml(run, **options)
    import unicodedata

    # Spacing is added after every character of a run: only the last (with its marks) takes it.
    cut = len(run.text) - 1
    while cut > 0 and unicodedata.combining(run.text[cut]):
        cut -= 1
    head = _run_xml(replace(run, text=run.text[:cut]), **options) if cut > 0 else ""
    return head + _run_xml(replace(run, text=run.text[cut:]), spacing=spacing, **options)


def _ink_of(run, palette, ink: str) -> str:
    """The colour a list or table run is written in: its own, a link's, or the text's."""

    from flexo.render_common import run_colour

    from flexo_talk.deck import link_colour

    if run.link and not run.color and (link := link_colour(palette)):
        return link
    return run_colour(run, palette) or ink


def _rtl_text(text: str) -> bool:
    from flexo.bidi import base_level

    return base_level(text) == 1


def _scaled(body: str, scale: float) -> str:
    import re

    return re.sub(r'sz="(\d+)"', lambda m: f'sz="{round(int(m.group(1)) * scale)}"', body)


# -- lists -----------------------------------------------------------------------------


@dataclass(frozen=True, slots=True)
class _ListRun:
    text: str
    size: float
    weight: int
    italic: bool
    face: object
    fill: str
    shift: float = 0.0
    link: str = ""


def _room_for(run, typography, face, size: float, weight: int) -> str:
    """Room in native words for a formula (``TextRun.math``): one space, spaced out to
    the formula's width. The formula itself is drawn over it, as outlines, where Flexo
    set it -- a slide program's words cannot hold a fraction or a matrix."""

    from flexo.fonts import hb_font
    from flexo.text import formula_of

    width = formula_of(run, typography, weight).width
    font = hb_font(face, weight)
    gid = font.get_nominal_glyph(ord(" ")) or 0
    space = font.get_glyph_h_advance(gid) / font.scale[0] * size
    spacing = max(-400000, min(400000, round((width - space) * 100)))
    room = _ListRun(" ", size, weight, False, face, "#000000")
    name = escape(_family_name(room), {'"': "&quot;"})
    # In the weight its width was measured at: a bold space is narrower than a regular one.
    bold = ' b="1"' if _bold(room) else ""
    return (
        f'<a:r><a:rPr lang="en-GB" sz="{round(size * 100)}"{bold} spc="{spacing}" dirty="0">'
        f'<a:latin typeface="{name}"/><a:ea typeface="{name}"/><a:cs typeface="{name}"/>'
        f"</a:rPr><a:t> </a:t></a:r>"
    )


def add_list(tree: etree._Element, deck, layout, *, formulas: list | None = None, editable: bool = False) -> int:
    """A bulleted list as one text box of bulleted paragraphs, wrapped by the slide program.

    The box is as wide as the list was set, so the words break where Flexo broke
    them when the fonts are the same, and reflow like any list when edited.

    Maths in it (``formulas``, the drawn formulas set over it) is written one of two
    ways: ``editable`` -- PowerPoint's own equations, in the words, with the drawn ones
    and words that leave room for them as the fallback for other slide programs --
    or else only the drawn ones over the words.
    """

    from flexo.text import font_stack

    existing = [int(item) for item in tree.xpath(".//@id") if str(item).isdigit()]
    ids = _Ids(max(existing, default=1))
    typography = deck.typography(layout.size)
    # flexo's shared stack: it holds the families measuring adopted for other scripts.
    stack = font_stack(typography)
    palette = layout.palette or deck.palette
    ink = palette.get("ink")
    accent = _colour(palette.get("tone-1-stroke"))
    muted = _colour(palette.get("muted-ink"))

    face = stack.face(400, False)

    def placed(native: bool) -> tuple[list[tuple[float, float]], float, float]:
        """Each item's line spacing and the room before it, and the box's top and height:
        each item's first baseline where Flexo set it."""

        rows: list[tuple[float, float]] = []
        top = end = 0.0
        for position, (_level, _runs, baseline) in enumerate(layout.items):
            rise, fall, count = layout.opened[position] if position < len(layout.opened) else (0.0, 0.0, 1)
            spacing = plain = layout.step(position) - rise - fall
            if rise + fall > 0.01 and (count > 1 or native):
                # Lines opened for a formula, set loosely (and lower) by the slide program.
                spacing += rise + fall
            # A single line opened for a drawn formula keeps its words' spacing, the room
            # the formula needs going before and after it.
            down = baseline_down(face, spacing, plain)
            if position == 0:
                top = end = baseline - down
            before = max(0.0, baseline - down - end)
            end += before + spacing * count
            rows.append((spacing, before))
        return rows, top, end - top

    def paragraphs(native: bool) -> str:
        written = []
        rows, _, _ = placed(native)
        for position, (level, runs, _baseline) in enumerate(layout.items):
            offset = layout.offset(level)
            mark_at = layout.mark_at(level)
            pieces = _runs_xml(runs, typography, stack, layout.size, palette, ink, native=native)
            spacing, before = rows[position]
            if layout.numbered and level == 0:
                # A native numbered list: the program numbers it, in the words' face.
                face = escape(_family_name(_ListRun("1", layout.size, 400, False, stack.face(400, False), ink)))
                mark = f'<a:buFont typeface="{face}"/><a:buAutoNum type="arabicPeriod"/>'
            else:
                mark = '<a:buFont typeface="Arial"/><a:buChar char="\u2022"/>'
            colour = accent if level == 0 else muted
            rtl = _rtl_text("".join(run.text for run in runs))
            written.append(
                f'<a:p><a:pPr marL="{round(offset * EMU_PER_POINT)}" '
                f'indent="{round((mark_at - offset) * EMU_PER_POINT)}" lvl="{min(level, 8)}"'
                f'{" rtl=\"1\" algn=\"r\"" if rtl else ""}>'
                f'<a:lnSpc><a:spcPts val="{round(spacing * 100)}"/></a:lnSpc>'
                f'<a:spcBef><a:spcPts val="{round(before * 100)}"/></a:spcBef>'
                f'<a:spcAft><a:spcPts val="0"/></a:spcAft>'
                f'<a:buClr><a:srgbClr val="{colour}"/></a:buClr><a:buSzPct val="{100000 if level == 0 else 85000}"/>'
                f"{mark}</a:pPr>{pieces}</a:p>"
            )
        return "".join(written)

    def shape(shape_id: int, native: bool) -> etree._Element:
        _, top, height = placed(native) if layout.items else ([], layout.y, layout.line_height)
        return etree.fromstring(
            f"<p:sp {_NS}><p:nvSpPr><p:cNvPr id=\"{shape_id}\" name=\"{escape(layout.id)}\"/>"
            f"<p:cNvSpPr txBox=\"1\"/><p:nvPr/></p:nvSpPr>"
            f'<p:spPr><a:xfrm><a:off x="{round(layout.x * EMU_PER_POINT)}" y="{round(top * EMU_PER_POINT)}"/>'
            f'<a:ext cx="{round((layout.width + 0.25) * EMU_PER_POINT)}" '
            f'cy="{round(height * EMU_PER_POINT)}"/></a:xfrm>'
            f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/></p:spPr>'
            f'<p:txBody><a:bodyPr wrap="square" lIns="0" tIns="0" rIns="0" bIns="0" anchor="t" rtlCol="0">'
            f"<a:noAutofit/></a:bodyPr><a:lstStyle/>{paragraphs(native)}</p:txBody></p:sp>"
        )

    maths = any(run.math for _, runs, _ in layout.items for run in runs)
    editable = editable and _expressible(run for _, runs, _ in layout.items for run in runs)
    shape_id = ids()
    drawn = _group(Group(None, list(formulas or [])), Placement(), ids) if formulas else None
    if maths and editable:
        fallback = [shape(ids(), native=False), *([drawn] if drawn is not None else [])]
        tree.append(alternate(shape(shape_id, native=True), _together(fallback, ids)))
        return shape_id
    tree.append(shape(shape_id, native=False))
    if drawn is not None:
        tree.append(drawn)
    return shape_id


def _runs_xml(runs, typography, stack, size: float, palette, ink: str, *, native: bool, bold: bool = False) -> str:
    """Runs of words as DrawingML runs: maths as PowerPoint's equations (``native``), or
    as room for a drawn formula."""

    from flexo.text import drawn_weight

    from flexo_talk.omml import omml

    pieces = []
    for run in runs:
        # Bold words' maths stays regular (TextRun.maths), as flexo draws it.
        weight = 700 if bold and run.weight == 400 and not run.maths else drawn_weight(run, None)
        if run.math:
            if native:
                pieces.append(omml(run.math, size=size, colour=_colour(_ink_of(run, palette, ink)),
                                   resolve=_resolver(palette), bold=weight >= 600))
            else:
                pieces.append(_room_for(run, typography, stack.face(weight, False), size, weight))
            continue
        for face, text in stack.segments(run.text, weight, run.italic, code=run.code):
            script = run.baseline_shift != "normal"
            run_size = size * 0.72 / SCRIPT_SCALE if script else size
            shift = {"super": 33000, "sub": -20000}.get(run.baseline_shift)
            pieces.append(
                _run_xml(
                    _ListRun(text, run_size, weight, run.italic, face, _ink_of(run, palette, ink), link=run.link),
                    baseline=shift,
                )
            )
    if native and runs and all(run.math for run in runs) and not any(
        run.math.startswith("\\displaystyle") for run in runs
    ):
        # A formula alone in its paragraph is displayed (full size) by PowerPoint; beside
        # a word -- a zero-width one -- it is set within the line, as flexo set it.
        face = stack.face(400, False)
        pieces.append(_run_xml(_ListRun("\u200b", size, 400, False, face, ink)))
    return "".join(pieces)




def _expressible(runs) -> bool:
    """Whether every formula among ``runs`` is one Office Math shows as flexo set it."""

    from flexo_talk.omml import expressible

    return all(expressible(run.math) for run in runs if run.math)


def _resolver(palette):
    """A colour a formula names (``\\color{accent}``, ``red``, ``#c0392b``) as ``rrggbb``."""

    from flexo.ir.semantic import TextRun
    from flexo.render_common import run_colour

    def resolve(colour: str) -> str | None:
        # A palette role, or a colour LaTeX names (as the drawing has it).
        found = run_colour(TextRun("", color=colour), palette)
        return _colour(found) if found else None

    return resolve


_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
_A14 = "http://schemas.microsoft.com/office/drawing/2010/main"


def alternate(choice: etree._Element, fallback: etree._Element | None) -> etree._Element:
    """``choice`` for PowerPoint (2010 on, which reads its own equations), ``fallback``
    for every other slide program -- as PowerPoint itself saves an equation. With no
    ``fallback``, its place (the last child) is left empty for one to be moved into."""

    element = etree.Element(f"{{{_MC}}}AlternateContent", nsmap={"mc": _MC})
    first = etree.SubElement(element, f"{{{_MC}}}Choice", nsmap={"a14": _A14})
    first.set("Requires", "a14")
    first.append(choice)
    second = etree.SubElement(element, f"{{{_MC}}}Fallback")
    if fallback is not None:
        second.append(fallback)
    return element


def _together(elements: list[etree._Element], ids: _Ids) -> etree._Element:
    """Shapes as one group (or the shape itself, when there is one)."""

    if len(elements) == 1:
        return elements[0]
    boxes = [_extent(child) for child in elements]
    left = min(box[0] for box in boxes)
    top = min(box[1] for box in boxes)
    right = max(box[0] + box[2] for box in boxes)
    bottom = max(box[1] + box[3] for box in boxes)
    group = etree.fromstring(
        f"<p:grpSp {_NS}><p:nvGrpSpPr><p:cNvPr id=\"{ids()}\" name=\"group\"/>"
        f"<p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr><p:grpSpPr><a:xfrm>"
        f'<a:off x="{left}" y="{top}"/><a:ext cx="{right - left}" cy="{bottom - top}"/>'
        f'<a:chOff x="{left}" y="{top}"/><a:chExt cx="{right - left}" cy="{bottom - top}"/>'
        f"</a:xfrm></p:grpSpPr></p:grpSp>"
    )
    for child in elements:
        group.append(child)
    return group


def editable_maths(tree: etree._Element, drawing: Drawing, worded: list, deck, palette) -> int:
    """Each displayed equation, and each passage of words with maths in it, written as
    PowerPoint's own -- its equations editable in its equation editor -- with the drawing
    of it as the fallback for other slide programs. How many were written."""

    from flexo_talk.omml import expressible, omml

    equations: dict[str, Group] = {}

    def find(group: Group) -> None:
        for item in group.items:
            if isinstance(item, Group):
                if item.id and item.data.get("data-flexo-talk") == "math":
                    equations[item.id] = item
                else:
                    find(item)

    find(drawing.root)
    passages = {words.id: words for words in worded}
    existing = [int(item) for item in tree.xpath(".//@id") if str(item).isdigit()]
    ids = _Ids(max(existing, default=1))
    count = 0
    for drawn in list(tree.iter(f"{{{_P}}}grpSp")):
        name_element = drawn.find(f"{{{_P}}}nvGrpSpPr/{{{_P}}}cNvPr")
        name = name_element.get("name") if name_element is not None else None
        if name in equations and not expressible(equations[name].data.get("data-flexo-math", "")):
            continue
        if name in passages and not _expressible(run for line in passages[name].lines for run in line):
            continue
        if name in equations:
            group = equations[name]
            size = float(group.data.get("data-flexo-size", "20"))
            align = group.data.get("data-flexo-align", "middle")
            fill = _first_fill(group) or deck.palette.get("ink")
            maths = omml(group.data.get("data-flexo-math", ""), size=size, colour=_colour(fill), display=True,
                         align=align, resolve=_resolver(palette))
            left, top, width, height = _extent(drawn)
            algn = {"start": "l", "end": "r"}.get(align, "ctr")
            choice = etree.fromstring(
                f"<p:sp {_NS}><p:nvSpPr><p:cNvPr id=\"{ids()}\" name=\"{escape(name)}\"/>"
                f"<p:cNvSpPr txBox=\"1\"/><p:nvPr/></p:nvSpPr>"
                f'<p:spPr><a:xfrm><a:off x="{left}" y="{top}"/><a:ext cx="{width}" cy="{height}"/></a:xfrm>'
                f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/></p:spPr>'
                f'<p:txBody><a:bodyPr wrap="none" lIns="0" tIns="0" rIns="0" bIns="0" anchor="ctr" rtlCol="0">'
                f'<a:noAutofit/></a:bodyPr><a:lstStyle/><a:p><a:pPr algn="{algn}"/>{maths}'
                f'<a:endParaRPr lang="en-GB" sz="{round(size * 100)}" dirty="0"/></a:p></p:txBody></p:sp>'
            )
        elif name in passages:
            choice = _passage(passages[name], deck, palette, ids, native=True)
        else:
            continue
        # The wrapper goes in first, and the drawing moves into it within the slide: a
        # drawing taken out of the slide alone is copied about (seconds, for a long one).
        wrapper = alternate(choice, None)
        drawn.addprevious(wrapper)
        wrapper[-1].append(drawn)
        count += 1
    return count


def _passage(words, deck, palette, ids: _Ids, *, native: bool, breaks: bool = False) -> etree._Element:
    """Words with maths in them (a ``WordsLayout``) as one text box where they were set:
    the maths PowerPoint's own equations (``native``), or room for the drawn formulas.
    With ``breaks``, one paragraph broken where Flexo broke it."""

    from flexo.text import font_stack

    typography = deck.typography(words.size)
    if words.family:
        typography = typography.with_family(words.family)
    bold = (words.weight or 400) >= 600
    stack = font_stack(typography)
    plain = words.size * typography.line_height
    # A line opened for a tall formula is set lower by some slide programs than by others: a
    # single line leaving room for a drawn formula keeps its words' own spacing.
    spacing = plain if not native and len(words.lines) == 1 else words.line_height
    lines = [
        _runs_xml(line, typography, stack, words.size, palette, words.fill, native=native, bold=bold)
        # Room for a formula ending a line is a space there, which a centred or right-aligned
        # line leaves out: a zero-width word after it keeps it in.
        + (_run_xml(_ListRun("\u200b", words.size, 400, False, stack.face(400, False), words.fill))
           if not native and line and line[-1].math else "")
        for line in words.lines
    ]
    start = (
        f'<a:p><a:pPr algn="{ {"start": "l", "middle": "ctr", "end": "r"}.get(words.align, "l") }">'
        f'<a:lnSpc><a:spcPts val="{round(spacing * 100)}"/></a:lnSpc>'
        f'<a:spcBef><a:spcPts val="0"/></a:spcBef><a:spcAft><a:spcPts val="0"/></a:spcAft></a:pPr>'
    )
    end = f'<a:endParaRPr lang="en-GB" sz="{round(words.size * 100)}" dirty="0"/></a:p>'
    if breaks:
        br = f'<a:br><a:rPr lang="en-GB" sz="{round(words.size * 100)}" dirty="0"/></a:br>'
        paragraphs = start + br.join(lines) + end
    else:
        paragraphs = "".join(start + line + end for line in lines)
    left = {"start": words.x, "middle": words.x - words.width / 2.0, "end": words.x - words.width}.get(
        words.align, words.x)
    top = words.baseline - baseline_down(stack.face(400, False), spacing, plain)
    height = spacing * len(words.lines)
    return etree.fromstring(
        f"<p:sp {_NS}><p:nvSpPr><p:cNvPr id=\"{ids()}\" name=\"{escape(words.id, {'"': '&quot;'})}\"/>"
        f"<p:cNvSpPr txBox=\"1\"/><p:nvPr/></p:nvSpPr>"
        f'<p:spPr><a:xfrm><a:off x="{round(left * EMU_PER_POINT)}" y="{round(top * EMU_PER_POINT)}"/>'
        f'<a:ext cx="{round((words.width + 0.5) * EMU_PER_POINT)}" '
        f'cy="{round(height * EMU_PER_POINT)}"/></a:xfrm>'
        f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/></p:spPr>'
        f'<p:txBody><a:bodyPr wrap="none" lIns="0" tIns="0" rIns="0" bIns="0" anchor="t" rtlCol="0">'
        f"<a:noAutofit/></a:bodyPr><a:lstStyle/>{paragraphs}</p:txBody></p:sp>"
    )


def _first_fill(group: Group) -> str | None:
    for item in group.items:
        if isinstance(item, Group):
            found = _first_fill(item)
            if found:
                return found
        elif isinstance(item, Shape) and item.paint.fill:
            return item.paint.fill
    return None


def add_reveals(slide_element: etree._Element, reveals: list[tuple[int, list[tuple[int, int]]]]) -> None:
    """Click-by-click builds: each ``(shape id, [(first, last) paragraph, ...])`` shows
    its paragraph ranges one click at a time, as PowerPoint's "Appear" by paragraph."""

    if not reveals:
        return
    ids = iter(range(3, 10_000))
    clicks = []
    for shape_id, ranges in reveals:
        for first, last in ranges:
            outer, inner, effect, behaviour = next(ids), next(ids), next(ids), next(ids)
            clicks.append(
                f'<p:par><p:cTn id="{outer}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst>'
                f'<p:childTnLst><p:par><p:cTn id="{inner}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst>'
                f'<p:childTnLst><p:par><p:cTn id="{effect}" presetID="1" presetClass="entr" presetSubtype="0" '
                f'fill="hold" grpId="0" nodeType="clickEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst>'
                f'<p:childTnLst><p:set><p:cBhvr><p:cTn id="{behaviour}" dur="1" fill="hold">'
                f'<p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{shape_id}">'
                f'<p:txEl><p:pRg st="{first}" end="{last}"/></p:txEl></p:spTgt></p:tgtEl>'
                f"<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr>"
                f'<p:to><p:strVal val="visible"/></p:to></p:set></p:childTnLst></p:cTn></p:par>'
                f"</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>"
            )
    builds = "".join(f'<p:bldP spid="{shape_id}" grpId="0" build="p"/>' for shape_id, _ in reveals)
    timing = etree.fromstring(
        f'<p:timing {_NS}><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">'
        f'<p:childTnLst><p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq">'
        f"<p:childTnLst>{''.join(clicks)}</p:childTnLst></p:cTn>"
        f'<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
        f'<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>'
        f"</p:seq></p:childTnLst></p:cTn></p:par></p:tnLst><p:bldLst>{builds}</p:bldLst></p:timing>"
    )
    # <p:timing> follows <p:clrMapOvr> (and <p:transition>) in a slide.
    slide_element.append(timing)


# -- placeholders ----------------------------------------------------------------------


PLACEHOLDER_KINDS = {
    "title": '<p:ph type="title"/>',
    "ctrTitle": '<p:ph type="ctrTitle"/>',
    "subTitle": '<p:ph type="subTitle" idx="1"/>',
    "body": '<p:ph type="body" idx="1"/>',
    "obj": '<p:ph idx="1"/>',
}
"""The placeholders a slide's headings fill, as the default template's layouts have them:
a title (a title slide's centred one), a subtitle, and the text of a section or content
layout."""


def placeholder(shape: etree._Element, kind: str) -> etree._Element:
    """A text box (``_text_box``'s, ``_passage``'s) made the slide's placeholder of ``kind``
    (``PLACEHOLDER_KINDS``) -- what PowerPoint's outline, navigation and accessibility checker,
    and a screen reader, take a slide's title from -- looking just as it did.

    A placeholder takes what its own words do not say from its layout and master: a
    section title's capitals and bold, a subtitle's bullet and indent. The box says
    everything a text box takes for granted, so none of it shows."""

    nv = shape.find(f"{{{_P}}}nvSpPr")
    locks = nv.find(f"{{{_P}}}cNvSpPr")
    locks.attrib.pop("txBox", None)
    etree.SubElement(locks, f"{{{_A}}}spLocks").set("noGrp", "1")
    nv.find(f"{{{_P}}}nvPr").insert(0, etree.fromstring(PLACEHOLDER_KINDS[kind].replace("<p:ph ", f"<p:ph {_NS} ")))
    body = shape.find(f"{{{_P}}}txBody")
    plain = etree.fromstring(
        f'<a:lstStyle {_NS}><a:lvl1pPr marL="0" indent="0"><a:buNone/>'
        f'<a:defRPr b="0" i="0" cap="none" spc="0" baseline="0"/>'
        f"</a:lvl1pPr></a:lstStyle>"
    )
    body.replace(body.find(f"{{{_A}}}lstStyle"), plain)
    later = {f"{{{_A}}}{tag}" for tag in ("tabLst", "defRPr", "extLst")}
    for properties in body.iter(f"{{{_A}}}pPr"):
        properties.set("marL", "0")
        properties.set("indent", "0")
        mark = etree.Element(f"{{{_A}}}buNone")
        follower = next((child for child in properties if child.tag in later), None)
        if follower is None:
            properties.append(mark)
        else:
            follower.addprevious(mark)
    for properties in body.iterfind(f".//{{{_A}}}r/{{{_A}}}rPr"):
        # Said on each run as well, for slide programs that read no list style.
        for name, value in (("b", "0"), ("i", "0"), ("cap", "none")):
            if properties.get(name) is None:
                properties.set(name, value)
    return shape


def add_heading(
    tree: etree._Element,
    heading: Text | Group,
    kind: str,
    *,
    words=None,
    formulas: list | None = None,
    deck=None,
    palette=None,
    editable: bool = False,
    span: tuple[float, float] | None = None,
) -> None:
    """A slide's heading -- its title, its subtitle -- appended to its shape tree as the
    slide's placeholder of ``kind`` (``PLACEHOLDER_KINDS``), where Flexo drew it and as it drew it.

    A ``Text`` is written as its text box was; a run that steps back over the words (an
    accent's mark, the second of a stacked pair of scripts) is drawn on its own, as it was.
    Words with maths in them (a ``Group``, with its ``words``, a ``WordsLayout``) are words
    leaving room for the drawn ``formulas``, and with ``editable`` (when Office Math can
    show them) PowerPoint's own equations in the words, with the former as the fallback
    for other slide programs (see ``alternate``).

    Some slide programs wrap a placeholder's words at its box, however it says not to:
    the box reaches across ``span`` (the slide's left and right margins, in points) on
    the side its words are not set against, so words a little wider there stay on one line."""

    existing = [int(item) for item in tree.xpath(".//@id") if str(item).isdigit()]
    ids = _Ids(max(existing, default=1))

    def made(shape: etree._Element) -> etree._Element:
        return _widened(placeholder(shape, kind), span) if span else placeholder(shape, kind)

    if isinstance(heading, Text):
        flowing, apart = _apart(heading)
        tree.append(made(_text_box(flowing, Placement(), ids, name=heading.id, breaks=True, steps=bool(apart))))
        if apart and (marks := _flat_text(apart, Placement(), ids)) is not None:
            tree.append(marks)
        return
    drawn = [made(_passage(words, deck, palette, ids, native=False, breaks=True))]
    if formulas and (shapes := _group(Group(None, list(formulas)), Placement(), ids)) is not None:
        drawn.append(shapes)
    if not (editable and _expressible(run for line in words.lines for run in line)):
        tree.extend(drawn)
        return
    wrapper = alternate(made(_passage(words, deck, palette, ids, native=True, breaks=True)), None)
    wrapper[-1].extend(drawn)
    tree.append(wrapper)


def _apart(text: Text) -> tuple[Text, Text | None]:
    """``text`` as words that run on, and the runs that step back over them (an accent's
    mark, the wider of a stacked pair of scripts), to be drawn apart, each where it was
    set -- one text box cannot step back."""

    if text.simple:
        return text, None
    lines, apart = [], []
    for line in text.lines:
        if line.logical:
            # Read right to left, ordered by the slide program: kept whole.
            lines.append(line)
            continue
        kept: list = []
        for run in line.runs:
            if kept and run.x < kept[-1].x + kept[-1].width - 0.01:
                apart.append(type(line)(line.baseline, (run,)))
            else:
                kept.append(run)
        lines.append(type(line)(line.baseline, tuple(kept)))
    flowing = replace(text, lines=tuple(lines), simple=True)
    if not apart:
        return flowing, None
    return flowing, replace(text, id=f"{text.id}.marks", lines=tuple(apart), simple=False)


def _widened(shape: etree._Element, span: tuple[float, float]) -> etree._Element:
    """``shape``'s box reaching to ``span`` (points) on the side its words are not set against."""

    off = shape.find(f"{{{_P}}}spPr/{{{_A}}}xfrm/{{{_A}}}off")
    ext = shape.find(f"{{{_P}}}spPr/{{{_A}}}xfrm/{{{_A}}}ext")
    left = int(off.get("x"))
    right = left + int(ext.get("cx"))
    low, high = (round(edge * EMU_PER_POINT) for edge in span)
    align = shape.find(f".//{{{_A}}}p/{{{_A}}}pPr").get("algn", "l")
    if align == "l":
        right = max(right, high)
    elif align == "r":
        left = min(left, low)
    else:
        middle = (left + right) // 2
        half = max(middle - left, min(middle - low, high - middle))
        left, right = middle - half, middle + half
    off.set("x", str(left))
    ext.set("cx", str(right - left))
    return shape


# -- tables ----------------------------------------------------------------------------


def add_table(tree: etree._Element, deck, layout, *, formulas: list | None = None, editable: bool = False) -> None:
    """A table as a native PowerPoint table: its columns, rows, rules, and words as set.

    Column widths and row heights are Flexo's; each cell's margins put its first
    baseline where Flexo put it; the rules are cell borders (booktabs: above,
    under the header, below), and every other border is off. Maths in its cells is
    written as in ``add_list``.
    """

    from flexo.text import font_stack

    existing = [int(item) for item in tree.xpath(".//@id") if str(item).isdigit()]
    ids = _Ids(max(existing, default=1))
    typography = deck.typography(layout.size)
    stack = font_stack(typography)
    palette = layout.palette or deck.palette
    ink = palette.get("ink")
    colour = _colour(ink)

    def table(frame_id: int, native: bool) -> etree._Element:
        rows = _table_rows(layout, typography, stack, palette, ink, colour, native)
        widths = list(reversed(layout.widths)) if layout.rtl else layout.widths
        grid = "".join(f'<a:gridCol w="{round(width * EMU_PER_POINT)}"/>' for width in widths)
        width = round(sum(layout.widths) * EMU_PER_POINT)
        height = round(sum(layout.heights) * EMU_PER_POINT)
        return etree.fromstring(
            f'<p:graphicFrame {_NS}><p:nvGraphicFramePr><p:cNvPr id="{frame_id}" name="{escape(layout.id)}"/>'
            f'<p:cNvGraphicFramePr><a:graphicFrameLocks noGrp="1"/></p:cNvGraphicFramePr><p:nvPr/>'
            f'</p:nvGraphicFramePr>'
            f'<p:xfrm><a:off x="{round(layout.x * EMU_PER_POINT)}" y="{round(layout.y * EMU_PER_POINT)}"/>'
            f'<a:ext cx="{width}" cy="{height}"/></p:xfrm>'
            f'<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/table">'
            f'<a:tbl><a:tblPr firstRow="{1 if layout.header else 0}" bandRow="0"/><a:tblGrid>{grid}</a:tblGrid>'
            f"{''.join(rows)}</a:tbl></a:graphicData></a:graphic></p:graphicFrame>"
        )

    maths = any(run.math for row in layout.cells for cell in row for run in cell)
    editable = editable and _expressible(run for row in layout.cells for cell in row for run in cell)
    frame_id = ids()
    drawn = _group(Group(None, list(formulas or [])), Placement(), ids) if formulas else None
    if maths and editable:
        fallback = [table(ids(), native=False), *([drawn] if drawn is not None else [])]
        tree.append(alternate(table(frame_id, native=True), _together(fallback, ids)))
        return
    tree.append(table(frame_id, native=False))
    if drawn is not None:
        tree.append(drawn)


def _table_rows(layout, typography, stack, palette, ink: str, colour: str, native: bool) -> list[str]:
    top_rule, mid_rule, bottom_rule = layout.rules
    rows = []
    last = len(layout.cells) - 1
    for r, row in enumerate(layout.cells):
        heading = layout.header and r == 0
        cells = []
        for c, cell in enumerate(row):
            pieces = [_runs_xml(cell, typography, stack, layout.size, palette, ink, native=native, bold=heading)]
            mirrored = {"start": "r", "middle": "ctr", "end": "l"}
            algn = (mirrored if layout.rtl else {"start": "l", "middle": "ctr", "end": "r"})[layout.align[c]]
            direction = ' rtl="1"' if _rtl_text("".join(run.text for run in cell)) else ""
            end = f'<a:endParaRPr lang="en-GB" sz="{round(layout.size * 100)}" dirty="0"/>'

            def border(tag: str, width: float | None) -> str:
                if not width:
                    return f"<a:{tag} w=\"0\"><a:noFill/></a:{tag}>"
                return (
                    f'<a:{tag} w="{round(width * EMU_PER_POINT)}" cap="flat" cmpd="sng">'
                    f'<a:solidFill><a:srgbClr val="{colour}"/></a:solidFill><a:prstDash val="solid"/></a:{tag}>'
                )

            above = top_rule if r == 0 else (mid_rule if layout.header and r == 1 else None)
            below = bottom_rule if r == last else (mid_rule if heading else None)
            share = ascent(stack.face(400, False))
            inset = round((layout.baseline - share * layout.line_height) * EMU_PER_POINT)
            margin = round(layout.pad * EMU_PER_POINT)
            cells.append(
                f"<a:tc><a:txBody><a:bodyPr/><a:lstStyle/><a:p><a:pPr algn=\"{algn}\"{direction}>"
                f'<a:lnSpc><a:spcPts val="{round(layout.line_height * 100)}"/></a:lnSpc>'
                f'<a:spcBef><a:spcPts val="0"/></a:spcBef><a:spcAft><a:spcPts val="0"/></a:spcAft></a:pPr>'
                f"{''.join(pieces)}{end}</a:p></a:txBody>"
                f'<a:tcPr marL="{margin}" marR="{margin}" marT="{max(inset, 0)}" marB="0" anchor="t">'
                f"{border('lnL', None)}{border('lnR', None)}{border('lnT', above)}{border('lnB', below)}"
                f"<a:noFill/></a:tcPr></a:tc>"
            )
        # A right-to-left table is written in the order it is seen, right column last,
        # rather than flagged rtl: not every slide program honours the flag.
        order = reversed(cells) if layout.rtl else cells
        rows.append(f'<a:tr h="{round(layout.heights[r] * EMU_PER_POINT)}">{"".join(order)}</a:tr>')
    return rows
