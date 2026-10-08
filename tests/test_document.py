"""A deck document makes the deck its Python makes, says where it is wrong, and is drawn
slide by slide in the studio."""

from __future__ import annotations

import importlib.util
import json
import re
import shutil
import subprocess
from pathlib import Path

import pytest
import yaml

from flexo_talk import Deck
from flexo_talk.cli import main
from flexo_talk.compose import render_slide
from flexo_talk.document import (
    DeckDocumentError,
    deck_document,
    deck_from_document,
    dump_document,
    read_deck,
)
from flexo_talk.studio import DeckKind

ROOT = Path(__file__).resolve().parents[1]


def _demo() -> Deck:
    spec = importlib.util.spec_from_file_location("demo", ROOT / "examples" / "demo.py")
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module.talk()


def _small() -> dict:
    return {
        "deck": {"id": "small", "look": "band", "footer": "A footer", "style": {"body_size": 22}},
        "slides": [
            {"layout": "title", "title": "A talk", "subtitle": "About things", "author": "Ada"},
            {"layout": "section", "title": "Part one"},
            {"title": "Points",
             "body": [{"bullets": ["One", "Two", ["under two"]], "numbered": True, "colour": "accent"}],
             "notes": "Say it slowly", "footnotes": ["[1] A source"]},
            {"layout": "two-columns", "title": "Both", "split": 0.4,
             "left": [{"text": "Words with $x$", "align": "middle"}],
             "right": [{"table": [["A", "B"], ["x", 1.5]]}]},
            {"layout": "columns", "title": "Three", "widths": [2, 1, 1],
             "columns": [[{"stats": [{"value": "93%", "label": "right"}]}], [{"quote": "Q", "by": "W"}],
                         [{"callout": "C", "title": "T", "colour": "accent2"}]]},
            {"layout": "statement", "words": "One [thing]{accent}.", "by": "Me"},
            {"layout": "agenda"},
            {"layout": "blank", "body": [{"code": "x = 1\n# a comment"}]},
        ],
    }


def test_a_slide_background_in_a_theme_colour_follows_the_theme() -> None:
    def drawn(palette: str) -> str:
        document = {"deck": {"id": "t", "palette": palette}, "slides": [{"title": "A", "background": "accent2"}]}
        deck = deck_from_document(document, ROOT)
        assert deck.slides[0].background == "accent2"
        svg = render_slide(deck, deck.slides[0]).svg
        return re.search(r'id="canvas.background"[^>]*fill="(#[0-9a-f]+)"', svg).group(1)

    # Named, not written as a colour: a palette changed later repaints it.
    assert drawn("Okabe-Ito") != drawn("Tableau")
    # In the theme's own colour, not the darker stroke its words and lines take.
    own = {"deck": {"id": "t", "palette": ["#222e50", "#007991", "#439a86", "#bcd8c1", "#e9d985"]},
           "slides": [{"title": "A", "background": "accent4"}]}
    assert deck_from_document(own, ROOT).slides[0].backdrop == "#bcd8c1"


def test_the_demo_as_a_document_draws_the_same_slides() -> None:
    deck = _demo()
    text = dump_document(deck_document(deck))
    again = deck_from_document(yaml.safe_load(text), ROOT / "examples")
    assert len(again.slides) == len(deck.slides)
    for original, copy in zip(deck.slides, again.slides, strict=True):
        assert render_slide(deck, original).svg == render_slide(again, copy).svg, original.id


def test_a_document_writes_itself_back_as_it_was_read() -> None:
    document = {"schema_version": 1, **_small()}
    deck = deck_from_document(document)
    assert deck.style.header == "band" and deck.style.body_size == 22
    assert deck_document(deck) == document


def test_words_over_several_lines_are_written_as_blocks() -> None:
    text = dump_document({"deck": {}, "slides": [{"body": [{"code": "a\nb"}]}]})
    assert "code: |-" in text


@pytest.mark.parametrize(
    ("slide", "where", "words"),
    [
        ({"layout": "sideways"}, "slides[0].layout", "Unknown layout"),
        ({"title": "T", "colour": "red"}, "slides[0]", "Unknown key colour"),
        ({"body": [{"table": "not rows"}]}, "slides[0].body[0] (table)", "list of rows"),
        ({"body": [{"text": "a", "bullets": ["b"]}]}, "slides[0].body[0]", "more than one kind (text, bullets)"),
        ({"body": [{"image": "nowhere.png"}]}, "slides[0].body[0] (image)", "Cannot find nowhere.png"),
        ({"layout": "columns", "columns": [[], []], "widths": [1]}, "slides[0].widths", "one share"),
        ({"body": [{"callout": "c", "colour": "red"}]}, "slides[0].body[0] (callout)", "accent"),
    ],
)
def test_a_wrong_document_says_where(tmp_path: Path, slide: dict, where: str, words: str) -> None:
    with pytest.raises(DeckDocumentError) as caught:
        deck_from_document({"deck": {}, "slides": [slide]}, tmp_path)
    assert caught.value.where == where
    assert words in caught.value.message


def test_a_wrong_slide_stands_aside_when_errors_are_collected(tmp_path: Path) -> None:
    errors: list[DeckDocumentError] = []
    document = {"deck": {}, "slides": [{"title": "Fine"}, {"body": [{"table": 3}]}, {"title": "Also fine"}]}
    deck = deck_from_document(document, tmp_path, errors=errors)
    assert [slide.index for slide in deck.slides] == [1, 2, 3]
    assert deck.slides[1].layout == "blank"
    assert [error.where for error in errors] == ["slides[1].body[0] (table)"]


def test_figures_and_plots_named_by_python_are_made_when_drawn(tmp_path: Path) -> None:
    pytest.importorskip("matplotlib")
    (tmp_path / "made.py").write_text(
        "import flexo\n"
        "calls = []\n"
        "def model():\n"
        "    calls.append('figure')\n"
        "    figure = flexo.Figure('model')\n"
        "    x = figure.root.text('x', 'Input')\n"
        "    figure.root.block('f', label='Model', input=x)\n"
        "    return figure\n"
        "def curve(deck):\n"
        "    import matplotlib.pyplot as plt\n"
        "    calls.append('plot')\n"
        "    figure, axes = plt.subplots()\n"
        "    axes.plot([0, 1], [1, 0])\n"
        "    return figure\n",
        encoding="utf-8",
    )
    document = {"deck": {}, "slides": [{"layout": "two-columns", "left": [{"figure": "made.py:model"}],
                                        "right": [{"plot": "made.py:curve"}]}]}
    deck = deck_from_document(document, tmp_path)
    from flexo_talk.document import import_file

    module = import_file(tmp_path / "made.py")
    assert module.calls == []
    svg = render_slide(deck, deck.slides[0]).svg
    assert module.calls == ["figure", "plot"]
    assert "slide1.left.0" in svg and "slide1.right.0" in svg


def test_the_command_line_builds_and_converts(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    target = tmp_path / "demo.yaml"
    assert main(["convert", str(ROOT / "examples" / "demo.py"), "-o", str(target)]) == 0
    assert read_deck(target).id == "demo"
    assert main(["build", str(target), "-o", str(tmp_path / "out"), "--formats", "svg"]) == 0
    assert len(list((tmp_path / "out").glob("demo-*.svg"))) == 5
    assert "ok" in capsys.readouterr().out


def test_the_studio_draws_the_slide_in_focus_first(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    import flexo_talk.studio as studio

    kind = DeckKind()
    document = {"schema_version": 1, **_small()}
    monkeypatch.setattr(studio, "BUDGET", 0.0)
    first = kind.draw(document, tmp_path, {"focus": 4})
    drawn = [page.id for page in first.pages if not page.pending]
    assert drawn == ["slide5"] and first.unfinished
    while (drawing := kind.draw(document, tmp_path, {"focus": 4})).unfinished:
        pass
    assert all(page.svg for page in drawing.pages)
    assert drawing.info["palette"]["accent"].startswith("#")
    # Nothing changed: nothing is drawn again, and every page is the same.
    again = kind.draw(document, tmp_path, {"focus": 0})
    assert [page.svg for page in again.pages] == [page.svg for page in drawing.pages]


def test_a_figure_written_in_the_deck_is_read_once_for_each_version(tmp_path: Path) -> None:
    figure = {"figure": {"id": "small"}, "nodes": [{"id": "a", "label": "A"}]}
    document = {"deck": {}, "slides": [{"body": [{"figure": figure}]}]}

    def read() -> object:
        return deck_from_document(document, tmp_path).slides[0].regions["body"].blocks[0].figure

    assert read() is read()
    figure["nodes"][0]["label"] = "B"
    assert read().nodes[0].label[0].text == "B"


def test_the_studio_reports_a_wrong_slide_on_its_page(tmp_path: Path) -> None:
    kind = DeckKind()
    document = {"deck": {}, "slides": [{"title": "Fine"}, {"body": [{"stats": []}]}]}
    drawing = kind.draw(document, tmp_path, {})
    assert [page.extra["error"] for page in drawing.pages] == [False, True]
    (message,) = [message for message in drawing.messages if message.severity == "error"]
    assert message.page == "slide2" and message.where == "slides[1].body[0] (stats)"
    assert message.place == "Slide 2 · Numbers"


def test_a_missing_picture_or_figure_file_stands_aside_and_its_slide_is_drawn(tmp_path: Path) -> None:
    from flexo.confine import folder_root

    from flexo_talk.document import MissingFile

    kind = DeckKind()
    document = {"deck": {}, "slides": [
        {"title": "Mine", "body": [{"bullets": ["One point"]}, {"image": "photo.png"}]},
        {"title": "Flow", "background": "paper.png", "body": [{"figure": "flow.yaml"}]},
    ]}
    root = folder_root.set(tmp_path)
    try:
        drawing = kind.draw(document, tmp_path, {})
    finally:
        folder_root.reset(root)
    assert [page.extra["error"] for page in drawing.pages] == [False, False]
    first, second = (page.svg for page in drawing.pages)
    assert "One point" in first and "Missing picture: photo.png" in first
    assert "Missing figure: flow.yaml" in second
    said = [(message.text, message.place) for message in drawing.messages if message.severity == "error"]
    assert said == [
        ("Can't find photo.png in the deck's folder.", "Slide 1 · Picture"),
        ("Can't find paper.png in the deck's folder.", "Slide 2 · Background"),
        ("Can't find flow.yaml in the deck's folder.", "Slide 2 · Figure"),
    ]
    # Watched while missing, so the slide is drawn again once the file is there.
    assert {tmp_path / "photo.png", tmp_path / "flow.yaml"} <= {Path(file) for file in drawing.files}
    # Built for real, a missing file is an error still.
    with pytest.raises(MissingFile, match=r"Cannot find photo\.png"):
        deck_from_document(document, tmp_path)


def test_a_figure_that_cannot_be_drawn_says_why_in_its_box_and_its_slide_is_drawn(tmp_path: Path) -> None:
    from flexo.confine import folder_root

    from flexo_talk.document import InvalidFigure

    timeline = {"id": "t", "kind": "timeline", "label": "Infection", "properties": {
        "events": [{"at": 0}, {"at": 10}], "spans": [{"start": 5, "end": 5, "label": "Span"}]}}
    kind = DeckKind()
    document = {"deck": {}, "slides": [
        {"title": "One infection", "body": [
            {"bullets": ["A point"]}, {"figure": {"figure": {"id": "f"}, "nodes": [timeline]}},
        ]},
    ]}
    root = folder_root.set(tmp_path)
    try:
        drawing = kind.draw(document, tmp_path, {})
    finally:
        folder_root.reset(root)
    # The slide is drawn, its title and words with it -- and the figure, the shape it can't
    # draw a plain box of its words; what is wrong is said under the slide, naming the shape
    # by its own words, its place naming it to choose it by.
    (page,) = drawing.pages
    assert not page.extra["error"]
    assert "One infection" in page.svg and "A point" in page.svg and "Infection" in page.svg
    assert 'data-flexo-talk="invalid"' not in page.svg
    (message,) = [message for message in drawing.messages if message.severity == "error"]
    assert message.text == (
        "\u201cInfection\u201d can\u2019t be drawn yet: span 1 ends where it starts (5): it needs a later end."
        " Choose it to set this in its panel."
    )
    assert "#t:" in message.where
    # Built for real, it is an error still.
    with pytest.raises(InvalidFigure):
        deck_from_document(document, tmp_path)


def test_a_message_says_where_as_a_person_would() -> None:
    from flexo_talk.studio import _place

    document = {"slides": [{"title": "A", "layout": "columns", "columns": [[{"text": "a"}], [{"table": [["x"]]}]]}]}
    assert _place("slides[0].columns[1][0] (table)", document) == "Slide 1 · Table"
    assert _place("slides[0] column2.0", document) == "Slide 1 · Table"
    assert _place("slides[0].title", document) == "Slide 1 · Title"
    assert _place("slides[0]", document) == "Slide 1"
    assert _place("deck.style.title_size", document) == "Design · Title Size"
    assert _place("deck", document) == "Design"
    assert _place("", document) == ""


def test_a_deck_its_editor_cannot_show_is_not_taken_in_or_written_over(tmp_path: Path) -> None:
    import time

    from flexo.studio.workspace import Workspace

    talk = tmp_path / "talk.yaml"
    talk.write_text(dump_document({"deck": {"id": "talk"}, "slides": [{"title": "A"}]}), encoding="utf-8")
    assert DeckKind().malformed({"slides": 5}) == "slides must be a list of slides, not a number (5)"
    assert "empty item" in DeckKind().malformed({"slides": [{"body": [None]}]})
    assert DeckKind().malformed({"slides": [{"title": "A", "body": [{"text": "x"}]}, "odd"]}) is None
    workspace = Workspace(tmp_path)
    try:
        doc = workspace.open("talk.yaml")
        talk.write_text("schema_version: 1\nslides: 5\n", encoding="utf-8")
        deadline = time.monotonic() + 5
        while not doc.held and time.monotonic() < deadline:
            time.sleep(0.05)
        assert doc.held and "slides must be a list" in (doc.problem or "")
        assert doc.document["slides"] == [{"title": "A"}]
        doc.update({**doc.document, "slides": [{"title": "B"}]}, doc.version, {"id": "p", "kind": "person"})
        workspace.flush()
        assert talk.read_text(encoding="utf-8") == "schema_version: 1\nslides: 5\n"
        with pytest.raises(ValueError, match="nothing was changed"):
            doc.update({"slides": 5}, doc.version, {"id": "agent", "kind": "agent"})
    finally:
        workspace.close()


def test_the_studio_catalog_offers_every_look_layout_and_block() -> None:
    catalog = DeckKind().catalog()
    assert {look["name"] for look in catalog["looks"]} == {"classic", "band", "editorial", "keynote", "margin"}
    assert "two-columns" in catalog["slide_keys"] and "split" in catalog["slide_keys"]["two-columns"]
    header = next(field for field in catalog["style"] if field["name"] == "header")
    assert header["kind"] == "choice" and header["choices"] == ["rule", "band", "line", "none"]


def test_the_studio_names_what_a_change_did_slide_by_slide() -> None:
    kind = DeckKind()
    before = {"deck": {"look": "classic"}, "slides": [
        {"title": "A", "body": [{"text": "long enough words to be recognised"}]}, {"title": "B"}, {"title": "C"}]}
    after = {"deck": {"look": "band"}, "slides": [
        {"title": "A!", "body": [{"text": "long enough words to be recognised"}]}, {"title": "New"}, {"title": "B"}]}
    notes = kind.describe(before, after)
    assert [note["text"] for note in notes] == [
        "changed the look", "edited the title on slide 1", "added slide 2, “New”", "deleted slide 3",
    ]
    assert notes[1]["where"] == {"page": 1, "label": "Slide 1"}
    # A slide's objects by name, and a slide changed through and through still itself.
    pasted = {"slides": [{"title": "Plan", "body": [{"bullets": ["One"]}]}]}
    grown = {"slides": [{"title": "Plan", "body": [{"bullets": ["Entirely", "different", "words", "pasted in"]}]}]}
    assert [note["text"] for note in kind.describe(pasted, grown)] == ["edited the list on slide 1"]
    added = {"slides": [{"title": "Plan", "body": [{"bullets": ["One"]}, {"table": [["a", "b"]]}], "date": "2026"}]}
    assert [note["text"] for note in kind.describe(pasted, added)] == ["added a table to slide 1"]
    # Added and deleted by the name the studio gives them, as said aloud: never "a numbers".
    numbers = {"slides": [{"title": "Plan",
                           "body": [{"bullets": ["One"]}, {"stats": [{"value": "93%", "label": "right"}]}]}]}
    assert [note["text"] for note in kind.describe(pasted, numbers)] == ["added numbers to slide 1"]
    coded = {"slides": [{"title": "Plan", "body": [{"bullets": ["One"]}, {"code": "print()"}]}]}
    assert [note["text"] for note in kind.describe(coded, pasted)] == ["deleted code from slide 1"]
    flow = {"slides": [{"title": "Plan", "body": [{"bullets": ["One"]}, {"figure": {"nodes": [
        {"id": "start", "kind": "terminal", "label": "Start"}, {"id": "end", "kind": "terminal", "label": "End"}]}}]}]}
    assert [note["text"] for note in kind.describe(pasted, flow)] == ["added a flow chart to slide 1"]
    dated = {"slides": [{"title": "Plan", "body": [{"bullets": ["One"]}], "date": "2026"}]}
    assert [note["text"] for note in kind.describe(pasted, dated)] == ["edited the date on slide 1"]
    # A setting by its name in the inspector, never its key.
    said = {"slides": [{"layout": "statement", "words": "Go", "by": "Ada"}]}
    resaid = {"slides": [{"layout": "statement", "words": "Go on", "by": "Ada L."}]}
    assert [note["text"] for note in kind.describe(said, resaid)] == [
        "edited the statement and the attribution on slide 1"]
    # Its typing undone, an object still there is emptied, not deleted; a quote's attribution
    # is the attribution.
    typed = {"slides": [{"title": "Plan", "body": [{"bullets": ["One"]}, {"table": [["a", "b"]]}]}]}
    emptied = {"slides": [{"title": "Plan", "body": [{"bullets": ["One"]}, {"table": [["", ""]]}]}]}
    assert [note["text"] for note in kind.describe(typed, emptied)] == ["emptied the table on slide 1"]
    quoted = {"slides": [{"title": "Q", "body": [{"quote": "Words", "by": "Ada"}]}]}
    requoted = {"slides": [{"title": "Q", "body": [{"quote": "Words", "by": "Ada Lovelace"}]}]}
    assert [note["text"] for note in kind.describe(quoted, requoted)] == ["edited the attribution on slide 1"]


def test_an_agent_gets_a_guide_and_a_quick_check() -> None:
    kind = DeckKind()
    guide = kind.guide()
    assert "two-columns" in guide and "bullets" in guide and "Look at each slide" in guide
    assert kind.check({"deck": {}, "slides": [{"body": [{"table": 1}]}]}, ROOT) == [
        "slides[0].body[0] (table): A table must be a list of rows, each a list of cells."
    ]
    document = {"schema_version": 1, **_small()}
    assert kind.parse(kind.dump(document)) == document


def test_the_theme_editor_shows_a_theme_on_slides_and_on_a_deck(tmp_path: Path) -> None:
    from flexo_talk.studio import SlideSamples

    samples = SlideSamples().pages("paper", tmp_path, {})
    assert len(samples) == 5 and "<svg" in samples[1][2]()
    (tmp_path / "talk.yaml").write_text(dump_document({"schema_version": 1, **_small()}), encoding="utf-8")
    own = SlideSamples().pages("tikz", tmp_path, {"deck": "talk.yaml"})
    assert [label for _, label, _ in own][:2] == ["A talk", "Part one"]


def test_a_deck_redraws_when_its_theme_file_changes(tmp_path: Path) -> None:
    import time

    theme = tmp_path / "lab.yaml"
    theme.write_text("theme: {name: lab-deck, base: paper, palette: ['#1d4e89']}\n", encoding="utf-8")
    stats = {"stats": [{"value": "9", "label": "x"}]}
    document = {"deck": {"theme": "lab.yaml"}, "slides": [{"title": "A", "body": [stats]}]}
    kind = DeckKind()
    first = kind.draw(document, tmp_path, {})
    time.sleep(0.01)
    theme.write_text("theme: {name: lab-deck, base: paper, palette: ['#8b1e3f']}\n", encoding="utf-8")
    second = kind.draw(document, tmp_path, {})
    assert first.pages[0].svg != second.pages[0].svg
    assert theme.resolve() in second.files


def test_a_theme_file_sets_the_slides_too(tmp_path: Path) -> None:
    from PIL import Image

    Image.new("RGB", (64, 36), "#f3ecdc").save(tmp_path / "paper.png")
    Image.new("RGB", (64, 36), "#101828").save(tmp_path / "night.png")
    (tmp_path / "notebook.yaml").write_text(
        "theme: {name: notebook-test, base: sketch, palette: ['#2b4c9b', '#c0392b']}\n"
        "slides:\n"
        "  look: margin\n"
        "  style: {header: none, body_size: 22}\n"
        "  background: paper.png\n"
        "  title_font: Caveat\n",
        encoding="utf-8",
    )
    document = {"deck": {"theme": "notebook.yaml"}, "slides": [
        {"title": "Paper"}, {"title": "Night", "background": "night.png"}]}
    deck = deck_from_document(document, tmp_path)
    assert deck.look == "margin" and deck.style.header == "none" and deck.style.body_size == 22
    assert deck.style.edge  # the look's, kept where the theme said nothing
    assert deck.title_font == "Caveat" and deck.background == str(tmp_path / "paper.png")
    # What the theme set is the theme's: the document is written back as it was.
    assert deck_document(deck)["deck"] == {"id": "talk", "theme": "notebook.yaml"}
    first, second = (render_slide(deck, slide).svg for slide in deck.slides)
    assert 'id="slide1.paper"' in first and "Caveat" in first
    # Words over a dark picture are set light, judged from the picture itself;
    # and a slide with a backdrop of its own is not covered by the paper.
    assert "#d4d0c8" in second and "#d4d0c8" not in first
    assert 'id="slide2.paper"' not in second and 'id="slide2.backdrop"' in second

    # The deck's own settings win over the theme's.
    mine = {"deck": {"theme": "notebook.yaml", "look": "band", "background": "#ffffff",
                     "style": {"body_size": 18}}, "slides": [{"title": "Mine"}]}
    deck = deck_from_document(mine, tmp_path)
    assert deck.style.header == "band" and deck.style.body_size == 18 and deck.background == "#ffffff"
    assert deck_document(deck)["deck"] == {"id": "talk", **mine["deck"]}
    assert 'id="slide1.paper"' not in render_slide(deck, deck.slides[0]).svg


def test_a_theme_file_says_where_its_slides_section_is_wrong(tmp_path: Path) -> None:
    (tmp_path / "wrong.yaml").write_text(
        "theme: {name: wrong-slides-test, base: paper}\nslides: {look: sideways}\n", encoding="utf-8"
    )
    with pytest.raises(DeckDocumentError, match="sideways"):
        deck_from_document({"deck": {"theme": "wrong.yaml"}, "slides": []}, tmp_path)


def test_a_table_alone_on_its_slide_is_centred(tmp_path: Path) -> None:
    rows = [["Model", "Active"], ["Ours", "38"]]
    alone = {"deck": {}, "slides": [{"title": "T", "body": [{"table": rows}]}]}
    beside = {"deck": {}, "slides": [{"title": "T", "body": [{"text": "Words above"}, {"table": rows}]}]}
    tables = []
    for document in (alone, beside):
        deck = deck_from_document(document, tmp_path)
        from flexo_talk.compose import render_slide as render

        (table,) = render(deck, deck.slides[0]).tables
        tables.append(table)
    style = deck.style
    body_middle = style.width / 2.0
    # Alone: across the slide and down it.
    assert abs(tables[0].x + sum(tables[0].widths) / 2.0 - body_middle) < 1.0
    assert tables[0].y > tables[1].y + 40
    # Under words, a table stands flush with them.
    assert abs(tables[1].x - style.margin) < 1.0


def test_a_table_or_picture_stands_at_the_left_middle_or_right_of_its_place_when_asked(tmp_path: Path) -> None:
    """Beside a list a table starts at the list's edge; asked (``place``, as the studio's
    Position and its drag write it) it is centred under the list, or set at the right --
    in the slide and in the PowerPoint's native table alike -- and kept so in its file."""

    rows = [["**mechazyme**", "Mechanisms to motifs"], ["**kaveh**", "Motifs to enzymes"]]

    def table_at(place: str | None) -> tuple[float, float]:
        block = {"table": rows, "header": False, **({"horizontal": place} if place is not None else {})}
        document = {"deck": {}, "slides": [{"title": "T", "body": [{"bullets": ["Build the tools"]}, block]}]}
        deck = deck_from_document(document, tmp_path)
        (table,) = render_slide(deck, deck.slides[0]).tables
        return table.x, table.x + sum(table.widths)

    deck = deck_from_document({"deck": {}, "slides": []}, tmp_path)
    margin, width = deck.style.margin, deck.style.width
    assert table_at(None)[0] == pytest.approx(margin, abs=1.0)
    left, right = table_at("middle")
    assert (left + right) / 2.0 == pytest.approx(width / 2.0, abs=1.0)
    assert table_at("centre") == (left, right)
    assert table_at("end")[1] == pytest.approx(width - margin, abs=1.0)
    assert table_at("left")[0] == pytest.approx(margin, abs=1.0)
    # Written as asked, and read back so.
    body = [{"bullets": ["A"]}, {"table": rows, "horizontal": "middle", "vertical": 0.25}]
    again = deck_document(deck_from_document({"deck": {}, "slides": [{"title": "T", "body": body}]}, tmp_path))
    assert again["slides"][0]["body"][1]["horizontal"] == "middle"
    assert again["slides"][0]["body"][1]["vertical"] == 0.25
    # Between: a quarter of the way from the left to the middle, and from the middle to the right.
    quarter, three = table_at(0.25), table_at(0.75)
    assert quarter[0] == pytest.approx(margin + (left - margin) / 2.0, abs=1.0)
    assert three[0] == pytest.approx(left + (left - margin) / 2.0, abs=1.0)
    # Only a side.
    def slide(*body: dict) -> dict:
        return {"deck": {}, "slides": [{"title": "T", "body": list(body)}]}

    with pytest.raises(DeckDocumentError, match="horizontal must be"):
        deck_from_document(slide({"table": rows, "horizontal": "top"}), tmp_path)
    with pytest.raises(DeckDocumentError, match="vertical must be"):
        deck_from_document(slide({"table": rows, "vertical": 2}), tmp_path)
    # In Python, as Build In is: after the object it places.
    made = Deck("placed")
    with made.slide("T") as slide:
        slide.bullets("Build the tools")
        slide.table(rows, header=False)
        slide.place("middle")
    (table,) = render_slide(made, made.slides[0]).tables
    assert table.x + sum(table.widths) / 2.0 == pytest.approx(width / 2.0, abs=1.0)
    with pytest.raises(ValueError, match="Add a block"):
        Deck("empty").slide("T").place("middle")


def test_every_object_stands_at_the_left_middle_or_right_of_its_place_when_asked(tmp_path: Path) -> None:
    """Words, a list, a listing, a quotation, a callout, numbers and an equation set at the
    middle (or the right) of their place stand there as a box as wide as they are, their
    lines as they were set; the PowerPoint's native list with them."""

    blocks = [
        {"text": "A short line of words."},
        {"bullets": ["One point", "Another point", ["A sub point"]]},
        {"code": "print('hello')"},
        {"quote": "Simplicity is prerequisite for reliability.", "by": "Edsger Dijkstra"},
        {"callout": "A short note.", "title": "Note"},
        {"stats": [{"value": "93%", "label": "accuracy"}, {"value": "4x", "label": "faster"}]},
        {"math": "E = mc^2", "align": "start"},
    ]

    def drawn(place: str | None) -> list[tuple[float, float]]:
        body = [{**block, **({"horizontal": place} if place else {})} for block in blocks]
        deck = deck_from_document({"deck": {}, "slides": [{"title": "T", "body": body}]}, tmp_path)
        rendered = render_slide(deck, deck.slides[0])
        svg = rendered.svg
        found = []
        for index in range(len(blocks)):
            xs = [float(x) for x in re.findall(rf'id="slide1\.body\.{index}(?:\.[^"]*)?"[^>]*?\bx="([-\d.]+)"', svg)]
            # (An equation is drawn where its group is moved to.)
            xs += [float(x) for x in re.findall(rf'translate\(([-\d.]+) [^)]*\)"[^>]*id="slide1\.body\.{index}"', svg)]
            found.append((min(xs), max(xs)) if xs else (0.0, 0.0))
        lists = [layout for layout in rendered.lists]
        return found, lists

    deck = deck_from_document({"deck": {}, "slides": []}, tmp_path)
    margin = deck.style.margin
    start, start_lists = drawn(None)
    middle, middle_lists = drawn("middle")
    end, _ = drawn("end")
    for index, block in enumerate(blocks):
        # Each moved right of where the slide sets it, further at the right than in the middle.
        assert start[index][0] < middle[index][0] < end[index][0], block
        assert start[index][0] < margin + 25, block
    # The native list moved with its drawing, its lines as they were.
    assert middle_lists[0].x > start_lists[0].x + 100
    assert [level for level, _, _ in middle_lists[0].items] == [level for level, _, _ in start_lists[0].items]


def test_an_object_stands_down_its_place_what_follows_it_following_it(tmp_path: Path) -> None:
    """Asked to stand at the bottom of its place, between the top and the middle, or in the
    middle, an object takes that share of the room its place has to spare above it; what
    comes after it follows it down, and what is before it stays where it was."""

    rows = [["mechazyme", "Mechanisms to motifs"], ["kaveh", "Motifs to enzymes"]]

    def drawn(vertical: object) -> tuple[float, float, float]:
        table = {"table": rows, "header": False, **({"vertical": vertical} if vertical is not None else {})}
        body = [{"text": "Before"}, table, {"text": "After"}]
        deck = deck_from_document({"deck": {}, "slides": [{"title": "T", "body": body}]}, tmp_path)
        rendered = render_slide(deck, deck.slides[0])
        before = float(re.search(r'id="slide1\.body\.0"[^>]*\by="([-\d.]+)"', rendered.svg).group(1))
        (table,) = rendered.tables
        return before, table.y, sum(table.heights)

    before, top, _ = drawn(None)
    for vertical, share in (("top", 0.0), (0.25, 0.25), ("middle", 0.5), ("bottom", 1.0)):
        now, y, _ = drawn(vertical)
        assert now == pytest.approx(before)
        assert y >= top - 0.01
        if share:
            assert y > top + 1.0
    # Each a share of the same spare room: the middle twice as far down as a quarter.
    _, quarter, _ = drawn(0.25)
    _, middle, _ = drawn("middle")
    assert middle - top == pytest.approx(2 * (quarter - top), rel=0.02)


def _deck_with_figures() -> dict:
    figure = {
        "figure": {"id": "inline"},
        "nodes": [{"id": "x", "kind": "text", "label": "x"}, {"id": "m", "label": "Model"}],
        "edges": [{"from": "x", "to": "m"}],
    }
    return {
        "schema_version": 1,
        "deck": {"id": "talk"},
        "slides": [
            {"title": "Here", "layout": "two-columns", "left": [{"text": "Words"}], "right": [{"figure": figure}]},
            {"title": "A file", "layout": "figure", "body": [{"figure": "model.yaml"}]},
            {"title": "Python", "body": [{"figure": "figures.py:model"}]},
        ],
    }


def test_a_figure_written_in_the_deck_is_edited_where_it_is_written(tmp_path: Path) -> None:
    kind = DeckKind()
    document = _deck_with_figures()
    at = {"slide": 0, "region": "right", "index": 0}
    read = kind.act(document, {"do": "figure", "at": at, "edit": {"do": "read"}}, tmp_path)
    assert [node["id"] for node in read["model"]["nodes"]] == ["x", "m"] and read["document"] == document
    edit = {"do": "add", "kind": "mlp", "after": "m", "source": "m"}
    made = kind.act(document, {"do": "figure", "at": at, "edit": edit}, tmp_path)
    assert made["select"] == ["mlp"]
    figure = made["document"]["slides"][0]["right"][0]["figure"]
    assert {"from": "m", "to": "mlp"} in figure["edges"] and len(figure["nodes"]) == 3
    assert len(document["slides"][0]["right"][0]["figure"]["nodes"]) == 2  # the one given is left alone
    deck = deck_from_document({**made["document"], "slides": made["document"]["slides"][:1]}, tmp_path)
    assert render_slide(deck, deck.slides[0]).svg


def test_a_figure_file_on_a_slide_is_edited_in_its_file(tmp_path: Path) -> None:
    from flexo.studio.figure_kind import SAMPLE_FIGURE

    (tmp_path / "model.yaml").write_text(SAMPLE_FIGURE, encoding="utf-8")
    kind = DeckKind()
    document = _deck_with_figures()
    at = {"slide": 1, "region": "body", "index": 0}
    kind.act(document, {"do": "figure", "at": at, "edit": {"do": "read"}}, tmp_path)
    assert (tmp_path / "model.yaml").read_text() == SAMPLE_FIGURE  # reading writes nothing
    rename = {"do": "rename", "id": "encoder", "to": "backbone"}
    result = kind.act(document, {"do": "figure", "at": at, "edit": rename}, tmp_path)
    assert result["file"] == "model.yaml" and result["document"] is document
    text = (tmp_path / "model.yaml").read_text()
    assert text.startswith("# A flexo figure") and "id: backbone" in text and "to: backbone" in text


def test_an_edit_to_a_figure_file_is_undone_by_putting_the_file_back(tmp_path: Path) -> None:
    from flexo.studio.figure_edit import EditError
    from flexo.studio.figure_kind import SAMPLE_FIGURE

    path = tmp_path / "model.yaml"
    path.write_text(SAMPLE_FIGURE, encoding="utf-8")
    kind = DeckKind()
    document = _deck_with_figures()
    at = {"slide": 1, "region": "body", "index": 0}
    read = kind.act(document, {"do": "figure", "at": at, "edit": {"do": "read"}}, tmp_path)
    assert "was" not in read  # nothing changed, nothing to undo
    rename = {"do": "rename", "id": "encoder", "to": "backbone"}
    result = kind.act(document, {"do": "figure", "at": at, "edit": rename}, tmp_path)
    assert result["was"] == SAMPLE_FIGURE and result["now"] == path.read_text()
    was, now = result["change"]
    assert "encoder" in [node["id"] for node in was["nodes"]] and "backbone" in [node["id"] for node in now["nodes"]]

    def restore(text: str, expect: str) -> None:
        kind.act(document, {"do": "figure-file", "file": "model.yaml", "text": text, "expect": expect}, tmp_path)

    restore(result["was"], result["now"])  # undone
    assert path.read_text() == SAMPLE_FIGURE
    restore(result["now"], result["was"])  # done again
    assert path.read_text() == result["now"]
    # Changed since by hand, elsewhere in the file: undone and done again around it.
    path.write_text(result["now"] + "\n# mine\n", encoding="utf-8")
    restore(result["was"], result["now"])
    assert path.read_text() == SAMPLE_FIGURE + "\n# mine\n"
    restore(result["now"], result["was"])
    assert path.read_text() == result["now"] + "\n# mine\n"
    # Changed since in the same place: left as it is.
    path.write_text(result["now"].replace("backbone", "trunk") + "\n# mine\n", encoding="utf-8")
    with pytest.raises(EditError, match="changed in the same place since"):
        restore(result["was"], result["now"])
    assert "trunk" in path.read_text() and path.read_text().endswith("# mine\n")
    with pytest.raises(EditError, match="Can't find the figure file"):
        kind.act(document, {"do": "figure-file", "file": "../elsewhere.yaml", "text": "", "expect": ""}, tmp_path)


def test_a_figure_made_in_python_or_gone_is_not_edited_on_its_slide(tmp_path: Path) -> None:
    from flexo.studio.figure_edit import EditError

    kind = DeckKind()
    document = _deck_with_figures()
    def read(slide: int, region: str, index: int) -> None:
        at = {"slide": slide, "region": region, "index": index}
        kind.act(document, {"do": "figure", "at": at, "edit": {"do": "read"}}, tmp_path)

    with pytest.raises(EditError, match="Python file"):
        read(2, "body", 0)
    with pytest.raises(EditError, match="no longer exists"):
        read(0, "right", 5)


def test_the_studio_offers_the_figure_editor_for_figures_on_slides() -> None:
    parts = DeckKind().catalog()["figure_editor"]["parts"]
    assert {"block", "protein", "plasmid", "attention"} <= set(parts)


def test_a_figure_on_a_slide_is_exported_by_itself_in_the_decks_look(tmp_path: Path) -> None:
    (tmp_path / "logo.svg").write_text(
        '<svg xmlns="http://www.w3.org/2000/svg" width="20" height="20"><circle cx="10" cy="10" r="8"/></svg>'
    )
    figure = {"figure": {"id": "pic"}, "nodes": [
        {"id": "logo", "kind": "image", "properties": {"source": "logo.svg"}}, {"id": "m", "label": "Model"}]}
    document = {"schema_version": 1, "deck": {"id": "t", "theme": "dark"},
                "slides": [{"title": "Logo", "body": [{"text": "Words"}, {"figure": figure}]}]}
    at = {"slide": 0, "region": "body", "index": 1}
    written = DeckKind().export_part(document, tmp_path, "talk", at, ["yaml", "editable"])
    assert sorted(path.name for path in written) == ["talk-pic.editable.svg", "talk-pic.yaml"]
    exported = yaml.safe_load((tmp_path / "build" / "talk-pic.yaml").read_text())
    # In the deck's theme, and naming its picture from where it is written.
    assert exported["figure"]["style"] == "dark"
    assert exported["nodes"][0]["properties"]["source"] == "../logo.svg"
    from flexo.serialization import load_figure

    assert load_figure(tmp_path / "build" / "talk-pic.yaml").id == "pic"
    from flexo.studio.figure_edit import EditError

    with pytest.raises(EditError, match="isn't a figure"):
        DeckKind().export_part(document, tmp_path, "talk", {**at, "index": 0}, ["yaml"])


def test_a_figure_file_is_written_into_the_deck_its_files_named_from_there(tmp_path: Path) -> None:
    (tmp_path / "figures").mkdir()
    (tmp_path / "logo.svg").write_text(
        '<svg xmlns="http://www.w3.org/2000/svg" width="20" height="20"><circle cx="10" cy="10" r="8"/></svg>'
    )
    text = "# its own\nfigure: {id: pic}\nnodes:\n- {id: logo, kind: image, properties: {source: ../logo.svg}}\n"
    (tmp_path / "figures" / "pic.yaml").write_text(text)
    document = {"schema_version": 1, "deck": {"id": "t"},
                "slides": [{"title": "Logo", "body": [{"figure": "figures/pic.yaml"}]}]}
    # Drawn from its file, its picture is found beside the file.
    deck = deck_from_document(document, tmp_path)
    assert render_slide(deck, deck.slides[0]).svg
    result = DeckKind().act(document, {"do": "inline", "at": {"slide": 0, "region": "body", "index": 0}}, tmp_path)
    inline = result["document"]["slides"][0]["body"][0]["figure"]
    assert inline["nodes"][0]["properties"]["source"] == "logo.svg"
    assert (tmp_path / "figures" / "pic.yaml").read_text() == text  # the file is left as it was
    deck = deck_from_document(result["document"], tmp_path)
    assert render_slide(deck, deck.slides[0]).svg


def test_a_picture_in_a_figure_written_in_the_deck_is_found_beside_the_deck(tmp_path: Path) -> None:
    (tmp_path / "logo.svg").write_text(
        '<svg xmlns="http://www.w3.org/2000/svg" width="20" height="20"><circle cx="10" cy="10" r="8"/></svg>'
    )
    figure = {"figure": {"id": "pic"}, "nodes": [{"id": "logo", "kind": "image", "properties": {"source": "logo.svg"}}]}
    document = {"schema_version": 1, "deck": {"id": "t"}, "slides": [{"title": "Logo", "body": [{"figure": figure}]}]}
    deck = deck_from_document(document, tmp_path)
    assert render_slide(deck, deck.slides[0]).svg


def test_one_theme_file_is_put_to_use_in_a_deck_and_a_figure_at_once(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from flexo.studio import theming
    from flexo.studio.figure_kind import SAMPLE_FIGURE
    from flexo.studio.workspace import Workspace

    monkeypatch.setenv("XDG_RUNTIME_DIR", str(tmp_path / "run"))
    (tmp_path / "themes").mkdir()
    (tmp_path / "themes" / "lab.theme.yaml").write_text(
        "theme: {name: lab-put-to-use, base: paper, palette: ['#8b1e3f']}\n", encoding="utf-8")
    (tmp_path / "talk.yaml").write_text(yaml.safe_dump({"deck": {"id": "talk"}, "slides": [{"title": "A"}]}))
    (tmp_path / "figure.yaml").write_text(SAMPLE_FIGURE, encoding="utf-8")
    workspace = Workspace(tmp_path)
    try:
        after = theming.use(workspace, "themes/lab.theme.yaml", ["talk.yaml", "figure.yaml"],
                            {"id": "page", "name": "Ada", "kind": "person"})
        cards = theming.cards(workspace, "talk.yaml")
        drawn = workspace.open("talk.yaml").kind.draw(workspace.open("talk.yaml").document, tmp_path, {})
    finally:
        workspace.close()
    assert {entry["file"]: entry["uses"] for entry in after} == {"figure.yaml": True, "talk.yaml": True}
    assert yaml.safe_load((tmp_path / "talk.yaml").read_text())["deck"]["theme"] == "themes/lab.theme.yaml"
    assert yaml.safe_load((tmp_path / "figure.yaml").read_text())["figure"]["style"] == "themes/lab.theme.yaml"
    assert cards[0]["value"] == "themes/lab.theme.yaml"
    assert drawn.info["tones"]["colours"][0]["stroke"].startswith("#")


def test_a_deck_s_python_is_kept_to_itself(tmp_path: Path) -> None:
    import os
    import time

    import matplotlib
    import matplotlib.pyplot as plt

    matplotlib.use("Agg")
    for folder, word in (("one", "first"), ("two", "second")):
        (tmp_path / folder).mkdir()
        (tmp_path / folder / "utils.py").write_text(f"WORD = {word!r}\n")
        (tmp_path / folder / "plots.py").write_text(
            "from dataclasses import dataclass\n"
            "import matplotlib.pyplot as plt\n"
            "from utils import WORD\n"
            "plt.rcParams['lines.linewidth'] = 9\n"
            "@dataclass\n"
            "class Point:\n"
            "    x: float\n"
            "def plot():\n"
            "    figure, axes = plt.subplots()\n"
            "    axes.plot([Point(0).x, 1])\n"
            "    axes.set_title(WORD)\n"
            "    return figure\n"
            "def leaves():\n"
            "    os.chdir('/')\n"
            "    return plot()\n"
            "def quits():\n"
            "    raise SystemExit(3)\n"
            "import os\n"
        )
    kind = DeckKind()
    width, figures, here = plt.rcParams["lines.linewidth"], len(plt.get_fignums()), os.getcwd()

    def draw(folder: str, function: str = "plot") -> str:
        document = {"slides": [{"title": "A", "body": [{"plot": f"plots.py:{function}"}]}]}
        drawing = kind.draw(document, tmp_path / folder, {})
        return drawing.pages[0].svg + " ".join(message.text for message in drawing.messages)

    assert "first" in draw("one") and "second" in draw("two")
    assert "stopped: SystemExit 3" in draw("one", "quits")
    draw("two", "leaves")
    assert os.getcwd() == here
    assert plt.rcParams["lines.linewidth"] == width and len(plt.get_fignums()) == figures
    time.sleep(0.01)
    (tmp_path / "one" / "utils.py").write_text("WORD = 'edited'\n")
    assert "edited" in draw("one")


def test_a_folder_not_trusted_runs_none_of_its_python_and_reads_only_itself(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from flexo.studio.workspace import Workspace

    monkeypatch.setenv("XDG_RUNTIME_DIR", str(tmp_path / "run"))
    folder = tmp_path / "sent"
    folder.mkdir()
    ran = tmp_path / "ran.txt"
    (folder / "plots.py").write_text(
        "from pathlib import Path\n"
        f"Path({str(ran)!r}).write_text('ran')\n"
        "import matplotlib.pyplot as plt\n"
        "def loss():\n    figure, axes = plt.subplots()\n    axes.plot([1, 2])\n    return figure\n"
    )
    (tmp_path / "private.png").write_bytes(b"not for the deck")
    (folder / "talk.yaml").write_text(yaml.safe_dump({"deck": {"id": "talk"}, "slides": [
        {"title": "Plot", "body": [{"plot": "plots.py:loss"}]},
        {"title": "Private", "body": [{"image": "../private.png"}]},
    ]}))
    workspace = Workspace(folder, trusted=False)
    try:
        doc = workspace.open("talk.yaml")
        first = workspace.draw("talk.yaml", doc.document, 1, {}, {})
        ran_before_trust = ran.exists()
        workspace.trust()
        second = workspace.draw("talk.yaml", doc.document, 2, {}, {})
    finally:
        workspace.close()
    codes = {message["code"] for message in first["messages"]}
    assert "code.untrusted" in codes and not ran_before_trust
    assert any("outside the folder" in message["text"] for message in first["messages"])
    assert ran.read_text() == "ran"
    assert "code.untrusted" not in {message["code"] for message in second["messages"]}
    assert any("outside the folder" in message["text"] for message in second["messages"])


def test_in_the_studio_a_decks_python_runs_apart_and_cannot_hang_or_take_down_the_app(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    pytest.importorskip("matplotlib")
    import os

    from flexo.studio.workspace import Workspace

    from flexo_talk import worker

    monkeypatch.setenv("XDG_RUNTIME_DIR", str(tmp_path / "run"))
    monkeypatch.setattr(worker, "TIMEOUT", 3.0)
    folder = tmp_path / "talk"
    folder.mkdir()
    (folder / "numbers.txt").write_text("3 1 2")
    (folder / "plots.py").write_text(
        "import os, sys, time\n"
        "from pathlib import Path\n"
        "import flexo\n"
        "import matplotlib.pyplot as plt\n"
        "def good(deck):\n"
        "    print('printed words go to the log, not among the answers')\n"
        "    numbers = [float(n) for n in Path('numbers.txt').read_text().split()]\n"
        "    figure, axes = plt.subplots()\n"
        "    axes.plot(numbers, color=deck.palette.get('tone-1-stroke'))\n"
        "    axes.set_title(f'pid {os.getpid()}')\n"
        "    return figure\n"
        "def model():\n"
        "    figure = flexo.Figure('model')\n"
        "    figure.root.block('f', label='Made apart')\n"
        "    return figure\n"
        "def forever():\n"
        "    while True:\n"
        "        time.sleep(0.05)\n"
        "def crash():\n"
        "    os._exit(3)\n"
        "def leave():\n"
        "    sys.exit(2)\n"
        "def nothing():\n"
        "    return 7\n"
        "def fault():\n"
        "    import ctypes\n"
        "    ctypes.string_at(0)\n"
        "def typo():\n"
        "    return figur\n"
    )
    slides = [{"title": name, "body": [{kind: f"plots.py:{name}"}]} for kind, name in (
        ("plot", "good"), ("figure", "model"), ("plot", "forever"), ("plot", "crash"), ("plot", "leave"),
        ("plot", "nothing"), ("plot", "good"), ("plot", "fault"), ("figure", "typo"),
    )]
    (folder / "talk.yaml").write_text(yaml.safe_dump({"deck": {"id": "talk"}, "slides": slides}))
    workspace = Workspace(folder, trusted=True)
    try:
        started = __import__("time").monotonic()
        drawing = workspace.drawing_of("talk.yaml")
        seconds = __import__("time").monotonic() - started
    finally:
        workspace.close()
        worker.stop_all()
    said = {message.page: message.text for message in drawing.messages if message.severity == "error"}
    pages = {page.id: page for page in drawing.pages}
    assert f"pid {os.getpid()}" not in pages["slide1"].svg and "pid " in pages["slide1"].svg
    assert "Made apart" in pages["slide2"].svg
    assert said["slide3"] == "plots.py:forever took longer than 3 seconds and was stopped."
    assert said["slide4"] == "plots.py:crash quit unexpectedly (exit code 3)."
    assert said["slide5"] == "plots.py:leave: sys.exit(2) was called (line 22)"
    assert said["slide6"] == "plots.py:nothing: nothing did not return a matplotlib figure"
    assert said["slide8"] == "plots.py:fault crashed (signal 11)."
    assert said["slide9"] == "plots.py:typo: name 'figur' is not defined (line 29)"
    assert "slide1" not in said and "slide7" not in said and "pid " in pages["slide7"].svg
    assert seconds < 30


def test_a_mechanism_block_is_read_and_its_mistakes_said_at_their_place(tmp_path: Path) -> None:
    document = {
        "deck": {"id": "chemistry"},
        "slides": [{"title": "SN2", "body": [{"mechanism": [
            {"smiles": "[OH-:1].[CH3:2][Br:3]", "arrows": ["1 -> 2", "2-3 -> 3"], "label": "backside attack"},
        ]}]}],
    }
    deck = deck_from_document(document, tmp_path)
    (rendered,) = deck.render()
    assert "backside attack" in rendered.svg
    # Half drawn, the step cannot be: it is drawn as far as it goes, and why is said.
    document["slides"][0]["body"][0]["mechanism"][0]["arrows"] = ["1 -> 2"]
    (rendered,) = deck_from_document(document, tmp_path).render()
    assert "backside attack" in rendered.svg
    assert any("C2 would have 10 electrons" in said for said in rendered.diagnostics)
    document["slides"][0]["body"][0]["mechanism"][0]["smiles"] = "[OH-:1].[CH3:2][Br:3"
    with pytest.raises(DeckDocumentError, match=r"slides\[0\]\.body\[0\] \(mechanism\).*\[ is never closed"):
        deck_from_document(document, tmp_path)


def test_the_studio_draws_a_mechanisms_arrows_from_two_clicks(tmp_path: Path) -> None:
    kind = DeckKind()
    document = {"deck": {"id": "c"}, "slides": [{"title": "SN2", "body": [
        {"mechanism": [{"smiles": "[OH-:1].[CH3:2][Br:3]"}]}]}]}
    at = {"slide": 0, "region": "body", "index": 0}
    seen = kind.act(document, {"do": "mechanism", "at": at, "step": 0, "holding": []}, tmp_path)
    assert seen["document"] is document and seen["sheet"]["arrows"] == [] and seen["sheet"]["svg"]
    # The hydroxide's lone pair to the carbon: drawn, and the five-bonded carbon said, not refused.
    made = kind.act(document, {"do": "mechanism", "at": at, "step": 0, "holding": [],
                               "add": {"tail": {"atom": 0}, "head": {"atom": 1}}}, tmp_path)
    assert made["arrow"] == "1 -> 2"
    assert made["document"]["slides"][0]["body"][0]["mechanism"] == [
        {"smiles": "[OH-:1].[CH3:2][Br:3]", "arrows": ["1 -> 2"]}]
    assert "10 electrons" in made["sheet"]["problem"]["message"]
    asked = kind.act(made["document"], {"do": "mechanism", "at": at, "step": 0,
                                        "add": {"tail": {"bond": [1, 2]}, "head": {"atom": 0}}}, tmp_path)
    assert [end["name"] for end in asked["ends"]] == ["C2", "Br3"]
    gone = kind.act(made["document"], {"do": "mechanism", "at": at, "step": 0, "remove": 0}, tmp_path)
    assert gone["document"]["slides"][0]["body"][0]["mechanism"] == [{"smiles": "[OH-:1].[CH3:2][Br:3]"}]
    # A molecule dragged on the sheet: kept on its step, by one of its atoms, and drawn there.
    placed = kind.act(made["document"], {"do": "mechanism", "at": at, "step": 0,
                                         "place": {"atom": 0, "move": [-1, 0.5]}}, tmp_path)
    step = placed["document"]["slides"][0]["body"][0]["mechanism"][0]
    assert step["place"] == {"1": {"move": [-1.0, 0.5]}}
    assert placed["sheet"]["molecules"][0]["placed"] is True
    (rendered,) = deck_from_document(placed["document"], tmp_path).render()
    assert "mechanism" in rendered.svg
    with pytest.raises(DeckDocumentError, match="place must map an atom"):
        bad = {"deck": {"id": "c"}, "slides": [{"body": [{"mechanism": [
            {"smiles": "[OH-:1].[CH3:2][Br:3]", "place": "1 sideways"}]}]}]}
        deck_from_document(bad, tmp_path)


def test_a_figure_keeps_its_layout_while_the_deck_is_changed_and_settles_after(tmp_path: Path, monkeypatch) -> None:
    import flexo

    monkeypatch.setenv("FLEXO_TALK_CACHE", "0")
    kind = DeckKind()
    nodes = [{"id": name, "label": name.title()} for name in ("one", "two", "three")]
    figure = {"figure": {"id": "chain"}, "nodes": nodes,
              "edges": [{"from": "one", "to": "two"}, {"from": "two", "to": "three"}]}
    document = {"deck": {"id": "d"}, "slides": [{"title": "Chain", "layout": "figure", "body": [{"figure": figure}]}]}
    first = kind.draw(document, tmp_path)
    assert first.info["unsettled"] is False  # drawn the first time, it is laid out at its best
    seen = []
    real = flexo.fit_in_box
    monkeypatch.setattr(flexo, "fit_in_box", lambda *a, **k: seen.append(k.get("keep")) or real(*a, **k))
    nodes[1]["label"] = "Second"
    edited = kind.draw(document, tmp_path)
    assert seen[-1] is not None and edited.info["unsettled"] is True  # kept its layout, in one compile
    settled = kind.draw(document, tmp_path, {"settle": True})
    assert seen[-1] is None and settled.info["unsettled"] is False
    assert kind.draw(document, tmp_path).info["unsettled"] is False  # and stays settled


def test_a_figure_whose_settling_was_given_up_is_not_laid_out_afresh_while_it_is_changed(
    tmp_path: Path, monkeypatch
) -> None:
    import itertools
    import time

    import flexo
    from flexo.draft import GivenUp, given_up_when

    monkeypatch.setenv("FLEXO_TALK_CACHE", "0")
    ids = [f"s{index}" for index in range(6)]
    nodes = [{"id": id, "label": f"Stage {index + 1}"} for index, id in enumerate(ids)]
    root = {"id": "root", "layout": {"kind": "row"}, "children": list(ids)}
    figure = {"figure": {"id": "chain"}, "nodes": nodes, "groups": [root],
              "edges": [{"from": a, "to": b} for a, b in itertools.pairwise(ids)]}
    document = {"deck": {"id": "given"}, "slides": [{"title": "Stages", "body": [{"figure": figure}]}]}
    kind = DeckKind()
    kind.draw(document, tmp_path, {"settle": True})  # laid out at its best, quickly
    # Changed, and its settling given up a while into its layouts -- a large figure, say -- for a
    # newer change.
    nodes[0]["label"] = "Stage one"
    real = flexo.fit_in_box
    begun = []
    monkeypatch.setattr(flexo, "fit_in_box", lambda *a, **k: begun.append(time.perf_counter()) or real(*a, **k))

    def newer() -> bool:
        if not begun:
            return False
        time.sleep(max(0.0, begun[0] + 1.2 - time.perf_counter()))
        return True

    with pytest.raises(GivenUp), given_up_when(newer):
        kind.draw(document, tmp_path, {"settle": True})
    # Changed again so that, kept as it is, it would be drawn much smaller (a shape added to a
    # long row): it is drawn as a draft, kept, in one compile -- not laid out every way at once,
    # which, for all its person knows, takes a while.
    calls = []
    monkeypatch.setattr(flexo, "fit_in_box", lambda *a, **k: calls.append(k.get("keep")) or real(*a, **k))
    nodes.extend({"id": f"x{index}", "label": f"Extra step {index}"} for index in range(6))
    root["children"] += [f"x{index}" for index in range(6)]
    drawn = kind.draw(document, tmp_path)
    assert len(calls) == 1 and calls[0] is not None and drawn.info["unsettled"] is True


def test_a_figure_changed_on_a_slide_is_drawn_as_a_draft_and_tried_other_ways_once_settled(
    tmp_path: Path, monkeypatch
) -> None:
    import itertools

    import flexo
    from flexo.draft import DRAFT

    monkeypatch.setenv("FLEXO_TALK_CACHE", "0")
    # A long row arranged by hand -- its last step put under the one before -- drawn small:
    # once settled, the slide says whether turning it would set its words larger.
    ids = [f"s{index}" for index in range(12)]
    nodes = [{"id": id, "label": f"Stage {index + 1}"} for index, id in enumerate(ids)]
    groups = [{"id": "root", "layout": {"kind": "row"}, "children": [*ids[:-2], "pair"]},
              {"id": "pair", "layout": {"kind": "column"}, "children": ids[-2:]}]
    edges = [{"from": a, "to": b} for a, b in itertools.pairwise(ids)]
    figure = {"figure": {"id": "long"}, "nodes": nodes, "groups": groups, "edges": edges}
    document = {"deck": {"id": "drafts"}, "slides": [{"title": "Stages", "body": [{"figure": figure}]}]}
    kind = DeckKind()
    kind.draw(document, tmp_path)
    seen = []
    real = flexo.fit_in_box
    monkeypatch.setattr(flexo, "fit_in_box", lambda *a, **k: seen.append((DRAFT.get(), k["turn"])) or real(*a, **k))
    nodes[3]["label"] = "Stage four"
    edited = kind.draw(document, tmp_path)
    # Drawn as a draft, once, as arranged: not laid out turned as well at every change.
    assert seen == [(True, False)] and edited.info["unsettled"] is True
    seen.clear()
    settled = kind.draw(document, tmp_path, {"settle": True})
    assert settled.info["unsettled"] is False
    assert (False, False) in seen and (False, True) in seen and not any(draft for draft, _ in seen)


def test_a_folded_chart_keeps_its_fold_when_its_words_are_changed(tmp_path: Path, monkeypatch) -> None:
    from flexo_talk import compose

    monkeypatch.setenv("FLEXO_TALK_CACHE", "0")
    names = {"start": ("terminal", "Message arrives"), "decision": ("decision", "Seen this ID?"),
             "skip": ("terminal", "Skip it"), "step": (None, "Update the order"),
             "check": ("decision", "Worked?"), "end": ("terminal", "Acknowledge")}
    nodes = [{"id": id, "label": label, **({"kind": kind} if kind else {})} for id, (kind, label) in names.items()]
    edges = [{"from": "start", "to": "decision"}, {"from": "decision", "to": "skip", "label": "yes"},
             {"from": "decision", "to": "step", "label": "no"}, {"from": "step", "to": "check"},
             {"from": "check", "to": "end", "label": "yes"}, {"from": "check", "to": "step", "label": "no"}]
    root = {"id": "root", "role": "canvas", "layout": {"kind": "row"}, "children": list(names)}
    figure = {"figure": {"id": "consumer"}, "nodes": nodes, "groups": [root], "edges": edges}
    document = {"deck": {"id": "folds"}, "slides": [{"title": "How a consumer handles a message",
                                                      "body": [{"figure": figure}]}]}
    kind = DeckKind()
    kind.draw(document, tmp_path, {"settle": True})
    where = next(key for key in compose._LAYOUTS if key[0] == "folds")
    assert compose._LAYOUTS[where] == "as written, folded"  # one long row, wrapped onto two
    folded = compose._SCALES[where]
    # Its words changed and settled: still folded, as large -- kept from turning (it was seen
    # as written), it is not kept from folding, and laid out as one long row at half the size.
    nodes[3]["label"] = "Update the order now"
    kind.draw(document, tmp_path)
    kind.draw(document, tmp_path, {"settle": True})
    assert compose._LAYOUTS[where] == "as written, folded"
    assert compose._SCALES[where] > folded * 0.8
    # Written as it is seen, each line a row of its own (as the studio writes them before a
    # part is put among them), it is drawn as written: not folded, nor turned, again.
    lines = [["start", "decision", "skip"], ["step", "check", "end"]]
    root.update(layout={"kind": "column"}, children=["line", "line-2"])
    figure["groups"] += [{"id": name, "role": "layout", "layout": {"kind": "row"}, "children": line}
                         for name, line in zip(("line", "line-2"), lines, strict=True)]
    kind.draw(document, tmp_path)
    assert compose._LAYOUTS[where] == "as written"


def test_a_chart_kept_from_turning_is_still_folded_to_fit(tmp_path: Path, monkeypatch) -> None:
    # Turn to Fit the Slide off (`turn: false`) keeps a chart from turning, not from folding:
    # the long row is wrapped onto two lines as it is left to the deck, never drawn unwrapped.
    from flexo_talk import compose

    monkeypatch.setenv("FLEXO_TALK_CACHE", "0")
    names = {"start": ("terminal", "Message arrives"), "decision": ("decision", "Seen this ID?"),
             "skip": ("terminal", "Skip it"), "step": (None, "Update the order"),
             "check": ("decision", "Worked?"), "end": ("terminal", "Acknowledge")}
    nodes = [{"id": id, "label": label, **({"kind": kind} if kind else {})} for id, (kind, label) in names.items()]
    edges = [{"from": "start", "to": "decision"}, {"from": "decision", "to": "skip", "label": "yes"},
             {"from": "decision", "to": "step", "label": "no"}, {"from": "step", "to": "check"},
             {"from": "check", "to": "end", "label": "yes"}, {"from": "check", "to": "step", "label": "no"}]
    root = {"id": "root", "role": "canvas", "layout": {"kind": "row"}, "children": list(names)}
    layouts = {}
    for turn in (None, False):
        figure = {"figure": {"id": "kept"}, "nodes": nodes, "groups": [root], "edges": edges}
        block = {"figure": figure, **({} if turn is None else {"turn": turn})}
        document = {"deck": {"id": f"kept-{turn}"}, "slides": [{"title": "Steps", "body": [block]}]}
        DeckKind().draw(document, tmp_path, {"settle": True})
        layouts[turn] = compose._LAYOUTS[next(key for key in compose._LAYOUTS if key[0] == f"kept-{turn}")]
    assert "folded" in layouts[None] and "turned" not in layouts[None]
    assert layouts[False] == layouts[None]


def test_a_turned_figure_written_as_it_was_drawn_is_not_turned_back(tmp_path: Path, monkeypatch) -> None:
    import itertools

    from flexo_talk import compose

    monkeypatch.setenv("FLEXO_TALK_CACHE", "0")
    names = ["Purify CA", "Mix CA with IP6", "Negative-stain EM", "Tubes formed?", "Cryo-EM grids",
             "Collect movies", "Motion correction", "3D refinement"]
    ids = list("abcdefgh")
    nodes = [{"id": id, "label": name} for id, name in zip(ids, names, strict=True)]
    nodes[3]["kind"] = "decision"
    root = {"id": "root", "layout": {"kind": "column"}, "children": list(ids)}
    figure = {"figure": {"id": "assay"}, "nodes": nodes, "groups": [root],
              "edges": [{"from": a, "to": b} for a, b in itertools.pairwise(ids)]}
    document = {"deck": {"id": "turned"}, "slides": [{"title": "Assay", "body": [{"figure": figure}]}]}
    kind = DeckKind()
    kind.draw(document, tmp_path)
    where = next(key for key in compose._LAYOUTS if key[0] == "turned")
    assert compose._LAYOUTS[where].startswith("turned")  # a column, drawn as a row to fit
    # The studio writes it as it is drawn, a row, to put the last part under the one before:
    # kept as it is drawn while it is changed, it is not turned back into a column (the part
    # put under the other drawn beside it).
    root["layout"]["kind"] = "row"
    root["children"] = [*ids[:-2], "pair"]
    figure["groups"].append({"id": "pair", "layout": {"kind": "column"}, "children": ids[-2:]})
    kind.draw(document, tmp_path)
    assert compose._LAYOUTS[where].startswith("as written")


def test_a_figure_arranged_by_hand_is_drawn_as_arranged_and_the_same_whatever_came_before(
    tmp_path: Path, monkeypatch
) -> None:
    import copy

    from flexo_talk import compose

    monkeypatch.setenv("FLEXO_TALK_CACHE", "0")
    names = {"u": "KaiC-U", "t": "KaiC-T", "st": "KaiC-ST", "s": "KaiC-S", "b": "KaiB binds"}
    nodes = [{"id": id, "label": name} for id, name in names.items()]
    groups = [
        {"id": "root", "layout": {"kind": "column"}, "children": ["row", "b"]},
        {"id": "row", "layout": {"kind": "row"}, "children": ["u", "t", "st", "s"]},
    ]
    edges = [{"from": "u", "to": "t"}, {"from": "t", "to": "st"}, {"from": "st", "to": "s"},
             {"from": "s", "to": "b"}]
    figure = {"figure": {"id": "kai"}, "nodes": nodes, "groups": groups, "edges": edges}
    document = {"deck": {"id": "kai"}, "slides": [{"title": "The clock", "body": [{"figure": figure}]}]}
    kind = DeckKind()
    first = kind.draw(document, tmp_path, {"settle": True}).pages[0].svg
    where = next(key for key in compose._LAYOUTS if key[0] == "kai")
    # Its part put on a line of its own by hand: drawn so, not turned round to fit.
    assert not compose._LAYOUTS[where].startswith("turned")
    # A line given words, drawn, then taken back: drawn as it was, whatever came between.
    labelled = copy.deepcopy(document)
    labelled["slides"][0]["body"][0]["figure"]["edges"][3]["label"] = "KaiB slows it"
    kind.draw(labelled, tmp_path)
    kind.draw(labelled, tmp_path, {"settle": True})
    assert kind.draw(document, tmp_path, {"settle": True}).pages[0].svg == first


def test_a_figure_arranged_by_hand_is_turned_to_fit_when_its_person_asks(tmp_path: Path, monkeypatch) -> None:
    from flexo_talk import compose

    monkeypatch.setenv("FLEXO_TALK_CACHE", "0")
    parts = {"customer": ("person", "Customer"), "form": ("io", "Order form"), "api": ("server", "Checkout API"),
             "queue": ("queue", "Order queue"), "worker": ("server", "Email worker"),
             "receipt": ("document", "Receipt email"), "feed": ("cloud", "Warehouse feed"),
             "db": ("database", "Orders DB")}
    nodes = [{"id": id, "kind": kind, "label": label} for id, (kind, label) in parts.items()]
    # As the studio writes them: the parts' places, which may be turned, not groups of meaning.
    groups = [
        {"id": "root", "layout": {"kind": "column"}, "role": "canvas", "children": ["row", "feed", "db"]},
        {"id": "row", "layout": {"kind": "row"}, "role": "layout",
         "children": ["customer", "form", "api", "queue", "worker", "receipt"]},
    ]
    lines = [("customer", "form", "fills in"), ("form", "api", "POST /orders"), ("api", "queue", "publishes"),
             ("queue", "worker", "consumes"), ("worker", "receipt", "sends"), ("queue", "feed", "streams"),
             ("api", "db", "writes")]
    edges = [{"from": a, "to": b, "label": label} for a, b, label in lines]

    def layout(**options: object) -> str:
        figure = {"figure": {"id": "orders"}, "nodes": nodes, "groups": groups, "edges": edges}
        document = {"deck": {"id": "orders"}, "slides": [
            {"layout": "figure", "title": "The new architecture", "body": [{"figure": figure, **options}]}]}
        DeckKind().draw(document, tmp_path, {"settle": True})
        return compose._LAYOUTS[next(key for key in compose._LAYOUTS if key[0] == "orders")]

    # Arranged by hand (a row with parts under it), it is drawn as arranged...
    assert not layout().startswith("turned")
    # ...unless its person turns "Turn to Fit the Slide" on for it.
    assert layout(turn=True).startswith("turned")
    assert not layout(turn=False).startswith("turned")


def test_a_figure_built_on_a_slide_keeps_the_way_it_is_drawn_until_its_person_asks(
    tmp_path: Path, monkeypatch
) -> None:
    import copy

    from flexo_talk import compose

    # (With the layouts kept between drawings, as the studio keeps them.)
    monkeypatch.setenv("FLEXO_TALK_CACHE", str(tmp_path / "cache"))
    nodes = [{"id": "one", "label": "One"}]
    figure = {"figure": {"id": "chart"}, "nodes": nodes}
    block = {"figure": figure}
    document = {"deck": {"id": "growing"}, "slides": [{"layout": "blank", "body": [block]}]}
    kind = DeckKind()
    kind.draw(document, tmp_path, {"settle": True})
    where = next(key for key in compose._LAYOUTS if key[0] == "growing")
    # Shapes added one at a time under the last, each drawn as it is made and then settled:
    # the column it is seen as stays a column, made smaller, never turned into a row.
    for word in ["Two", "Three", "Four", "Five", "Six", "Seven"]:
        nodes.append({"id": word.lower(), "label": word})
        kind.draw(document, tmp_path)
        assert not compose._LAYOUTS[where].startswith("turned"), word
        kind.draw(document, tmp_path, {"settle": True})
        assert not compose._LAYOUTS[where].startswith("turned"), word
    # Opened again another day (nothing remembered but the layouts kept): drawn as it was.
    for memory in (compose._LAYOUTS, compose._ASKED, compose._TURNS, compose._SCALES, compose._SHOWN, compose._PLACES):
        memory.pop(where, None)
    DeckKind().draw(copy.deepcopy(document), tmp_path, {"settle": True})
    assert not compose._LAYOUTS[where].startswith("turned")
    # Turned only when its person asks (Turn to Fit the Slide), and back when they ask again.
    block["turn"] = True
    kind.draw(document, tmp_path, {"settle": True})
    assert compose._LAYOUTS[where].startswith("turned")
    nodes.append({"id": "eight", "label": "Eight"})
    kind.draw(document, tmp_path, {"settle": True})
    assert compose._LAYOUTS[where].startswith("turned")
    block["turn"] = False
    kind.draw(document, tmp_path, {"settle": True})
    assert not compose._LAYOUTS[where].startswith("turned")


@pytest.mark.skipif(shutil.which("node") is None, reason="needs node")
def test_a_part_dragged_on_a_slide_swaps_goes_between_or_goes_home() -> None:
    script = Path(__file__).parents[1] / "src/flexo_talk/studio/static/slidedrop.js"

    def box(left, top, right, bottom):
        return {"left": left, "top": top, "right": right, "bottom": bottom}

    # Three parts down a body; two columns, the left holding two parts, the right one or none.
    body = [
        {
            "key": "body",
            "room": box(0, 0, 400, 300),
            "blocks": [
                {"index": 0, "box": box(0, 0, 300, 40)},
                {"index": 1, "box": box(0, 60, 300, 140)},
                {"index": 2, "box": box(0, 160, 200, 172)},  # a line of words, 12 tall
            ],
        }
    ]
    left = {
        "key": "left",
        "room": box(0, 0, 200, 300),
        "blocks": [{"index": 0, "box": box(0, 0, 180, 40)}, {"index": 1, "box": box(0, 60, 180, 100)}],
    }
    right_one = {"key": "right", "room": box(220, 0, 420, 300), "blocks": [{"index": 0, "box": box(220, 0, 400, 20)}]}
    right_none = {"key": "right", "room": box(220, 0, 420, 300), "blocks": []}
    drops = [
        [body, {"region": "body", "index": 0}, {"x": 100, "y": 166}],  # the middle of the line
        [body, {"region": "body", "index": 2}, {"x": 100, "y": 3}],  # the top edge of the first
        [body, {"region": "body", "index": 0}, {"x": 100, "y": 62}],  # just after itself: home
        [body, {"region": "body", "index": 0}, {"x": 100, "y": 250}],  # below them all
        [body, {"region": "body", "index": 1}, {"x": 900, "y": 900}],  # off the slide
        [[left, right_none], {"region": "left", "index": 1}, {"x": 300, "y": 200}],  # empty column
        [[left, right_one], {"region": "left", "index": 0}, {"x": 300, "y": 250}],  # under its one part
    ]
    code = (
        f"import {{ blockDrop, blockPlan }} from {json.dumps(script.as_uri())};\n"
        f"const drops = {json.dumps(drops)};\n"
        "const out = drops.map(([regions, from, point]) => blockDrop(regions, from, point));\n"
        "const at = (kind, region, index) => ({ kind, region, index });\n"
        "out.push(blockPlan({ body: 3 }, { region: 'body', index: 0 }, at('swap', 'body', 2)));\n"
        "out.push(blockPlan({ left: 2, right: 1 }, { region: 'left', index: 0 }, at('between', 'right', 1)));\n"
        "console.log(JSON.stringify(out));\n"
    )
    done = subprocess.run(["node", "--input-type=module", "-e", code], capture_output=True, text=True, check=True)
    found = json.loads(done.stdout)
    assert found[:7] == [
        {"kind": "swap", "region": "body", "index": 2},
        {"kind": "between", "region": "body", "index": 0},
        {"kind": "home"},
        {"kind": "between", "region": "body", "index": 3},
        None,
        {"kind": "into", "region": "right", "index": 0},
        {"kind": "between", "region": "right", "index": 1},
    ]
    place = lambda region, index: {"region": region, "index": index}  # noqa: E731
    # Swapped: the first and the last trade places, the middle stays.
    assert found[7] == [
        [place("body", 2), place("body", 0)],
        [place("body", 1), place("body", 1)],
        [place("body", 0), place("body", 2)],
    ]
    # Moved across: the left's second part rises to first, the moved one goes after the right's.
    assert found[8] == [
        [place("left", 1), place("left", 0)],
        [place("right", 0), place("right", 0)],
        [place("left", 0), place("right", 1)],
    ]


@pytest.mark.skipif(shutil.which("node") is None, reason="needs node")
def test_several_parts_dragged_together_land_together_in_their_order() -> None:
    script = Path(__file__).parents[1] / "src/flexo_talk/studio/static/slidedrop.js"

    def box(left, top, right, bottom):
        return {"left": left, "top": top, "right": right, "bottom": bottom}

    # Four parts down a body, the first and the third dragged: the second and the fourth,
    # left behind, closed up (their boxes where they show while the two are dragged).
    body = [
        {
            "key": "body",
            "room": box(0, 0, 400, 400),
            "blocks": [
                {"index": 0, "box": box(0, 0, 300, 60)},
                {"index": 1, "box": box(0, 0, 300, 60)},
                {"index": 2, "box": box(0, 200, 300, 260)},
                {"index": 3, "box": box(0, 100, 300, 160)},
            ],
        }
    ]
    # A column holding only a part dragged out of it.
    right = [{"key": "right", "room": box(500, 0, 700, 400), "blocks": [{"index": 0, "box": box(500, 0, 700, 60)}]}]
    chosen = [{"region": "body", "index": 0}, {"region": "body", "index": 2}]
    drops = [
        [body, chosen, {"x": 100, "y": 35}],  # the second's lower half: after it, never swapped
        [body, chosen, {"x": 100, "y": 105}],  # the fourth's upper half: before it
        [body, chosen, {"x": 100, "y": 395}],  # below them all
        [right, [{"region": "right", "index": 0}], {"x": 600, "y": 30}],  # its column, emptied
        [body, chosen, {"x": 100, "y": 450}],  # under the region's room, on the slide: at its end
    ]
    code = (
        f"import {{ groupDrop, gatherPlan, blockDrop }} from {json.dumps(script.as_uri())};\n"
        f"const out = {json.dumps(drops)}.map(([regions, all, point]) => groupDrop(regions, all, point));\n"
        # One part, under the room: at the end too.
        f"out.push(blockDrop({json.dumps(body)}, {{ region: 'body', index: 0 }}, {{ x: 100, y: 450 }}));\n"
        f"const chosen = {json.dumps(chosen)};\n"
        "out.push(gatherPlan({ body: 4 }, chosen, { kind: 'between', region: 'body', index: 4 }));\n"
        "const across = [{ region: 'left', index: 0 }, { region: 'right', index: 0 }];\n"
        "out.push(gatherPlan({ left: 2, right: 1 }, across, { kind: 'between', region: 'left', index: 2 }));\n"
        "console.log(JSON.stringify(out));\n"
    )
    done = subprocess.run(["node", "--input-type=module", "-e", code], capture_output=True, text=True, check=True)
    found = json.loads(done.stdout)
    assert found[:6] == [
        {"kind": "between", "region": "body", "index": 2},
        {"kind": "between", "region": "body", "index": 3},
        {"kind": "between", "region": "body", "index": 4},
        {"kind": "into", "region": "right", "index": 0},
        {"kind": "between", "region": "body", "index": 4},
        {"kind": "between", "region": "body", "index": 4},
    ]
    place = lambda region, index: {"region": region, "index": index}  # noqa: E731
    # Let go below them all: the two not chosen rise, the chosen two follow in their order.
    assert found[6] == [
        [place("body", 1), place("body", 0)],
        [place("body", 3), place("body", 1)],
        [place("body", 0), place("body", 2)],
        [place("body", 2), place("body", 3)],
    ]
    # From two columns to the end of the left: the right's part follows the left's.
    assert found[7] == [
        [place("left", 1), place("left", 0)],
        [place("left", 0), place("left", 1)],
        [place("right", 0), place("left", 2)],
    ]


def test_each_region_of_a_slide_says_where_its_room_is_empty_or_not(tmp_path: Path) -> None:
    deck = deck_from_document(
        {
            "deck": {"id": "rooms"},
            "slides": [{"layout": "two-columns", "title": "Two", "left": [{"text": "words"}], "right": []}],
        },
        tmp_path,
    )
    svg = render_slide(deck, deck.slides[0]).svg
    rooms = dict(re.findall(r'<g id="slide1\.(left|right)"[^>]*data-flexo-box="([^"]+)"', svg))
    assert set(rooms) == {"left", "right"}  # the empty column too: a part can be dropped in it
    (lx, ly, lw, lh), (rx, ry, rw, rh) = (map(float, rooms[key].split()) for key in ("left", "right"))
    assert lw > 0 and lh > 0 and (rw, rh, ry) == (lw, lh, ly) and rx > lx + lw


def test_a_structure_on_a_slide_says_the_mol_sketch_settings_it_is_drawn_with(tmp_path: Path) -> None:
    pytest.importorskip("molsketch")
    import shutil

    data = Path(__file__).resolve().parents[2] / "flexo" / "tests" / "unit" / "data" / "1a7g.cif"
    if not data.is_file():
        pytest.skip("no structure file beside flexo")
    shutil.copy(data, tmp_path / "1a7g.cif")
    figure = {
        "figure": {"id": "mol"},
        "nodes": [{"id": "m", "kind": "structure",
                   "properties": {"source": "1a7g.cif", "style": {"line": {"width": 2.5}}}}],
    }
    document = {"schema_version": 1, "deck": {"id": "talk", "theme": "sketch"},
                "slides": [{"title": "A molecule", "body": [{"figure": figure}]}]}
    at = {"slide": 0, "region": "body", "index": 0}
    asked = {"do": "figure", "at": at, "edit": {"do": "structure-settings", "id": "m"}}
    result = DeckKind().act(document, asked, tmp_path)
    assert result["document"] == document
    # In the deck's look: a sketched theme draws in watercolour, the molecule's own pen over it.
    assert result["settings"]["look"] == "watercolour"
    assert result["settings"]["style"]["line.width"] == 2.5


def test_an_empty_title_or_text_holds_its_place_and_shows_only_in_the_studio() -> None:
    from flexo_talk.compose import PLACEHOLDERS

    document = yaml.safe_load(
        'deck: {id: holds}\nslides:\n  - title: ""\n    body: [{bullets: [""]}, {text: ""}, {text: Kept}]\n'
        '  - layout: title\n    title: ""\n    subtitle: ""\n'
    )
    deck = deck_from_document(document, Path("."))
    exported = [render_slide(deck, slide).svg for slide in deck.slides]
    token = PLACEHOLDERS.set(True)
    try:
        studio = [render_slide(deck, slide).svg for slide in deck.slides]
    finally:
        PLACEHOLDERS.reset(token)
    placed = re.findall(r'id="([^"]+)"[^>]*data-flexo-placeholder="([^"]+)"', studio[0])
    assert placed == [("slide1.title", "Title"), ("slide1.body.0", "Text"), ("slide1.body.1", "Text")]
    assert re.findall(r'data-flexo-placeholder="([^"]+)"', studio[1]) == ["Title", "Subtitle"]
    assert "data-flexo-placeholder" not in exported[0] + exported[1]
    assert "Title" not in exported[1] and ">Text<" not in exported[0]
    # Words left empty keep their place while edited (they are typed in there), and take
    # none when presented or exported: what follows them is where it would be without them.
    kept = r'<text[^>]*y="([\d.]+)"[^>]*>(?:(?!</text>).)*Kept'
    alone = deck_from_document(yaml.safe_load(
        'deck: {id: holds}\nslides:\n  - title: ""\n    body: [{text: Kept}]\n'), Path("."))
    shown = re.findall(kept, exported[0], re.S)
    assert shown == re.findall(kept, render_slide(alone, alone.slides[0]).svg, re.S)
    assert float(shown[0]) < float(re.findall(kept, studio[0], re.S)[0])
    assert not [layout for layout in render_slide(deck, deck.slides[0]).lists if layout.id == "slide1.body.0"]


def test_an_agenda_before_any_section_shows_its_rows_faintly_while_editing() -> None:
    from flexo_talk.compose import PLACEHOLDERS

    document = yaml.safe_load("deck: {id: early}\nslides:\n  - {layout: agenda, title: Outline}\n")
    deck = deck_from_document(document, Path("."))
    exported = render_slide(deck, deck.slides[0])
    token = PLACEHOLDERS.set(True)
    try:
        editing = render_slide(deck, deck.slides[0])
    finally:
        PLACEHOLDERS.reset(token)
    # Put in before its sections, as often: faint rows and a note, not a warning at once.
    assert not editing.diagnostics and editing.notes
    assert re.findall(r'id="slide1\.agenda\d"[^>]*data-flexo-placeholder="Section"', editing.svg)
    # Exported so, it is empty: that is said.
    assert exported.diagnostics and "slide1.agenda" not in exported.svg


def test_a_slide_with_no_title_has_no_band_when_presented(tmp_path: Path) -> None:
    from pptx import Presentation

    from flexo_talk.compose import PLACEHOLDERS
    from flexo_talk.export import write_pptx

    document = yaml.safe_load(
        "deck: {id: banded, look: band}\nslides:\n"
        "  - {title: '', body: [{text: Words}]}\n  - {title: Titled, body: [{text: Words}]}\n"
    )
    deck = deck_from_document(document, tmp_path)

    def marks(svg: str) -> list[str]:
        return re.findall(r'id="slide\d\.(band|title)"', svg)

    token = PLACEHOLDERS.set(True)
    try:
        drawn = [render_slide(deck, slide).svg for slide in deck.slides]
    finally:
        PLACEHOLDERS.reset(token)
    # While editing, the empty title's place is shown on its band; presented, neither is:
    # the band is marked as a placeholder is, so Present (which shows this drawing) hides it.
    assert [marks(svg) for svg in drawn] == [["band", "title"], ["band", "title"]]
    hidden = [re.findall(r'id="slide\d\.band"[^>]*data-flexo-placeholder', svg) != [] for svg in drawn]
    assert hidden == [True, False]
    assert [marks(render_slide(deck, slide).svg) for slide in deck.slides] == [[], ["band", "title"]]
    slides = Presentation(write_pptx(deck, deck.render(), tmp_path / "talk.pptx")).slides
    assert ["Band" in [shape.name for shape in slide.shapes] for slide in slides] == [False, True]


def test_a_powerpoint_file_carries_the_talks_title_and_author_not_the_templates(tmp_path: Path) -> None:
    from pptx import Presentation

    from flexo_talk.export import write_pptx

    document = yaml.safe_load(
        'deck: {id: cache}\nslides:\n  - {layout: title, title: "Making it **fast**", author: Ada Lovelace}\n'
    )
    deck = deck_from_document(document, tmp_path)
    properties = Presentation(write_pptx(deck, deck.render(), tmp_path / "talk.pptx")).core_properties
    assert (properties.title, properties.author, properties.last_modified_by) == (
        "Making it fast", "Ada Lovelace", "Ada Lovelace"
    )
    assert "python-pptx" not in (properties.comments or "") and properties.created.year >= 2026
    # The PDF is named for the talk too, not for its file.
    from flexo_talk.export import build_deck

    info = build_deck(deck, tmp_path / "out", ("pdf",)).pdf.read_bytes()
    assert b"/Title (Making it fast)" in info and b"/Author (Ada Lovelace)" in info


def test_a_pictures_description_is_read_out_and_is_powerpoints_alt_text(tmp_path: Path) -> None:
    from PIL import Image
    from pptx import Presentation

    from flexo_talk.export import write_pptx

    Image.new("RGB", (64, 36), "#2b4c9b").save(tmp_path / "chart.png")
    (tmp_path / "plot.svg").write_text(
        '<svg xmlns="http://www.w3.org/2000/svg" width="80" height="40" viewBox="0 0 80 40">'
        '<rect x="0" y="0" width="80" height="40" fill="#c0392b"/></svg>'
    )
    document = yaml.safe_load(
        "deck: {id: alt}\nslides:\n"
        "  - title: Results\n    body:\n"
        "      - {image: chart.png, description: \"  Latency   fell by half  \"}\n"
        "      - {image: plot.svg, description: A red bar}\n"
    )
    deck = deck_from_document(document, tmp_path)
    svg = render_slide(deck, deck.slides[0]).svg
    assert 'aria-label="Latency fell by half"' in svg and 'aria-label="A red bar"' in svg
    slide = Presentation(write_pptx(deck, deck.render(), tmp_path / "talk.pptx")).slides[0]
    described = [(element.get("name"), element.get("descr"))
                 for element in slide._element.iter() if element.get("descr")]
    assert described == [("Picture", "Latency fell by half"), ("Picture", "A red bar")]


def test_the_studio_draws_a_deck_and_says_the_order_its_tones_take(tmp_path: Path) -> None:
    from flexo.themes import palette_order

    document = {"deck": {"id": "o", "palette": "Deep Sea Harvest"}, "slides": [{"title": "One"}, {"title": "Two"}]}
    drawing = DeckKind().draw(document, tmp_path, {"focus": 1})
    assert len(drawing.pages) == 2 and all(page.svg for page in drawing.pages)
    assert drawing.info["order"] == palette_order("paper", "Deep Sea Harvest")


def test_a_new_quote_callout_code_or_numbers_is_a_placeholder_until_typed_in() -> None:
    from flexo_talk.compose import PLACEHOLDERS

    document = yaml.safe_load(
        'deck: {id: holds}\nslides:\n  - title: Empty objects\n    body:\n'
        '      - {quote: ""}\n      - {callout: ""}\n      - {code: ""}\n'
        '      - {stats: [{value: "", label: ""}, {value: "", label: ""}]}\n'
        '      - {callout: "", title: Kept}\n'
    )
    deck = deck_from_document(document, Path("."))
    exported = render_slide(deck, deck.slides[0]).svg
    token = PLACEHOLDERS.set(True)
    try:
        studio = render_slide(deck, deck.slides[0]).svg
    finally:
        PLACEHOLDERS.reset(token)
    placed = re.findall(r'id="([^"]+)"[^>]*data-flexo-placeholder="([^"]+)"', studio)
    assert placed == [("slide1.body.0", "Quote"), ("slide1.body.1", "Text"), ("slide1.body.2", "Code"),
                      ("slide1.body.3", "Numbers")]
    assert "data-flexo-placeholder" not in exported and ">Quote<" not in exported and ">Kept<" in exported


def test_an_attribution_or_caption_put_there_empty_holds_its_place_while_it_is_typed() -> None:
    from flexo_talk.compose import PLACEHOLDERS

    def drawn(by: str | None, caption: str | None) -> tuple[str, str]:
        quote = {"quote": "Waiting is the cost.", **({} if by is None else {"by": by})}
        table = {"table": [["Year", "Orders"], ["2024", "12"]], **({} if caption is None else {"caption": caption})}
        slide = {"title": "Own lines", "body": [quote, table, {"text": "Under"}]}
        deck = deck_from_document({"deck": {"id": "own"}, "slides": [slide]}, Path("."))
        exported = render_slide(deck, deck.slides[0]).svg
        token = PLACEHOLDERS.set(True)
        try:
            return render_slide(deck, deck.slides[0]).svg, exported
        finally:
            PLACEHOLDERS.reset(token)

    held, held_exported = drawn("", "")
    written, _ = drawn("Sam Rivera", "Table 1. Orders by year.")
    none, none_exported = drawn(None, None)
    # Emptied (or put there) while it is typed, each shows its placeholder where its words go...
    placed = re.findall(r'id="([^"]+)"[^>]*data-flexo-placeholder="([^"]+)"', held)
    assert placed == [("slide1.body.0.by", "Who said it"), ("slide1.body.1.caption", "Caption")]
    assert "\u2014 Who said it<" in held
    # ... taking the room its words will: what is under it stays where it is.
    under = r'<text id="slide1\.body\.2"[^>]*\by="([^"]+)"'
    assert re.findall(under, held) == re.findall(under, written) != re.findall(under, none)
    # Presented or exported, it is not there at all.
    assert held_exported == none_exported and "Who said it" not in held_exported


def test_a_sample_table_figure_or_equation_is_a_placeholder_until_changed(tmp_path: Path) -> None:
    from flexo_talk.compose import PLACEHOLDERS

    document = yaml.safe_load(
        'deck: {id: samples}\nslides:\n  - title: Samples\n    body:\n'
        '      - {placeholder: true, table: [[Item, Value], [First, "1.0"]]}\n'
        '      - {placeholder: true, math: "a^2 + b^2 = c^2"}\n'
        '      - {placeholder: true, figure: {figure: {id: f}, nodes: [{id: input, label: Input}]}}\n'
        '      - {table: [[Kept, Too]]}\n'
    )
    deck = deck_from_document(document, Path("."))
    exported = render_slide(deck, deck.slides[0])
    token = PLACEHOLDERS.set(True)
    try:
        studio = render_slide(deck, deck.slides[0]).svg
    finally:
        PLACEHOLDERS.reset(token)
    placed = re.findall(r'id="([^"]+)"[^>]*data-flexo-placeholder=', studio)
    assert placed == ["slide1.body.0", "slide1.body.1", "slide1.body.2"]
    # Drawn as itself while editing; neither drawn nor a native table when exported.
    assert ">Item<" in studio and ">Item<" not in exported.svg and ">Kept<" in exported.svg
    assert [table.id for table in exported.tables] == ["slide1.body.3"]
    # Written back as it was read; a figure changed is the person's own.
    assert deck_document(deck)["slides"][0]["body"][0]["placeholder"] is True
    changed = DeckKind().act(document, {"do": "figure", "at": {"slide": 0, "region": "body", "index": 2},
                                        "edit": {"do": "update", "target": {"type": "node", "id": "input"},
                                                 "values": {"label": "Load"}}}, tmp_path)
    assert "placeholder" not in changed["document"]["slides"][0]["body"][2]
    assert changed["document"]["slides"][0]["body"][0]["placeholder"] is True


def test_a_placeholder_takes_no_room_from_the_words_and_a_fit_note_names_what_is_there() -> None:
    words = {"bullets": ["a sentence that goes on and on across the slide"] * 8}
    sample = {"figure": {"figure": {"id": "f"}, "nodes": [{"id": "a", "label": "Input"}]}}
    document = {"deck": {"id": "fit"}, "slides": [
        {"title": "Samples", "body": [words, {"placeholder": True, **sample},
                                      {"placeholder": True, "math": "a^2 + b^2 = c^2"}]},
        {"title": "A figure", "body": [words, sample]},
    ]}
    deck = deck_from_document(document, Path("."))
    samples, figure = (render_slide(deck, slide).diagnostics for slide in deck.slides)
    # Samples are never presented, so the words keep their size and nothing is said.
    assert samples == []
    # A figure is there, and is named: not "the picture".
    assert figure == [
        "slide2: Text size reduced to 85% to fit the slide. For full-size text, give the figure "
        "a column of its own (Two Columns) or a slide of its own."
    ]


def test_an_empty_placeholder_takes_no_room_from_what_is_there_when_presented() -> None:
    from flexo_talk.compose import PLACEHOLDERS

    nodes = [{"id": "a", "label": "Repressor"}, {"id": "b", "label": "Operator"}]
    real = {"figure": {"figure": {"id": "r"}, "nodes": nodes, "edges": [{"from": "a", "to": "b"}]}}
    empty = {"figure": {"figure": {"id": "e"}, "nodes": [{"id": "shape", "label": ""}]}}
    document = {"deck": {"id": "room"}, "slides": [
        {"layout": "figure", "title": "With", "body": [empty, real]},
        {"layout": "figure", "title": "Alone", "body": [real]},
    ]}
    deck = deck_from_document(document, Path("."))

    def placed(slide, editing: bool) -> list[tuple[str, str]]:
        token = PLACEHOLDERS.set(editing)
        try:
            svg = render_slide(deck, slide).svg
        finally:
            PLACEHOLDERS.reset(token)
        return re.findall(r'<g id="slide\d\.body\.(\d)" transform="([^"]+)"', svg)

    alone = placed(deck.slides[1], False)
    # Exported, the figure is as large as it is alone, with no empty band where the sample was.
    assert placed(deck.slides[0], False) == [("1", alone[0][1])]
    # While editing too (Present shows that drawing, the sample hidden): the sample is drawn
    # after it, taking none of its room.
    assert placed(deck.slides[0], True)[0] == ("1", alone[0][1])
    assert [index for index, _ in placed(deck.slides[0], True)] == ["1", "0"]


def test_a_slide_holding_empty_words_is_presented_as_it_is_exported(tmp_path: Path) -> None:
    nodes = [{"id": "a", "label": "Repressor"}, {"id": "b", "label": "Operator"}]
    figure = {"figure": {"figure": {"id": "r"}, "nodes": nodes, "edges": [{"from": "a", "to": "b"}]}}
    document = {"deck": {"id": "shown"}, "slides": [
        {"title": "Held", "body": [{"text": ""}, figure]},
        {"title": "Plain", "body": [figure]},
    ]}
    held, plain = (page.svg for page in DeckKind().draw(document, tmp_path, {"settle": True}).pages)
    deck = deck_from_document(document, tmp_path)
    exported = render_slide(deck, deck.slides[0])

    def placed(svg: str, prefix: str = "") -> list[str]:
        return re.findall(rf'<g id="{prefix}slide\d\.body\.1" transform="([^"]+)"', svg)

    # Edited, the empty text holds its place; presented, the slide as exported is shown
    # (present.js shows the drawing marked presented), the figure where the PDF has it.
    exported_at = re.findall(r'<g id="slide1\.body\.1" transform="([^"]+)"', exported.svg)
    assert placed(held) != exported_at and placed(held, "presented-") == exported_at
    assert 'data-flexo-editing=""' in held and 'display="none" data-flexo-presented=""' in held
    # A slide that presents as it is edited is drawn once.
    assert "data-flexo-presented" not in plain


def test_one_rule_places_a_body_s_things_and_an_emptied_object_changes_nothing(tmp_path: Path) -> None:
    from PIL import Image

    Image.new("RGB", (64, 36), "#2b4c9b").save(tmp_path / "gel.png")
    table = {"table": [["Variant", "Families"], ["R122H", "62"], ["N29I", "19"]]}
    document = {"deck": {"id": "placed"}, "slides": [
        {"title": "Lone", "body": [table]},
        {"title": "Captioned", "body": [table, {"text": "Families in the registry, 2024"}]},
        {"title": "Listed", "body": [{"image": "gel.png"}, {"bullets": ["Lag", "Burst"]}, {"text": "n = 3 gels"}]},
        {"title": "Quote", "body": [{"quote": "Words", "by": "Ada"}, {"callout": ""}]},
        {"title": "Quote", "body": [{"quote": "Words", "by": "Ada"}]},
    ]}
    deck = deck_from_document(document, tmp_path)
    lone, captioned, listed, emptied, alone = (render_slide(deck, slide) for slide in deck.slides)
    # A table keeps its place across the slide when a caption is added: the two are one
    # group, centred, the caption centred under it.
    assert captioned.tables[0].x == pytest.approx(lone.tables[0].x)
    assert re.search(r'<text id="slide2\.body\.1" x="480"[^>]*text-anchor="middle"', captioned.svg)
    # Beside a list, a picture starts where the list's words do.
    assert re.search(r'<image id="slide3\.body\.0" x="48"', listed.svg)
    # A slide holding an emptied callout is exported as the slide without it: the quote
    # stands where it stands alone.
    shift = r'<g id="slide\d\.body" [^>]*transform="([^"]+)"'
    assert re.findall(shift, emptied.svg) == re.findall(shift, alone.svg) != []


def test_a_new_table_equation_or_figure_starts_empty_and_shows_only_while_editing() -> None:
    from flexo_talk.compose import PLACEHOLDERS

    document = yaml.safe_load(
        'deck: {id: empty}\nslides:\n  - title: New objects\n    body:\n'
        '      - {table: [["", ""], ["", ""]]}\n'
        '      - {math: ""}\n'
        '      - {figure: {figure: {id: f}, nodes: [{id: start, kind: terminal, label: ""}]}}\n'
        '      - {table: [[Year, ""]]}\n'
    )
    deck = deck_from_document(document, Path("."))
    exported = render_slide(deck, deck.slides[0])
    token = PLACEHOLDERS.set(True)
    try:
        studio = render_slide(deck, deck.slides[0]).svg
    finally:
        PLACEHOLDERS.reset(token)
    placed = re.findall(r'id="([^"]+)"[^>]*data-flexo-placeholder="([^"]+)"', studio)
    # (And in a table typed in, its empty header cell hints its column, as it is typed in.)
    # A sample figure is drawn after what is really there, taking none of its room.
    assert sorted(placed) == [("slide1.body.0", "Table"), ("slide1.body.1", "Equation"),
                              ("slide1.body.2", "Placeholder"), ("slide1.body.3.0.1.hint", "Column")]
    # Hints while editing ("Column 1", the first step's "Start"); none of it presented.
    assert ">Column 1<" in studio and ">Start<" in studio
    assert ">Column 1<" not in exported.svg and ">Start<" not in exported.svg and ">Column 2<" not in exported.svg
    # A table with one cell typed in is the person's own, its empty cells empty.
    assert ">Year<" in exported.svg and [table.id for table in exported.tables] == ["slide1.body.3"]


def test_a_placeholder_added_and_taken_away_is_not_activity() -> None:
    kind = DeckKind()
    plain = {"slides": [{"title": "A", "body": [{"bullets": ["x"]}]}]}
    empty = {"slides": [{"title": "A", "body": [{"bullets": ["x"]}, {"code": ""}]}]}
    typed = {"slides": [{"title": "A", "body": [{"bullets": ["x"]}, {"code": "print()"}]}]}
    assert kind.describe(plain, empty) == [] and kind.describe(empty, plain) == []
    # Typed in, the code that was there empty is filled in -- not added: it was there.
    assert [note["text"] for note in kind.describe(empty, typed)] == ["filled in the code on slide 1"]
    # An object made another kind, a layout changed: said so, not as objects added or deleted.
    texted = {"slides": [{"title": "A", "body": [{"text": "-"}]}]}
    listed = {"slides": [{"title": "A", "body": [{"bullets": ["x"]}]}]}
    assert [note["text"] for note in kind.describe(texted, listed)] == ["made the text a list on slide 1"]
    columns = {"slides": [{"layout": "two-columns", "title": "A", "left": [{"bullets": ["x"]}]}]}
    assert [note["text"] for note in kind.describe(listed, columns)] == ["changed the layout on slide 1"]


HAND = """schema_version: 1
deck:
  id: talk
  theme: paper   # chosen by Kiarash
slides:
# --- opening ---
- layout: title
  title: "Making a web service fast"
# --- the middle ---
- title: The pipeline
  body:
  - bullets:
    - Generate   # step one
    - Design
  - figure:
      edges:
      - {from: start, to: step}
- title: Voices
  body:
  - quote: Simple things should be simple.
# end of deck
"""


def test_a_deletion_or_a_copy_undone_through_the_studio_puts_the_file_back_word_for_word(tmp_path: Path) -> None:
    import copy

    from flexo.studio.workspace import Workspace

    talk = tmp_path / "talk.yaml"
    talk.write_text(HAND, encoding="utf-8")
    me = {"id": "p", "kind": "person"}
    workspace = Workspace(tmp_path)
    try:
        doc = workspace.open("talk.yaml")
        whole = copy.deepcopy(doc.document)
        # Each written as the studio writes it (beside the file, then over it), then undone.
        for change in ("delete", "duplicate first", "duplicate middle"):
            changed = copy.deepcopy(whole)
            if change == "delete":
                del changed["slides"][1]
            else:
                at = 0 if change == "duplicate first" else 1
                changed["slides"].insert(at + 1, copy.deepcopy(changed["slides"][at]))
            doc.update(changed, doc.version, me)
            workspace.flush()
            assert talk.read_text(encoding="utf-8") != HAND
            doc.update(copy.deepcopy(whole), doc.version, me)
            workspace.flush()
            assert talk.read_text(encoding="utf-8") == HAND, change
    finally:
        workspace.close()


def test_a_deck_moved_into_a_new_folder_or_up_out_of_one_is_followed_never_copied(tmp_path: Path) -> None:
    import copy
    import time

    from flexo.studio.workspace import Workspace

    def waited(condition) -> None:
        deadline = time.monotonic() + 5
        while not condition():
            assert time.monotonic() < deadline, "waited in vain"
            time.sleep(0.05)

    (tmp_path / "talk.yaml").write_text(HAND, encoding="utf-8")
    me = {"id": "p", "kind": "person"}
    workspace = Workspace(tmp_path)
    try:
        doc = workspace.open("talk.yaml")
        # The Finder's New Folder with Selection (the folder and the move at once), then back up.
        moves = [("talk.yaml", "Lab meeting/talk.yaml"), ("Lab meeting/talk.yaml", "talk.yaml")]
        for at, (old, new) in enumerate(moves):
            time.sleep(0.3)
            (tmp_path / new).parent.mkdir(exist_ok=True)
            (tmp_path / old).rename(tmp_path / new)
            waited(lambda new=new: doc.name == new)
            edited = copy.deepcopy(doc.document)
            edited["slides"][2]["title"] = f"Voices {at}"
            doc.update(edited, doc.version, me)
            waited(lambda new=new, at=at: f"Voices {at}" in (tmp_path / new).read_text(encoding="utf-8"))
            assert [found.relative_to(tmp_path).as_posix() for found in tmp_path.rglob("*.yaml")] == [new]
            # Written over its words as they are: its comments kept where it went.
            assert "# chosen by Kiarash" in (tmp_path / new).read_text(encoding="utf-8")
    finally:
        workspace.close()


def test_activity_says_a_slide_moved_and_follows_an_object_where_it_goes() -> None:
    import copy

    kind = DeckKind()

    def said(before: list, after: list) -> list[str]:
        return [note["text"] for note in kind.describe({"slides": before}, {"slides": after})]

    slides = [{"title": f"Slide {n}", "body": [{"text": f"words of slide {n}"}]} for n in range(1, 6)]
    assert said(slides, [*slides[:2], slides[3], slides[2], slides[4]]) == ["moved slide 3 after slide 4"]
    assert said(slides, [slides[0], slides[4], *slides[1:4]]) == ["moved slide 5 before slide 2"]
    # A long title is cut at a word, and said to be.
    longer = [*slides, {"title": "Part 1 · The signalling pathway and its last long step"}]
    assert said(slides, longer) == ["added slide 6, “Part 1 · The signalling pathway and its…”"]
    # An object put in an empty placeholder's place is added.
    empty, filled = [{"title": "A", "body": [{"bullets": []}]}], [{"title": "A", "body": [{"image": "a.png"}]}]
    assert said(empty, filled) == ["added a picture to slide 1"]
    # A picture edited, then dragged to the slide before: its note goes to it there.
    one = {"slides": [{"title": "A", "body": [{"text": "a"}]},
                      {"title": "B", "body": [{"text": "b"}, {"image": "a.png"}]}]}
    two = copy.deepcopy(one)
    two["slides"][1]["body"][1]["image"] = "b.png"
    where = kind.describe(one, two)[0]["where"]
    assert where == {"page": 2, "label": "Slide 2", "object": {"region": "body", "index": 1}}
    three = copy.deepcopy(two)
    three["slides"][0]["body"].append(three["slides"][1]["body"].pop())
    assert [note["text"] for note in kind.describe(two, three)] == ["moved the picture from slide 2 to slide 1"]
    assert kind.follow(two, three)(where) == {
        "page": 1, "label": "Now on slide 1", "object": {"region": "body", "index": 1}, "first": 2}
    # Edited again where it is, it stays put; deleted, it is said to be gone.
    again = copy.deepcopy(two)
    again["slides"][1]["body"][1]["image"] = "c.png"
    assert kind.follow(two, again)(where) is where
    gone = copy.deepcopy(two)
    del gone["slides"][1]["body"][1]
    assert kind.follow(two, gone)(where)["gone"] == "That picture has since been deleted."


def test_a_powerpoint_figure_is_described_and_its_shapes_named_by_their_words(tmp_path: Path) -> None:
    from pptx import Presentation

    from flexo_talk.export import write_pptx

    document = yaml.safe_load(
        "deck: {id: named}\nslides:\n  - title: Flow\n    body:\n      - figure:\n"
        "          figure: {id: f}\n"
        "          nodes: [{id: a, label: Customer, kind: person}, {id: b, label: Orders API}]\n"
        "          edges: [{from: a, to: b}]\n"
    )
    deck = deck_from_document(document, tmp_path)
    slide = Presentation(write_pptx(deck, deck.render(), tmp_path / "talk.pptx")).slides[0]
    names = [element.get("name") for element in slide._element.iter() if element.tag.endswith("}cNvPr")]
    described = {element.get("name"): element.get("descr") for element in slide._element.iter() if element.get("descr")}
    # Read along its line, and every shape named as a person would: no ids.
    assert described == {"Figure": "A figure: Customer → Orders API"}
    assert {"Title", "Figure", "Customer", "Orders API", "Line from Customer to Orders API",
            "Label “Customer”", "Slide Number"} <= set(names)
    assert not any(name.startswith(("slide", "path", "rect")) for name in names)


def test_a_powerpoint_slide_s_shapes_are_named_by_what_they_are(tmp_path: Path) -> None:
    from collections import Counter

    from pptx import Presentation

    from flexo_talk.export import write_pptx

    document = yaml.safe_load(
        "deck: {id: parts}\nslides:\n  - title: Parts\n    body:\n"
        "      - {quote: To be or not to be, by: Hamlet}\n      - callout: Mind the gap\n"
        "      - math: 'f = 1 + \\frac{a}{4c} - \\sqrt{b}'\n"
    )
    deck = deck_from_document(document, tmp_path)
    slide = Presentation(write_pptx(deck, deck.render(), tmp_path / "talk.pptx")).slides[0]
    names = Counter(element.get("name") for element in slide._element.iter()
                    if element.tag.endswith("}cNvPr") and element.get("name"))
    # A quote by its words (not its large quotation mark), its parts and a drawn formula's
    # glyphs and rules by what they are: no ids, no "path" or "rect".
    assert {"Title", "Quote “To be or not to be”", "Quotation Mark", "Attribution", "Callout “Mind the gap”",
            "Panel", "Bar", "Equation", "Equation Glyph", "Equation Line", "Slide Number"} <= set(names)
    assert not [name for name in names if "." in name or name in {"path", "rect", "Shape"}]


def test_a_drawn_protein_is_named_by_its_name_and_its_parts_by_what_they_are(tmp_path: Path) -> None:
    from pptx import Presentation

    from flexo_talk.export import write_pptx

    document = {"deck": {"id": "protein"}, "slides": [{"title": "Barrel", "body": [{"figure": {
        "figure": {"id": "f"},
        "nodes": [
            {"id": "gas", "label": "Oxygen"},
            {"id": "sfgfp", "kind": "protein", "label": "sfGFP", "properties": {"length": 238, "features": [
                {"type": "domain", "label": "β-barrel", "start": 1, "end": 230},
                {"type": "mutation", "label": "S65T", "at": 65},
            ]}},
        ],
        "edges": [{"from": "gas", "to": "sfgfp"}],
    }}]}]}
    deck = deck_from_document(document, tmp_path)
    slide = Presentation(write_pptx(deck, deck.render(), tmp_path / "talk.pptx")).slides[0]
    names = [element.get("name") for element in slide._element.iter()
             if element.tag.endswith("}cNvPr") and element.get("name")]
    # By its own name, not every word on it (its axis's numbers); its parts plainly.
    assert {"Protein sfGFP", "Feature β-barrel", "Site S65T", "Axis", "Tick 50", "Label “S65T”",
            "Line from Oxygen to Protein sfGFP"} <= set(names)
    assert "Shape" not in names and not any("50 100" in name for name in names)


def test_a_drawn_things_words_are_read_its_name_first_then_its_parts_in_order(tmp_path: Path) -> None:
    pdfium = pytest.importorskip("pypdfium2")
    from pptx import Presentation

    protein = {"id": "chea", "kind": "protein", "label": "CheA", "properties": {"length": 654, "features": [
        {"type": "domain", "label": "P1", "start": 1, "end": 134},
        {"type": "domain", "label": "P2", "start": 159, "end": 227},
        {"type": "domain", "label": "P4 kinase", "start": 355, "end": 507},
        {"type": "phosphorylation", "label": "His48", "at": 48}]}}
    construct = {"id": "fret", "kind": "construct", "label": "Reporter", "properties": {"parts": [
        {"type": "promoter", "label": "Ptrc"}, {"type": "cds", "label": "cheY"},
        {"type": "terminator", "label": "T1"}]}}
    timeline = {"id": "assay", "kind": "timeline", "label": "Assay", "properties": {
        "unit": "s", "events": [{"at": 0, "label": "Added"}, {"at": 30, "label": "Back"}],
        "spans": [{"start": 10, "end": 30, "label": "Adapted"}, {"start": 0, "end": 10, "label": "Adapting"}]}}
    document = {"deck": {"id": "order"}, "slides": [
        {"title": "Parts", "body": [{"figure": {"figure": {"id": "f"}, "nodes": [protein, construct]}}]},
        {"title": "Times", "body": [{"figure": {"figure": {"id": "g"}, "nodes": [timeline]}}]},
    ]}
    deck = deck_from_document(document, tmp_path)
    built = deck.build(tmp_path, formats=("pdf", "pptx"))
    pages = pdfium.PdfDocument(built.pdf)  # type: ignore[arg-type]
    said = [" ".join(pages[index].get_textpage().get_text_range().split()) for index in range(2)]
    # Its name, then its parts from the N-terminus (a site among its domains) and 5' to 3';
    # a timeline's in time order, then the times along its axis.
    assert "CheA P1 His48 P2 P4 kinase 1 100" in said[0]
    assert "Reporter Ptrc cheY T1" in said[0]
    assert "Assay Added Adapting Adapted Back 0 s" in said[1]
    # So in the PowerPoint: each drawn thing's name first, every piece of it named plainly.
    slide = Presentation(str(built.pptx)).slides[0]
    groups = {group.name: [shape.name for shape in group.shapes] for group in slide.shapes[-2].shapes}
    chea, fret = groups["Protein CheA"], groups["Construct Reporter"]
    assert chea[0] == "Label “CheA”" and fret[0] == "Label “Reporter”"
    features = [name for name in chea if name.startswith("Feature")]
    assert features == ["Feature P1", "Feature P2", "Feature P4 kinase"]
    assert {"Chain", "Stem of Site His48", "Leader to Site His48"} <= set(chea)
    assert {"Stem of Part Ptrc", "Arrowhead of Part Ptrc", "Part cheY", "Part T1"} <= set(fret)
    assert not {"Shape", "Stem", "Head", "Leader"} & {*chea, *fret}


def test_a_part_named_by_a_label_on_two_lines_is_named_with_a_space_where_it_wrapped() -> None:
    from flexo.drawing import Group, read_drawing

    from flexo_talk.pptx import _shown_name

    drawing = read_drawing(
        '<svg xmlns="http://www.w3.org/2000/svg" width="200" height="60">'
        '<g id="f.lcd" data-flexo-entity="component" data-flexo-kind="structure">'
        '<text id="f.lcd.label" x="100" y="20" text-anchor="middle" font-family="Figtree" font-size="7.5">'
        '<tspan x="100">Lac repressor headpiece</tspan><tspan x="100" dy="8.85">on its operator (1LCD)</tspan>'
        "</text></g></svg>"
    )
    (part,) = (item for item in drawing.root.items if isinstance(item, Group))
    assert _shown_name(part) == "Structure Lac repressor headpiece on its operator (1LCD)"


def test_a_figure_says_what_its_lines_do_and_a_drawn_thing_its_parts(tmp_path: Path) -> None:
    from flexo_talk.export import _figure_said

    nodes = [{"id": "ep", "label": "Enteropeptidase", "kind": "terminal"}, {"id": "tg", "label": "Trypsinogen"},
             {"id": "t", "label": "Trypsin"}, {"id": "s", "label": "SPINK1"}, {"id": "ca", "label": "Ca²⁺"}]
    edges = [{"from": "ep", "to": "tg", "head": "catalysis"}, {"from": "tg", "to": "t"},
             {"from": "s", "to": "t", "head": "inhibition"}, {"from": "ca", "to": "t", "head": "modulation"}]
    timeline = {"id": "assay", "kind": "timeline", "label": "Assay", "properties": {
        "unit": "min", "events": [{"at": 14, "label": "EGTA added"}, {"at": 0, "label": "Start"}],
        "spans": [{"start": 0, "end": 8, "label": "Lag"}]}}
    protein = {"id": "p", "kind": "protein", "label": "PRSS1", "properties": {"length": 247, "features": [
        {"type": "domain", "label": "Protease", "start": 24, "end": 247},
        {"type": "mutation", "label": "R122H", "at": 122},
        {"type": "domain", "label": "Signal peptide", "start": 1, "end": 15}]}}
    document = {"deck": {"id": "said"}, "slides": [
        {"title": "Brakes", "body": [{"figure": {"figure": {"id": "f"}, "nodes": nodes, "edges": edges}}]},
        {"title": "Assay", "body": [{"figure": {"figure": {"id": "g"}, "nodes": [timeline]}}]},
        {"title": "PRSS1", "body": [{"figure": {"figure": {"id": "h"}, "nodes": [protein]}}]},
    ]}
    deck = deck_from_document(document, tmp_path)
    said = [_figure_said(slide.regions["body"].blocks[0]) for slide in deck.slides]
    # A line that inhibits is not an arrow onward: what each regulating line does is said.
    assert said[0] == ("A flow chart: Trypsinogen → Trypsin; Enteropeptidase catalyses Trypsinogen; "
                       "SPINK1 inhibits Trypsin; Ca²⁺ modulates Trypsin")
    # A timeline in time order, a protein from its first residue to its last.
    assert said[1] == "A timeline, Assay, in min: 0 Start; 0\u20138 Lag; 14 EGTA added."
    assert said[2] == ("A protein, PRSS1, 247 residues: Signal peptide 1\u201315; Protease 24\u2013247; "
                       "R122H at 122.")


def test_a_flow_chart_is_read_in_its_order_and_a_description_given_is_said_instead(tmp_path: Path) -> None:
    from pptx import Presentation

    from flexo_talk.export import build_deck, write_pptx

    document = yaml.safe_load(
        "deck: {id: flows}\nslides:\n  - title: Assay\n    body:\n      - figure:\n"
        "          figure: {id: f}\n          nodes:\n"
        "            - {id: check, label: 'Tubes formed?', kind: decision}\n"
        "            - {id: purify, label: Purify CA, kind: terminal}\n"
        "            - {id: mix, label: Mix CA with IP6}\n"
        "            - {id: grids, label: Cryo-EM grids, kind: terminal}\n"
        "          edges:\n"
        "            - {from: mix, to: check}\n            - {from: purify, to: mix}\n"
        "            - {from: check, to: grids, label: 'yes'}\n            - {from: check, to: mix, label: 'no'}\n"
        "  - title: Given\n    body:\n      - figure:\n          figure: {id: g}\n"
        "          nodes: [{id: a, label: Customer}, {id: b, label: Orders API}]\n"
        "          edges: [{from: a, to: b}]\n"
        "        description: '  The orders   path '\n"
    )
    deck = deck_from_document(document, tmp_path)
    slides = Presentation(write_pptx(deck, deck.render(), tmp_path / "talk.pptx")).slides
    described = [
        {element.get("name"): element.get("descr") for element in slide._element.iter() if element.get("descr")}
        for slide in slides
    ]
    # Read from where it starts, along its lines, its branches said by their labels (not
    # in the order its file lists them).
    flow = "A flow chart: Purify CA → Mix CA with IP6 → Tubes formed? (yes: Cryo-EM grids; no: back to Mix CA with IP6)"
    assert described == [{"Figure": flow}, {"Figure": "The orders path"}]
    pdf = build_deck(deck, tmp_path / "out", ("pdf",)).pdf.read_bytes()
    assert b"/Alt (The orders path)" in pdf
    # Written back as it was read.
    assert deck_document(deck)["slides"][1]["body"][0]["description"] == "  The orders   path "


def test_a_merged_deck_keeps_no_line_to_a_shape_deleted() -> None:
    from flexo.studio.merge import merge3

    def deck(nodes: list[str], edges: list[tuple[str, str]]) -> dict:
        figure = {
            "figure": {"id": "f"},
            "nodes": [{"id": node} for node in nodes],
            "groups": [{"id": "root", "layout": {"kind": "row"}, "children": nodes}],
        }
        if edges:
            figure["edges"] = [{"from": a, "to": b} for a, b in edges]
        return {"deck": {"id": "d"}, "slides": [{"title": "T", "body": [{"figure": figure}]}]}

    base = deck(["a", "b", "c"], [])
    ours = deck(["a", "b", "c"], [("a", "b"), ("b", "c")])
    theirs = deck(["a", "b"], [])
    merged = DeckKind().mended(merge3(base, ours, theirs))
    figure = merged["slides"][0]["body"][0]["figure"]
    assert figure["edges"] == [{"from": "a", "to": "b"}]
    assert figure["groups"][0]["children"] == ["a", "b"]
    # A line written to a shape the figure never had is no merge's doing: it stays, for the
    # slide to say (below).
    stray = deck(["a", "b"], [("a", "b"), ("b", "nowhere")])
    merged = DeckKind().mended(merge3(base, stray, base), [], base)
    assert {"from": "b", "to": "nowhere"} in merged["slides"][0]["body"][0]["figure"]["edges"]


def test_a_line_to_a_shape_there_is_none_of_is_left_out_of_the_slide_and_said(tmp_path: Path) -> None:
    figure = {
        "figure": {"id": "f"},
        "nodes": [{"id": "a", "label": "A"}, {"id": "b", "label": "B"}],
        "edges": [{"from": "a", "to": "b"}, {"from": "b", "to": "nowhere"}],
    }
    deck = deck_from_document({"slides": [{"title": "Lines", "body": [{"figure": figure}]}]}, tmp_path)
    drawn = render_slide(deck, deck.slides[0])
    assert any(
        message.endswith("A line to \u201cnowhere\u201d has no shape to go to.") for message in drawn.diagnostics
    )
    assert 'data-flexo-talk="invalid"' not in drawn.svg
    # Exported, it is said with its slide: not dropped without a word.
    from flexo_talk.studio import DeckKind

    kind = DeckKind()
    document = {"slides": [{"title": "Lines", "body": [{"figure": figure}]}]}
    assert kind.export(document, tmp_path, "t", ["pdf"])
    assert kind.export_notes == ["Slide 1: A line to \u201cnowhere\u201d has no shape to go to. It is left out."]


def test_a_paragraph_made_a_list_while_typed_in_is_one_list_with_the_words_typed() -> None:
    from flexo.studio.merge import merge3

    def deck(*body: dict) -> dict:
        return {"slides": [{"title": "Q", "body": [{"text": "First."}, *body]}]}

    base = deck({"text": "Second paragraph."})
    typed = deck({"text": "Second paragraph. typed by Alice"})
    listed = deck({"bullets": ["Second paragraph."], "numbered": True})
    for ours, theirs in ((typed, listed), (listed, typed)):
        notes: list = []
        merged = DeckKind().mended(merge3(base, ours, theirs, notes), notes, base)
        # Not the paragraph kept beside the list made of it: one list, with every word.
        assert merged["slides"][0]["body"] == [
            {"text": "First."},
            {"bullets": ["Second paragraph. typed by Alice"], "numbered": True},
        ]
        assert notes == []  # settled: no one is told it was deleted
    # A list made words, as its items were typed in.
    base = deck({"bullets": ["One", "Two"]})
    notes = []
    worded, typed = deck({"text": "One\nTwo"}), deck({"bullets": ["One", "Two, typed"]})
    merged = DeckKind().mended(merge3(base, worded, typed, notes), notes, base)
    assert merged["slides"][0]["body"] == [{"text": "First."}, {"text": "One\nTwo, typed"}]


def test_a_column_taken_away_while_a_cell_in_it_is_typed_in_leaves_the_rows_aligned() -> None:
    from flexo.studio.merge import merge3

    def deck(rows: list[list[str]]) -> dict:
        return {"slides": [{"title": "Q", "body": [{"table": rows}]}]}

    base = deck([["Model", "Params"], ["Baseline", "25.6M"], ["Ours", "24.0M"]])
    typed = deck([["Model", "Params"], ["Baseline, typed", "25.6M"], ["Ours", "24.0M"]])
    taken = deck([["Params"], ["25.6M"], ["24.0M"]])
    for ours, theirs in ((typed, taken), (taken, typed)):
        notes: list = []
        merged = DeckKind().mended(merge3(base, ours, theirs, notes), notes, base)
        # The taking away wins: no row left a cell longer than the others.
        assert merged["slides"][0]["body"] == [{"table": [["Params"], ["25.6M"], ["24.0M"]]}]
        assert notes == []


def test_an_object_written_with_two_kinds_is_made_the_one_it_became() -> None:
    def deck(*body: dict) -> dict:
        return {"slides": [{"title": "Q", "body": [{"text": "First."}, *body]}]}

    # A paragraph's last words written into it as it became a list: one list, with them.
    base = deck({"text": "Second paragraph."})
    both = deck({"bullets": ["Second paragraph.", "agent item"], "text": "Second paragraph. alice"})
    merged = DeckKind().mended(both, [], base)
    assert merged["slides"][0]["body"][1] == {
        "bullets": ["Second paragraph. alice", "agent item"]
    }
    # A list's words as they are, nested, where the paragraph's added nothing.
    both = deck({"bullets": ["Second paragraph.", ["under it"]], "text": "Second paragraph."})
    merged = DeckKind().mended(both, [], base)
    assert merged["slides"][0]["body"][1] == {"bullets": ["Second paragraph.", ["under it"]]}


def test_an_object_left_empty_by_someone_whose_window_went_is_taken_away() -> None:
    deck = {"slides": [{"title": "A"}, {"title": "Q", "body": [{"text": "Written."}, {"text": ""}]}]}
    kind = DeckKind()
    # Where they were typing: not the written one, nor a slide not there.
    assert not kind.abandoned(deck, {"page": 2, "block": "body[0]", "editing": True})
    assert not kind.abandoned(deck, {"page": 5, "block": "body[1]", "editing": True})
    assert kind.abandoned(deck, {"page": 2, "block": "body[1]", "editing": True})
    assert deck["slides"][1]["body"] == [{"text": "Written."}]


def test_a_merged_slide_keeps_its_objects_in_places_its_layout_has() -> None:
    def deck(slide: dict) -> dict:
        return {"slides": [slide]}

    # A body left beside the columns the slide was set out in: its stale copy goes, and an
    # object only it had is put in the first column.
    merged = deck({
        "layout": "two-columns",
        "left": [{"text": "First."}],
        "right": [{"text": "Second paragraph. alice keeps"}],
        "body": [{"text": "Second paragraph. ali"}, {"code": "x = 1"}],
    })
    slide = DeckKind().mended(merged, [], merged)["slides"][0]
    assert "body" not in slide
    assert slide["left"] == [{"text": "First."}, {"code": "x = 1"}]
    assert slide["right"] == [{"text": "Second paragraph. alice keeps"}]


HAND_WRITTEN = (
    "# Deck notes: three slides, short.\n"
    "schema_version: 1\n"
    "deck:\n  theme: paper\n  id: talk\n"
    "slides:\n"
    '- title: "The pipeline"\n  body:\n  - bullets:\n    - Generate   # agent: tighten this\n    - Design\n'
    "\n# The question comes last.\n"
    "- title: 'The question'\n  body:\n  - text: Why\n"
)


def test_a_deck_written_again_keeps_its_comments_quoting_and_order(tmp_path: Path) -> None:
    from flexo_talk.document import load_document, save_document

    path = tmp_path / "talk.yaml"
    path.write_text(HAND_WRITTEN, encoding="utf-8")
    document = load_document(path)
    document["slides"][1]["body"][0]["text"] = "Why it matters"
    save_document(document, path)
    # Only the words changed are written otherwise: the comments (a line's own, one over a
    # slide, one over the whole), the quotes and the deck's keys in their order stay.
    assert path.read_text(encoding="utf-8") == HAND_WRITTEN.replace("text: Why\n", "text: Why it matters\n")
    # A slide added is written as the studio writes one; a slide taken out takes its own
    # lines, and the comment over the next one stays with it.
    document["slides"].append({"title": "Last", "body": [{"text": "End"}]})
    del document["slides"][0]
    save_document(document, path)
    written = path.read_text(encoding="utf-8")
    assert written.startswith(HAND_WRITTEN.split("slides:")[0])
    assert "# The question comes last.\n- title: 'The question'" in written
    assert "The pipeline" not in written and written.endswith("- title: Last\n  body:\n  - text: End\n")
    assert yaml.safe_load(written)["slides"] == document["slides"]


def test_the_studio_writes_a_deck_over_the_words_it_read(tmp_path: Path) -> None:
    from flexo.studio.workspace import Workspace

    (tmp_path / "talk.yaml").write_text(HAND_WRITTEN, encoding="utf-8")
    workspace = Workspace(tmp_path)
    try:
        doc = workspace.open("talk.yaml")
        changed = json.loads(json.dumps(doc.document))
        changed["slides"][0]["title"] = "The pipeline, in short"
        doc.update(changed, doc.version, {"id": "ada", "kind": "person"})
        workspace.flush()
        written = (tmp_path / "talk.yaml").read_text(encoding="utf-8")
        # Its words changed in the quotes they were written in; the rest as it was.
        assert written == HAND_WRITTEN.replace('"The pipeline"', '"The pipeline, in short"')
    finally:
        workspace.close()


def test_a_figure_with_a_shape_that_cannot_be_drawn_draws_the_rest_and_says_which(tmp_path: Path) -> None:
    from flexo_talk.studio import DeckKind

    figure = {
        "figure": {"id": "f"},
        "nodes": [{"id": "p", "kind": "protein", "label": "Spike"}, {"id": "q", "label": "Next"}],
        "edges": [{"from": "p", "to": "q"}],
    }
    document = {"slides": [{"title": "Spike", "body": [{"figure": figure}]}]}
    drawing = DeckKind().draw(document, tmp_path, {"settle": True})
    (message,) = [item for item in drawing.messages if item.code == "deck.figure"]
    # Said of the shape by its words, its place naming it, to choose it by: "#p:length".
    assert message.text.startswith("“Spike” can\u2019t be drawn yet: a protein needs its length")
    assert message.where.endswith("#p:length") and message.place == "Slide 1 · Figure"
    # The rest of the figure is drawn, the protein a plain box of its words: no empty box.
    svg = drawing.pages[0].svg
    assert 'id="slide1.body.0.q"' in svg and 'id="slide1.body.0.p"' in svg
    assert 'data-flexo-talk="invalid"' not in svg


def test_a_figure_that_fails_as_it_is_laid_out_leaves_the_rest_of_its_slide(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    import flexo_talk.compose as compose
    from flexo_talk.studio import DeckKind

    fit = compose._fit_in_box

    def failing(spec, *args, **options):
        if any(node.id == "boom" for node in spec.nodes):
            raise KeyError("boom")
        return fit(spec, *args, **options)

    monkeypatch.setattr(compose, "_fit_in_box", failing)
    figure = {"figure": {"id": "f"}, "nodes": [{"id": "boom", "label": "Log it"}]}
    document = {"slides": [{"title": "Kept", "body": [{"figure": figure}, {"text": "Words stay"}]}]}
    drawing = DeckKind().draw(document, tmp_path, {"settle": True})
    svg = drawing.pages[0].svg
    assert "Kept" in svg and "Words stay" in svg and 'data-flexo-talk="invalid"' in svg
    assert any("“Log it”" in message.text for message in drawing.messages)
    kind = DeckKind()
    assert kind.export(document, tmp_path, "t", ["pdf"])
    assert kind.export_notes == ["Slide 1 has a figure left empty: it can\u2019t be drawn as it is."]


def test_a_figure_edited_on_a_slide_is_said_as_the_figure_says_it() -> None:
    from flexo_talk.studio import DeckKind

    figure = {"figure": {"id": "f"}, "nodes": [{"id": "a", "label": "Start"}]}
    before = {"slides": [{"title": "Flow", "body": [{"figure": figure}]}]}
    grown = {**figure, "nodes": [*figure["nodes"], {"id": "x", "label": "Extra step"}]}
    after = {"slides": [{"title": "Flow", "body": [{"figure": grown}]}]}
    (note,) = DeckKind().describe(before, after)
    assert note["text"] == "added “Extra step” to the figure on slide 1"


def _captioned(tmp_path: Path) -> dict:
    from PIL import Image

    Image.new("RGB", (400, 300), (30, 90, 160)).save(tmp_path / "tracks.png")
    figure = {"figure": {"id": "f"}, "nodes": [{"id": "r", "label": "Receptor"}, {"id": "k", "label": "Kinase"}],
              "edges": [{"from": "r", "to": "k"}]}
    return {"deck": {"id": "captions"}, "slides": [
        {"title": "Tracks", "body": [{"image": "tracks.png", "caption": "Figure 1. Tracks", "description": "Tracks"}]},
        {"title": "Cues", "body": [{"table": [["Receptor", "Cue"], ["Tar", "Aspartate"]], "caption": "Table 1. Cues"}]},
        {"title": "Path", "body": [{"figure": figure, "caption": "Figure 2. Receptor to kinase"}]},
    ]}


def test_a_picture_table_or_figure_has_a_caption_under_it_tagged_and_grouped_with_it(tmp_path: Path) -> None:
    pytest.importorskip("PIL")
    import xml.etree.ElementTree as ET

    from pptx import Presentation

    deck = deck_from_document(_captioned(tmp_path), tmp_path)
    for index, (slide, rendered) in enumerate(zip(deck.slides, deck.render(), strict=True)):
        block = f"{slide.id}.body.0"
        # Under it, centred under it and a size smaller -- and part of it: in its group, so it
        # is chosen and moved with it.
        group = next(item for item in ET.fromstring(rendered.svg).iter() if item.get("id") == block)
        caption = next((item for item in group.iter() if item.get("id") == f"{block}.caption"), None)
        assert caption is not None, index
        assert caption.get("text-anchor") == "middle" and float(caption.get("font-size")) == deck.style.small_size
    built = deck.build(tmp_path / "out", formats=("pdf", "pptx"))
    # A caption of what it is under, in the PDF's tags.
    pdf = built.pdf.read_bytes()  # type: ignore[union-attr]
    parents = [re.search(rb"\b%s 0 obj\s*<< /Type /StructElem /S /(\w+)" % parent, pdf).group(1)
               for parent in re.findall(rb"/S /Caption /P (\d+) 0 R", pdf)]
    assert parents == [b"Figure", b"Table", b"Figure"]
    # In the PowerPoint, a picture and a figure are grouped with theirs; a table (which no
    # slide program groups) is followed by it.
    slides = Presentation(str(built.pptx)).slides
    picture = next(shape for shape in slides[0].shapes if shape.name == "Picture")
    assert [shape.name for shape in picture.shapes] == ["Picture", "Caption “Figure 1. Tracks”"]
    # Described once, by the group: the picture in it is marked decorative, not read again.
    inner = picture.shapes[0]._element.find(".//{*}cNvPr")
    assert picture._element.find(".//{*}cNvPr").get("descr") == "Tracks" and inner.get("descr") is None
    assert inner.find(".//{*}decorative").get("val") == "1"
    names = [shape.name for shape in slides[1].shapes]
    assert names[names.index("Table") + 1] == "Caption “Table 1. Cues”"
    figure = next(shape for shape in slides[2].shapes if shape.name == "Figure")
    assert "Caption “Figure 2. Receptor to kinase”" in [shape.name for shape in figure.shapes]


def test_words_set_as_a_caption_say_they_are_centred_and_keep_an_alignment_given_them(tmp_path: Path) -> None:
    pytest.importorskip("PIL")
    from PIL import Image

    Image.new("RGB", (400, 300), (30, 90, 160)).save(tmp_path / "tracks.png")

    def drawn(text: dict) -> str:
        document = {"deck": {"id": "aligned"}, "slides": [{"title": "Tracks", "body": [{"image": "tracks.png"}, text]}]}
        deck = deck_from_document(document, tmp_path)
        return re.search(r'<text id="slide1\.body\.1"[^>]*>', render_slide(deck, deck.slides[0]).svg).group(0)

    # Under a picture, words are centred as a caption is, and say so (the inspector shows
    # Centre, as drawn); an alignment written for them is kept, Left too.
    centred = drawn({"text": "Swimming tracks"})
    assert 'text-anchor="middle"' in centred and 'data-flexo-align="middle"' in centred
    left = drawn({"text": "Swimming tracks", "align": "start"})
    assert "text-anchor" not in left and "data-flexo-align" not in left


def test_any_object_builds_in_on_a_click_of_its_own_in_order(tmp_path: Path) -> None:
    pytest.importorskip("PIL")
    import zipfile

    document = _captioned(tmp_path)
    figure = document["slides"][2]["body"][0]
    document["slides"].append({"title": "Builds", "layout": "two-columns", "left": [
        {"bullets": ["One", "Two"], "reveal": True}], "right": [
        {**figure, "build": True}, {"text": "Then this", "build": True}]})
    document["slides"][1]["body"][0]["build"] = True
    deck = deck_from_document(document, tmp_path)
    rendered = deck.render()
    builds = rendered[3]
    # The list's items a click each, then the figure, then the words: as Keynote builds in.
    assert builds.steps == 5 and builds.builds == [(4, "slide4.right.0"), (5, "slide4.right.1")]
    assert 'id="slide4.right.0"' not in builds.at_step(3) and 'id="slide4.right.0"' in builds.at_step(4)
    assert 'id="slide4.right.1"' not in builds.at_step(4)
    built = deck.build(tmp_path / "out", formats=("pdf", "pptx"))
    import pypdfium2 as pdfium

    assert len(pdfium.PdfDocument(built.pdf)) == 1 + 2 + 1 + 5  # type: ignore[arg-type]
    with zipfile.ZipFile(built.pptx) as archive:  # type: ignore[arg-type]
        table, last = (archive.read(f"ppt/slides/slide{n}.xml").decode() for n in (2, 4))

    def clicks(xml: str) -> list[list[tuple[str, str]]]:
        names = dict(re.findall(r'<p:cNvPr id="(\d+)" name="([^"]*)"', xml))
        timing = xml[xml.index("<p:timing"):]
        found = re.split(r'<p:cond delay="indefinite"/>', timing)[1:]
        return [[(names[spid], kind) for kind, spid in re.findall(r'nodeType="(\w+)".*?spid="(\d+)"', click)]
                for click in found]

    # In PowerPoint, an Appear on each click: a table with its caption together.
    assert clicks(table) == [[("Table", "clickEffect"), ("Caption “Table 1. Cues”", "withEffect")]]
    assert [[name for name, _ in click] for click in clicks(last)] == [
        ["List “One Two”"], ["List “One Two”"], ["Figure"], ["Text “Then this”"]]
    assert '<p:bldGraphic spid="' in table


def test_a_figure_arranged_by_hand_drawn_small_says_turning_would_draw_it_larger(tmp_path: Path) -> None:
    import itertools

    from flexo_talk.studio import DeckKind

    steps = [f"n{index}" for index in range(9)]
    figure = {
        "figure": {"id": "row"},
        "nodes": [{"id": step, "label": f"Step number {index}"} for index, step in enumerate(steps)],
        "edges": [{"from": a, "to": b} for a, b in itertools.pairwise(steps)],
        "groups": [
            {"id": "root", "layout": {"kind": "column"}, "role": "canvas", "children": ["line", "n8"]},
            {"id": "line", "layout": {"kind": "row"}, "role": "layout", "children": steps[:8]},
        ],
    }
    drawing = DeckKind().draw({"slides": [{"title": "Long", "body": [{"figure": figure}]}]}, tmp_path, {"settle": True})
    (note,) = [message for message in drawing.messages if message.code.startswith("figure.small")]
    # Coded for the studio to offer the switch in one click.
    assert note.code == "figure.small.turn" and note.text.endswith("Turned to fit the slide, it would be larger.")
    assert note.where == "slides[0] row"


def test_a_slide_of_a_layout_there_is_none_of_is_content_with_every_object_in_its_body(tmp_path: Path) -> None:
    from flexo_talk.document import UnknownLayout

    errors: list[DeckDocumentError] = []
    slide = {"layout": "two-column", "title": "Typed wrong", "left": [{"text": "Words"}],
             "right": [{"bullets": ["One", "Two"]}]}
    deck = deck_from_document({"deck": {}, "slides": [slide]}, tmp_path, errors=errors)
    (drawn,) = deck.slides
    # Drawn as Content -- its objects in its body, in order -- not stood aside as blank.
    assert drawn.layout == "content" and [type(block).__name__ for block in drawn.regions["body"].blocks] == [
        "_Words", "_Bullets"]
    assert [(type(error), error.where) for error in errors] == [(UnknownLayout, "slides[0].layout")]
    assert errors[0].message == "\u201ctwo-column\u201d isn\u2019t a layout: drawn as Content. Choose one from Layout."


def test_a_figure_of_one_molecule_is_described_once_in_powerpoint(tmp_path: Path) -> None:
    pytest.importorskip("molsketch")
    import shutil

    from pptx import Presentation

    from flexo_talk.export import write_pptx

    data = Path(__file__).resolve().parents[2] / "flexo" / "tests" / "unit" / "data" / "1a7g.cif"
    if not data.is_file():
        pytest.skip("no structure file beside flexo")
    shutil.copy(data, tmp_path / "1a7g.cif")
    figure = {"figure": {"id": "mol"}, "nodes": [{"id": "m", "kind": "structure", "label": "E2 domain",
                                                   "properties": {"source": "1a7g.cif"}}]}
    document = {"deck": {"id": "talk"}, "slides": [
        {"title": "A molecule", "body": [{"figure": figure, "description": "A dimeric beta-barrel"}]}]}
    deck = deck_from_document(document, tmp_path)
    slide = Presentation(write_pptx(deck, deck.render(), tmp_path / "talk.pptx")).slides[0]
    described = [(element.get("name"), element.get("descr"), element.find(".//{*}decorative") is not None)
                 for element in slide._element.iter() if element.tag.endswith("}cNvPr")
                 and (element.get("descr") or element.find(".//{*}decorative") is not None)]
    # Its description on the figure; the molecule's picture in it says the same, so it is
    # marked decorative -- a screen reader says it once.
    assert described == [("Figure", "A dimeric beta-barrel", False), ("Structure E2 domain", None, True)]
