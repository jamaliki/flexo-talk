"""A slide of columns set in rows (``rows``): block *i* of every column in row *i*, every
block of a row starting on one line, and a label beside each row in a gutter at the slide's
edge -- in the slide, the PowerPoint, the file and the studio alike. Without ``rows`` each
column stacks its blocks on its own, as it always did."""

from __future__ import annotations

import json
import re
import shutil
import subprocess
import xml.etree.ElementTree as ET
from pathlib import Path

import pytest
import yaml

from flexo_talk import Deck, compose
from flexo_talk.compose import PLACEHOLDERS, ROW_GAP, ROW_GUTTER, ROW_LABEL, ROW_WEIGHT, _Canvas, render_slide
from flexo_talk.deck import RenderedSlide, inline
from flexo_talk.document import DeckDocumentError, deck_document, deck_from_document, dump_document
from flexo_talk.studio import DeckKind

LONG = "A paragraph of words long enough to wrap onto a second line in its half of the slide, and so on."


def _picture(folder: Path, name: str = "pic.png", size: tuple[int, int] = (320, 200)) -> str:
    from PIL import Image

    Image.new("RGB", size, (91, 141, 239)).save(folder / name)
    return name


def _deck(slide: dict, folder: Path = Path("."), **deck: object) -> Deck:
    return deck_from_document({"deck": {"id": "rows", **deck}, "slides": [slide]}, folder)


def _drawn(slide: dict, folder: Path = Path("."), *, editing: bool = False, **deck: object) -> RenderedSlide:
    made = _deck(slide, folder, **deck)
    token = PLACEHOLDERS.set(editing)
    try:
        return render_slide(made, made.slides[0])
    finally:
        PLACEHOLDERS.reset(token)


def _part(svg: str, identifier: str) -> ET.Element | None:
    """What a slide drew under an id, where it is on the slide: its ``y`` with how far the
    groups it is in were moved down (a region centred, a block stood down its row)."""

    root = ET.fromstring(svg)
    parents = {child: parent for parent in root.iter() for child in parent}
    node = next((item for item in root.iter() if item.get("id") == identifier), None)
    if node is None:
        return None
    above = node
    down = 0.0
    while above in parents:
        above = parents[above]
        moved = re.fullmatch(r"translate\(0 ([-\d.]+)\)", above.get("transform") or "")
        down += float(moved.group(1)) if moved else 0.0
    node.set("y", str(float(node.get("y", 0.0)) + down))
    return node


def _y(svg: str, identifier: str) -> float:
    node = _part(svg, identifier)
    assert node is not None, identifier
    return float(node.get("y"))


def _room(svg: str, region: str) -> tuple[float, float, float, float]:
    found = re.search(rf'<g id="slide1\.{region}"[^>]*data-flexo-box="([^"]+)"', svg)
    assert found, region
    x, y, width, height = map(float, found.group(1).split())
    return x, y, width, height


def _two(left: list, right: list, **settings: object) -> dict:
    return {"title": "Rows", "layout": "two-columns", "left": left, "right": right, **settings}


def test_every_block_of_a_row_starts_on_one_line_and_a_row_is_as_tall_as_its_tallest() -> None:
    left = [{"text": LONG}, {"text": "Left, second"}, {"text": "Left, third"}]
    right = [{"text": "Short"}, {"text": LONG}, {"text": "Right, third"}]
    stacked = _drawn(_two(left, right)).svg
    rowed = _drawn(_two(left, right, rows=[])).svg
    # Each column on its own: the second of one starts higher than the second of the other.
    assert _y(stacked, "slide1.right.1") < _y(stacked, "slide1.left.1") - 10
    for row in range(3):
        assert _y(rowed, f"slide1.left.{row}") == pytest.approx(_y(rowed, f"slide1.right.{row}"))
    # A row is as tall as its tallest block, the next after the usual space: the second row
    # where the left column's own second block was, under its two lines; the third under the
    # right's two lines, as far again.
    assert _y(rowed, "slide1.left.1") == pytest.approx(_y(stacked, "slide1.left.1"))
    step = _y(stacked, "slide1.left.1") - _y(stacked, "slide1.left.0")
    assert _y(rowed, "slide1.left.2") - _y(rowed, "slide1.left.1") == pytest.approx(step)
    # No label, no gutter: only the rows line up, the columns as wide as they were.
    assert "slide1.row" not in rowed
    assert _room(rowed, "left") == _room(stacked, "left") and _room(rowed, "right") == _room(stacked, "right")


def test_each_label_stands_in_a_gutter_beside_its_row_centred_down_it(tmp_path: Path) -> None:
    picture = {"image": _picture(tmp_path)}
    slide = _two([{"text": "**One**"}, picture, picture], [{"text": "**Other**"}, picture, picture],
                 rows=["", "Mechanism", "Geometry", "Kept, and not drawn"], split=0.3)
    deck = _deck(slide, tmp_path)
    rendered = render_slide(deck, deck.slides[0])
    svg, style = rendered.svg, deck.style
    # A row without a label has none, and one past the last row is not drawn.
    assert _part(svg, "slide1.row0") is None and _part(svg, "slide1.row3") is None
    canvas = _Canvas(deck, deck.slides[0])
    size = style.body_size * ROW_LABEL
    widest = max(canvas.measure(inline(words), size, None, ROW_WEIGHT).width for words in ("Mechanism", "Geometry"))
    for row, words in ((1, "Mechanism"), (2, "Geometry")):
        label = _part(svg, f"slide1.row{row}")
        assert label is not None and "".join(label.itertext()) == words
        # The deck's face, a step smaller than its body's words, semibold, in the muted ink.
        assert float(label.get("font-size")) == pytest.approx(size) and label.get("font-weight") == "600"
        assert label.get("data-flexo-fill") == "muted-ink" and label.get("fill") == deck.palette.get("muted-ink")
        # At the slide's left, where its title starts; centred down its row, which is as tall
        # as the taller of its pictures (the wider column's).
        assert float(label.get("x")) == pytest.approx(style.margin)
        metrics = canvas.measure(inline(words), size, None, ROW_WEIGHT)
        first, second = (_part(svg, f"slide1.{name}.{row}") for name in ("left", "right"))
        assert float(second.get("y")) == pytest.approx(float(first.get("y")))
        assert float(second.get("height")) > float(first.get("height")) + 5.0
        middle = float(second.get("y")) + float(second.get("height")) / 2.0
        assert float(label.get("y")) - metrics.baseline + metrics.height / 2.0 == pytest.approx(middle, abs=0.01)
    # The gutter is as wide as the widest label; the columns share what is left by the split.
    (left, top, wide, _), (other, _, rest, _) = _room(svg, "left"), _room(svg, "right")
    assert left == pytest.approx(style.margin + widest + 0.01 + style.column_gap * ROW_GAP, abs=0.01)
    assert wide / (wide + rest) == pytest.approx(0.3)
    assert other + rest == pytest.approx(style.width - style.margin) and other == pytest.approx(left + wide + 36.0)
    # The rows follow one another from the body's top, the usual space between them.
    assert _y(svg, "slide1.left.1") > top + style.block_gap
    under = float(_part(svg, "slide1.right.1").get("y")) + float(_part(svg, "slide1.right.1").get("height"))
    assert _y(svg, "slide1.left.2") == pytest.approx(under + style.block_gap, abs=0.01)
    assert not rendered.diagnostics


def test_a_label_longer_than_a_fifth_of_the_body_wraps_in_its_gutter() -> None:
    long = "A label of a good many words that will not fit on one line"
    deck = _deck(_two([{"text": "a"}, {"text": "b"}], [{"text": "c"}, {"text": "d"}], rows=["Short", long]))
    svg = render_slide(deck, deck.slides[0]).svg
    body = deck.style.width - 2 * deck.style.margin
    left, *_ = _room(svg, "left")
    assert left - deck.style.margin <= body * ROW_GUTTER + deck.style.column_gap * ROW_GAP + 0.02
    label = _part(svg, "slide1.row1")
    assert label is not None and len([line for line in label.iter() if line.tag.endswith("tspan")]) > 1
    # Its row is as tall as its lines, a line of words beside it or not: nothing under it is
    # set over them.
    lines = [float(line.get("y") or 0.0) for line in label.iter() if line.tag.endswith("tspan") and line.get("y")]
    assert _y(svg, "slide1.left.1") == pytest.approx(_y(svg, "slide1.right.1"))
    three = _drawn(_two([{"text": "a"}, {"text": "b"}, {"text": "e"}], [{"text": "c"}, {"text": "d"}],
                        rows=["Short", long, "Under"])).svg
    assert _y(three, "slide1.left.2") - _y(three, "slide1.left.1") > (len(lines) or 2) * 17.0


def test_labels_read_from_the_right_stand_at_the_right() -> None:
    slide = _two([{"text": "متن"}, {"text": "دیگر"}], [{"text": "یک"}, {"text": "دو"}], rows=["سازوکار", "هندسه"])
    deck = _deck({**slide, "title": "عنوان"})
    svg = render_slide(deck, deck.slides[0]).svg
    style = deck.style
    label = _part(svg, "slide1.row0")
    assert label is not None and label.get("text-anchor") == "end"
    assert float(label.get("x")) == pytest.approx(style.width - style.margin, abs=0.02)
    # The columns start at the left margin, the gutter after them.
    (left, *_), (other, _, wide, _) = _room(svg, "left"), _room(svg, "right")
    assert left == pytest.approx(style.margin) and other + wide < float(label.get("x")) - 40.0
    # One not yet written beside it (its placeholder, read from the left) moves the gutter
    # nowhere while the slide is edited; nor, on a slide read from the right, with none written.
    for labels in (["", "هندسه"], ["", ""]):
        editing = _drawn({**slide, "title": "عنوان", "rows": labels}, editing=True).svg
        assert _room(editing, "left")[0] == pytest.approx(style.margin)
    assert _room(_drawn({**slide, "rows": ["", ""]}, editing=True).svg, "left")[0] > style.margin + 40.0


def test_a_label_follows_the_theme_and_a_dark_slide() -> None:
    slide = _two([{"text": "a"}], [{"text": "b"}], rows=["Label"])
    light = _drawn(slide).svg
    dark = _drawn({**slide, "background": "#1b2a41"}).svg
    other = _drawn(slide, theme="dark").svg
    fills = [_part(svg, "slide1.row0").get("fill") for svg in (light, dark, other)]
    assert fills[0] != fills[1] and fills[0] != fills[2]
    assert all(_part(svg, "slide1.row0").get("data-flexo-fill") == "muted-ink" for svg in (light, dark, other))


def test_a_picture_shorter_than_the_words_beside_it_is_centred_down_its_row(tmp_path: Path) -> None:
    words = {"bullets": [f"Point {index}" for index in range(7)]}
    picture = {"image": _picture(tmp_path), "width": 150}
    slide = _two([{"text": "Head"}, words], [{"text": "Head"}, picture], rows=["", "Row"])
    svg = _drawn(slide, tmp_path).svg

    def drawn(block: dict, **settings: object) -> str:
        return _drawn(_two([{"text": "Head"}, words], [{"text": "Head"}, block], rows=["", "Row"], **settings),
                      tmp_path).svg

    def top(svg: str) -> float:
        return _y(svg, "slide1.right.1")

    # Its middle is its row's, where the row's label is centred too.
    middle = top(svg) + float(_part(svg, "slide1.right.1").get("height")) / 2.0
    assert abs(float(_part(svg, "slide1.row1").get("y")) - middle) < 12.0
    # With the slide's content at its top, it starts on its row's line, as the words do; so
    # too asked to (vertical), wherever the content stands.
    line = top(drawn(picture, align="top"))
    assert "data-flexo-down" not in drawn(picture, align="top")
    assert top(drawn({**picture, "vertical": "top"})) == pytest.approx(line) and top(svg) > line + 40.0
    # Asked to stand at the bottom of its place, it stands at the bottom of its row: as far
    # under the row's middle as its top is over it.
    assert top(drawn({**picture, "vertical": "bottom"})) - top(svg) == pytest.approx(top(svg) - line)


def test_the_rows_stand_down_the_body_together_as_the_content_does(tmp_path: Path) -> None:
    slide = _two([{"text": "One"}, {"text": "Two"}], [{"text": "Three"}, {"text": "Four"}], rows=["A", "B"])
    top = _drawn(slide).svg
    middle = _drawn({**slide, "align": "middle"}).svg
    moved = _y(middle, "slide1.left.0") - _y(top, "slide1.left.0")
    assert moved > 50.0
    # Every column, and every label, by as much.
    for part in ("left.1", "right.0", "right.1", "row0", "row1"):
        assert _y(middle, f"slide1.{part}") - _y(top, f"slide1.{part}") == pytest.approx(moved)
    # Pictures alone stand a little above the body's middle, and share one top in a row.
    small, large = {"image": _picture(tmp_path), "width": 100}, {"image": _picture(tmp_path)}
    alone = _drawn(_two([small, large], [large, small], rows=["A", "B"]), tmp_path).svg
    assert _y(alone, "slide1.left.0") == pytest.approx(_y(alone, "slide1.right.0"))
    assert _y(alone, "slide1.left.1") == pytest.approx(_y(alone, "slide1.right.1"))
    at_top = _drawn(_two([small, large], [large, small], rows=["A", "B"], align="top"), tmp_path).svg
    assert 0.0 <= _y(alone, "slide1.left.0") - _y(at_top, "slide1.left.0") < 60.0


def _shaped(folder: Path, left: dict | None = None, right: dict | None = None, **settings: object) -> dict:
    """A slide of a heading, a mechanism and a picture in each column, a row each, under a
    title of two lines: a wide picture and a squarer one, each asked to be wider than its
    column (``left``, ``right``: what else each picture is asked)."""

    def mechanism(smiles: str, bond: str) -> dict:
        return {"mechanism": [{"smiles": smiles, "arrows": ["1 -> 2", f"{bond} -> 3"]}], "horizontal": "middle"}

    wide, squarer = _picture(folder, "wide.png", (660, 300)), _picture(folder, "squarer.png", (600, 300))
    return _two(
        [{"text": "**Backside attack**", "align": "center"}, mechanism("[X-:1].[Y:2][Z:3]", "2-3"),
         {"image": wide, "width": 400, **(left or {})}],
        [{"text": "**Attack from above**", "align": "center"}, mechanism("[X-:1].[Y:2]=[Z:3]", "2=3"),
         {"image": squarer, "width": 360, **(right or {})}],
        rows=[" ", "Mechanism", "Geometry"],
        title="A title long enough to take two lines of the slide, as a claim does", **settings,
    )


def _box(svg: str, identifier: str) -> tuple[float, float, float, float]:
    part = _part(svg, identifier)
    assert part is not None, identifier
    return tuple(float(part.get(name)) for name in ("x", "y", "width", "height"))


def test_a_row_that_cannot_use_its_share_leaves_the_rest_to_the_rows_that_can(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    deck = _deck(_shaped(tmp_path), tmp_path)
    svg, gap = render_slide(deck, deck.slides[0]).svg, deck.style.block_gap
    _, top, wide, room = _room(svg, "left")
    heading = _Canvas(deck, deck.slides[0]).measure(inline("**Backside attack**"), deck.style.body_size, wide).height
    left, right = _box(svg, "slide1.left.2"), _box(svg, "slide1.right.2")
    # The mechanisms need a good deal less than half the height the two rows have: their row
    # is as tall as they are, the pictures starting on one line under it.
    equal = (room - heading - 2 * gap) / 2.0
    mechanisms = left[1] - gap - (top + heading + gap)
    assert left[1] == pytest.approx(right[1]) and 60.0 < mechanisms < equal - 30.0
    # The pictures' row has the rest, to the body's foot -- no band of the slide left empty
    # under it -- and the squarer picture, the taller, is as tall as it: narrower than it was
    # asked to be, and than its column.
    rest = room - heading - mechanisms - 2 * gap
    assert right[3] == pytest.approx(rest, abs=0.01) and right[1] + right[3] == pytest.approx(top + room, abs=0.01)
    assert right[2] == pytest.approx(2.0 * rest, abs=0.01) and right[2] < min(360.0, wide) - 20.0
    # The wide one is as wide as its column, the most it can be however wide it is asked.
    assert left[2] == pytest.approx(wide) and left[3] == pytest.approx(wide / 2.2) and left[3] < rest
    # Each a good deal taller than in half the height, which is what they had -- the rows
    # sharing it equally whatever they used, and what the mechanisms left of theirs left
    # under the pictures.
    monkeypatch.setattr(compose, "ROW_TRIES", 0)
    before = render_slide(deck, deck.slides[0]).svg
    monkeypatch.undo()
    assert _box(before, "slide1.left.2")[3] == pytest.approx(equal) == _box(before, "slide1.right.2")[3]
    assert left[3] > equal + 30.0 and right[3] > equal + 35.0
    assert _box(before, "slide1.right.2")[1] + equal < top + room - 35.0
    # Asked nothing, they are drawn the same: as wide as its column, as tall as its row.
    assert _drawn(_shaped(tmp_path, {"width": None}, {"width": None}), tmp_path).svg == svg
    # A width its row now has the height for is the picture's: each 300 wide, where half the
    # height let the wide one be 280 and the squarer one 255.
    asked = _drawn(_shaped(tmp_path, {"width": 300}, {"width": 300}), tmp_path).svg
    assert _box(asked, "slide1.left.2")[2:] == pytest.approx((300.0, 300.0 / 2.2))
    assert _box(asked, "slide1.right.2")[2:] == pytest.approx((300.0, 150.0))
    assert 2.2 * equal < 300.0 < 2.0 * rest


def test_what_a_row_leaves_is_shared_out_until_no_row_can_use_more(tmp_path: Path) -> None:
    wide, tall = _picture(tmp_path, "wide.png", (660, 300)), _picture(tmp_path, "tall.png", (200, 400))
    words = [{"text": "One"}, {"text": "Two"}, {"text": "Three"}]

    def drawn(*pictures: dict, **settings: object) -> tuple[str, list[tuple[float, float, float, float]], float]:
        svg = _drawn(_two(list(pictures), words[:len(pictures)], rows=[], **settings), tmp_path).svg
        return svg, [_box(svg, f"slide1.left.{row}") for row in range(len(pictures))], _room(svg, "left")[3]

    # Rows that use all they have share the height equally, as they did.
    svg, boxes, room = drawn({"image": tall}, {"image": tall}, {"image": tall})
    assert [box[3] for box in boxes] == pytest.approx([(room - 36.0) / 3.0] * 3)
    # A small picture's row is as tall as it; the two others share the rest, equally.
    svg, boxes, room = drawn({"image": wide, "width": 110}, {"image": tall}, {"image": tall})
    assert boxes[0][3] == pytest.approx(50.0)
    assert [box[3] for box in boxes[1:]] == pytest.approx([(room - 36.0 - 50.0) / 2.0] * 2)
    # One of those is as wide as it was asked to be before it has its half of the rest: the
    # last row has what that leaves too, to the body's foot.
    svg, boxes, room = drawn({"image": wide, "width": 110}, {"image": wide, "width": 264}, {"image": tall})
    assert (room - 36.0) / 3.0 < 120.0 < (room - 36.0 - 50.0) / 2.0
    assert [box[3] for box in boxes] == pytest.approx([50.0, 120.0, room - 36.0 - 170.0])
    assert boxes[2][1] + boxes[2][3] == pytest.approx(_room(svg, "left")[1] + room, abs=0.01)
    # None able to use more, the rows are as tall as their pictures and the rest is left: at
    # the body's foot, the rows at its top.
    svg, boxes, room = drawn({"image": wide, "width": 110}, {"image": wide, "width": 264}, align="top")
    assert [box[3] for box in boxes] == pytest.approx([50.0, 120.0])
    assert boxes[0][1] == pytest.approx(_room(svg, "left")[1]) and boxes[1][1] == pytest.approx(boxes[0][1] + 68.0)
    # A row whose words need more than its share is as tall as they are, as it was: what the
    # small picture's row leaves goes to the row that can use it.
    long = {"bullets": [f"Point {index}" for index in range(5)]}
    svg = _drawn(_two([{"image": wide, "width": 110}, {"image": tall}, {"image": tall}],
                      [{"text": "One"}, long, {"text": "Three"}], rows=[]), tmp_path).svg
    listed = _y(svg, "slide1.left.2") - 18.0 - _y(svg, "slide1.left.1")
    assert listed > (room - 36.0 - 50.0) / 2.0 + 5.0
    assert _box(svg, "slide1.left.2")[3] == pytest.approx(room - 36.0 - 50.0 - listed, abs=0.01)


def test_the_height_a_row_would_have_at_most_is_what_the_rest_leave_it() -> None:
    # Rows that take all they are given share the room equally; one that takes less than an
    # equal share leaves the rest to the others, and one that would take more has a share
    # like theirs.
    assert compose._level(300.0, [], 2) == pytest.approx(150.0)
    assert compose._level(300.0, [50.0], 1) == pytest.approx(250.0)
    assert compose._level(300.0, [50.0, 200.0], 1) == pytest.approx(125.0)
    assert compose._level(300.0, [120.0, 50.0], 2) == pytest.approx(250.0 / 3.0)


def test_figures_at_one_scale_down_a_column_grow_together_and_the_rows_still_fit(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    def chain(name: str, count: int) -> dict:
        nodes = [{"id": f"n{index}", "label": f"Step {index}"} for index in range(count)]
        edges = [{"from": f"n{index}", "to": f"n{index + 1}"} for index in range(count - 1)]
        return {"figure": {"figure": {"id": name}, "nodes": nodes, "edges": edges}, "turn": False}

    # A short figure and a long one down a column, at one scale, over a tall picture: the
    # short one's row leaves room, the others take it -- and the short one grows with the
    # long one, which the rows then allow for.
    slide = _two([chain("a", 2), chain("b", 6), {"image": _picture(tmp_path, "tall.png", (200, 400))}],
                 [{"text": "One"}, {"text": "Two"}, {"text": "Three"}], rows=["A", "B", "C"])
    deck = _deck(slide, tmp_path)

    def drawn() -> tuple[list[float], tuple[float, float, float, float]]:
        svg = render_slide(deck, deck.slides[0]).svg
        scales = re.findall(r'<g id="slide1\.left\.[01]" transform="translate\([^)]*\) scale\(([\d.]+)\)"', svg)
        assert _box(svg, "slide1.left.2")[1] + _box(svg, "slide1.left.2")[3] <= sum(_room(svg, "left")[1::2]) + 0.02
        return [float(scale) for scale in scales], _box(svg, "slide1.left.2")

    (short, long), picture = drawn()
    monkeypatch.setattr(compose, "ROW_TRIES", 0)
    (was_short, was_long), was = drawn()
    assert short == pytest.approx(long) and was_short == pytest.approx(was_long)
    assert long > was_long * 1.1 and picture[3] > was[3] * 1.1
    # To the body's foot, where the short figure's row left a band empty under the picture.
    foot = sum(_room(render_slide(deck, deck.slides[0]).svg, "left")[1::2])
    assert picture[1] + picture[3] == pytest.approx(foot, abs=0.02) and was[1] + was[3] < foot - 25.0
    # A figure that cannot be laid out for the height it would be given keeps the rows as
    # they fitted: the slide is drawn as it was, and nothing is said of it.
    monkeypatch.undo()
    plan, asked = compose._plan_figures, []

    def failing(canvas: object, blocks: list, box: object, words: object, each: float | None = None) -> tuple:
        asked.append(each)
        if len(asked) > 1:
            raise compose._FigureFailed(blocks[0], ValueError("no"), compose.FIGURE_FAILED, compose.FIGURE_FAILED)
        return plan(canvas, blocks, box, words, each=each)

    monkeypatch.setattr(compose, "_plan_figures", failing)
    kept = render_slide(deck, deck.slides[0])
    assert len(asked) == 2 and not any(compose.FIGURE_FAILED in said for said in kept.diagnostics)
    monkeypatch.setattr(compose, "_plan_figures", plan)
    monkeypatch.setattr(compose, "ROW_TRIES", 0)
    assert kept.svg == render_slide(deck, deck.slides[0]).svg


def test_a_picture_its_row_makes_narrower_than_its_place_stands_where_it_is_asked(tmp_path: Path) -> None:
    square = _picture(tmp_path, "square.png", (300, 300))

    def drawn(**asked: object) -> tuple[float, float, float]:
        """Its left edge and width, and the room its column has to spare beside it."""

        picture = {"image": square, **asked}
        svg = _drawn(_two([{"text": "**Head**"}, picture, {"image": square}],
                          [{"text": "**Head**"}, {"text": "Beside"}, {"text": "Beside"}], rows=[]), tmp_path).svg
        (x, _, wide, _), (left, _, width, _) = _room(svg, "left"), _box(svg, "slide1.left.1")
        return left - x, width, wide - width

    # In a row half the body's height a square picture is that wide, whatever width it was
    # asked to be -- none, wider than its column, narrower than it but wider than its row is
    # tall -- with room beside it.
    _, width, spare = drawn()
    assert spare > 150.0
    for asked in ({}, {"width": 900}, {"width": width + 60.0}):
        assert drawn(**asked)[1:] == pytest.approx((width, spare))
        # Under words it starts at their edge; asked, it stands at the left, the middle or
        # the right of its place, or any share of the way across.
        assert drawn(**asked)[0] == pytest.approx(0.0)
        for where, share in (("left", 0.0), ("middle", 0.5), ("right", 1.0), (0.25, 0.25), (0.8, 0.8)):
            assert drawn(**asked, horizontal=where)[0] == pytest.approx(spare * share)
    # As one narrower than its place by its own width does.
    assert drawn(width=120, horizontal="right") == pytest.approx((spare + width - 120.0, 120.0, spare + width - 120.0))
    # A mechanism too, asked to be wider than its column in a row too short for that.
    step = {"mechanism": [{"smiles": "[X-:1].[Y:2][Z:3]", "arrows": ["1 -> 2", "2-3 -> 3"]}], "width": 900}

    def mechanism(**asked: object) -> tuple[float, float]:
        svg = _drawn(_two([{**step, **asked}, step, step], [{"text": "a"}, {"text": "b"}, {"text": "c"}], rows=[]),
                     tmp_path).svg
        found = re.search(r'<g id="slide1\.left\.0" transform="translate\(([-\d.]+),[^>]*data-flexo-width="([\d.]+)"',
                          svg)
        return float(found.group(1)), float(found.group(2))

    (start, wide), (middle, _), (end, _) = mechanism(), mechanism(horizontal="middle"), mechanism(horizontal="right")
    room = _room(_drawn(_two([step], [{"text": "a"}], rows=[]), tmp_path).svg, "left")[2]
    assert wide < room - 60.0
    assert middle == pytest.approx(start, abs=0.01) and end - start == pytest.approx((room - wide) / 2.0, abs=0.01)
    assert mechanism(horizontal="left")[0] == pytest.approx(start - (room - wide) / 2.0, abs=0.01)


def test_a_picture_in_a_row_says_how_much_larger_its_row_lets_it_be(tmp_path: Path) -> None:
    def most(svg: str, identifier: str) -> float:
        found = re.search(rf'id="{re.escape(identifier)}"[^>]*data-flexo-most="([\d.]+)"', svg)
        assert found, identifier
        return float(found.group(1))

    editing = _drawn(_shaped(tmp_path), tmp_path, editing=True).svg
    assert "data-flexo-most" not in _drawn(_shaped(tmp_path), tmp_path).svg
    # The squarer picture's row has given it all it can: it can be no larger. (The wide one
    # is as wide as its column: its row would let it be a hair taller, its column no wider.)
    assert most(editing, "slide1.right.2") == pytest.approx(1.0)
    left, right = _box(editing, "slide1.left.2"), _box(editing, "slide1.right.2")
    assert most(editing, "slide1.left.2") == pytest.approx(right[3] / left[3], abs=1e-4)
    # Made smaller, it says how far it can be made larger again: as large as it was.
    small = _drawn(_shaped(tmp_path, right={"width": 200}), tmp_path, editing=True).svg
    assert _box(small, "slide1.right.2")[2:] == pytest.approx((200.0, 100.0))
    assert most(small, "slide1.right.2") == pytest.approx(right[3] / 100.0, abs=1e-4)
    # So too where another row has taken the height it left, and the rows fill the body
    # with no room to be seen: it could have half of it again, and has when it asks.
    tall = _picture(tmp_path, "tall.png", (200, 400))

    def stacked(**asked: object) -> str:
        return _drawn(_two([{"image": tall, **asked}, {"image": tall}], [{"text": "a"}, {"text": "b"}], rows=[]),
                      tmp_path, editing=True).svg

    even, uneven = stacked(), stacked(width=50)
    half = _box(even, "slide1.left.0")[3]
    assert _box(uneven, "slide1.left.0")[3] == pytest.approx(100.0) and most(even, "slide1.left.0") == 1.0
    assert _box(uneven, "slide1.left.1")[3] == pytest.approx(2.0 * half - 100.0)
    assert most(uneven, "slide1.left.1") == 1.0
    assert most(uneven, "slide1.left.0") == pytest.approx(half / 100.0, abs=1e-4)
    assert _box(stacked(width=900), "slide1.left.0")[3] == pytest.approx(half)
    # A mechanism, its words as large as words are, can be given a width that makes it as
    # tall as an equal share of the height -- and, its column wide enough, is then drawn so.
    width = re.compile(r'id="slide1\.left\.1"[^>]*data-flexo-width="([\d.]+)"')
    grown = _shaped(tmp_path, split=0.7)
    editing = _drawn(grown, tmp_path, editing=True).svg
    grown["left"][1]["width"] = 900
    larger = float(width.search(_drawn(grown, tmp_path).svg).group(1)) / float(width.search(editing).group(1))
    assert larger > 1.4 and most(editing, "slide1.left.1") == pytest.approx(larger, abs=1e-3)
    assert _room(editing, "left")[2] / float(width.search(editing).group(1)) > larger + 0.2


def _node(script: str, names: str, code: str) -> object:
    """What a few lines of JavaScript print, run in node with ``names`` of a static script."""

    path = Path(__file__).parents[1] / "src/flexo_talk/studio/static" / script
    code = f"import {{ {names} }} from {json.dumps(path.as_uri())};\n{code}"
    done = subprocess.run(["node", "--input-type=module", "-e", code], capture_output=True, text=True, check=True)
    return json.loads(done.stdout)


CARRY_JS = """
const stops = [0, 0.25, 0.5, 0.75, 1];
console.log(JSON.stringify(carried.map(({ room, home, dx, placed }) =>
  ({ room: roomAcross(room, home), drop: acrossDrop(room, home, dx, stops, placed) }))));
"""


@pytest.mark.skipif(shutil.which("node") is None, reason="needs node")
def test_a_picture_is_carried_across_the_room_it_is_drawn_with_whatever_width_it_was_asked(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    def places(svg: str) -> list[dict]:
        """Each column's picture as the page finds it: its place and itself, as drawn, and
        where the slide says it stood it."""

        found = []
        for name in ("left", "right"):
            (x, _, wide, _), (left, _, width, _) = _room(svg, name), _box(svg, f"slide1.{name}.2")
            placed = float(re.search(rf'id="slide1\.{name}\.2"[^>]*data-flexo-share="([\d.]+)"', svg).group(1))
            found.append({"room": {"left": x, "right": x + wide}, "home": {"left": left, "right": left + width},
                          "placed": placed})
        return found

    def carried(pictures: list[dict], *moves: float) -> list[dict]:
        tried = [{**picture, "dx": dx} for picture in pictures for dx in moves]
        return _node("slidedrop.js", "roomAcross, acrossDrop", f"const carried = {json.dumps(tried)};{CARRY_JS}")

    # Each picture is asked to be wider than its column. The squarer one is drawn narrower,
    # its row no taller: carried half the room it has, it would stand in the middle of its
    # column -- not where it stands, so that is written -- and carried a hair, where it is.
    wide, squarer = places(_drawn(_shaped(tmp_path), tmp_path, editing=True).svg)
    spare = squarer["room"]["right"] - squarer["home"]["right"]
    assert spare > 20.0 and squarer["placed"] == 0.0
    middle, hair, right = carried([squarer], spare / 2.0, 1.0, spare)
    assert middle["room"] is True and middle["drop"]["place"] == 0.5 != squarer["placed"]
    assert middle["drop"]["left"] == pytest.approx(squarer["home"]["left"] + spare / 2.0)
    assert hair["drop"] == {"place": 0, "left": squarer["home"]["left"]} and right["drop"]["place"] == 1
    # The wide one is as wide as its column: it has nowhere across it to go.
    assert carried([wide], 10.0) == [{"room": False, "drop": None}]
    # In half the height each, as the rows were, both are a good deal narrower than they
    # were asked to be, and than their columns: each goes to the middle of its own.
    monkeypatch.setattr(compose, "ROW_TRIES", 0)
    for picture in places(_drawn(_shaped(tmp_path), tmp_path, editing=True).svg):
        spare = picture["room"]["right"] - picture["home"]["right"]
        assert spare > 70.0 and picture["placed"] == 0.0
        assert carried([picture], spare / 2.0 - 5.0)[0]["drop"]["place"] == 0.5


@pytest.mark.skipif(shutil.which("node") is None, reason="needs node")
def test_the_editor_knows_where_the_slide_stands_a_picture_it_is_not_asked_about(tmp_path: Path) -> None:
    picture = _picture(tmp_path)
    blocks = {
        "text": {"text": "Words"}, "bullets": {"bullets": ["One", "Two"]}, "code": {"code": "x = 1"},
        "image": {"image": picture}, "table": {"table": [["A", "B"], ["1", "2"]]},
        "figure": {"figure": {"figure": {"id": "f"}, "nodes": [{"id": "a", "label": "A"}]}},
        "mechanism": {"mechanism": [{"smiles": "[X-:1].[Y:2][Z:3]", "arrows": ["1 -> 2", "2-3 -> 3"]}]},
    }
    columns = [
        ["text", "mechanism", "image"], ["mechanism"], ["text", "mechanism"], ["mechanism", "text"], ["image"],
        ["text", "image"], ["image", "text"], ["bullets", "image"], ["text", "figure", "table"],
        ["text", "code", "image"], ["image", "mechanism"], ["text", "text", "figure"],
    ]
    said = _node("slidedrop.js", "autoAcross",
                 f"console.log(JSON.stringify({json.dumps(columns)}.map((kinds) => autoAcross(kinds))));")
    for kinds, share in zip(columns, said, strict=True):
        # Where the slide stands each figure, mechanism, picture or table of the column, asked
        # nothing: said on its drawing while it is edited, a column of its own or one in rows.
        for settings in ({}, {"rows": []}):
            svg = _drawn(_two([blocks[kind] for kind in kinds], [{"text": "Beside"}], **settings), tmp_path,
                         editing=True).svg
            stood = [
                float(re.search(rf'id="slide1\.left\.{index}"[^>]*data-flexo-share="([\d.]+)"', svg).group(1))
                for index, kind in enumerate(kinds) if kind in {"figure", "mechanism", "image", "table"}
            ]
            assert stood and all(each == share for each in stood), kinds
    # A heading, a mechanism and a picture: under the words, at their edge -- not centred,
    # where a picture carried to the middle would be taken to stand already.
    assert said[0] == 0 and said[1] == 0.5


def test_the_editor_sizes_and_carries_a_picture_by_what_the_slide_draws() -> None:
    editor = (Path(__file__).parents[1] / "src/flexo_talk/studio/static/editor.js").read_text("utf-8")
    # Carried across its place: the room it has is its drawing's, the stop it stands at is
    # the slide's own word for it, and it is let go at the stop nearest.
    assert "const room = own && roomAcross(own, home) ? own : null;" in editor
    assert "placedAt(slideAt(), carry.from.region, block, carry.from)" in editor
    assert "acrossDrop(carry.room, carry.home, dx, ACROSS.map(([share]) => share), carry.placed)" in editor
    assert 'PLACEABLE.has(kind)) ? attribute(blockElement(at.region, at.index), "data-flexo-share")' in editor
    assert 'const PLACEABLE = new Set(["figure", "mechanism", "image", "table"]);' in editor
    # Sized by a corner on a slide set in rows: no larger than its row would let it be.
    assert 'const taller = attribute(element, "data-flexo-most") ?? (region.room.bottom - region.room.top' in editor
    assert "Math.min(across / box.width, taller)" in editor


def test_builds_footnotes_notes_and_shrinking_work_in_rows(tmp_path: Path) -> None:
    picture = {"image": _picture(tmp_path)}
    slide = _two(
        [{"text": "First", "build": True}, picture],
        [{"bullets": ["a", "b"], "reveal": True}, {**picture, "build": True}],
        rows=["Words", "Pictures"], footnotes=["A footnote"], notes="Say it.",
    )
    deck = _deck(slide, tmp_path)
    rendered = render_slide(deck, deck.slides[0])
    # Each builds in on a click of its own, in the slide's order: the left column, then the right.
    assert [name for _, name in rendered.builds] == ["slide1.left.0", "slide1.right.1"]
    assert rendered.steps == 5 and deck.slides[0].notes_text == "Say it."
    # The rows end above the footnote.
    low = _part(rendered.svg, "slide1.left.1")
    assert float(low.get("y")) + float(low.get("height")) < _y(rendered.svg, "slide1.footnote0") - 10.0
    assert _y(rendered.svg, "slide1.left.1") == pytest.approx(_y(rendered.svg, "slide1.right.1"))
    # Too many words for the slide are set smaller together, the labels with them, and said.
    many = [{"bullets": [f"Point {index}" for index in range(6)]}, {"text": "More"}]
    full = _drawn(_two(many, [{"text": "Beside"}, {"bullets": [f"Item {index}" for index in range(6)]}],
                       rows=["One", "Two"]))
    assert full.diagnostics == ["slide1: Text size reduced to 79% to fit the slide."]
    assert float(_part(full.svg, "slide1.row0").get("font-size")) < deck.style.body_size * ROW_LABEL - 1.0
    sizes = {_part(full.svg, f"slide1.{name}").get("font-size") for name in ("left.1", "right.0")}
    assert len(sizes) == 1 and float(sizes.pop()) < deck.style.body_size


def test_an_empty_label_shows_its_placeholder_in_the_studio_only(tmp_path: Path) -> None:
    picture = {"image": _picture(tmp_path)}
    slide = _two([{"text": "Head"}, picture], [{"text": ""}, {"figure": {"figure": {"id": "f"}, "nodes": [
        {"id": "shape", "label": ""}]}}], rows=["", ""])
    exported = _drawn(slide, tmp_path).svg
    editing = _drawn(slide, tmp_path, editing=True).svg
    held = re.findall(r'id="(slide1\.row\d+)"[^>]*data-flexo-placeholder="([^"]+)"', editing)
    assert held == [("slide1.row0", "Label"), ("slide1.row1", "Label")]
    assert "slide1.row" not in exported and "Label" not in exported
    # They hold a gutter while edited (said, for the studio to present the slide as it is
    # exported); with none written there is none, and the columns start at the margin.
    assert 'data-flexo-gutter="held"' in editing and "data-flexo-gutter" not in exported
    assert _room(editing, "left")[0] > 48.0 + 40.0 and _room(exported, "left")[0] == pytest.approx(48.0)
    # One written beside one not: its gutter is the written one's, held by no placeholder.
    one = _drawn({**slide, "rows": ["", "Mechanism"]}, tmp_path, editing=True).svg
    assert "data-flexo-gutter" not in one and _room(one, "left") == _room(
        _drawn({**slide, "rows": ["", "Mechanism"]}, tmp_path).svg, "left")
    # A row a hand-written list has no label for has its placeholder too, to be written.
    short = _drawn({**slide, "rows": []}, tmp_path, editing=True).svg
    assert len(re.findall(r'id="slide1\.row\d+"[^>]*data-flexo-placeholder="Label"', short)) == 2
    # While edited every block keeps its cell (empty words, a new figure's sample), and says
    # the room its row has to spare for it; exported, neither is there.
    assert _part(editing, "slide1.right.0") is not None and _part(exported, "slide1.right.0") is None
    assert _y(editing, "slide1.left.1") == pytest.approx(float(_room(editing, "left")[1]) + 24.0 + 18.0, abs=1.0)
    assert re.search(r'id="slide1\.left\.0"[^>]*data-flexo-spare="0"', editing)
    assert "data-flexo-spare" not in re.search(r'<image[^>]*id="slide1\.left\.1"[^>]*>', exported).group(0)


def test_a_slide_set_in_rows_is_written_back_as_it_was_read() -> None:
    document = {"schema_version": 1, "deck": {"id": "rows"}, "slides": [
        {"layout": "two-columns", "title": "Two", "rows": ["", "Mechanism", "Geometry", "Past the last"],
         "left": [{"text": "a"}, {"text": "b"}], "right": [{"text": "c"}]},
        {"layout": "columns", "title": "Three", "widths": [2, 1, 1], "rows": [],
         "columns": [[{"text": "a"}], [{"text": "b"}], [{"text": "c"}]]},
        {"layout": "two-columns", "title": "Stacked", "left": [{"text": "a"}], "right": [{"text": "b"}]},
    ]}
    deck = deck_from_document(document)
    assert [len(slide.rows) if slide.rows is not None else None for slide in deck.slides] == [4, 0, None]
    assert deck.slides[0].rows[0] == () and deck.slides[0].rows[1][0].text == "Mechanism"
    assert deck_document(deck) == document
    # In the file, with its other settings, before its columns.
    text = dump_document(document)
    assert text.index("rows:") < text.index("left:") and yaml.safe_load(text) == document


def test_python_sets_a_slide_in_rows_where_it_gives_its_split_or_widths() -> None:
    deck = Deck("rows")
    with deck.slide("Two", layout="two-columns", split=0.4, rows=["", "Mechanism"]) as slide:
        slide.left.text("a")
        slide.right.text("b")
    with deck.slide("Three", layout="columns", widths=(2, 1), rows=("One", 2)) as slide:
        slide.columns[0].text("a")
        slide.columns[1].text("b").text("c")
    written = deck_document(deck)
    assert written["slides"][0]["rows"] == ["", "Mechanism"] and written["slides"][1]["rows"] == ["One", "2"]
    again = deck_from_document(written)
    for first, second in zip(deck.slides, again.slides, strict=True):
        assert render_slide(deck, first).svg == render_slide(again, second).svg
    assert "slide2.row1" in render_slide(deck, deck.slides[1]).svg
    # Only a slide of columns is set in rows; and its labels are a list of words.
    with pytest.raises(ValueError, match="Only a two-columns or columns slide"):
        deck.slide("One body", rows=["a"])
    with pytest.raises(ValueError, match="list of labels"):
        deck.slide("Two", layout="two-columns", rows="Mechanism")


@pytest.mark.parametrize(
    ("slide", "where", "words"),
    [
        ({"layout": "two-columns", "rows": "Mechanism"}, "slides[0].rows", "list of labels"),
        ({"layout": "columns", "columns": [[]], "rows": [["a"]]}, "slides[0].rows", "a label must be text"),
        ({"layout": "two-columns", "rows": [True]}, "slides[0].rows", "a label must be text"),
        ({"title": "T", "rows": ["a"]}, "slides[0]", "Unknown key rows"),
    ],
)
def test_wrong_rows_say_where(slide: dict, where: str, words: str) -> None:
    with pytest.raises(DeckDocumentError) as caught:
        deck_from_document({"deck": {}, "slides": [slide]})
    assert caught.value.where == where and words in caught.value.message


def test_a_label_is_a_text_box_in_the_powerpoint_where_it_is_drawn(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Emu

    from flexo_talk.export import write_pptx

    picture = {"image": _picture(tmp_path)}
    slide = _two([{"text": "**One**"}, picture], [{"text": "**Other**"}, {"bullets": ["a", "b"]}],
                 rows=["", "Mechanism"])
    deck = _deck(slide, tmp_path)
    rendered = deck.render()
    svg = rendered[0].svg
    shapes = list(Presentation(str(write_pptx(deck, rendered, tmp_path / "rows.pptx"))).slides[0].shapes)
    names = [shape.name.split(" “")[0] for shape in shapes]
    # Read after the title, before the columns: the left top to bottom, then the right.
    assert names == ["Title", "Rule", "Row Label", "Text", "Picture", "Text", "List", "Slide Number"]
    label = shapes[2]
    assert label.name == "Row Label “Mechanism”" and label.has_text_frame
    assert label.text_frame.text == "Mechanism"
    # (A semibold is a family of its own to a slide program: "Figtree SemiBold".)
    run = label.text_frame.paragraphs[0].runs[0]
    assert run.font.size.pt == pytest.approx(deck.style.body_size * ROW_LABEL) and "SemiBold" in run.font.name
    # Where the slide draws it: at the margin, its words' middle its row's.
    drawn = _part(svg, "slide1.row1")
    assert Emu(label.left).pt == pytest.approx(float(drawn.get("x")), abs=0.5)
    row = _part(svg, "slide1.left.1")
    middle = float(row.get("y")) + float(row.get("height")) / 2.0
    assert Emu(label.top).pt + Emu(label.height).pt / 2.0 == pytest.approx(middle, abs=1.5)
    # And every block where the slide draws it: the picture in its row, the list beside it.
    picture_shape, listed = shapes[4], shapes[6]
    assert Emu(picture_shape.top).pt == pytest.approx(float(row.get("y")), abs=0.5)
    assert Emu(picture_shape.left).pt == pytest.approx(float(row.get("x")), abs=0.5)
    assert Emu(listed.top).pt == pytest.approx(rendered[0].lists[0].y, abs=6.0)
    assert rendered[0].lists[0].y == pytest.approx(float(row.get("y")), abs=0.01)


def test_the_studio_draws_names_and_presents_a_slide_set_in_rows(tmp_path: Path) -> None:
    kind = DeckKind()
    catalog = kind.catalog()
    assert "rows" in catalog["slide_keys"]["two-columns"] and "rows" in catalog["slide_keys"]["columns"]
    assert "rows" not in catalog["slide_keys"]["content"]
    assert "rows" in kind.guide() and "Mechanism" in kind.guide()
    picture = {"image": _picture(tmp_path)}
    stacked = {"schema_version": 1, "deck": {"id": "rows"}, "slides": [
        _two([{"text": "One"}, picture], [{"text": LONG}, picture])]}
    switched = {**stacked, "slides": [{**stacked["slides"][0], "rows": ["", ""]}]}
    labelled = {**stacked, "slides": [{**stacked["slides"][0], "rows": ["", "Geometry"]}]}
    assert not kind.check(switched, tmp_path)
    assert kind.check({**stacked, "slides": [{**stacked["slides"][0], "rows": "no"}]}, tmp_path)
    # Row Labels switched on: the labels' placeholders hold a gutter on the stage, and the
    # slide is drawn beside it as it is presented and exported, with none.
    page = kind.draw(switched, tmp_path, {"settle": True}).pages[0].svg
    editing, presented = page.split('id="presented-slide1.content"')
    assert editing.count('data-flexo-placeholder="Label"') == 2 and "data-flexo-presented" in presented
    assert "presented-slide1.row" not in presented and "Label" not in presented
    assert _room(editing, "left")[0] > 48.0
    assert re.search(r'id="presented-slide1\.left"[^>]*data-flexo-box="48 ', presented)
    # A label written, none held: one drawing, the label on it.
    page = kind.draw(labelled, tmp_path, {"settle": True}).pages[0].svg
    assert "data-flexo-presented" not in page and ">Geometry<" in page
    # What the change did is said by the setting's name.
    said = kind.describe(stacked, switched) + kind.describe(switched, labelled)
    assert [note["text"] for note in said] == ["edited the row labels on slide 1"] * 2


ROWS_JS = """
const block = (words) => ({ text: words });
const two = () => ({ layout: "two-columns", left: [block("a"), block("b"), block("c")], right: [block("d")] });
const three = { layout: "columns", columns: [[block("a")], [block("b"), block("c")], []] };
const out = {};
out.can = [canRow(two()), canRow(three), canRow({ body: [] }), canRow({ layout: "figure" })];
out.count = [rowCount(two()), rowCount(three), rowCount({ layout: "two-columns" }), rowCount({ body: [block("a")] })];
// Row Labels on: an empty label a row; off: gone.
const on = two();
setRows(on, true);
out.on = structuredClone(on.rows);
out.rowed = [rowed(on), rowed(two())];
setRows(on, true);
out.again = on.rows;
const off = { ...two(), rows: ["x", "", "z"] };
setRows(off, false);
out.off = "rows" in off;
const body = { body: [block("a")] };
setRows(body, true);
out.body = "rows" in body;
// Switched on over labels a hand wrote: those kept, the rows without one given one.
const kept = { ...two(), rows: ["Head"] };
setRows(kept, true);
out.kept = kept.rows;
// A label written: those before it it had none of are empty.
const few = { ...two(), rows: [] };
setLabel(few, 2, "Third");
out.few = few.rows;
const plain = two();
setLabel(plain, 0, "No");
out.plain = "rows" in plain;
out.labels = labelsOf({ ...two(), rows: [null, 7, "x"] });
// After a change: the labels in step with the rows.
const step = (slide, change) => {
  const was = structuredClone(slide);
  change(slide);
  keepRows(was, slide);
  return slide.rows;
};
const full = () => ({ ...two(), rows: ["A", "B", "C"] });
out.added = step(full(), (s) => s.left.push(block("e")));
out.beside = step(full(), (s) => s.right.push(block("e")));
out.deleted = step({ ...two(), rows: ["A", "B", ""] }, (s) => s.left.pop());
out.written = step(full(), (s) => s.left.pop());
out.both = step({ ...two(), rows: ["A", "", ""] }, (s) => s.left.splice(1, 2));
out.middle = step({ ...two(), rows: ["A", "", "C"] }, (s) => s.left.splice(1, 2));
out.unrowed = step(two(), (s) => s.left.push(block("e")));
out.switched = step(two(), (s) => { s.left.push(block("e")); s.rows = ["only"]; });
console.log(JSON.stringify(out));
"""


@pytest.mark.skipif(shutil.which("node") is None, reason="needs node")
def test_the_editor_keeps_a_slides_labels_in_step_with_its_rows() -> None:
    script = Path(__file__).parents[1] / "src/flexo_talk/studio/static/rows.js"
    names = "canRow, rowed, rowCount, labelsOf, setRows, setLabel, keepRows"
    code = f"import {{ {names} }} from {json.dumps(script.as_uri())};\n{ROWS_JS}"
    done = subprocess.run(["node", "--input-type=module", "-e", code], capture_output=True, text=True, check=True)
    found = json.loads(done.stdout)
    assert found["can"] == [True, True, False, False] and found["count"] == [3, 2, 0, 0]
    # Row Labels switched on writes an empty label for each row the slide has (once, however
    # often); off, none -- and only a slide of columns takes them.
    assert found["on"] == ["", "", ""] and found["again"] == ["", "", ""] and found["rowed"] == [True, False]
    assert found["off"] is False and found["body"] is False
    assert found["kept"] == ["Head", "", ""]
    assert found["few"] == ["", "", "Third"] and found["plain"] is False
    assert found["labels"] == ["", "7", "x"]
    # An object added that makes a new row adds an empty label for it; one beside others, none.
    assert found["added"] == ["A", "B", "C", ""] and found["beside"] == ["A", "B", "C"]
    # The last row's last object deleted: its label goes with it if it is empty (those of
    # every row gone), and one written stays -- with any before it, empty or not.
    assert found["deleted"] == ["A", "B"] and found["written"] == ["A", "B", "C"]
    assert found["both"] == ["A"] and found["middle"] == ["A", "", "C"]
    # A slide not set in rows, or switched on by the change itself, is left as it is.
    assert "unrowed" not in found and found["switched"] == ["only"]


def test_the_editor_offers_row_labels_by_name_and_types_them_on_the_slide() -> None:
    editor = (Path(__file__).parents[1] / "src/flexo_talk/studio/static/editor.js").read_text("utf-8")
    # A label is one of the slide's lines of words: chosen, typed over and typed in place as
    # they are, by the id the slide draws it under.
    assert r"footnote\d+|row\d+)$/" in editor and 'ROW.test(key) ? "Row Label"' in editor
    assert 'ROW.test(key) ? "Label"' in editor
    # Its switch in the Format panel, its command for the palette, and their one step each.
    assert 'ui.toggle({ value: rowed(slide), label: "Row Labels"' in editor
    assert 'label: rowed(slideAt()) ? "Hide Row Labels" : "Show Row Labels", also: ["Row Labels"]' in editor
    assert '{ label: on ? "Show Row Labels" : "Hide Row Labels" }' in editor
    # And every change made on the page keeps the labels in step with the rows.
    assert "keepRows(was, slide)" in editor
