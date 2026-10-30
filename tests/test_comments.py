"""Comments on a deck's slides, the objects on them and the parts of their figures: kept in its
YAML, never drawn, merged as any edit is, read and answered by an agent's tools, said in the
activity, and carried into its PowerPoint as PowerPoint's own comments
(flexo_talk.studio.comments, flexo.studio.commenting, flexo_talk.pptx_comments)."""

from __future__ import annotations

import copy
import json
import shutil
import subprocess
import zipfile
from pathlib import Path

import pytest
from flexo.studio.agent import Tools
from flexo.studio.merge import merge3
from flexo.studio.workspace import Workspace
from pptx import Presentation

from flexo_talk.document import deck_from_document, load_document, save_document
from flexo_talk.export import build_deck
from flexo_talk.studio import DeckKind, _slide_keys
from flexo_talk.studio.comments import note, uncommented

AGENT = {"id": "mcp:claude-code", "name": "Claude Code"}
PERSON = {"id": "page-a", "name": "Ada", "kind": "person"}


def _deck() -> dict:
    return {
        "schema_version": 1,
        "deck": {"id": "review", "theme": "paper"},
        "slides": [
            {"layout": "title", "title": "Review", "subtitle": "Comments"},
            {
                "title": "The model",
                "body": [
                    {"text": "Write what the model *is*"},
                    {
                        "figure": {
                            "figure": {"id": "model"},
                            "nodes": [{"id": "encoder", "label": "Encoder"}, {"id": "decoder", "label": "Decoder"}],
                            "edges": [{"from": "encoder", "to": "decoder"}],
                        },
                        "comments": [{"id": "c000001", "author": "Ana Ruiz", "time": "2026-10-29T09:12:00Z",
                                      "text": "What does it take in?", "part": "encoder"}],
                    },
                ],
                "comments": [{"id": "c000002", "author": "Ana Ruiz", "time": "2026-10-28T16:40:00Z",
                              "text": "A takeaway in the title, please.",
                              "replies": [{"author": "Ada", "time": "2026-10-28T17:02:00Z", "text": "Will do."}]}],
            },
            {"title": "Skipped", "skip": True, "comments": [{"id": "c000003", "author": "Bo", "text": "Not sent"}]},
            {
                "title": "Results",
                "layout": "two-columns",
                "left": [{"bullets": ["First", "Second"]}],
                "right": [{"text": "More", "comments": [
                    {"id": "c000004", "author": "Bo", "text": "Done with", "resolved": True},
                    {"id": "c000005", "author": "Claude Code", "text": "Is “More” the word?"},
                ]}],
            },
        ],
    }


@pytest.fixture(autouse=True)
def _apart(tmp_path_factory: pytest.TempPathFactory, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("XDG_RUNTIME_DIR", str(tmp_path_factory.mktemp("run")))
    monkeypatch.setenv("FLEXO_ACTS", "0")


@pytest.fixture
def workspace(tmp_path: Path):
    save_document(_deck(), tmp_path / "review.yaml")
    opened = Workspace(tmp_path)
    try:
        yield opened
    finally:
        opened.close()


def test_a_deck_with_comments_reads_and_draws_none_of_them(tmp_path: Path) -> None:
    document = _deck()
    kind = DeckKind()
    assert kind.check(document, tmp_path) == []
    deck = deck_from_document(document, tmp_path)
    svgs = "".join(item.svg for item in deck.render())
    assert "takeaway" not in svgs and "take in" not in svgs and "word?" not in svgs
    # A comment added, answered or resolved draws nothing again: the slide is kept as it was drawn.
    keys = _slide_keys(document, tmp_path, set())
    before = [keys(index) for index in range(len(document["slides"]))]
    changed = copy.deepcopy(document)
    changed["slides"][1]["comments"][0]["resolved"] = True
    changed["slides"][1]["body"][0]["comments"] = [{"id": "c9", "author": "Ada", "text": "New"}]
    keys = _slide_keys(changed, tmp_path, set())
    assert [keys(index) for index in range(len(changed["slides"]))] == before
    # But the slide's words changed are another drawing, as ever.
    changed["slides"][1]["title"] = "The model, plainly"
    keys = _slide_keys(changed, tmp_path, set())
    assert keys(1) != before[1]


def test_a_comment_written_wrongly_is_said_and_draws_the_slide_all_the_same(tmp_path: Path) -> None:
    document = _deck()
    document["slides"][1]["comments"].append({"words": "no text"})
    said = DeckKind().check(document, tmp_path)
    assert said == ["slides[1].comments[1]: a comment is a mapping with its words in text"]
    deck_from_document(document, tmp_path)


def test_uncommented_takes_comments_out_and_leaves_the_rest_the_very_same() -> None:
    slide = _deck()["slides"][1]
    plain = uncommented(slide)
    assert "comments" not in plain and "comments" not in plain["body"][1]
    assert plain["body"][0] is slide["body"][0]
    assert uncommented(_deck()["slides"][0]) == _deck()["slides"][0]


def test_the_activity_says_what_a_change_to_comments_did() -> None:
    old = _deck()["slides"][1]
    new = copy.deepcopy(old)
    new["comments"][0]["replies"].append({"author": "Bo", "text": "Agreed"})
    assert note(old, new, 1) == {"text": "replied to a comment on slide 2", "where": {"page": 2, "label": "Slide 2"}}
    new = copy.deepcopy(old)
    new["body"][1]["comments"][0]["resolved"] = True
    said = note(old, new, 1)
    assert said["text"] == "resolved a comment on the figure on slide 2"
    assert said["where"]["object"] == {"region": "body", "index": 1}
    new = copy.deepcopy(old)
    new["body"][0]["comments"] = [{"id": "c7", "author": "Ada", "text": "Hm"}]
    assert note(old, new, 1)["text"] == "commented on the text on slide 2"
    new["title"] = "Other"
    assert note(old, new, 1) is None
    # As the deck's activity says it.
    after = _deck()
    after["slides"][1]["comments"].append({"id": "c8", "author": "Ada", "text": "One more"})
    assert [item["text"] for item in DeckKind().describe(_deck(), after)] == ["commented on slide 2"]


def test_an_agent_works_through_the_comments(workspace: Workspace) -> None:
    tools = Tools(workspace, AGENT)
    listed, failed = tools.call("list_comments", {"file": "review.yaml"})
    assert not failed
    text = listed[0]["text"]
    # Open ones, where each is, who wrote it and its replies; resolved ones not shown.
    assert "4 comment threads (1 resolved not shown)" in text
    assert "c000001 · Slide 2 · Figure, part “encoder” · open" in text
    assert "Ada (28 Oct 2026 17:02 UTC): Will do." in text
    assert "Done with" not in text
    assert "Done with" in tools.call("list_comments", {"file": "review.yaml", "all": True})[0][0]["text"]
    answer, failed = tools.call("reply_to_comment", {"file": "review.yaml", "id": "c000002", "text": "Retitled it.",
                                                     "resolve": True, "label": "Work through comments"})
    assert not failed and "Replied and resolved" in answer[0]["text"]
    answer, failed = tools.call("add_comment", {
        "file": "review.yaml", "text": "Which data set is this?", "slide": 2, "object": "body[1]", "part": "decoder",
        "label": "Work through comments"})
    assert not failed
    _, failed = tools.call("resolve_comment", {"file": "review.yaml", "id": "c000005"})
    assert not failed
    # Wrong: said, and nothing changed.
    for wrong, why in [
        ({"slide": 9}, "There is no slide 9"),
        ({"slide": 2, "object": "body[5]"}, "has no object body[5]"),
        ({"slide": 2, "object": "body[1]", "part": "nope"}, "has no part 'nope'"),
        ({"slide": 2, "object": "body[0]", "part": "x"}, "no figure written in the deck"),
        ({"slide": 2, "object": "middle"}, "names no object"),
    ]:
        answer, failed = tools.call("add_comment", {"file": "review.yaml", "text": "x", **wrong})
        assert failed and why in answer[0]["text"], answer
    _, failed = tools.call("reply_to_comment", {"file": "review.yaml", "id": "c404", "text": "x"})
    assert failed
    doc = workspace.open("review.yaml")
    slide = doc.document["slides"][1]
    assert slide["comments"][0]["resolved"] is True
    assert slide["comments"][0]["replies"][-1]["author"] == "Claude Code"
    added = slide["body"][1]["comments"][-1]
    assert added["part"] == "decoder" and added["author"] == "Claude Code" and added["time"].endswith("Z")
    assert doc.document["slides"][3]["right"][0]["comments"][1]["resolved"] is True
    # One turn, said as what it did.
    (row,) = [entry for entry in workspace.activity if "turn" in entry]
    assert row["text"] == "Work through comments"
    assert row["turn"]["done"] == ["resolved a comment on slide 2", "commented on the figure on slide 2",
                                   "resolved a comment on the text on slide 4"]
    # Written to its file as YAML, the writer's words kept round them.
    workspace.flush()
    written = load_document(workspace.root / "review.yaml")
    assert written["slides"][1]["body"][1]["comments"][-1]["text"] == "Which data set is this?"
    assert "time: '20" in (workspace.root / "review.yaml").read_text(encoding="utf-8")


def test_two_people_answering_one_thread_at_once_keep_both_replies() -> None:
    base = _deck()
    ours, theirs = copy.deepcopy(base), copy.deepcopy(base)
    ours["slides"][1]["comments"][0]["replies"].append({"author": "Ada", "text": "Mine"})
    theirs["slides"][1]["comments"][0]["replies"].append({"author": "Bo", "text": "Theirs"})
    theirs["slides"][1]["comments"][0]["resolved"] = True
    merged = merge3(base, ours, theirs)
    thread = merged["slides"][1]["comments"][0]
    assert [reply["text"] for reply in thread["replies"]] == ["Will do.", "Mine", "Theirs"]
    assert thread["resolved"] is True


def test_a_powerpoint_carries_the_open_comments_as_its_own(tmp_path: Path) -> None:
    deck = deck_from_document(_deck(), tmp_path)
    # (A slide skipped is neither presented nor exported, as the studio's export leaves it out.)
    deck.slides = [slide for slide in deck.slides if slide.source.get("skip") is not True]
    built = build_deck(deck, tmp_path / "build", ("pptx",))
    with zipfile.ZipFile(built.pptx) as package:
        names = set(package.namelist())
        authors = package.read("ppt/commentAuthors.xml").decode()
        second = package.read("ppt/comments/comment2.xml").decode()
        third = package.read("ppt/comments/comment3.xml").decode()
        rels = package.read("ppt/slides/_rels/slide2.xml.rels").decode()
        types = package.read("[Content_Types].xml").decode()
    assert "ppt/comments/comment1.xml" not in names
    assert 'name="Ana Ruiz" initials="AR" lastIdx="2"' in authors and 'name="Ada"' in authors
    # The slide's thread, its reply threaded under it (PowerPoint 2013's), the part's named.
    assert "A takeaway in the title, please." in second and "Will do." in second
    assert '<p15:parentCm authorId="0" idx="1"/>' in second
    assert "On “Encoder”: What does it take in?" in second
    assert 'dt="2026-10-29T09:12:00.000"' in second
    # Resolved: left out. The skipped slide's: not there.
    assert "Done with" not in third and "Is “More” the word?" in third
    assert "Not sent" not in "".join([second, third])
    assert "relationships/comments" in rels and "comments+xml" in types and "commentAuthors+xml" in types
    # Placed at the top right of what it is about: the figure, in eighths of a point.
    positions = [chunk.split('"')[1:4:2] for chunk in second.split("<p:pos ")[1:]]
    slide_spot, figure_spot = positions[0], positions[-1]
    assert int(slide_spot[0]) > 7000 and int(slide_spot[1]) < 200
    assert int(figure_spot[0]) < 2000 and int(figure_spot[1]) > 1000
    # And the package reads as a presentation.
    assert len(Presentation(str(built.pptx)).slides) == 3


def test_a_deck_with_no_open_comments_is_written_as_ever(tmp_path: Path) -> None:
    document = _deck()
    for slide in document["slides"]:
        slide.pop("comments", None)
        for key in ("body", "left", "right"):
            for block in slide.get(key) or []:
                block.pop("comments", None)
    built = build_deck(deck_from_document(document, tmp_path), tmp_path / "build", ("pptx",))
    with zipfile.ZipFile(built.pptx) as package:
        assert not [name for name in package.namelist() if "comment" in name.lower()]


# -- in the editor (static/comments.js), run in Node ---------------------------------------

STUB = """
import { register } from "node:module";
const studio = "data:text/javascript," + encodeURIComponent(
  "const none = () => ({}); export const h = none, clear = none, icon = none, popover = none, closeMenu = none,"
  + " toast = none, avatar = none; export const ui = {};");
register("data:text/javascript," + encodeURIComponent(`
export async function resolve(specifier, context, next) {
  if (specifier === "/static/studio/studio.js") return { url: ${JSON.stringify(studio)}, shortCircuit: true };
  return next(specifier, context);
}`));
"""


def test_the_editor_finds_each_thread_and_counts_the_open_ones_on_each_slide() -> None:
    node = shutil.which("node")
    if node is None:
        pytest.skip("Node is not installed")
    script = Path(__file__).parents[1] / "src/flexo_talk/studio/static/comments.js"
    code = STUB + (
        f"const C = await import({json.dumps(script.as_uri())});\n"
        f"const deck = {json.dumps(_deck())};\n"
        "const threads = C.deckThreads(deck)\n"
        "  .map((one) => [one.slide, one.region, one.index, one.part, one.id, one.open]);\n"
        "const now = new Date(Date.now() - 5 * 60000).toISOString();\n"
        "const when = [C.when(now), C.when('nonsense')];\n"
        "console.log(JSON.stringify({ threads, counts: C.openCounts(deck), when }));\n"
    )
    done = subprocess.run([node, "--input-type=module", "-e", code], capture_output=True, text=True, check=True)
    said = json.loads(done.stdout.strip().splitlines()[-1])
    assert said["threads"] == [
        [1, None, None, None, "c000002", True],
        [1, "body", 1, "encoder", "c000001", True],
        [2, None, None, None, "c000003", True],
        [3, "right", 0, None, "c000004", False],
        [3, "right", 0, None, "c000005", True],
    ]
    assert said["counts"] == [0, 2, 1, 1]
    assert said["when"] == ["5 min ago", ""]
