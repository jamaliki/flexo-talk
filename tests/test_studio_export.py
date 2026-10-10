"""The studio's exports of a deck, as Keynote's Export To makes them: a PDF with a
page per slide (each stage of its builds only when asked), and an image of each slide
in a folder of the deck's name and theirs, made aside for the page to hand over, a PNG
as wide as chosen; the CLI's build/ is as it was."""

from __future__ import annotations

import re
import zlib
from pathlib import Path

import pytest

from flexo_talk.studio import DeckKind

DECK = {
    "schema_version": 1,
    "deck": {"id": "My Talk", "theme": "paper"},
    "slides": [
        {"layout": "title", "title": "A talk"},
        {"title": "Three things", "body": [{"bullets": ["One", "Two", "Three"], "reveal": True}]},
    ],
}


def _pages(pdf: Path) -> int:
    return pdf.read_bytes().count(b"/Type /Page ")


def _streams(data: bytes) -> list[bytes]:
    """A PDF's content streams, as drawn: each inflated."""

    found = []
    for stream in re.findall(rb"stream\r?\n(.*?)\r?\nendstream", data, re.DOTALL):
        try:
            found.append(zlib.decompress(stream))
        except zlib.error:
            found.append(stream)
    return found


def test_the_pdf_has_a_page_per_slide_showing_it_whole(tmp_path: Path) -> None:
    written = DeckKind().export(DECK, tmp_path, "talk", ["pdf"], into=tmp_path / "out")
    assert [path.name for path in written] == ["My Talk.pdf"]
    assert _pages(written[0]) == 2


def test_each_stage_of_builds_is_a_page_when_asked(tmp_path: Path) -> None:
    written = DeckKind().export(DECK, tmp_path, "talk", ["pdf"], into=tmp_path / "out", steps=True)
    # The title slide, then the list's slide bare and with each of its three items.
    assert _pages(written[0]) == 1 + 4


def test_images_made_aside_are_a_folder_of_the_deck_s_name(tmp_path: Path) -> None:
    written = DeckKind().export(DECK, tmp_path, "talk", ["png", "svg"], into=tmp_path / "out")
    # (Named for what they are, too: a folder of PNGs is told apart from one of SVGs.)
    folder = "My Talk \u2013 PNG 1920 px \u2013 SVG"
    assert sorted(path.relative_to(tmp_path / "out").as_posix() for path in written) == [
        f"{folder}/slide-01.png", f"{folder}/slide-01.svg", f"{folder}/slide-02.png", f"{folder}/slide-02.svg",
    ]
    assert not (tmp_path / "build").exists()


def test_build_is_written_beside_the_deck_as_before(tmp_path: Path) -> None:
    written = DeckKind().export(DECK, tmp_path, "talk", ["pdf", "png"])
    assert sorted(path.relative_to(tmp_path).as_posix() for path in written) == [
        "build/My Talk-01.png", "build/My Talk-02.png", "build/My Talk.pdf",
    ]


def test_a_figure_that_cannot_be_drawn_is_an_empty_box_and_the_export_says_so(tmp_path: Path) -> None:
    protein = {"id": "p", "kind": "protein", "label": "Spike", "properties": {}}
    deck = {**DECK, "slides": [
        *DECK["slides"], {"title": "The spike", "body": [{"figure": {"figure": {"id": "f"}, "nodes": [protein]}}]},
    ]}
    kind = DeckKind()
    (pdf,) = kind.export(deck, tmp_path, "talk", ["pdf"], into=tmp_path / "out")
    assert _pages(pdf) == 3
    # (Its other shapes drawn, the one it can't draw a plain box of its words.)
    assert kind.export_notes == [
        "Slide 3 · Figure: \u201cSpike\u201d is drawn as a plain box: a protein needs its length in residues."
    ]
    # Its slide as exported (and presented) shows no words of what is wrong.
    svgs = kind.export(deck, tmp_path, "talk", ["svg"], into=tmp_path / "svg")
    last = sorted(path for path in svgs if path.suffix == ".svg")[-1].read_text()
    assert "can\u2019t be drawn" not in last and "length" not in last and "The spike" in last


def test_a_slide_of_a_layout_there_is_none_of_is_drawn_and_exported_as_content(tmp_path: Path) -> None:
    quote = {"layout": "quote", "title": "Why it matters",
             "body": [{"quote": "The queue lets the customer stop waiting."}]}
    deck = {**DECK, "slides": [*DECK["slides"], quote]}
    kind = DeckKind()
    drawing = kind.draw(deck, tmp_path)
    while drawing.unfinished:
        drawing = kind.draw(deck, tmp_path)
    page = drawing.pages[-1]
    # Its objects drawn as on a Content slide, not a blank slide with a warning sign.
    assert "The queue lets the customer" in page.svg and "Why it matters" in page.svg
    said = [message for message in drawing.messages if message.where == f"slides[{len(deck['slides']) - 1}].layout"]
    assert [(message.text, message.severity) for message in said] == [
        ("\u201cquote\u201d isn\u2019t a layout: drawn as Content. Choose one from Layout.", "warning")]
    # Exported the same, and said so.
    (pdf,) = kind.export(deck, tmp_path, "talk", ["pdf"], into=tmp_path / "out")
    assert _pages(pdf) == len(deck["slides"])
    number = len(deck["slides"])
    assert kind.export_notes == [f"Slide {number}: \u201cquote\u201d isn\u2019t a layout: drawn as Content."]


def test_shapes_found_wanting_as_they_are_drawn_are_plain_boxes_each_said(tmp_path: Path) -> None:
    nodes = [
        {"id": "p", "kind": "plasmid", "label": "pUC19", "properties": {"length": -5}},
        {"id": "t", "kind": "tree", "label": "Tree", "properties": {"newick": "((A,B"}},
        {"id": "q", "label": "Next"},
    ]
    bad = {"title": "Bad", "body": [{"figure": {"figure": {"id": "f"}, "nodes": nodes}}]}
    deck = {**DECK, "slides": [*DECK["slides"], bad]}
    kind = DeckKind()
    drawing = kind.draw(deck, tmp_path)
    while drawing.unfinished:
        drawing = kind.draw(deck, tmp_path)
    # The rest of the figure is drawn; each shape that can't be is said where it is, to be chosen.
    assert "Next" in drawing.pages[-1].svg and 'data-flexo-talk="invalid"' not in drawing.pages[-1].svg
    assert {message.where for message in drawing.messages if "can\u2019t be drawn yet" in message.text} == {
        "slides[2] f#p:length", "slides[2] f#t:newick"}
    kind.export(deck, tmp_path, "talk", ["pdf"], into=tmp_path / "out")
    assert kind.export_notes == [
        "Slide 3 \u00b7 Figure: \u201cpUC19\u201d is drawn as a plain box: a plasmid of -5 bp is too short to draw."
        " \u201cTree\u201d is drawn as a plain box: the tree's Newick does not read: a ( is never closed."
    ]


def test_a_skipped_slide_is_kept_but_neither_exported_nor_counted_among_the_sections(tmp_path: Path) -> None:
    import pytest

    from flexo_talk.document import DeckDocumentError, deck_from_document

    deck = {**DECK, "slides": [
        {"layout": "agenda"},
        {"layout": "section", "title": "Skipped part", "skip": True},
        {"layout": "section", "title": "Shown part"},
        {"title": "Last", "body": [{"text": "Words"}]},
    ]}
    kind = DeckKind()
    # Drawn to be edited, as Keynote's slide navigator shows a slide skipped.
    drawing = kind.draw(deck, tmp_path, {"settle": True})
    while drawing.unfinished:
        drawing = kind.draw(deck, tmp_path, {"settle": True})
    agenda, skipped, *_ = (page.svg for page in drawing.pages)
    assert "Skipped part" in skipped
    # The agenda lists the sections presented, numbered as they are.
    assert "Shown part" in agenda and "Skipped part" not in agenda
    (pdf,) = kind.export(deck, tmp_path, "talk", ["pdf"], into=tmp_path / "out")
    assert _pages(pdf) == 3
    with pytest.raises(DeckDocumentError, match="skip must be true or false"):
        deck_from_document({**DECK, "slides": [{"title": "T", "skip": "yes"}]}, tmp_path)


def _numbers(svgs: list[str]) -> list[str | None]:
    import re

    found = (re.search(r'id="slide\d+\.number"[^>]*><tspan[^>]*>(\d+)<', svg) for svg in svgs)
    return [match and match[1] for match in found]


def test_the_slides_after_one_skipped_are_numbered_as_they_are_shown(tmp_path: Path) -> None:
    # As Keynote numbers them: the show counts the slides it shows, and the numbers drawn on
    # them, the PDF's and the PowerPoint's agree with it -- no gap where one was skipped.
    deck = {**DECK, "slides": [
        {"title": "One", "body": [{"text": "a"}]},
        {"title": "Two", "skip": True, "body": [{"text": "b"}]},
        {"title": "Three", "body": [{"text": "c"}]},
    ]}
    kind = DeckKind()

    def drawn(document: dict) -> list[str]:
        drawing = kind.draw(document, tmp_path, {"settle": True})
        while drawing.unfinished:
            drawing = kind.draw(document, tmp_path, {"settle": True})
        return [page.svg for page in drawing.pages]

    # The slide skipped has none; the one after it is the second shown.
    assert _numbers(drawn(deck)) == ["1", None, "2"]
    # Shown again, the slides after it are drawn again with their numbers.
    shown = {**deck, "slides": [{**slide, "skip": False} for slide in deck["slides"]]}
    assert _numbers(drawn(shown)) == ["1", "2", "3"]
    svgs = sorted(kind.export(deck, tmp_path, "talk", ["svg"], into=tmp_path / "out"))
    assert [path.name for path in svgs] == ["slide-01.svg", "slide-02.svg"]
    assert _numbers([path.read_text(encoding="utf-8") for path in svgs]) == ["1", "2"]


def test_a_deck_s_logos_are_in_every_export_as_pictures(tmp_path: Path) -> None:
    import re
    import zipfile

    import pytest

    pytest.importorskip("PIL")
    from PIL import Image

    Image.new("RGBA", (300, 100), (120, 30, 40, 255)).save(tmp_path / "institute.png")
    Image.new("RGB", (80, 80), (10, 60, 140)).save(tmp_path / "crest.jpg")
    deck = {**DECK, "deck": {**DECK["deck"], "logos": ["institute.png", "crest.jpg"], "logos_on": "every"}}
    written = DeckKind().export(deck, tmp_path, "talk", ["pdf", "pptx", "svg"], into=tmp_path / "out")
    pptx = next(path for path in written if path.suffix == ".pptx")
    archive = zipfile.ZipFile(pptx)
    slides = [archive.read(f"ppt/slides/slide{n}.xml").decode() for n in (1, 2)]
    # Native pictures on every slide, named as PowerPoint's Selection Pane shows them.
    assert [len(re.findall(r'<p:cNvPr id="\d+" name="Logo"', slide)) for slide in slides] == [2, 2]
    pdf = next(path for path in written if path.suffix == ".pdf")
    assert pdf.read_bytes().count(b"/Subtype /Image") >= 2
    svgs = sorted(path for path in written if path.suffix == ".svg")
    assert 'id="slide1.logo0"' in svgs[0].read_text() and 'id="slide2.logo1"' in svgs[1].read_text()


def test_a_png_is_as_wide_as_chosen_and_named_apart_from_an_svg(tmp_path: Path) -> None:
    from PIL import Image

    pngs = DeckKind().export(DECK, tmp_path, "talk", ["png"], into=tmp_path / "png", png_width=1280)
    assert {path.parent.name for path in pngs} == {"My Talk \u2013 PNG 1280 px"}
    assert {Image.open(path).size for path in pngs} == {(1280, 720)}
    (wide,) = DeckKind().export({**DECK, "slides": DECK["slides"][:1]}, tmp_path, "talk", ["png"],
                                into=tmp_path / "big", png_width=3840)
    assert Image.open(wide).size == (3840, 2160)
    # Not chosen, 1920 pixels across, as the CLI's build draws a 16:9 slide.
    (plain,) = DeckKind().export({**DECK, "slides": DECK["slides"][:1]}, tmp_path, "talk", ["png"], into=tmp_path / "x")
    assert Image.open(plain).size == (1920, 1080) and plain.parent.name == "My Talk \u2013 PNG 1920 px"
    svgs = DeckKind().export(DECK, tmp_path, "talk", ["svg"], into=tmp_path / "svg")
    assert {path.parent.name for path in svgs} == {"My Talk \u2013 SVG"}
    with pytest.raises(ValueError, match="A PNG is 1280, 1920 or 3840 pixels wide, not 1000"):
        DeckKind().export(DECK, tmp_path, "talk", ["png"], into=tmp_path / "bad", png_width=1000)


def test_a_cropped_picture_is_cropped_natively_in_the_powerpoint_and_its_pdf_shows_only_the_part_kept(
    tmp_path: Path,
) -> None:
    from PIL import Image
    from pptx import Presentation

    picture = Image.new("RGB", (400, 200), "#ffffff")
    picture.paste(Image.new("RGB", (200, 200), "#c0392b"), (100, 0))
    picture.save(tmp_path / "wide.png")
    (tmp_path / "drawing.svg").write_text(
        '<svg xmlns="http://www.w3.org/2000/svg" width="80" height="40" viewBox="0 0 80 40">'
        '<rect x="0" y="0" width="80" height="40" fill="#2b4c9b"/></svg>'
    )
    deck = {"deck": {"id": "Crops"}, "slides": [
        {"title": "Kept", "body": [{"image": "wide.png", "crop": [0.25, 0.1, 0.5, 0.8]}]},
        {"title": "Round", "body": [{"image": "wide.png", "crop": [0.25, 0, 0.5, 1], "mask": "circle"}]},
        {"title": "Vectors", "body": [{"image": "drawing.svg", "crop": [0, 0, 0.5, 1]}]},
    ]}
    (pptx,) = DeckKind().export(deck, tmp_path, "talk", ["pptx"], into=tmp_path / "out")
    pictures = [[shape for shape in slide.shapes if shape.shape_type == 13] for slide in Presentation(pptx).slides]
    kept, round_, vectors = (found[0] for found in pictures)
    assert [round(value, 3) for value in (kept.crop_left, kept.crop_top, kept.crop_right, kept.crop_bottom)] == [
        0.25, 0.1, 0.25, 0.1]
    assert abs(kept.width / kept.height - (0.5 * 400) / (0.8 * 200)) < 0.01
    # Round: PowerPoint's own oval picture, of the square kept.
    assert round_._element.spPr.prstGeom.get("prst") == "ellipse" and abs(round_.width - round_.height) < 2
    assert round(round_.crop_left, 3) == round(round_.crop_right, 3) == 0.25
    # A cropped SVG is a picture of its vectors, cropped as a photograph is.
    assert round(vectors.crop_right, 3) == 0.5 and "svgBlip" in vectors._element.xml
    # The PDF and the PNG show the part kept: the picture's red middle, none of its white.
    (pdf,) = DeckKind().export(deck, tmp_path, "talk", ["pdf"], into=tmp_path / "pdf")
    drawn = b"".join(_streams(pdf.read_bytes()))
    assert drawn.count(b" re W n ") == 2 and drawn.count(b" c h W n ") == 1
    (png, *_) = DeckKind().export(deck, tmp_path, "talk", ["png"], into=tmp_path / "png")
    shown = Image.open(png).convert("RGB")
    reds = [x for x in range(shown.width) if shown.getpixel((x, shown.height // 2))[0] > 150
            and shown.getpixel((x, shown.height // 2))[1] < 100]
    whites = [x for x in range(min(reds), max(reds)) if shown.getpixel((x, shown.height // 2)) == (255, 255, 255)]
    assert reds and not whites
