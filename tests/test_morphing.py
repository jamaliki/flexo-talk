"""Morph (``flexo_talk.morphing``, the studio's morph.js): a slide that says ``transition: morph``
arrives by moving what it shares with the slide shown before it into place -- read and written as
a deck says it, wrong said where it is; what each object is said in the slide's drawing, so the
studio pairs them as it presents, never by a file's path; paired in the order the steps go,
named, framed, the same, placed, as the studio's script pairs them; and written into a PowerPoint
as its own Morph, a fade where it has none, each object either side named ``!!`` as PowerPoint
pairs them."""

from __future__ import annotations

import copy
import json
import re
import shutil
import subprocess
import xml.etree.ElementTree as ET
import zipfile
from itertools import pairwise
from pathlib import Path
from types import SimpleNamespace

import pytest

from flexo_talk.compose import render_slide
from flexo_talk.document import DeckDocumentError, deck_from_document, dump_document, load_document
from flexo_talk.export import build_deck
from flexo_talk.morphing import (
    DURATION,
    Transition,
    Unit,
    add_transition,
    page_fades,
    paired,
    plan,
    transition_of,
    transition_written,
    transition_xml,
    units_of,
)

PAIRS_JS = Path(__file__).parents[1] / "src/flexo_talk/studio/static/morphpairs.js"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"


def figure(count: int, ident: str = "smc") -> dict:
    """A pipeline of ``count`` steps, as each slide of a talk shows a little more of it."""

    names = ["prior", "weight", "resample", "move"][:count]
    return {
        "figure": {"id": ident},
        "nodes": [{"id": name, "label": name.title()} for name in names],
        "edges": [{"from": a, "to": b} for a, b in pairwise(names)],
        "groups": [{"id": "root", "children": names, "layout": {"kind": "row"}, "role": "canvas"}],
    }


def deck(*slides: dict) -> dict:
    return {"schema_version": 1, "deck": {"id": "morphs", "theme": "paper"}, "slides": list(slides)}


def talk() -> dict:
    return deck(
        {
            "title": "Sampling",
            "layout": "two-columns",
            "left": [{"bullets": ["Sample the prior"]}],
            "right": [{"figure": figure(2), "frame": "smc"}],
            "shapes": [
                {"kind": "arrow", "id": "pointer", "from": [100, 460], "to": [300, 460]},
                {"kind": "ellipse", "at": [700, 380], "size": [80, 60]},
            ],
        },
        {
            "title": "Sampling",
            "layout": "two-columns",
            "transition": "morph",
            "left": [{"bullets": ["Sample the prior", "Weight the samples"]}],
            "right": [{"figure": figure(3), "frame": "smc"}],
            "shapes": [{"kind": "arrow", "id": "pointer", "from": [500, 300], "to": [700, 250]}],
        },
        {
            "title": "An equation",
            "transition": {"kind": "morph", "duration": 2},
            "body": [{"text": "Words above it"}, {"math": "f(x) = a + b", "morph": "eq"}],
        },
        {"title": "Rearranged", "transition": "morph", "body": [{"math": "f(x) = b + a", "morph": "eq"}]},
    )


def drawn(document: dict) -> list[str]:
    made = deck_from_document(document, Path("."))
    return [render_slide(made, slide).svg for slide in made.slides]


# -- the format --


def test_a_slide_morphs_from_the_one_before_in_a_second_unless_it_says_how_long() -> None:
    assert transition_of("morph") == Transition("morph", DURATION) == Transition()
    assert transition_of({"kind": "morph", "duration": 2}) == Transition("morph", 2.0)
    assert transition_of({"duration": 0.5}) == Transition("morph", 0.5)
    assert transition_of(None) is transition_of("none") is transition_of(False) is None
    # Written as briefly as it can be: its kind alone, where it takes the time it takes unless told.
    assert transition_written(Transition()) == "morph"
    assert transition_written(Transition("morph", 2.0)) == {"kind": "morph", "duration": 2}
    assert transition_written(Transition("morph", 1.5)) == {"kind": "morph", "duration": 1.5}


@pytest.mark.parametrize(
    ("value", "said"),
    [
        ("dissolve", "transition must be morph"),
        ({"kind": "morph", "speed": 2}, "transition must be morph"),
        ({"kind": "morph", "duration": "slow"}, "a number of seconds"),
        ({"kind": "morph", "duration": 60}, "between 0.1 and 10 seconds"),
    ],
)
def test_a_transition_that_cannot_be_is_said_under_its_slide(value: object, said: str) -> None:
    document = deck({"title": "One"}, {"title": "Two", "transition": value})
    with pytest.raises(DeckDocumentError) as error:
        deck_from_document(document, Path("."))
    assert error.value.where == "slides[1].transition" and said in error.value.message


def test_an_object_or_a_shape_names_its_partner_with_a_name_and_nothing_else() -> None:
    for wrong, where in (
        ({"text": "Hello", "morph": True}, "slides[0].body[0] (text).morph"),
        ({"text": "Hello", "morph": ""}, "slides[0].body[0] (text).morph"),
    ):
        with pytest.raises(DeckDocumentError) as error:
            deck_from_document(deck({"title": "One", "body": [wrong]}), Path("."))
        assert error.value.where == where
    with pytest.raises(DeckDocumentError) as error:
        deck_from_document(
            deck({"title": "One", "shapes": [{"kind": "ellipse", "at": [10, 10], "size": [20, 20], "morph": ["a"]}]}),
            Path("."),
        )
    assert "morph must be a name" in error.value.message
    # A number is a name too, as YAML reads one.
    deck_from_document(
        deck(
            {
                "title": "One",
                "body": [{"text": "Hi", "morph": 3}],
                "shapes": [{"kind": "ellipse", "at": [10, 10], "size": [20, 20], "morph": "ring"}],
            }
        ),
        Path("."),
    )


def test_a_deck_written_again_keeps_its_transitions_where_a_person_reads_them(tmp_path: Path) -> None:
    written = dump_document(talk())
    assert "transition: morph" in written and "morph: eq" in written
    (tmp_path / "talk.yaml").write_text(written)
    again = load_document(tmp_path / "talk.yaml")
    assert [slide.get("transition") for slide in again["slides"]] == [
        None,
        "morph",
        {"kind": "morph", "duration": 2},
        "morph",
    ]
    # After its look, before its content.
    assert list(again["slides"][1]) == ["layout", "title", "transition", "left", "right", "shapes"]


# -- what a slide's drawing says --


def test_each_object_says_what_it_is_in_the_drawing_and_never_a_files_path(tmp_path: Path) -> None:
    (tmp_path / "cell.png").write_bytes(_png())
    document = deck(
        {
            "title": "One",
            "body": [
                {"text": "Hello", "morph": "greeting"},
                {"image": "cell.png"},
                {"figure": figure(2), "frame": "smc"},
            ],
            "shapes": [
                {"kind": "arrow", "id": "pointer", "from": [10, 10], "to": [90, 10]},
                {"kind": "ellipse", "at": [10, 10], "size": [20, 20], "morph": "ring"},
            ],
        }
    )
    made = deck_from_document(document, tmp_path)
    svg = render_slide(made, made.slides[0]).svg
    assert 'data-flexo-morph="greeting"' in svg and 'data-flexo-morph="ring"' in svg
    assert 'data-flexo-framed="smc"' in svg and 'data-flexo-same="shape:pointer"' in svg
    assert re.search(r'data-flexo-same="picture:[0-9a-f]{12}"', svg)
    assert re.search(r'data-flexo-same="text:[0-9a-f]{12}"', svg)
    assert str(tmp_path) not in re.sub(r'(?:href|xlink:href)="[^"]*"', "", svg)


def test_a_copy_of_a_slide_shows_the_same_figure_as_the_one_it_copies() -> None:
    """The studio's Duplicate Slide gives a figure written in the slide an id of its own
    (``smc-copy-x3k1``): the copy's figure is still the same figure, for a morph."""

    first = {"title": "One", "body": [{"figure": figure(2)}]}
    copied = copy.deepcopy(first)
    copied["body"][0]["figure"]["figure"]["id"] = "smc-copy-x3k1"
    other = {"title": "Three", "body": [{"figure": figure(2, "other")}]}
    found = [re.search(r'id="slide\d+\.body\.0"[^>]*', svg).group(0) for svg in drawn(deck(first, copied, other))]
    keys = [re.search(r'data-flexo-same="([^"]+)"', element).group(1) for element in found]
    assert keys[0] == keys[1] != keys[2] and keys[0].startswith("figure:")


# -- pairing --


def unit(key: str, kind: str = "text", box=(0, 0, 10, 10), region: str = "", **said) -> Unit:
    made = Unit(key, f"slide1.{key}", kind, tuple(box), region, **said)
    if not region and not made.same and not made.owner:
        made.same = f"furniture:{key}"
    return made


PAIRING = {
    "old": [
        {"key": "title", "kind": "text", "texty": True, "box": [48, 50, 300, 30]},
        {
            "key": "body.0",
            "kind": "text",
            "texty": True,
            "box": [48, 130, 200, 20],
            "region": "body",
            "same": "text:aaa",
        },
        {"key": "body.1", "kind": "figure", "box": [100, 200, 300, 100], "region": "body", "framed": "smc"},
        {"key": "body.1.caption", "kind": "text", "texty": True, "box": [100, 310, 100, 12], "owner": "body.1"},
        {"key": "body.2", "kind": "image", "box": [500, 200, 100, 100], "region": "body", "same": "picture:ccc"},
        {"key": "shapes.0", "kind": "shape:arrow", "box": [10, 10, 80, 2], "region": "shapes", "same": "shape:pointer"},
        {"key": "shapes.1", "kind": "shape:ellipse", "box": [700, 400, 40, 40], "region": "shapes"},
        {"key": "shapes.2", "kind": "shape:ellipse", "box": [800, 400, 40, 40], "region": "shapes", "morph": "ring"},
        {
            "key": "body.3",
            "kind": "text",
            "texty": True,
            "box": [48, 400, 200, 20],
            "region": "body",
            "same": "text:ddd",
        },
    ],
    "now": [
        {"key": "title", "kind": "text", "texty": True, "box": [48, 50, 420, 30]},
        {"key": "body.0", "kind": "image", "box": [600, 300, 200, 200], "region": "body", "same": "picture:ccc"},
        {
            "key": "body.1",
            "kind": "text",
            "texty": True,
            "box": [48, 132, 260, 20],
            "region": "body",
            "same": "text:bbb",
        },
        {"key": "left.0", "kind": "figure", "box": [400, 200, 400, 150], "region": "left", "framed": "smc"},
        {"key": "left.0.caption", "kind": "text", "texty": True, "box": [400, 360, 100, 12], "owner": "left.0"},
        {"key": "shapes.0", "kind": "shape:ellipse", "box": [703, 404, 41, 38], "region": "shapes"},
        {
            "key": "shapes.1",
            "kind": "shape:arrow",
            "box": [300, 300, 80, 40],
            "region": "shapes",
            "same": "shape:pointer",
        },
        {"key": "shapes.2", "kind": "shape:rectangle", "box": [100, 100, 40, 40], "region": "shapes", "morph": "ring"},
        {
            "key": "body.2",
            "kind": "text",
            "texty": True,
            "box": [48, 160, 200, 20],
            "region": "body",
            "same": "text:ddd",
        },
    ],
}
"""Two slides' units, as both pairings take them: named, framed, the same, and placed."""

PAIRED = [[0, 0], [4, 1], [1, 2], [2, 3], [3, 4], [6, 5], [5, 6], [7, 7], [8, 8]]


def test_what_two_slides_share_is_found_by_name_then_frame_then_what_it_is_then_where_it_is() -> None:
    old = [
        unit(**{key: tuple(value) if key == "box" else value for key, value in item.items()}) for item in PAIRING["old"]
    ]
    now = [
        unit(**{key: tuple(value) if key == "box" else value for key, value in item.items()}) for item in PAIRING["now"]
    ]
    pairs = paired(old, now)
    assert [list(pair) for pair in pairs] == PAIRED
    # Words that changed, where words were, of the same kind: the same words (body.0 -> body.1); a
    # shape of another kind named the same (the ring, now a rectangle); a caption with its figure.
    # Nothing paired for looking alike: the old body.0's words went nowhere else.


def _node(code: str) -> object:
    node = shutil.which("node")
    if node is None:
        pytest.skip("Node is not installed")
    script = f"const M = await import({json.dumps(PAIRS_JS.as_uri())});\n{code}"
    done = subprocess.run([node, "--input-type=module", "-e", script], capture_output=True, text=True, check=True)
    return json.loads(done.stdout)


def test_the_studio_pairs_as_a_powerpoint_is_written() -> None:
    """morph.js's ``paired`` (which the studio presents with) and flexo_talk.morphing's (which a
    PowerPoint is written with) pair the same slides alike."""

    made = _node(f"""
      const units = {json.dumps(PAIRING)};
      const fill = (u) => ({{ morph: "", framed: "", same: "", owner: "", region: "", texty: false, ...u,
        same: u.same || (!u.region && !u.owner ? `furniture:${{u.key}}` : "") }});
      console.log(JSON.stringify(M.paired(units.old.map(fill), units.now.map(fill))));
    """)
    assert made == PAIRED


def test_the_glyphs_two_equations_share_travel_and_those_that_commute_cross() -> None:
    made = _node("""
      const glyphs = (keys, xs) => keys.map((key, i) => ({ key, box: [xs[i], 0, 8, 10] }));
      // f = a + b  ->  f = b + a
      const before = glyphs(["f", "=", "a", "+", "b"], [0, 10, 20, 30, 40]);
      const after = glyphs(["f", "=", "b", "+", "a"], [0, 10, 20, 30, 40]);
      console.log(JSON.stringify({
        swapped: M.matchGlyphs(before, after),
        // One more term: what is kept is kept in order; the new one has no partner.
        grown: M.matchGlyphs(before, glyphs(["f", "=", "a", "+", "b", "+", "c"], [0, 10, 20, 30, 40, 50, 60])),
        none: M.matchGlyphs([], after),
        transitions: [M.transitionOf("morph"), M.transitionOf({ kind: "morph", duration: 99 }), M.transitionOf("fade"),
          M.transitionWritten({ kind: "morph", duration: 1 }), M.transitionWritten({ kind: "morph", duration: 2.5 })],
      }));
    """)
    assert made["swapped"] == [[0, 0], [1, 1], [4, 2], [3, 3], [2, 4]]
    assert made["grown"] == [[0, 0], [1, 1], [2, 2], [3, 3], [4, 4]]
    assert made["none"] == []
    assert made["transitions"] == [
        {"kind": "morph", "duration": 1},
        {"kind": "morph", "duration": 10},
        None,
        "morph",
        {"kind": "morph", "duration": 2.5},
    ]


def test_the_units_of_a_slide_are_what_it_shows_its_figure_paired_by_its_frame() -> None:
    svgs = drawn(talk())
    first, second = units_of(svgs[0]), units_of(svgs[1], entering=True)
    keys = {unit.key for unit in second}
    assert {"title", "left.0", "right.0", "shapes.0", "number", "canvas.background"} <= keys
    pairs = {(first[i].key, second[j].key) for i, j in paired(first, second)}
    assert ("right.0", "right.0") in pairs and ("shapes.0", "shapes.0") in pairs and ("title", "title") in pairs
    # The ellipse slide 2 has none of fades: it is paired with nothing.
    assert not any(old == "shapes.1" for old, _ in pairs)


# -- PowerPoint --


def test_a_morphing_slide_is_written_with_powerpoints_morph_and_a_fade_for_one_without(tmp_path: Path) -> None:
    made = deck_from_document(talk(), Path("."))
    pptx = build_deck(made, tmp_path, ("pptx",)).pptx
    with zipfile.ZipFile(pptx) as archive:
        slides = [archive.read(f"ppt/slides/slide{n}.xml").decode() for n in range(1, 5)]
    assert "p159:morph" not in slides[0]
    for xml, ms in ((slides[1], 1000), (slides[2], 2000), (slides[3], 1000)):
        found = re.search(r"<mc:AlternateContent [^>]*><mc:Choice xmlns:p159=.*?</mc:AlternateContent>", xml)
        assert found, xml[-600:]
        said = found.group(0)
        assert 'Requires="p159"' in said and '<p159:morph option="byObject"/>' in said and f'p14:dur="{ms}"' in said
        assert "<mc:Fallback><p:transition" in said and "<p:fade/>" in said
        # After the slide's colours, as PowerPoint writes it.
        assert xml.index("clrMapOvr") < xml.index(said)


def test_what_a_morph_pairs_is_named_alike_for_powerpoint_to_pair_it(tmp_path: Path) -> None:
    made = deck_from_document(talk(), Path("."))
    pptx = build_deck(made, tmp_path, ("pptx",)).pptx
    from pptx import Presentation

    shown = Presentation(str(pptx))
    names = [[shape.name for shape in slide.shapes] for slide in shown.slides]
    for before, after in ((0, 1), (1, 2), (2, 3)):
        assert all(name.startswith("!!") for name in names[before] + names[after]), names
    # What the first two share is named alike; what each has alone, apart.
    assert "!!smc" in names[0] and "!!smc" in names[1]
    assert "!!pointer" in names[0] and "!!pointer" in names[1]
    assert "!!title" in names[0] and "!!title" in names[1]
    for slide in names:
        assert len(slide) == len(set(slide))
    alone = set(names[0]) - set(names[1])
    assert alone and not alone & set(names[1])
    # The equation named for its name, from slide 3 to 4: PowerPoint's own and flexo's drawing.
    with zipfile.ZipFile(pptx) as archive:
        for number in (3, 4):
            assert archive.read(f"ppt/slides/slide{number}.xml").decode().count('name="!!eq"') == 2


def test_a_deck_without_morphs_is_named_as_it_always_was(tmp_path: Path) -> None:
    document = talk()
    for slide in document["slides"]:
        slide.pop("transition", None)
    made = deck_from_document(document, Path("."))
    assert plan([]).transitions == {}
    pptx = build_deck(made, tmp_path, ("pptx",)).pptx
    from pptx import Presentation

    names = [shape.name for slide in Presentation(str(pptx)).slides for shape in slide.shapes]
    assert not any(name.startswith("!!") for name in names)
    with zipfile.ZipFile(pptx) as archive:
        assert not any(
            "p159" in archive.read(name).decode() for name in archive.namelist() if name.startswith("ppt/slides/slide")
        )


def test_a_slide_kept_as_written_is_written_again_with_its_transition_as_it_is_now(tmp_path: Path) -> None:
    """A PowerPoint written again from slides drawn before (the studio's, kept by their key) gives
    a slide whose transition, or whose neighbour, changed its transition and names as they are."""

    from flexo_talk.export import write_pptx

    made = deck_from_document(talk(), Path("."))
    rendered = made.render()
    for number, item in enumerate(rendered):
        item.key = f"kept-{number}"
    first = write_pptx(made, rendered, tmp_path / "a.pptx")
    rendered[1].slide.source.pop("transition")
    again = write_pptx(made, rendered, tmp_path / "b.pptx")
    with zipfile.ZipFile(first) as one, zipfile.ZipFile(again) as two:
        assert "p159:morph" in one.read("ppt/slides/slide2.xml").decode()
        assert "p159:morph" not in two.read("ppt/slides/slide2.xml").decode()
        assert "!!" not in two.read("ppt/slides/slide1.xml").decode()


def test_a_pdf_presented_full_screen_fades_into_a_slide_that_morphs(tmp_path: Path) -> None:
    """A PDF can do no more than fade: the first page of a slide that morphs is faded to, over
    the morph's time; every other page is cut to."""

    made = deck_from_document(talk(), Path("."))
    data = build_deck(made, tmp_path, ("pdf",)).pdf.read_bytes()
    pages = re.findall(rb"/Type /Page /Parent[^\n]*", data)
    fades = [re.search(rb"/Trans << /Type /Trans /S /Fade /D ([\d.]+) >>", page) for page in pages]
    assert [fade.group(1) if fade else None for fade in fades] == [None, b"1", b"2", b"1"]


def test_a_transition_goes_after_the_slides_colours_and_before_its_builds() -> None:
    from lxml import etree

    slide = etree.fromstring(f'<p:sld xmlns:p="{P[1:-1]}"><p:cSld/><p:clrMapOvr/><p:timing/></p:sld>')
    add_transition(slide, Transition("morph", 0.5))
    add_transition(slide, Transition("morph", 1.5))
    names = [etree.QName(child).localname for child in slide]
    assert names == ["cSld", "clrMapOvr", "AlternateContent", "timing"]
    assert 'p14:dur="1500"' in etree.tostring(slide).decode()
    assert 'spd="fast"' in transition_xml(Transition("morph", 0.5))


def _png() -> bytes:
    import base64

    return base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAIAAAACCAIAAAD91JpzAAAAFklEQVR4nGP8z8DAwMDAxMDAwMDAAAAMBgEB6X0K0gAAAABJRU5ErkJggg=="
    )


# -- with clicks (flexo_talk.onclick) --


def clicking() -> dict:
    """Three slides that morph, each walked through on clicks: on the first, words that appear on
    its first click and a stage hidden from its second; the second shows those words from the
    start, hides the same stage on its first click, and has a ring that appears on a click."""

    return deck(
        {
            "title": "Sampling",
            "layout": "two-columns",
            "left": [
                {"bullets": ["Prior", "Weight"], "click 1": {"highlight": [1]}, "click 2": {}},
                {"text": "Shown on a click", "hide": True, "click 1": {}},
            ],
            "right": [{"figure": figure(3), "frame": "smc", "click 2": {"hide": ["prior"]}}],
            "shapes": [{"kind": "ellipse", "id": "ring", "at": [700, 400], "size": [80, 60]}],
        },
        {
            "title": "Sampling",
            "layout": "two-columns",
            "transition": "morph",
            "left": [{"bullets": ["Prior", "Weight"]}, {"text": "Shown on a click"}],
            "right": [{"figure": figure(3), "frame": "smc", "click 1": {"hide": ["prior"]}}],
            "shapes": [
                {"kind": "ellipse", "id": "ring", "at": [700, 400], "size": [80, 60], "hide": True, "click 1": {}}
            ],
        },
        {
            "title": "Sampling",
            "layout": "two-columns",
            "transition": {"kind": "morph", "duration": 2},
            "left": [{"bullets": ["Prior", "Weight"]}],
            "right": [{"figure": figure(3), "frame": "smc"}],
        },
    )


def _top_names(xml: str) -> list[tuple[str, int]]:
    tree = next(ET.fromstring(xml).iter(f"{P}spTree"))
    found = []
    for child in tree:
        properties = next(child.iter(f"{P}cNvPr"), None)
        if properties is not None and child.tag != f"{P}nvGrpSpPr":
            found.append((properties.get("name"), int(properties.get("id"))))
    return found


def _named_with(xml: str, words: str) -> str | None:
    """The name of the slide's shape whose words are ``words``."""

    tree = next(ET.fromstring(xml).iter(f"{P}spTree"))
    for child in tree:
        if (
            "".join(t.text or "" for t in child.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}t"))
            == words
        ):
            return next(child.iter(f"{P}cNvPr")).get("name")
    return None


def test_a_morph_goes_from_the_slide_before_after_its_last_click_to_this_one_before_its_first() -> None:
    """The slide before is left as its last click shows it -- words that appeared on a click
    there, and are on this slide from the start, are paired -- and this one is arrived at as it is
    before its first click: what appears on a click here is not paired with what was there."""

    made = deck_from_document(clicking(), Path("."))
    rendered = made.render()
    morphs = plan(rendered)
    assert morphs.transitions == {1: Transition(), 2: Transition("morph", 2.0)}
    words = {
        index: next(ident for ident in names if ident.endswith("left.1"))
        for index, names in morphs.names.items()
        if any(ident.endswith("left.1") for ident in names)
    }
    assert morphs.names[0][words[0]] == morphs.names[1][words[1]]
    # The ring is on the first slide; on the second it appears on a click: not paired, not named so.
    assert morphs.names[0]["slide1.shapes.0"] == "!!ring"
    assert "slide2.shapes.0" not in morphs.names[1] or morphs.names[1]["slide2.shapes.0"] != "!!ring"
    # The figure, the same frame, goes on from slide to slide under one name.
    assert morphs.names[0]["slide1.right.0"] == morphs.names[1]["slide2.right.0"] == morphs.names[2]["slide3.right.0"]


def test_a_powerpoint_with_clicks_and_morphs_keeps_its_transition_before_its_timing_and_pairs_lifted_parts(
    tmp_path: Path,
) -> None:
    made = deck_from_document(clicking(), Path("."))
    pptx = build_deck(made, tmp_path, ("pptx",)).pptx
    with zipfile.ZipFile(pptx) as archive:
        slides = [archive.read(f"ppt/slides/slide{n}.xml").decode() for n in range(1, 4)]
    for xml in slides[1:]:
        assert xml.index("<p159:morph") < xml.index("<p:timing") if "<p:timing" in xml else "<p159:morph" in xml
    assert "<p:timing" in slides[0] and "<p:timing" in slides[1]
    for xml in slides:
        names = _top_names(xml)
        # Every shape on a slide either side of a morph named for PowerPoint to pair as the morph does.
        assert all(name.startswith("!!") for name, _ in names), names
        assert len({name for name, _ in names}) == len(names)
        # What the clicks' animations act on is still there, by its id.
        ids = {int(properties.get("id")) for properties in ET.fromstring(xml).iter(f"{P}cNvPr")}
        assert {int(spid) for spid in re.findall(r'spTgt spid="(\d+)"', xml)} <= ids
    first, second, third = (dict(_top_names(xml)) for xml in slides)
    # The figure's part hidden on a click, lifted out of it on both slides: named alike.
    lifted = [name for name in first if name.startswith("!!smc ") and "prior" in name]
    assert lifted and lifted[0] in second
    # Words that appear on a click on the first slide are paired with those on the second.
    assert _named_with(slides[0], "Shown on a click") == _named_with(slides[1], "Shown on a click")
    # The ring that appears on a click on the second is not the first's.
    assert "!!ring" in first and "!!ring" not in second
    assert "!!smc" in first and "!!smc" in second and "!!smc" in third


def test_a_pdf_of_a_deck_walked_through_on_clicks_fades_only_into_the_first_page_of_a_morph(tmp_path: Path) -> None:
    made = deck_from_document(clicking(), Path("."))
    rendered = made.render()
    assert [item.steps for item in rendered] == [3, 2, 1]
    assert page_fades(rendered) == [None, None, None, 1.0, None, 2.0]
    assert page_fades(rendered, handout=True) == [None, 1.0, 2.0]
    data = build_deck(made, tmp_path, ("pdf",)).pdf.read_bytes()
    pages = re.findall(rb"/Type /Page /Parent[^\n]*", data)
    fades = [re.search(rb"/Trans << /Type /Trans /S /Fade /D ([\d.]+) >>", page) for page in pages]
    assert [fade.group(1) if fade else None for fade in fades] == [None, None, None, b"1", None, b"2"]
    handout = build_deck(made, tmp_path / "handout", ("pdf",), handout=True).pdf.read_bytes()
    pages = re.findall(rb"/Type /Page /Parent[^\n]*", handout)
    fades = [re.search(rb"/Trans << /Type /Trans /S /Fade /D ([\d.]+) >>", page) for page in pages]
    assert [fade.group(1) if fade else None for fade in fades] == [None, b"1", b"2"]


def test_a_slide_that_morphs_is_drawn_at_each_step_as_it_would_be_without(tmp_path: Path) -> None:
    """A morph is how a slide arrives: its drawing, at each of its clicks (as an agent looks at it,
    and the PDF draws it), is the slide's own."""

    from flexo_talk.studio.deck_agent import at_step

    document = clicking()
    still = copy.deepcopy(document)
    for slide in still["slides"]:
        slide.pop("transition", None)
    moving, plain = deck_from_document(document, Path(".")).render(), deck_from_document(still, Path(".")).render()
    for one, other in zip(moving, plain, strict=True):
        assert one.svg == other.svg
        for step in range(1, one.steps + 1):
            assert one.at_step(step) == other.at_step(step)
            looked = at_step(SimpleNamespace(svg=one.svg, steps=one.steps), step)
            assert looked == at_step(SimpleNamespace(svg=other.svg, steps=other.steps), step)
