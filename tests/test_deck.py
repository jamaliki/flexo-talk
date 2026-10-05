"""A deck builds, and its PowerPoint keeps every word, list, figure, and note."""

from __future__ import annotations

import importlib.util
import re
import zipfile
from pathlib import Path

import pytest

from flexo_talk import Deck
from flexo_talk.deck import inline

ROOT = Path(__file__).resolve().parents[1]


def _demo() -> Deck:
    spec = importlib.util.spec_from_file_location("demo", ROOT / "examples" / "demo.py")
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module.talk()


def _slides(path: Path) -> list[str]:
    archive = zipfile.ZipFile(path)
    names = sorted(
        (n for n in archive.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)),
        key=lambda n: int(re.search(r"(\d+)", n).group(1)),  # type: ignore[union-attr]
    )
    return [archive.read(name).decode() for name in names]


def test_the_demo_builds_every_format(tmp_path: Path) -> None:
    result = _demo().build(tmp_path)
    assert result.pptx is not None and result.pptx.exists()
    assert len(result.svgs) == len(result.pngs) == 5
    assert not result.diagnostics, result.summary()
    assert result.pdf is not None
    pdf = result.pdf.read_bytes()
    assert pdf.count(b"/Type /Page ") == 5 and b"/FontFile2" in pdf


def test_every_word_on_a_slide_is_live_text_in_the_pptx(tmp_path: Path) -> None:
    deck = _demo()
    result = deck.build(tmp_path, formats=("pptx",))
    slides = _slides(result.pptx)  # type: ignore[arg-type]
    assert len(slides) == 5
    text = "".join(re.findall(r"<a:t>([^<]*)</a:t>", slides[1]))
    for words in ("Why another figure tool?", "boxes drift out of line", "flexo-talk · a demo"):
        assert words.replace("·", "·") in text
    # The list is one box of bulleted paragraphs, the second level a level in.
    assert slides[1].count("<a:buChar") == 5 and 'lvl="1"' in slides[1]
    # The figure is native shapes: a rounded box per component, a freeform per line.
    assert slides[4].count('prst="roundRect"') >= 5 and "<a:cubicBezTo>" in slides[4]
    # Dashes are presets every slide program draws; scripts are written larger
    # than drawn, since slide programs shrink raised runs.
    assert '<a:prstDash val="dash"/>' in slides[4] and "custDash" not in slides[4]
    assert re.search(r'sz="\d+" i="1" baseline="', slides[4]) or 'baseline="' in slides[4]
    notes = [n for n in zipfile.ZipFile(result.pptx).namelist() if "notesSlide" in n]  # type: ignore[arg-type]
    assert notes, "the speaker notes are kept"


def test_slide_text_takes_emphasis_and_math() -> None:
    runs = inline("A *model* is **not** $x^2$")
    assert [(r.text, r.italic, r.weight) for r in runs][:4] == [
        ("A ", False, 400),
        ("model", True, 400),
        (" is ", False, 400),
        ("not", False, 700),
    ]
    assert any(r.baseline_shift == "super" for r in runs)


def test_a_deck_takes_any_flexo_theme_and_checks_its_layouts() -> None:
    for theme in ("tikz", "sketch", "dark"):
        with Deck("t", theme=theme) as deck, deck.slide("Hello") as slide:
            slide.bullets("one", "two")
        assert deck.render()[0].svg.startswith("<?xml")
    with pytest.raises(ValueError, match="Available layouts"), Deck("t") as deck:
        deck.slide("x", layout="three-columns")  # type: ignore[arg-type]


def test_a_matplotlib_plot_is_native_shapes_and_text(tmp_path: Path) -> None:
    matplotlib = pytest.importorskip("matplotlib")
    matplotlib.use("svg")
    import matplotlib.pyplot as plt

    deck = Deck("plots", theme="dark")
    with deck.plotting():
        figure, axes = plt.subplots()
        axes.plot([0, 1, 2], [2, 1, 1.5], label="loss")
        axes.set_ylabel(r"Loss $\theta$")
        axes.imshow([[0, 1], [1, 0]], extent=(0, 2, 1, 2))
    with deck.slide("A plot") as slide:
        slide.plot(figure)
    result = deck.build(tmp_path, formats=("pptx", "pdf"))
    slide = _slides(result.pptx)[0]  # type: ignore[arg-type]
    # Tick labels are live text in the deck's font.
    assert 'typeface="Figtree"' in slide and ">0.00<" in slide
    # Lines are freeforms; the image is a picture, flipped as matplotlib stores it.
    assert "<a:custGeom>" in slide and "<p:pic>" in slide and 'flipV="1"' in slide
    # Words with maths in them are set by flexo, as on a slide, and drawn as outlines.
    assert ">Loss" not in slide and "STIX" not in slide


def test_an_illustrator_file_is_placed_as_shapes(tmp_path: Path) -> None:
    pytest.importorskip("pypdfium2")
    from flexo.pdf import write_pdf

    # An Illustrator file is a PDF with Illustrator's data beside it; flexo's own PDF
    # stands in for one: a box, and words in an embedded face.
    source = tmp_path / "panel.ai"
    write_pdf(
        ['<svg xmlns="http://www.w3.org/2000/svg" width="200" height="80" viewBox="0 0 200 80">'
         '<rect x="10" y="10" width="60" height="60" fill="#d55e00"/>'
         '<text x="90" y="50" font-size="20">Panel A</text></svg>'],
        source,
    )
    deck = Deck("illustrator")
    with deck.slide("A colleague's panel") as slide:
        slide.image(source)
    result = deck.build(tmp_path, formats=("pptx", "pdf"))
    xml = _slides(result.pptx)[0]  # type: ignore[arg-type]
    assert "<p:pic>" not in xml
    assert xml.count("<a:custGeom>") >= 2  # the box, and the words' outlines
    assert 'srgbClr val="D55E00"' in xml  # in the box's own colour


def test_an_overfull_slide_is_set_smaller_and_reported(tmp_path: Path) -> None:
    deck = Deck("full")
    with deck.slide("Too much") as slide:
        slide.bullets(*(["a long sentence that goes on and on across the whole slide width"] * 12))
    result = deck.build(tmp_path, formats=("svg",))
    assert any("to fit" in message for message in result.diagnostics)


def test_a_table_cell_is_drawn_on_the_lines_its_row_was_sized_for(monkeypatch: pytest.MonkeyPatch) -> None:
    from flexo_talk import compose

    # A line measured again in the width it was found to take can come out a hair wider
    # (kerning across a space) and wrap: here, any width given falls half a point short.
    measure = compose._Canvas.measure

    def short(self, runs, size, width, *args, **kwargs):  # type: ignore[no-untyped-def]
        return measure(self, runs, size, width - 0.5 if width else width, *args, **kwargs)

    monkeypatch.setattr(compose._Canvas, "measure", short)
    deck = Deck("tables")
    with deck.slide("Half-times") as slide:
        slide.table([["[O2] (µM)", "t1/2 EGFP (min)", "t1/2 sfGFP (min)"], ["250", "25.1", "13.6"]])
    rendered = compose.render_slide(deck, deck.slides[0])
    header = re.search(r'id="slide1\.body\.0\.0\.1"[^>]*>(.*?)</text>', rendered.svg, re.S).group(1)
    # Set as the table was planned: one line, in a row one line tall, over nothing below it.
    assert header.count("<tspan") == 1 and "t1/2 EGFP (min)" in header
    (table,) = rendered.tables
    assert table.heights[0] == pytest.approx(table.heights[1]) and table.measured is None


def test_words_set_a_few_per_cent_smaller_to_fit_go_unsaid() -> None:
    from flexo_talk import compose

    deck = Deck("snug")
    with deck.slide("Snug") as slide:
        slide.bullets(*(["a sentence that goes on and on across the slide"] * 11))
    rendered = compose.render_slide(deck, deck.slides[0])
    # Set a little smaller, so it fits; too little to be seen, so nothing is said of it.
    (words,) = rendered.lists
    assert compose.QUIET_SHRINK * deck.style.body_size <= words.size < deck.style.body_size
    assert rendered.diagnostics == []


def test_a_text_that_is_one_formula_is_set_in_display_style() -> None:
    deck = Deck("alone")
    with deck.slide("Alone") as slide:
        slide.text(r"$F = \frac{k_2}{k_2 - k_1}$")
        slide.text(r"Words and $\frac{a}{b}$ among them")
        slide.text(r"$x^2$")
    alone, among, simple = (block.runs for block in deck.slides[0].regions["body"].blocks)
    # Alone, its fraction is full size, as a displayed one is; among words, set to sit in them.
    assert [run.math for run in alone] == [r"\displaystyle F = \frac{k_2}{k_2 - k_1}"]
    assert [run.math for run in among if run.math] == [r"\frac{a}{b}"]
    assert not any(run.math for run in simple)


def test_a_formula_displayed_in_words_is_centred_in_a_list_after_its_bullet_set_apart_and_tagged(
    tmp_path: Path,
) -> None:
    from flexo.drawing import Drawing, Group, ink_bounds, read_drawing

    from flexo_talk import compose

    def build(gap: float) -> tuple[list[tuple[float, float]], list[float], float, bytes | None, str | None]:
        compose.DISPLAY_GAP, kept = gap, compose.DISPLAY_GAP
        try:
            deck = Deck("display")
            with deck.slide("Model") as slide:
                slide.bullets("Let p be the fraction phosphorylated:", r"$$\frac{dp}{dt} = k_B (1 - p)$$",
                              "Delay makes a cycle.")
                slide.text(r"The period is $$T = \frac{2\pi}{\omega}$$ about a day.")
            rendered = compose.render_slide(deck, deck.slides[0])
            built = deck.build(tmp_path / str(gap), formats=("pdf", "pptx")) if gap else None
        finally:
            compose.DISPLAY_GAP = kept

        def formulas(group: Group) -> list[Group]:
            found = [group] if "data-flexo-math" in group.data else []
            return found + [inner for item in group.items if isinstance(item, Group) for inner in formulas(item)]

        spans = []
        for formula in formulas(read_drawing(rendered.svg).root):
            left, _, right, _ = ink_bounds(Drawing(0.0, 0.0, Group(None, list(formula.items))))
            spans.append((left, right))
        words = float(re.search(r'<text id="[^"]*\.0" x="([\d.]+)"', rendered.svg).group(1))
        pdf = built.pdf.read_bytes() if built else None  # type: ignore[union-attr]
        xml = _slides(built.pptx)[0] if built else None  # type: ignore[arg-type]
        return spans, [baseline for _, _, baseline in rendered.lists[0].items], words, pdf, xml

    spans, baselines, words, pdf, xml = build(compose.DISPLAY_GAP)
    # In a list's item, after its bullet where the item's words start; among words, centred in
    # the room they are set in (the region, 48 to 912), as LaTeX centres one.
    assert spans[0][0] == pytest.approx(words, abs=3.0)
    assert (spans[1][0] + spans[1][1]) / 2.0 == pytest.approx(480.0, abs=1.0)
    # So in PowerPoint's equations too: at the item's start, and centred among words.
    assert xml is not None and re.findall(r'<m:jc m:val="(\w+)"/>', xml) == ["left", "centerGroup"]
    # Set apart from the items either side of it, by the same room above and below.
    _, close, _, _, _ = build(0.0)
    apart = [now - before for now, before in zip(baselines, close, strict=True)]
    assert apart[0] == 0.0 and apart[1] == pytest.approx(8.0) and apart[2] == pytest.approx(16.0)
    # A formula in the PDF, described by its words, within the item and the paragraph it is in.
    assert pdf is not None
    found = re.findall(rb"/S /Formula /P (\d+) 0 R /Pg \d+ 0 R /Alt <FEFF([0-9A-F]+)>", pdf)
    alts = [bytes.fromhex(alt.decode()).decode("utf-16-be") for _, alt in found]
    assert alts == ["dp/dt = k_B(1 \u2212 p)", "T = 2π/ω"]
    parents = [re.search(rb"\b%s 0 obj\s*<< /Type /StructElem /S /(\w+)" % parent, pdf).group(1) for parent, _ in found]
    assert parents == [b"LBody", b"P"]


def test_maths_in_a_powerpoint_list_reads_with_single_spaces_and_breaks_round_a_display(tmp_path: Path) -> None:
    deck = Deck("spaced")
    with deck.slide("Period") as slide:
        slide.bullets(r"A period $T \approx 24$ h", r"The rate $$\frac{dp}{dt} = k$$ falls")
    xml = _slides(deck.build(tmp_path, formats=("pptx",)).pptx)[0]  # type: ignore[arg-type]
    # An italic letter's lean before a sign is not a second space where the words are read.
    first = re.search(r"<a:p>(?:(?!</a:p>).)*A period (?:(?!</a:p>).)*</a:p>", xml, re.S).group(0)
    assert "".join(re.findall(r"<a:t>([^<]*)</a:t>", first)) == "A period T ≈ 24 h"
    # A formula displayed among an item's words is on a line of its own there too.
    item = re.search(r"<a:p>(?:(?!</a:p>).)*The rate (?:(?!</a:p>).)*</a:p>", xml, re.S).group(0)
    assert item.count("<a:br>") == 2


def test_a_line_that_inhibits_or_catalyses_keeps_its_head_in_powerpoint(tmp_path: Path) -> None:
    from flexo_talk.document import deck_from_document

    heads = ["arrow", "inhibition", "catalysis", "stimulation", "necessary", "modulation"]
    nodes = [{"id": "src", "label": "Source"}] + [{"id": head, "label": head.title()} for head in heads]
    edges = [{"from": "src", "to": head, "head": head} for head in heads]
    figure = {"figure": {"id": "f"}, "nodes": nodes, "edges": edges}
    deck = deck_from_document({"deck": {"id": "heads"}, "slides": [{"title": "Heads", "body": [{"figure": figure}]}]},
                              tmp_path)
    xml = _slides(deck.build(tmp_path, formats=("pptx",)).pptx)[0]  # type: ignore[arg-type]
    # Only the arrow is a slide program's own line end. A bar (represses), a circle (catalyses)
    # or an open triangle has none: each is drawn as Flexo drew it, never turned into an arrow.
    assert re.findall(r'<a:(?:head|tail)End type="(\w+)"', xml) == ["stealth"]
    for head in heads[1:]:
        line = re.search(rf'<p:grpSp><p:nvGrpSpPr><p:cNvPr id="\d+" name="Line from Source to {head.title()}"/>'
                         r"(?:(?!</p:grpSp>).)*</p:grpSp>", xml, re.S)
        assert line is not None and line.group(0).count("<a:custGeom>") >= 2, head


def test_a_tall_formula_opens_only_its_own_line_and_its_words_read_as_one_paragraph(tmp_path: Path) -> None:
    pdfium = pytest.importorskip("pypdfium2")
    from flexo_talk import compose

    deck = Deck("tall")
    with deck.slide("Poisson", layout="two-columns") as slide:
        slide.left.text(
            "Infect at multiplicity $m$: the number of phages per cell, $k$, is Poisson,\n"
            "$$P(k) = \\frac{m^k e^{-m}}{k!}$$\nso the share of cells with $k \\ge 2$ is $1 - e^{-m}(1+m)$."
        )
        slide.right.text("Words beside it")
    rendered = compose.render_slide(deck, deck.slides[0])
    (words,) = (words for words in rendered.worded if words.id == "slide1.left.0")
    plain = deck.typography(words.size).line_height * words.size
    # Each line of words is spaced as words are; only the formula's own line is opened, and
    # set apart from them -- in the PowerPoint as on the slide.
    assert words.heights is not None and words.downs is not None
    assert [height == pytest.approx(plain) for height in words.heights] == [True, True, False, True, True]
    assert words.downs[1] == pytest.approx(plain) and words.downs[4] - words.downs[3] == pytest.approx(plain)
    pdf = deck.build(tmp_path, formats=("pdf",)).pdf.read_bytes()  # type: ignore[union-attr]
    # One paragraph with its formula in it, every script in it read by one rule.
    assert re.findall(rb"/Type /StructElem /S /(\w+)", pdf) == [b"Document", b"Sect", b"H1", b"P", b"Formula", b"P"]
    said = pdfium.PdfDocument(pdf)[0].get_textpage().get_text_range()
    assert "P(k) = (m^k e^(\u2212m))/k!" in said and "k \u2265 2 is 1 \u2212 e^(\u2212m)(1 + m)." in said


def test_a_slide_is_read_title_body_footnote_then_its_number(tmp_path: Path) -> None:
    from pptx import Presentation

    deck = Deck("order")
    with deck.slide("Why a phage decides") as slide:
        slide.bullets("A temperate phage chooses", "The choice depends on dose")
        slide.footnote("Zeng L. et al. (2010) Cell 141:682.")
    result = deck.build(tmp_path, formats=("pdf", "pptx"))
    pdf = result.pdf.read_bytes()  # type: ignore[union-attr]
    # In the PDF: the heading, the list, then the footnote; the page number is decoration.
    kinds = re.findall(rb"/Type /StructElem /S /(\w+)", pdf)
    assert kinds == [b"Document", b"Sect", b"H1", b"L", b"LI", b"LBody", b"LI", b"LBody", b"P"]
    # In the PowerPoint's order (its Selection Pane, a screen reader) the same.
    slide_xml = Presentation(result.pptx).slides[0]._element  # type: ignore[arg-type]
    names = [element.get("name") for element in slide_xml.iter()
             if element.tag.endswith("}cNvPr") and element.get("name")]
    assert [name.split(" ")[0] for name in names] == ["Title", "Rule", "List", "Footnote", "Slide"]


def test_a_slide_s_shapes_are_written_in_the_order_they_are_read(tmp_path: Path) -> None:
    from pptx import Presentation

    deck = Deck("order")
    with deck.slide("Families", layout="two-columns") as slide:
        slide.left.table([["Family", "Carriers"], ["A", "3"]])
        slide.left.text("Families in the registry")
        slide.right.bullets("Lag: enteropeptidase", "Burst: trypsin")
        slide.right.text("n = 3 gels")
    pptx = deck.build(tmp_path, formats=("pptx",)).pptx
    shapes = Presentation(pptx).slides[0]._element.iter()  # type: ignore[arg-type]
    names = [element.get("name") for element in shapes if element.tag.endswith("}cNvPr") and element.get("name")]
    # The left column top to bottom, then the right: a table and a list where they are drawn,
    # not written after the words below them.
    assert [name.split(" ")[0] for name in names] == ["Title", "Rule", "Table", "Text", "List", "Text", "Slide"]


def test_a_figure_s_heading_is_never_as_large_as_the_slide_title() -> None:
    from flexo_talk import compose

    deck = Deck("heading")
    with deck.slide("Where the brakes sit", layout="figure") as slide, slide.figure() as figure:
        figure.protein("prss1", 247, [{"type": "domain", "label": "Protease", "start": 24, "end": 247}],
                       label="PRSS1 (cationic trypsinogen)")
    svg = compose.render_slide(deck, deck.slides[0]).svg
    scale = float(re.search(r'<g id="slide1\.body\.0" transform="[^"]*scale\(([\d.]+)\)"', svg).group(1))
    size = float(re.search(r'<text id="[^"]*prss1\.label"[^>]*font-size="([\d.]+)"', svg).group(1))
    # Its bold name no larger than the body's words, however large the figure is let be.
    assert size * scale <= deck.style.body_size + 0.01 < deck.style.title_size


def test_a_png_picture_is_embedded(tmp_path: Path) -> None:
    import struct
    import zlib

    def chunk(tag: bytes, payload: bytes) -> bytes:
        return struct.pack(">I", len(payload)) + tag + payload + struct.pack(">I", zlib.crc32(tag + payload))

    rows = b"".join(b"\x00" + bytes([30, 120, 200] * 4) for _ in range(3))
    source = tmp_path / "p.png"
    source.write_bytes(
        b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 4, 3, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(rows)) + chunk(b"IEND", b"")
    )
    deck = Deck("pictures")
    with deck.slide("A picture") as slide:
        slide.image(source)
    result = deck.build(tmp_path, formats=("pptx",))
    assert "<p:pic>" in _slides(result.pptx)[0]  # type: ignore[arg-type]
    assert any(name.startswith("ppt/media/") for name in zipfile.ZipFile(result.pptx).namelist())  # type: ignore[arg-type]


def test_a_table_is_a_native_table_ruled_as_in_a_paper(tmp_path: Path) -> None:
    deck = Deck("tables")
    with deck.slide("Results", layout="two-columns", split=0.6) as slide:
        slide.left.table([["Model", "Top-1"], ["ViT", "**81.8**"], ["Ours ($\\lambda$)", "81.2"]])
        slide.right.bullets("Numbers are set flush right")
    result = deck.build(tmp_path, formats=("pptx", "svg"))
    slide = _slides(result.pptx)[0]  # type: ignore[arg-type]
    assert "<a:tbl>" in slide and slide.count("<a:tr ") == 3 and slide.count("<a:gridCol ") == 2
    # Booktabs: a rule above, one under the header, one below; nothing else.
    assert slide.count("<a:lnT w=\"13970\"") + slide.count("<a:lnB w=\"13970\"") >= 2
    # Italic Greek the words' face lacks is the maths font's (TeX's) letter.
    assert 'algn="r"' in slide and "\U0001d706" in slide
    # The drawn copy of the table is not also in the PowerPoint.
    assert slide.count(">Model<") == 1


def test_bold_in_a_bold_statement_shows_in_the_accent(tmp_path: Path) -> None:
    deck = Deck("strong")
    deck.statement("A queue lets the customer **stop waiting** for us.")
    result = deck.build(tmp_path, formats=("pptx", "svg"))
    svg = result.svgs[0].read_text()
    assert re.search(r'data-flexo-fill="tone-1-stroke"[^>]*>stop waiting<', svg)
    accent = re.search(r'fill="#([0-9a-f]{6})"[^>]*>stop waiting<', svg).group(1).upper()
    slide = _slides(result.pptx)[0]  # type: ignore[arg-type]
    assert re.search(rf'<a:srgbClr val="{accent}"/>.{{0,200}}<a:t>stop waiting</a:t>', slide)
    # A title already in the accent (the Margin look's) has its strong words in the ink.
    margin = Deck("edge", look="margin")
    margin.slide("Why **we** moved")
    drawn = margin.render()[0].svg
    assert re.search(r'data-flexo-fill="ink"[^>]*>we<', drawn)


def test_a_change_column_of_percentages_and_multiples_is_numbers() -> None:
    from flexo_talk.deck import _numeric

    # Set flush right with the numbers: a multiple before its number as well as after it.
    assert all(_numeric(cell) for cell in ["\u221298%", "-97%", "\u00d77", "x7", "7\u00d7", "+5%"])
    assert not any(_numeric(cell) for cell in ["x", "\u00d7", "Change"])


def test_titles_words_and_figures_each_take_their_own_family(tmp_path: Path) -> None:
    from flexo_talk import DeckStyle

    deck = Deck(
        "type", font="IBM Plex Sans", title_font="Latin Modern Roman", figure_font="Figtree",
        style=DeckStyle(title_align="middle", title_role="tone-1-stroke"),
    )
    with deck.slide("A title", layout="two-columns") as slide:
        slide.left.bullets("Words and `code`")
        with slide.right.figure() as figure:
            figure.block("b", label="Block")
    result = deck.build(tmp_path, formats=("pptx", "svg"))
    svg = result.svgs[0].read_text()
    assert 'font-family="Latin Modern Roman"' in svg and 'text-anchor="middle"' in svg
    assert 'font-family="Figtree"' in svg
    slide = _slides(result.pptx)[0]  # type: ignore[arg-type]
    assert 'typeface="LM Roman 10"' in slide and 'typeface="IBM Plex Sans"' in slide


def test_a_tall_figure_is_laid_out_for_a_wide_slide(tmp_path: Path) -> None:
    deck = Deck("tall")
    with deck.slide("A stack", layout="figure") as slide, slide.figure() as figure, figure.column(
        "layers", reverse=True
    ) as layers:
        previous = layers.text("x", "$x$")
        for index in range(6):
            previous = layers.block(f"b{index}", label=f"Layer {index}", input=previous)
    result = deck.build(tmp_path, formats=("svg",))
    # Said as a person sees it: the chart turned to run across, not rows and columns swapped.
    assert any(note.endswith("To fit the slide, the chart was turned to run across.") for note in result.notes)
    assert not any("too small" in message for message in result.diagnostics)


def test_a_figure_leaves_its_own_page_behind(tmp_path: Path) -> None:
    # Its page would reach past what it draws (off the slide, in a column), and be
    # what the studio frames and clicks.
    deck = Deck("page")
    with deck.slide("Two sides", layout="two-columns") as slide:
        slide.left.bullets("A point")
        with slide.right.figure() as figure:
            x = figure.root.text("x", "Input $x$")
            figure.root.block("model", label="Model", input=x)
    result = deck.build(tmp_path, formats=("svg",))
    svg = result.svgs[0].read_text()
    assert 'id="slide1.right.0.root"' in svg
    assert "slide1.right.0.layer.background" not in svg
    assert "slide1.right.0.canvas.background" not in svg


def test_a_code_listing_is_set_in_monospace(tmp_path: Path) -> None:
    deck = Deck("code")
    with deck.slide("Code") as slide:
        slide.code("""
            def f(x):
                # a comment
                return x + 1
        """)
    svg = deck.build(tmp_path, formats=("svg",)).svgs[0].read_text()
    assert "def f(x):" in svg and 'xml:space="preserve"' in svg
    assert 'data-flexo-fill="muted-ink"' in svg


def test_a_list_in_another_script_names_a_face_that_has_it(tmp_path: Path) -> None:
    from flexo.fonts import family_covering

    if family_covering("한국어") is None:
        pytest.skip("no installed font has Hangul")
    deck = Deck("scripts")
    with deck.slide("Scripts") as slide:
        slide.bullets("한국어 문장도 됩니다", "See https://github.com/jamaliki/flexo-talk/blob/main/src/flexo_talk/compose.py")
    result = deck.build(tmp_path, formats=("pptx", "svg"))
    slide = _slides(result.pptx)[0]  # type: ignore[arg-type]
    run = re.search(r'<a:r>(?:(?!</a:r>).)*한국어(?:(?!</a:r>).)*</a:r>', slide)
    assert run is not None and 'typeface="Figtree"' not in run.group(0)


def test_a_numbered_list_is_numbered_natively(tmp_path: Path) -> None:
    deck = Deck("numbers")
    with deck.slide("Steps") as slide:
        slide.bullets("First", ["a detail"], "Second", numbered=True)
    result = deck.build(tmp_path, formats=("pptx", "svg"))
    slide = _slides(result.pptx)[0]  # type: ignore[arg-type]
    # Every level numbered in its tier: the detail is "a.", not a bullet.
    assert slide.count('<a:buAutoNum type="arabicPeriod"/>') == 2
    assert slide.count('<a:buAutoNum type="alphaLcPeriod"/>') == 1
    assert "<a:buChar" not in slide
    svg = result.svgs[0].read_text()
    assert ">1.<" in svg and ">a.<" in svg and ">2.<" in svg


def test_a_plain_list_has_no_marks_and_keeps_its_levels(tmp_path: Path) -> None:
    from flexo_talk.document import deck_document, deck_from_document

    deck = Deck("plain")
    with deck.slide("Points") as slide:
        slide.bullets("First", ["a detail"], "Second", plain=True)
    result = deck.build(tmp_path, formats=("pptx", "svg"))
    svg = result.svgs[0].read_text()
    # No bullets drawn; the detail still a level in.
    assert "<circle" not in svg and ">1.<" not in svg
    xs = {n: float(x) for n, x in re.findall(r'id="slide1\.body\.0\.(\d)" x="([\d.]+)"', svg)}
    assert xs["1"] > xs["0"] == xs["2"]
    slide_xml = _slides(result.pptx)[0]  # type: ignore[arg-type]
    assert "<a:buChar" not in slide_xml and "<a:buAutoNum" not in slide_xml and 'lvl="1"' in slide_xml
    # Written as it was read.
    document = {"schema_version": 1, "deck": {"id": "plain"}, "slides": [
        {"title": "Points", "body": [{"bullets": ["First", ["a detail"]], "plain": True}]}]}
    assert deck_document(deck_from_document(document))["slides"] == document["slides"]


def test_a_list_in_a_colour_of_its_own_keeps_it_in_the_pptx(tmp_path: Path) -> None:
    deck = Deck("coloured")
    with deck.slide("Points") as slide:
        slide.bullets("First", ["a detail"], colour="#c0392b")
    result = deck.build(tmp_path, formats=("pptx", "svg"))
    svg = result.svgs[0].read_text()
    # The words in its colour, the bullets in the theme's.
    assert 'fill="#c0392b"' in svg.lower()
    slide_xml = _slides(result.pptx)[0]  # type: ignore[arg-type]
    assert slide_xml.count('<a:srgbClr val="C0392B"/>') >= 2


def test_a_revealed_list_builds_click_by_click(tmp_path: Path) -> None:
    deck = Deck("reveal")
    with deck.slide("Steps") as slide:
        slide.bullets("First", ["under the first"], "Second", "Third", reveal=True)
    with deck.slide("Plain") as slide:
        slide.bullets("No steps")
    result = deck.build(tmp_path, formats=("pptx", "pdf"))
    slide = _slides(result.pptx)[0]  # type: ignore[arg-type]
    assert slide.count('nodeType="clickEffect"') == 3 and '<p:pRg st="0" end="1"/>' in slide
    # Title alone, then one item at a time: four pages, and one for the plain slide.
    assert result.pdf is not None and result.pdf.read_bytes().count(b"/Type /Page ") == 5
    handout = deck.build(tmp_path / "handout", formats=("pdf",), handout=True)
    assert handout.pdf is not None and handout.pdf.read_bytes().count(b"/Type /Page ") == 2


def test_footnotes_sit_above_the_footer_and_shorten_the_body(tmp_path: Path) -> None:
    deck = Deck("notes", footer="Footer")
    with deck.slide("Cited") as slide:
        slide.bullets(*(["a line of text that takes its room"] * 14))
        slide.footnote("[1] A reference.")
    result = deck.build(tmp_path, formats=("svg",))
    svg = result.svgs[0].read_text()
    assert "[1] A reference." in svg
    assert any("fit" in message for message in result.diagnostics)


def test_links_are_native_hyperlinks_in_the_pptx(tmp_path: Path) -> None:
    deck = Deck("links")
    with deck.slide("Links") as slide:
        slide.bullets("Code: [the repository](https://github.com/jamaliki/flexo)")
        slide.footnote("[1] [*A paper*](https://arxiv.org/abs/1706.03762)")
    result = deck.build(tmp_path, formats=("pptx", "pdf"))
    archive = zipfile.ZipFile(result.pptx)  # type: ignore[arg-type]
    relations = archive.read("ppt/slides/_rels/slide1.xml.rels").decode()
    assert "https://github.com/jamaliki/flexo" in relations and "https://arxiv.org/abs/1706.03762" in relations
    assert archive.read("ppt/slides/slide1.xml").decode().count("<a:hlinkClick") == 2
    assert result.pdf is not None and result.pdf.read_bytes().count(b"/Subtype /Link") == 2


def test_a_persian_list_is_set_right_to_left(tmp_path: Path) -> None:
    deck = Deck("fa")
    with deck.slide("سلام دنیا") as slide:
        slide.bullets("این یک جمله است", "English و فارسی")
    result = deck.build(tmp_path, formats=("pptx", "svg"))
    slide = _slides(result.pptx)[0]  # type: ignore[arg-type]
    # The title and the Persian item read right to left; the item starting in English does not.
    assert slide.count('rtl="1"') == 2 and 'algn="r"' in slide
    svg = result.svgs[0].read_text()
    assert 'text-anchor="end"' in svg


def test_an_acknowledgement_slide_has_columns_colours_and_a_gallery(tmp_path: Path) -> None:
    pytest.importorskip("PIL")
    from PIL import Image

    photo = tmp_path / "face.jpg"
    Image.new("RGB", (300, 400), (200, 150, 120)).save(photo)
    deck = Deck("thanks")
    with deck.slide("Thanks", layout="columns", widths=(2, 1, 1)) as slide:
        first, second, third = slide.columns
        first.text("**Lab** [— PI]{muted}", colour="accent")
        second.bullets("[Emmy]{accent2}", "[Bob]{#c0392b}")
        third.gallery([(photo, "**Ada**\\nPI")], crop="circle")
    result = deck.build(tmp_path, formats=("pptx", "svg"))
    svg = result.svgs[0].read_text()
    assert 'data-flexo-talk="gallery"' in svg and "data:image/png;base64" in svg
    slide = _slides(result.pptx)[0]  # type: ignore[arg-type]
    assert "C0392B" in slide and "<p:pic>" in slide and ">Ada<" in slide


def test_a_figure_laid_out_once_is_reused(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    import flexo

    monkeypatch.setenv("FLEXO_TALK_CACHE", str(tmp_path / "cache"))
    calls = []
    real = flexo.fit_in_box

    def counting(*args, **kwargs):
        calls.append(1)
        return real(*args, **kwargs)

    monkeypatch.setattr(flexo, "fit_in_box", counting)
    for _ in range(2):
        deck = Deck("cached")
        with deck.slide("A figure") as slide, slide.figure() as figure:
            figure.block("b", label="Block", input=figure.text("x", "$x$"))
        deck.render()
    assert len(calls) == 1 and list((tmp_path / "cache" / "fits").glob("*.json"))


def test_a_persian_table_reads_from_the_right(tmp_path: Path) -> None:
    deck = Deck("fa-table")
    with deck.slide("Table") as slide:
        slide.table([["مدل", "دقت"], ["الف", "81.2"]])
    result = deck.build(tmp_path, formats=("pptx", "svg"))
    slide = _slides(result.pptx)[0]  # type: ignore[arg-type]
    header = re.search(r"<a:tr .*?</a:tr>", slide)
    assert header is not None
    cells = re.findall(r"<a:t>([^<]*)</a:t>", header.group(0))
    assert cells == ["دقت", "مدل"]  # the first column is written last


def test_a_slide_takes_its_own_background(tmp_path: Path) -> None:
    pytest.importorskip("PIL")
    from PIL import Image

    photo = tmp_path / "wide.jpg"
    Image.new("RGB", (1600, 600), (40, 80, 120)).save(photo)
    deck = Deck("backgrounds")
    deck.title("On a photograph", background=photo, shade=0.5)
    with deck.slide("On a colour", background="#1b2a41") as slide:
        slide.bullets("Light words, lighter accents")
    result = deck.build(tmp_path, formats=("pptx", "svg", "pdf"))
    first, second = _slides(result.pptx)  # type: ignore[arg-type]
    # The photograph is cropped to the slide natively, not stretched.
    assert "<a:srcRect" in first and "<p:pic>" in first
    assert "1B2A41" in second
    assert 'fill="#f7f5f0"' in result.svgs[1].read_text()
    # The native list on a dark slide is set in the slide's light ink, not the page's.
    assert "F7F5F0" in second


def _texts(svg: str, pattern: str) -> list[tuple[float, float]]:
    """The ``(x, y)`` of every text element whose id matches ``pattern``."""

    import xml.etree.ElementTree as ET

    found = []
    for item in ET.fromstring(svg).iter():
        if item.tag.endswith("text") and re.search(pattern, item.get("id", "")):
            found.append((float(item.get("x", 0)), float(item.get("y", 0))))
    return found


@pytest.mark.parametrize("look", ["classic", "band", "editorial", "keynote", "margin"])
def test_every_look_builds_every_kind_of_slide(tmp_path: Path, look: str) -> None:
    deck = Deck("looks", look=look)
    deck.title("A talk", subtitle="In one look", author="Ada")
    deck.agenda()
    deck.section("First part", subtitle="Why")
    deck.statement("One claim, [large]{accent}.", by="someone")
    with deck.slide("Numbers") as slide:
        slide.stats(("93%", "accuracy"), ("4x", "faster"))
        slide.callout("The *key* point.", title="Idea", colour="accent2")
    with deck.slide("Words") as slide:
        slide.quote("A thing worth saying.", by="a reader")
    deck.section("Second part")
    result = deck.build(tmp_path, formats=("pptx", "svg"))
    assert not result.diagnostics, result.summary()
    agenda = result.svgs[1].read_text()
    assert "First part" in agenda and "Second part" in agenda and ">02<" in agenda
    slides = _slides(result.pptx)  # type: ignore[arg-type]
    assert ">93%<" in slides[4] and ">Idea<" in slides[4]
    assert "“" in slides[5] and "a reader" in slides[5]


def test_a_filled_section_takes_the_accent_and_light_words(tmp_path: Path) -> None:
    deck = Deck("fill", look="band")
    deck.section("A part")
    with deck.slide("After it") as slide:
        slide.bullets("An item")
    result = deck.build(tmp_path, formats=("pptx", "svg"))
    section, _ = _slides(result.pptx)  # type: ignore[arg-type]
    accent = deck.palette.get("tone-1-stroke").lstrip("#").upper()
    assert re.search(rf'<p:bg>.*?<a:srgbClr val="{accent}"/>', section, re.S)
    from flexo.colour import contrast

    # Its words set light, and readable on the accent.
    words = re.findall(r'<a:t>A part</a:t>', section)
    run = r'<a:rPr[^>]*>(?:(?!</a:rPr>).)*?<a:srgbClr val="([0-9A-F]{6})"/>(?:(?!</a:rPr>).)*?</a:rPr><a:t>A part'
    colours = re.findall(run, section, re.S)
    assert words and colours and contrast(f"#{colours[0]}", f"#{accent}") >= 4.5
    # The content slide's title sits on a band of the accent across the top.
    assert 'id="slide2.band"' in result.svgs[1].read_text()


def test_words_start_at_the_top_and_a_shorter_picture_is_centred_against_them(tmp_path: Path) -> None:
    def tall(figure) -> None:
        previous = figure.block("a", label="A")
        for name in "bcdef":
            previous = figure.block(name, label=name.upper(), input=previous)

    deck = Deck("aligned")
    with deck.slide("Beside a taller picture", layout="two-columns") as slide:
        slide.left.bullets("One", "Two")
        with slide.right.figure(turn=False) as figure:
            tall(figure)
    with deck.slide("Beside a shorter picture", layout="two-columns") as slide:
        slide.left.bullets(*(f"Point {index}" for index in range(8)))
        with slide.right.figure() as figure:
            figure.block("a", label="A")
    with deck.slide("Middle", layout="two-columns", align="middle") as slide:
        slide.left.bullets("One", "Two")
        with slide.right.figure(turn=False) as figure:
            tall(figure)
    result = deck.build(tmp_path, formats=("pptx", "svg"))
    taller, shorter, middle = (path.read_text() for path in result.svgs)

    def moved(svg: str, region: str) -> bool:
        tag = re.search(rf'<g id="{re.escape(region)}"[^>]*>', svg)
        return tag is not None and 'transform="translate(0' in tag.group(0)

    # The words start where the body starts, whatever is beside them.
    assert not moved(taller, "slide1.left") and not moved(shorter, "slide2.left")
    assert moved(shorter, "slide2.right") and moved(middle, "slide3.left")
    # The native list moves with its region: its box starts lower on the middle slide.
    first, _, third = _slides(result.pptx)  # type: ignore[arg-type]
    box = r'name="List[^"]*".*?<a:off x="\d+" y="(\d+)"'
    offsets = [int(re.search(box, s, re.S).group(1)) for s in (first, third)]  # type: ignore[union-attr]
    assert offsets[1] > offsets[0]


def test_a_single_column_gallery_is_flush_with_the_words(tmp_path: Path) -> None:
    pytest.importorskip("PIL")
    from PIL import Image

    logo = tmp_path / "logo.png"
    Image.new("RGB", (200, 80), (30, 90, 160)).save(logo)
    deck = Deck("logos")
    with deck.slide("Funding") as slide:
        slide.text("**Funded by**")
        slide.gallery([logo, logo], columns=1, height=30)
    result = deck.build(tmp_path, formats=("svg",))
    svg = result.svgs[0].read_text()
    xs = {float(x) for x in re.findall(r'<image id="slide1\.body\.1\.\d" x="([\d.]+)"', svg)}
    assert xs == {deck.style.margin}


def test_a_look_is_a_style_preset() -> None:
    from flexo_talk import DeckStyle

    assert DeckStyle.look("band").header == "band"
    assert Deck(look="editorial").style.sections == "number"
    assert DeckStyle.look("keynote", body_size=24).body_size == 24
    with pytest.raises(ValueError, match="Available looks"):
        Deck(look="fancy")
    with pytest.raises(ValueError, match="callout"):
        Deck().slide("x").callout("words", colour="red")


def _word_sizes(svg: str) -> dict[str, set[float]]:
    """The sizes words are drawn at, by the region they are in."""

    from flexo.drawing import Text, read_drawing

    sizes: dict[str, set[float]] = {}
    for item in read_drawing(svg).walk():
        if isinstance(item, Text) and item.id and item.id.count(".") > 2:
            region = ".".join(item.id.split(".")[:2])
            sizes.setdefault(region, set()).update(round(run.size, 2) for line in item.lines for run in line.runs)
    return sizes


def test_figures_side_by_side_set_their_words_at_one_size(tmp_path: Path) -> None:
    deck = Deck("shared")
    with deck.slide("Two figures", layout="two-columns") as slide:
        with slide.left.figure() as figure:
            a = figure.block("a", label="Short")
            figure.block("b", label="Pair", input=a)
        with slide.right.figure() as figure:
            previous = figure.text("x", "$x$")
            for index in range(6):
                previous = figure.block(f"b{index}", label=f"Layer {index}", input=previous)
    result = deck.build(tmp_path, formats=("svg",))
    sizes = _word_sizes(result.svgs[0].read_text())
    # The short pair alone would be set at the body size; beside the tall stack it matches it.
    assert sizes["slide1.left"] == sizes["slide1.right"]
    assert max(sizes["slide1.left"]) < deck.style.body_size


def test_a_figure_given_a_width_is_drawn_that_wide_as_far_as_its_place_allows(tmp_path: Path) -> None:
    def sizes(width: float | None) -> set[float]:
        deck = Deck(f"wide{width}")
        with deck.slide("Sized", layout="two-columns") as slide:
            slide.left.bullets("A point")
            with slide.right.figure(width=width) as figure:
                a = figure.block("a", label="Short")
                figure.block("b", label="Pair", input=a)
        result = deck.build(tmp_path / str(width), formats=("svg",))
        return _word_sizes(result.svgs[0].read_text())["slide1.right"]

    (narrow,), (twice,), (auto,), (huge,) = sizes(100), sizes(200), sizes(None), sizes(4000)
    # Its words are as large as that width makes them: past the body size too.
    assert twice == pytest.approx(2 * narrow, rel=0.02)
    assert auto <= Deck("x").style.body_size < huge
    # No wider than its column.
    assert huge == sizes(3000).pop()


def test_a_column_of_numbers_with_gaps_is_still_set_flush_right() -> None:
    deck = Deck("gaps")
    with deck.slide("Results") as slide:
        slide.table([["", "Ours", "Theirs"], ["ImageNet", "88.55", "88.4"], ["CIFAR-100", "94.55", "\N{EN DASH}"]])
    (table,) = slide.body.blocks
    assert table.align == ("start", "end", "end")


def test_a_callout_fills_its_panel_line_by_line(tmp_path: Path) -> None:
    from flexo.drawing import Text, read_drawing

    words = "DeiT trained ViTs on ImageNet alone, with strong augmentation and distillation from a CNN teacher."
    deck = Deck("panel")
    with deck.slide("Next", layout="two-columns") as slide:
        slide.left.bullets("One")
        slide.right.callout(words, title="What came next")
    result = deck.build(tmp_path, formats=("svg",))
    (text,) = [
        item for item in read_drawing(result.svgs[0].read_text()).walk()
        if isinstance(item, Text) and item.id == "slide1.right.0.words"
    ]
    first = "".join(run.text for run in text.lines[0].runs)
    # Greedy: the first line takes all it can, as a paragraph does; balanced lines leave the panel ragged.
    assert len(first) > len(words) / len(text.lines) + 5


def test_an_equation_is_displayed_on_its_own_line_and_fits_its_place(tmp_path: Path) -> None:
    from flexo_talk.document import deck_document, deck_from_document

    long = r" + ".join(rf"\frac{{a_{i}}}{{b_{i}}}" for i in range(40))
    document = {"deck": {"id": "maths"}, "slides": [
        {"title": "Loss", "body": [
            {"text": "We minimise"},
            {"math": r"\mathcal{L} = -\frac{1}{N}\sum_{i=1}^{N} \log p(y_i \mid x_i)"},
            {"text": r"$$\begin{pmatrix} a & b \\ c & d \end{pmatrix}$$"},
            {"math": long},
        ]},
    ]}
    deck = deck_from_document(document, tmp_path)
    kinds = [type(block).__name__ for block in deck.slides[0].regions["body"].blocks]
    assert kinds == ["_Words", "_Math", "_Math", "_Math"]
    # The document keeps what was written.
    assert deck_document(deck)["slides"][0]["body"][2] == {"text": r"$$\begin{pmatrix} a & b \\ c & d \end{pmatrix}$$"}
    result = deck.build(tmp_path / "out", formats=("svg", "pptx"))
    # Only the long one is said: set that small, it would read better broken into lines.
    assert len(result.diagnostics) == 1 and "Equation reduced to" in result.diagnostics[0], result.summary()
    svg = result.svgs[0].read_text()
    groups = re.findall(r'<g [^>]*data-flexo-math="[^"]*"[^>]*>', svg)
    assert len(groups) == 3 and all('data-flexo-talk="math"' in group for group in groups)
    # The long one is set smaller to fit the slide, not run off it.
    from flexo.drawing import read_drawing

    def formulas(group):
        for item in group.items:
            if hasattr(item, "items"):
                if "data-flexo-math" in item.data:
                    yield item
                else:
                    yield from formulas(item)

    def shapes(group):
        for item in group.items:
            yield from shapes(item) if hasattr(item, "items") else [item]

    last = list(formulas(read_drawing(svg).root))[-1]
    right = max(item.x + item.width for item in shapes(last))
    assert right <= deck.style.width - deck.style.margin + 1.0
    slide = _slides(result.pptx)[0]
    assert slide.count("<a:custGeom>") > 50  # formulas are native shapes, glyph by glyph


def test_maths_in_a_list_leaves_room_in_the_native_words_and_is_drawn_over_it(tmp_path: Path) -> None:
    deck = Deck("maths")
    deck.slide("Rates").bullets(
        r"The rate $k = A e^{-E_a/RT}$ rises with temperature",
        r"A plain one: $x_t$",
    )
    deck.slide("Table").table([["what", "value"], ["half", r"$\frac{1}{2}$"]])
    result = deck.build(tmp_path, formats=("pptx",))
    listing, table = _slides(result.pptx)
    # One space, spaced out to the formula's width, keeps its place in the words...
    gaps = re.findall(r'<a:rPr [^>]*spc="(\d+)"[^>]*>.*?</a:rPr><a:t> </a:t>', listing)
    assert len(gaps) == 1 and int(gaps[0]) > 1000
    assert "rises with temperature" in listing and "<a:t>t</a:t>" in listing
    # ...and the formula is drawn as shapes where it was set.
    assert listing.count("<a:custGeom>") > 5
    assert re.search(r'spc="\d+"', table) and table.count("<a:custGeom>") >= 2


@pytest.mark.parametrize("editable", [False, True])
def test_a_list_opened_for_tall_maths_keeps_its_baselines_in_powerpoint(tmp_path: Path, editable: bool) -> None:
    from flexo.text import font_stack

    from flexo_talk.pptx import LOOSE, ascent

    deck = Deck("tall")
    deck.slide("Tall").bullets(
        r"First: $\frac{\partial \mathcal{L}}{\partial \theta} = \sum_i \frac{1}{p_i}$ is tall",
        "A plain second item",
        r"A long item that wraps onto a second line because it goes on and on, with $\frac{a}{b}$ in "
        "the middle of it and more words after",
    )
    (rendered,) = deck.render()
    (layout,) = rendered.lists
    assert [count for _, _, count in layout.opened] == [1, 1, 2]
    xml = _slides(deck.build(tmp_path, formats=("pptx",), editable_maths=editable).pptx)[0]
    share = ascent(font_stack(deck.typography(layout.size)).face(400, False))
    shapes = re.findall(r'<p:sp>(?:(?!</p:sp>).)*name="List “First.*?</p:sp>', xml, re.S)
    assert len(shapes) == (2 if editable else 1)
    for shape in shapes:
        # PowerPoint sets an exactly spaced line's baseline at the face's share of single
        # spacing, and at LOOSE of any wider spacing: replayed, each item's first baseline
        # lands where Flexo set it.
        y = int(re.search(r'<a:off x="\d+" y="(\d+)"/>', shape).group(1)) / 12700
        spacings = [int(v) / 100 for v in re.findall(r'<a:lnSpc><a:spcPts val="(\d+)"/>', shape)]
        befores = [int(v) / 100 for v in re.findall(r'<a:spcBef><a:spcPts val="(\d+)"/>', shape)]
        for index, ((_, _, baseline), (rise, fall, count), spacing, before) in enumerate(
            zip(layout.items, layout.opened, spacings, befores, strict=True)
        ):
            y += before
            plain = layout.step(index) - rise - fall
            down = LOOSE * spacing if spacing > plain + 0.01 else share * spacing
            assert y + down == pytest.approx(baseline, abs=0.02)
            y += spacing * count


def test_a_text_box_starts_where_its_words_were_set(tmp_path: Path) -> None:
    deck = Deck("edges", footer="A footer")
    deck.slide("A title").text("Words.")
    xml = _slides(deck.build(tmp_path, formats=("pptx",)).pptx)[0]
    for words in ("A title", "Words.", "A footer"):
        shape = re.search(rf'<p:sp>(?:(?!</p:sp>).)*<a:t>{re.escape(words)}</a:t>', xml, re.S).group(0)
        # Left-aligned words start at the margin; the room to spare is on their right.
        assert int(re.search(r'<a:off x="(\d+)"', shape).group(1)) / 12700 == pytest.approx(48.0, abs=0.01)
    number = re.search(r'<p:sp>(?:(?!</p:sp>).)*algn="r"(?:(?!</p:sp>).)*<a:t>1</a:t>', xml, re.S).group(0)
    left, width = (int(v) / 12700 for v in re.search(r'<a:off x="(\d+)".*?<a:ext cx="(\d+)"', number, re.S).groups())
    assert left + width == pytest.approx(960.0 - 48.0, abs=0.5)


def test_maths_that_cannot_be_read_is_reported_on_its_slide(tmp_path: Path) -> None:
    deck = Deck("broken")
    deck.slide("Oops").text(r"Here: $\frac{1}{2} + \foo{x}$").math(r"\sqrt{x")
    # Maths simple enough to be set as words is read for its mistakes all the same.
    deck.slide("Squared").text(r"Then $x^{2$ is squared.")
    result = deck.build(tmp_path, formats=("svg",))
    said = " ".join(result.diagnostics)
    assert r"\foo is not a maths command flexo knows" in said
    assert "a { is not closed" in said and "slide1" in said
    assert "slide2: a { is not closed, in the maths \u201cx^{2\u201d" in said
    assert "Traceback" not in said and "Error" not in said


def test_escaped_backticks_and_brackets_are_themselves() -> None:
    runs = inline(r"Use \`ls\` and [box\](x) and [red\]{accent}, but `code`.")
    assert [(run.text, run.code, run.link, run.color) for run in runs] == [
        ("Use `ls` and [box](x) and [red]{accent}, but ", False, "", ""),
        ("code", True, "", ""),
        (".", False, "", ""),
    ]


def test_code_as_the_studio_writes_it_keeps_its_backticks_and_its_look() -> None:
    # A backtick in code is written \` in it; code bold, italic, coloured or linked is
    # written round the backticks.
    runs = inline(r"Run `a\`b`, **`bold`**, *`slanted`*, [`red`]{accent} and [`see`](https://example.com).")
    assert [(run.text, run.weight, run.italic, run.color, run.link) for run in runs if run.code] == [
        ("a`b", 400, False, "", ""), ("bold", 700, False, "", ""), ("slanted", 400, True, "", ""),
        ("red", 400, False, "accent", ""), ("see", 400, False, "", "https://example.com")]


def test_bold_and_italic_that_meet_inside_a_word_are_read_as_the_studio_writes_them() -> None:
    # "Second para" bold and "paragraph." italic: the studio closes the bold alone there,
    # never a close and a reopen run together (which read as ****).
    runs = inline("**Second *para**graph.*")
    assert [(run.text, run.weight, run.italic) for run in runs] == [
        ("Second ", 700, False), ("para", 700, True), ("graph.", 400, True)]
    runs = inline("**bold *both*** *italic*")
    assert [(run.text, run.weight, run.italic) for run in runs if run.text.strip()] == [
        ("bold ", 700, False), ("both", 700, True), ("italic", 400, True)]


def test_a_typed_backslash_before_a_bracket_is_no_maths() -> None:
    runs = inline(r"Type \\(x\) or \\[y\\] for maths, and \(z\) is maths.")
    assert "".join(run.text for run in runs if not run.maths).startswith(r"Type \(x\) or \[y\] for maths, and ")
    assert [run.text for run in runs if run.maths] == ["z"]


def test_a_link_and_a_colour_can_hold_each_other_and_maths() -> None:
    for words in ("[[First]{accent}](https://example.com)", "[[First](https://example.com)]{accent}"):
        runs = inline(words)
        assert [(run.text, run.link, run.color) for run in runs] == [("First", "https://example.com", "accent")]
    runs = inline("[the loss $L$]{accent2} and [see $x$](https://example.com)")
    assert all(run.color == "accent2" for run in runs[:2]) and runs[1].maths
    assert [run.link for run in runs if run.maths] == ["", "https://example.com"]


def test_prices_are_prices_and_escaped_dollars_are_dollars() -> None:
    runs = inline(r"It costs $5 and $10, a sample is \$20, and $x^2$ is maths.")
    assert "".join(run.text for run in runs if not run.italic).startswith("It costs $5 and $10, a sample is $20")
    assert [run.text for run in runs if run.italic] == ["x"]


def test_a_plot_sets_its_maths_as_a_slide_does_and_says_what_it_cannot_read() -> None:
    pytest.importorskip("matplotlib")
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    from flexo_talk.compose import plot_svg

    # Maths matplotlib cannot set at all: a matrix, \le, a fraction in a legend.
    figure, axes = plt.subplots()
    axes.plot([0, 1], [0, 1], label=r"$\beta = \frac{1}{2}$")
    axes.set_title(r"$\begin{pmatrix} a & b \\ c & d \end{pmatrix}$ and $x \le y$")
    axes.set_ylabel(r"$\frac{1}{$")
    axes.legend()
    said: list[tuple[str, str]] = []
    svg = plot_svg(figure, 300.0, 200.0, "Figtree", "p", maths="Fira Math", said=said)
    assert said == [(r"$\frac{1}{$", "a { is not closed")]
    assert "STIX" not in svg and "<path" in svg


def test_a_plot_annotation_with_maths_keeps_its_arrow_and_box() -> None:
    import matplotlib.pyplot as plt

    from flexo_talk.compose import plot_svg

    def drawn(maths: bool) -> str:
        figure, axes = plt.subplots()
        label = r"peak $\alpha$" if maths else "peak"
        axes.annotate(label, xy=(0.5, 0.5), xytext=(0.1, 0.8), arrowprops={"arrowstyle": "->"})
        axes.text(0.6, 0.2, r"box $\beta$" if maths else "box", bbox={"boxstyle": "round", "fc": "#ffff00"})
        return plot_svg(figure, 300.0, 200.0, "Figtree", "a", maths="Fira Math")

    with_maths, without = drawn(True), drawn(False)
    assert "#ffff00" in with_maths.lower()
    # The arrow is there as it is for plain words (one path more for the formula's glyphs at most).
    assert with_maths.count("stroke-linecap") >= without.count("stroke-linecap")


def test_a_log_axis_ticks_are_matplotlibs_and_say_nothing() -> None:
    import matplotlib.pyplot as plt
    import numpy as np

    from flexo_talk.compose import plot_svg

    figure, (left, right) = plt.subplots(1, 2)
    left.loglog(np.logspace(-2, 4, 20), np.logspace(-2, 4, 20), label=r"$\kappa \le 1$")
    left.set_xlim(1e-3, 1e5)
    figure.canvas.draw()  # ticks made once (as tight_layout would), some then out of view
    left.set_xlim(1e-2, 1e3)
    left.legend()
    right.semilogy([1, 2], [1, 1e5])
    right.secondary_yaxis("right")
    said: list[tuple[str, str]] = []
    plot_svg(figure, 300.0, 200.0, "Figtree", "p", maths="Fira Math", said=said)
    assert said == []


@pytest.mark.parametrize(
    ("source", "elements"),
    [
        (r"\frac{a}{b}", ["m:f", "m:num", "m:den"]),
        (r"\sqrt{x} + \sqrt[3]{y}", ["m:rad", "m:degHide", "m:deg"]),
        (r"\sum_{i=1}^{n} x_i", ["m:nary", 'm:chr m:val="∑"', 'm:limLoc m:val="undOvr"', "m:sSub"]),
        (r"\sum_{i=1}^{n}", ["m:nary", "<m:e><m:r>", "\u200b"]),
        (r"\int_0^1 f\,dx", ['m:chr m:val="∫"', 'm:limLoc m:val="subSup"']),
        (r"\left( \frac{a}{b} \right)", ["m:d", 'm:begChr m:val="("', 'm:endChr m:val=")"']),
        (r"\begin{pmatrix} a & b \\ c & d \end{pmatrix}", ["m:m", "m:mr", 'm:begChr m:val="("']),
        (r"\begin{cases} x & x \ge 0 \\ -x & x < 0 \end{cases}", ['m:begChr m:val="{"', 'm:endChr m:val=""']),
        (r"a &= b \\ &= c", ["m:eqArr", "m:aln"]),
        (r"\begin{gathered} a \\ b \end{gathered}", ["m:eqArr"]),
        (r"\log p(x) + \sin x", ["m:func", "m:fName"]),
        (r"\hat{x} + \overline{AB}", ["m:acc", "m:bar", 'm:pos m:val="top"']),
        (r"\overbrace{a+b}^{n} \xrightarrow{k}", ["m:limUpp", "m:groupChr", 'm:chr m:val="⏞"']),
        (r"\mathcal{L} + \mathbb{R} + \mathbf{x} + \text{if}", ['m:scr m:val="script"', 'm:scr m:val="double-struck"',
                                                               'm:sty m:val="b"', "m:nor"]),
        (r"\lim_{n \to \infty} a_n", ["m:limLow"]),
        (r"\mathop{\mathrm{Res}}_{z=0} f", ["m:func", "m:limLow", "<m:t>R</m:t>"]),
        (r"\binom{n}{k}", ['m:type m:val="noBar"']),
        (r"\foo{x}", ['val="C0392B"']),
    ],
)
def test_maths_is_written_as_powerpoints_own_equations(source: str, elements: list[str]) -> None:
    from lxml import etree

    from flexo_talk.omml import omml

    written = omml(source, size=20.0, colour="242126", display=True)
    root = etree.fromstring(written)
    assert root.tag == "{http://schemas.microsoft.com/office/drawing/2010/main}m"
    for element in elements:
        assert f"<{element}" in written or element in written, element


def test_powerpoints_equations_keep_tex_sizes_and_spaces() -> None:
    from flexo_talk.omml import omml

    # \big( to \Bigg(: a delimiter grown to a hidden, zero-width bar of TeX's size, alone.
    sizes = []
    for command in (r"\big", r"\Big", r"\bigg", r"\Bigg"):
        written = omml(command + "( x " + command + ")", size=20.0)
        assert written.count("<m:d>") == 2 and '<m:endChr m:val=""/>' in written
        assert '<m:begChr m:val=""/><m:endChr m:val=")"/>' in written
        assert '<m:show m:val="0"/><m:zeroWid m:val="1"/>' in written
        sizes.append(int(re.search(r'sz="(\d+)"[^>]*>(?:(?!</m:r>).)*<m:t>\|</m:t>', written).group(1)))
    assert sizes == sorted(sizes) and sizes[0] > 2000
    # amsmath's \pmod: an em from what it follows in display, 8mu (a three-per-em space) in words.
    assert "\u2003" in omml(r"a \equiv b \pmod{n}", size=20.0, display=True)
    inline = omml(r"a \equiv b \pmod{n}", size=20.0)
    assert "\u2004" in inline and "\u2003" not in inline


def test_a_bold_headers_maths_is_regular_all_of_it(tmp_path: Path) -> None:
    import xml.etree.ElementTree as ET

    from flexo_talk.compose import render_slide

    # Bold is a meaning in maths, and a maths font has no bold for its Greek, its h-bar
    # or its blackboard E: a bold header's maths stays regular, all of it, as LaTeX's
    # does -- and so does strong words' (**...**), unless it asks (\mathbf).
    deck = Deck("header")
    deck.slide("Table").table([[r"Mean $\mathbb{E}[X]$", r"$H_n(\xi)$ over $E_n / \hbar\omega$"], ["1", "2"]])
    deck.slide("Strong").text(r"**Strong $\xi + x$ and $\mathbf{v}$**")
    strong = {run.text.strip(): run.weight for run in inline(r"**a $\xi + x$ $\mathbf{v}$**")}
    assert strong["a"] == strong["v"] == 700 and strong["x"] == 400
    weights = {}
    for slide in deck.slides:
        for text in ET.fromstring(render_slide(deck, slide).svg).iter("{http://www.w3.org/2000/svg}text"):
            for span in text.findall("{http://www.w3.org/2000/svg}tspan"):  # a run each
                weight = span.get("font-weight") or text.get("font-weight") or "400"
                weights["".join(span.itertext()).strip()] = weight
    blackboard, hbar = "\U0001d53c[", "\u210f"
    assert weights["Mean"] == weights["over"] == weights["Strong"] == weights["v"] == "700"
    assert {weights[maths] for maths in (blackboard, "H", hbar, "x")} == {"400"}
    # The PowerPoint's table says the same: its words bold, its maths not.
    table = re.search(r"<a:tbl>.*?</a:tr>", _slides(deck.build(tmp_path, formats=("pptx",)).pptx)[0], re.S)
    runs = {text: attributes for attributes, text in re.findall(
        r"<a:rPr([^>]*)>(?:(?!</a:r>).)*?<a:t>([^<]*)</a:t>", table.group(0), re.S)}
    assert 'b="1"' in runs["Mean "] and 'b="1"' in runs[" over "]
    assert 'b="1"' not in runs[blackboard] and 'b="1"' not in runs[hbar]
    # Its equations' \text is the words' weight, as amsmath's is; their maths is not.
    from flexo_talk.omml import omml

    assert re.findall(r'b="(\d)"', omml(r"x \text{ if }", size=20.0, bold=True)) == ["0", "1"]


def test_a_formula_with_rules_stays_drawn_in_the_powerpoint(tmp_path: Path) -> None:
    from flexo_talk.omml import expressible

    ruled = r"\left[\begin{array}{cc|c} 1 & 0 & b_1 \\ \hline 0 & 0 & 1 \end{array}\right]"
    assert not expressible(ruled) and not expressible(r"x^{\begin{array}{c} a \\ \hline b \end{array}}")
    assert expressible(r"\begin{pmatrix} a & b \end{pmatrix}") and expressible(r"\frac{1}{2}")
    # Office Math has no rules in a matrix: the formula keeps its drawing, a picture for
    # every slide program, rather than becoming an equation without them.
    deck = Deck("ruled")
    slide = deck.slide("Rules")
    slide.math(ruled).math(r"\frac{a}{b}")
    slide.text(r"Inline $\begin{array}{c|c} a & b \end{array}$ too.")
    slide.bullets(r"A list $\begin{array}{c} a \\ \hline b \end{array}$", "and more")
    deck.slide("Table").table([["what", "value"], ["ruled", r"$\begin{array}{c|c} a & b \end{array}$"]])
    first, second = _slides(deck.build(tmp_path, formats=("pptx",)).pptx)
    assert first.count("<mc:AlternateContent") == 1 and first.count("<a14:m") == 1
    assert "<a:custGeom>" in first  # the ruled ones, drawn
    assert "AlternateContent" not in second and "<a:tbl>" in second


def test_the_powerpoint_has_editable_equations_and_the_drawing_for_other_programs(tmp_path: Path) -> None:
    from pptx import Presentation

    deck = Deck("equations")
    slide = deck.slide("Maths")
    slide.text(r"We minimise $\frac{1}{N}\sum_i \ell_i$ here.")
    slide.math(r"\begin{pmatrix} a & b \\ c & d \end{pmatrix}")
    slide.bullets(r"A rate $k = A e^{-E_a/RT}$")
    deck.slide("Table").table([["what", "value"], ["half", r"$\frac{1}{2}$"]])
    result = deck.build(tmp_path / "editable", formats=("pptx",))
    first, second = _slides(result.pptx)
    for xml, count in ((first, 3), (second, 1)):
        blocks = re.findall(r"<mc:AlternateContent.*?</mc:AlternateContent>", xml, re.S)
        assert len(blocks) == count
        for block in blocks:
            choice, fallback = block.split("<mc:Fallback>")
            assert 'Requires="a14"' in choice and "<a14:m" in choice and "<a14:m" not in fallback
            assert "<a:custGeom>" in fallback  # the formula, drawn
    Presentation(str(result.pptx))  # python-pptx still reads it
    # Without editable maths, only the drawing, and no Office Math anywhere.
    drawn = deck.build(tmp_path / "drawn", formats=("pptx",), editable_maths=False)
    assert all("<a14:m" not in xml and "AlternateContent" not in xml for xml in _slides(drawn.pptx))


def _placeholders(slide) -> dict:
    """A slide's placeholders by kind, each with its words (a broken line read as a space)."""

    return {
        shape.placeholder_format.type: " ".join(shape.text_frame.text.replace("\v", " ").split())
        for shape in slide.placeholders
    }


def test_every_title_is_its_slide_s_title_placeholder(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.enum.shapes import PP_PLACEHOLDER as PLACEHOLDER

    deck = Deck("titled", footer="A footer")
    deck.title("A **bold** talk", subtitle="Its *subtitle*", author="Ada", date="2026")
    deck.agenda()
    with deck.slide("Why [this](https://example.com) matters", subtitle="With `code`") as slide:
        slide.bullets("One", "Two")
    deck.section("Part one", subtitle="Its theme")
    deck.statement("One thing, said large", by="Someone")
    deck.slide("", layout="blank").text("Nothing on top.")
    deck.slide("").text("No title.")
    deck.slide("A title long enough that it has to break onto a second line of the slide to fit").text("Words.")
    deck.slide("Broken here\nby hand").text("Words.")
    result = deck.build(tmp_path, formats=("pptx",))
    slides = list(Presentation(str(result.pptx)).slides)  # type: ignore[arg-type]
    assert [slide.slide_layout.name for slide in slides] == [
        "Title Slide", "Title Only", "Title Only", "Section Header", "Title Only", "Blank", "Blank",
        "Title Only", "Title Only",
    ]
    assert _placeholders(slides[0]) == {PLACEHOLDER.CENTER_TITLE: "A bold talk", PLACEHOLDER.SUBTITLE: "Its subtitle"}
    assert _placeholders(slides[1]) == {PLACEHOLDER.TITLE: "Outline"}
    assert _placeholders(slides[2]) == {PLACEHOLDER.TITLE: "Why this matters", PLACEHOLDER.SUBTITLE: "With code"}
    assert _placeholders(slides[3]) == {PLACEHOLDER.TITLE: "Part one", PLACEHOLDER.BODY: "Its theme"}
    # A statement's words are its title, as on Keynote's Statement layout.
    assert _placeholders(slides[4]) == {PLACEHOLDER.TITLE: "One thing, said large"}
    # No title, no title placeholder; and none of the layout's empty ones.
    assert _placeholders(slides[5]) == _placeholders(slides[6]) == {}
    from pptx.util import Pt

    # A title Flexo wrapped is one paragraph that PowerPoint wraps, as wide as its longest
    # line; only a break typed in it is kept.
    title = slides[7].shapes.title
    assert len(title.text_frame.paragraphs) == 1 and "\v" not in title.text_frame.text
    assert title.text_frame.word_wrap is True and title.width < Pt(deck.style.width - 2 * deck.style.margin)
    assert _placeholders(slides[7])[PLACEHOLDER.TITLE] == "".join(run.text for run in deck.slides[7].title_runs)
    typed = slides[8].shapes.title.text_frame
    assert len(typed.paragraphs) == 1 and typed.text == "Broken here\vby hand"
    # Read first, its link still a link, and its words written once: no drawn copy.
    xml = _slides(result.pptx)  # type: ignore[arg-type]
    for slide in (slides[1], slides[2], slides[3], slides[4], slides[7]):
        assert slide.shapes[0] == slide.shapes.title
    assert "<a:hlinkClick" in xml[2].split("</p:sp>")[0]
    assert xml[2].count(">Why <") == xml[0].count(">Its <") == xml[4].count(">One thing, said large<") == 1


def test_a_title_placeholder_is_where_the_title_was_drawn_and_says_how_it_looks(tmp_path: Path) -> None:
    from flexo.drawing import Text, read_drawing

    from flexo_talk.pptx import ascent

    # Centred titles, and a section title, whose layout would set it bold and in capitals.
    deck = Deck("look", look="keynote")
    deck.slide("A centred title").text("Words.")
    deck.section("A quiet section")
    result = deck.build(tmp_path, formats=("pptx", "svg"))
    for svg, xml in zip(result.svgs, _slides(result.pptx), strict=True):  # type: ignore[arg-type]
        drawn = next(item for item in read_drawing(svg.read_text()).walk()
                      if isinstance(item, Text) and (item.id or "").endswith(".title"))
        shape = re.search(r'<p:sp>(?:(?!</p:sp>).)*<p:ph type="title"/>.*?</p:sp>', xml, re.S).group(0)
        left, top = (int(value) / 12700 for value in re.search(r'<a:off x="(-?\d+)" y="(-?\d+)"/>', shape).groups())
        width = int(re.search(r'<a:ext cx="(\d+)"', shape).group(1)) / 12700
        spacing = int(re.search(r'<a:lnSpc><a:spcPts val="(\d+)"/>', shape).group(1)) / 100
        (line,) = drawn.lines
        # Its first baseline where Flexo set it, and centred where it was centred.
        assert top + ascent(line.runs[0].face) * spacing == pytest.approx(line.baseline, abs=0.02)
        assert left + width / 2 == pytest.approx((line.left + line.right) / 2, abs=0.05)
        assert 'algn="ctr"' in shape and f'sz="{round(drawn.size * 100)}"' in shape and ' b="1"' in shape
        # What it would take from its layout and master is said: no capitals, bullet,
        # indent, autofit or inset.
        assert 'cap="none"' in shape and "<a:buNone/>" in shape and 'marL="0" indent="0"' in shape
        assert "<a:noAutofit/>" in shape and 'lIns="0" tIns="0"' in shape and 'wrap="square"' in shape


def test_a_title_with_words_under_it_grows_upward_so_another_font_cannot_cover_them(tmp_path: Path) -> None:
    deck = Deck("opening")
    deck.title("Building the HIV-1 capsid from its parts", subtitle="Structure, assembly and maturation")
    deck.title("Only a title")
    deck.section("A section", subtitle="With words under it")
    result = deck.build(tmp_path, formats=("pptx",))

    def anchors(xml: str) -> dict[str, str]:
        found = {}
        for shape in re.findall(r"<p:sp>.*?</p:sp>", xml, re.S):
            if kind := re.search(r'<p:ph type="(\w+)"', shape):
                found[kind.group(1)] = re.search(r'<a:bodyPr[^>]* anchor="(\w)"', shape).group(1)
        return found

    # Its box ends where its last line does, so the same font draws it where it was, and a
    # wider one sets its extra line above it rather than over the subtitle.
    assert [anchors(xml) for xml in _slides(result.pptx)] == [  # type: ignore[arg-type]
        {"ctrTitle": "b", "subTitle": "t"}, {"ctrTitle": "t"}, {"title": "b", "body": "t"},
    ]


def test_a_title_is_first_on_its_slide_and_over_what_it_was_drawn_on(tmp_path: Path) -> None:
    from pptx import Presentation

    for look, first in (("classic", ["Title"]), ("band", ["Band", "Title"])):
        deck = Deck(look, look=look)
        deck.slide("On the slide").bullets("A point")
        slide = Presentation(str(deck.build(tmp_path / look, formats=("pptx",)).pptx)).slides[0]
        # The band it is set on stays under it; nothing else comes before it.
        assert [shape.name for shape in slide.shapes][: len(first)] == first


def test_a_title_with_maths_is_a_placeholder_with_its_equation_or_room_for_it(tmp_path: Path) -> None:
    deck = Deck("maths")
    deck.slide(r"Halves $\frac{1}{2}$ and more").text("Words.")
    xml = _slides(deck.build(tmp_path / "editable", formats=("pptx",)).pptx)[0]  # type: ignore[arg-type]
    (block,) = re.findall(r"<mc:AlternateContent.*?</mc:AlternateContent>", xml, re.S)
    choice, fallback = block.split("<mc:Fallback>")
    # PowerPoint's own equation in the title; for other programs, the title's words
    # leaving room for the formula, drawn over them.
    assert '<p:ph type="title"/>' in choice and "<a14:m" in choice
    assert '<p:ph type="title"/>' in fallback and "<a:custGeom>" in fallback and re.search(r'spc="\d+"', fallback)
    assert xml.index("<mc:AlternateContent") < xml.index('name="Text')
    drawn = _slides(deck.build(tmp_path / "drawn", formats=("pptx",), editable_maths=False).pptx)[0]  # type: ignore[arg-type]
    assert "AlternateContent" not in drawn and drawn.count('<p:ph type="title"/>') == 1 and "<a:custGeom>" in drawn


def test_a_title_s_marks_that_step_back_are_drawn_beside_its_placeholder(tmp_path: Path) -> None:
    from pptx import Presentation

    deck = Deck("marks")
    deck.slide(r"The vector $\vec{v}$ and $x_i^2$ here").text("Words.")
    result = deck.build(tmp_path, formats=("pptx",))
    slide = Presentation(str(result.pptx)).slides[0]  # type: ignore[arg-type]
    # One placeholder for the words; the arrow and the stacked script, each where it was set.
    assert slide.shapes.title.text_frame.text.startswith("The vector v and x")
    assert [shape.name for shape in slide.shapes][:2] == ["Title", "Title Accents"]
    # Nothing steps back by negative spacing, which slide programs set differently.
    assert not re.search(r'spc="-', _slides(result.pptx)[0])  # type: ignore[arg-type]


def test_no_formula_however_broken_makes_office_math_powerpoint_cannot_read() -> None:
    import random

    from lxml import etree

    from flexo_talk.omml import omml

    pieces = [r"\frac", r"\sqrt", "{", "}", "^", "_", "&", r"\\", r"\left(", r"\right)", r"\begin{pmatrix}",
              r"\end{pmatrix}", "x", "2", r"\alpha", r"\sum", r"\middle|", r"\over", "'", r"\hat", r"\ce{",
              r"\text{", "<", ">", "&amp;", r"\color{red}", r"\overbrace", r"\not", r"\big", r"\mathbb", " "]
    rng = random.Random(9)
    for _ in range(500):
        source = "".join(rng.choice(pieces) for _ in range(rng.randint(1, 14)))
        etree.fromstring(omml(source, size=18.0, colour="000000", display=rng.random() < 0.5))


def test_a_mechanism_is_drawn_as_shapes_with_its_arrows_and_a_wrong_one_is_said(tmp_path: Path) -> None:
    deck = Deck("chemistry")
    with deck.slide("Acyl substitution") as slide:
        slide.mechanism([
            {"smiles": "[OH-:5].[CH3:1][C:2](=[O:3])[Cl:4]", "arrows": ["5 -> 2", "2=3 -> 3"],
             "reagents": "NaOH"},
            {"arrows": ["3 -> 2", "2-4 -> 4"], "label": "tetrahedral intermediate"},
        ])
    result = deck.build(tmp_path, formats=("pptx", "svg"))
    xml = _slides(result.pptx)[0]  # type: ignore[arg-type]
    assert "<p:pic>" not in xml
    assert xml.count("<a:cubicBezTo>") >= 4  # the curly arrows, as curves
    assert 'srgbclr val="d466d6"' in xml.lower()  # in their one magenta ink
    assert ">tetrahedral intermediate<" in xml and ">NaOH<" in xml
    with pytest.raises(ValueError, match="arrow_colour"):
        deck.slide("Pink").mechanism("[OH-:1]", arrow_colour="pink")
    # The deck's own colours: the arrows follow its theme.
    deck.slide("Accent").mechanism([{"smiles": "[OH-:1].[CH3:2][Br:3]", "arrows": "1 -> 2; 2-3 -> 3",
                                     "place": {1: {"move": [-1, 0]}}}], arrow_colour="accent")
    # A step that cannot be is drawn as far as it goes, its arrows on it, and said.
    deck.slide("Half drawn").mechanism([{"smiles": "[OH-:5].[CH3:1][C:2](=[O:3])[Cl:4]", "arrows": ["5 -> 2"]}])
    result = deck.build(tmp_path, formats=("svg",))
    assert any("C2 would have 10 electrons" in said for said in result.diagnostics)
    with pytest.raises(ValueError, match="never closed"):
        deck.slide("Wrong").mechanism([{"smiles": "[OH-:5].[CH3:1][C:2](=[O:3])[Cl:4", "arrows": ["5 -> 2"]}])


@pytest.mark.parametrize("look", [{"theme": "paper"}, {"theme": "paper", "font": "IBM Plex Sans"}, {"theme": "tikz"}])
def test_a_mechanism_is_lettered_as_its_deck_is(look) -> None:
    deck = Deck("chemistry", **look)
    with deck.slide("Acyl substitution") as slide:
        slide.text("Hydroxide attacks the carbonyl.").mechanism([
            {"smiles": "[OH-:5].[CH3:1][C:2](=[O:3])[Cl:4]", "arrows": ["5 -> 2", "2=3 -> 3"], "reagents": "NaOH"},
            {"arrows": ["3 -> 2", "2-4 -> 4"], "label": "tetrahedral intermediate"},
        ])
    (rendered,) = deck.render()
    texts = re.findall(r'<text[^>]*id="([^"]+)"[^>]*font-family="([^"]+)"[^>]*font-size="([^"]+)"', rendered.svg)
    face = {ident: family.split(",")[0] for ident, family, _ in texts}
    words = next(face[ident] for ident in face if ident.endswith("body.0"))
    lettered = [ident for ident in face if ".mechanism." in ident]
    assert lettered and {face[ident] for ident in lettered} == {words}
    # Its atoms are set at the size of the words beside it, once scaled onto the slide.
    scale = float(re.search(r'body\.1[^>]*transform="[^"]*scale\(([\d.]+)\)', rendered.svg).group(1))
    size = {ident: float(value) for ident, _, value in texts}
    atom = next(ident for ident in lettered if ".atom" in ident)
    assert size[atom] * scale == pytest.approx(size[next(i for i in face if i.endswith("body.0"))], rel=0.02)


def test_what_a_figure_check_found_is_said_in_words_under_the_slide() -> None:
    from flexo_talk.compose import _figure_check

    assert _figure_check("routing.connector.crossing") == "Lines cross in this figure."
    assert _figure_check("routing.net.target.arrow") == "A line or arrowhead is cramped in this figure."
    assert _figure_check("routing.source.direction").startswith("A line doesn't meet its shape")
    # Never a raw code: a check of the drawing's file (nothing to put right on the slide) is
    # not said at all, and one with no words of its own is said in plain ones.
    assert _figure_check("svg.id.missing") is None
    assert _figure_check("svg.hierarchy.deep") is None
    for code in ("layout.size.grown", "cells.something.new", "label.unheard.of"):
        said = _figure_check(code)
        assert said and code not in said and "(" not in said


def test_a_code_listing_is_one_text_box_in_powerpoint(tmp_path: Path) -> None:
    from pptx import Presentation

    deck = Deck("listing")
    with deck.slide("Code") as slide:
        slide.code("def greet(name):\n    # say hello\n\n    return name")
    result = deck.build(tmp_path, formats=("pptx",))
    def every(shapes):
        for shape in shapes:
            yield shape
            if shape.shape_type == 6:  # a group
                yield from every(shape.shapes)

    slide = Presentation(str(result.pptx)).slides[0]  # type: ignore[arg-type]
    listing = [shape for shape in every(slide.shapes) if shape.has_text_frame and "greet" in shape.text_frame.text]
    assert len(listing) == 1
    # A paragraph a line, the blank one kept, so each line is where it was drawn.
    lines = [paragraph.text for paragraph in listing[0].text_frame.paragraphs]
    assert lines == ["def greet(name):", "    # say hello", "", "    return name"]


def test_the_pdf_is_tagged_as_a_screen_reader_reads_each_slide(tmp_path: Path) -> None:
    import struct
    import zlib

    def chunk(tag: bytes, payload: bytes) -> bytes:
        return struct.pack(">I", len(payload)) + tag + payload + struct.pack(">I", zlib.crc32(tag + payload))

    source = tmp_path / "p.png"
    source.write_bytes(
        b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 2, 2, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(b"\x00" + bytes(6) + b"\x00" + bytes(6))) + chunk(b"IEND", b"")
    )
    deck = Deck("tagged")
    deck.slide("Points").bullets("One", ["Under one"], "Two")
    deck.slide("Numbers").table([["State", "Count"], ["Open", "3"]])
    deck.slide("A picture").image(source, description="Two dark pixels")
    deck.slide("A ratio").math(r"\frac{a+b}{2}")
    pdf = deck.build(tmp_path, formats=("pdf",)).pdf.read_bytes()  # type: ignore[union-attr]
    assert b"/MarkInfo << /Marked true >>" in pdf and b"/StructTreeRoot" in pdf
    kinds = re.findall(rb"/Type /StructElem /S /(\w+)", pdf)
    # A section a slide, its title a heading; a list's items in their levels, a table's rows
    # and header cells, a picture by its description and a formula by its words.
    assert kinds.count(b"Sect") == 4 and kinds.count(b"H1") == 4
    assert kinds.count(b"L") == 2 and kinds.count(b"LI") == 3 and kinds.count(b"LBody") == 3
    assert kinds.count(b"TR") == 2 and kinds.count(b"TH") == 2 and kinds.count(b"TD") == 2
    assert b"/S /Figure" in pdf and b"/Alt (Two dark pixels)" in pdf
    assert re.search(rb"/S /Formula /P \d+ 0 R /Pg \d+ 0 R /Alt \(\\\(a \+ b\\\)/2\)", pdf)


def test_a_molecule_s_picture_says_what_it_shows() -> None:
    from types import SimpleNamespace

    from flexo.ir.semantic import NodeSpec, TextRun

    from flexo_talk.export import _structures_said

    nodes = [
        NodeSpec("pocket", "structure", (TextRun("Trypsin with its inhibitor"),),
                 properties=(("source", "assets/1gbt.cif"),)),
        NodeSpec("fetched", "structure", properties=(("source", "4hhb"),)),
        NodeSpec("box", "block", (TextRun("Not a molecule"),)),
    ]
    said = _structures_said(SimpleNamespace(figure=SimpleNamespace(nodes=nodes)))  # type: ignore[arg-type]
    assert said == {
        "pocket.molecule": "Molecular structure: Trypsin with its inhibitor (1GBT).",
        "fetched.molecule": "Molecular structure 4HHB.",
    }


def test_a_molecule_s_picture_is_named_from_its_file_s_title_not_its_id(tmp_path: Path) -> None:
    from types import SimpleNamespace

    from flexo.ir.semantic import NodeSpec

    from flexo_talk.export import _structures_said

    (tmp_path / "1a8o.pdb").write_text(
        "HEADER    VIRAL PROTEIN\nTITLE     HIV CAPSID C-TERMINAL DOMAIN\n"
        "COMPND    MOL_ID: 1;\nCOMPND   2 MOLECULE: HIV CAPSID;\n"
    )
    (tmp_path / "2xyz.cif").write_text(
        "data_2XYZ\n_struct.title   'CRYSTAL STRUCTURE OF THE E2 PROTEIN BOUND TO DNA'\n"
    )
    nodes = [
        NodeSpec("pdb", "structure", properties=(("source", str(tmp_path / "1a8o.pdb")),)),
        NodeSpec("cif", "structure", properties=(("source", str(tmp_path / "2xyz.cif")),)),
    ]
    said = _structures_said(SimpleNamespace(figure=SimpleNamespace(nodes=nodes)))  # type: ignore[arg-type]
    # Its file's title, in sentence case with names kept as they are, and its id said once.
    assert said == {
        "pdb.molecule": "Molecular structure: HIV capsid C-terminal domain (1A8O).",
        "cif.molecule": "Molecular structure: Crystal structure of the E2 protein bound to DNA (2XYZ).",
    }


def test_a_figure_s_lines_are_connectors_joined_to_what_they_join(tmp_path: Path) -> None:
    from lxml import etree

    deck = Deck("lines")
    with deck.slide("A flow", layout="figure") as slide, slide.figure() as figure:
        with figure.row("steps") as row:
            start = row.block("start", label="Start")
            middle = row.block("middle", label="Middle", input=start)
            row.block("end", label="End", input=middle)
        figure.connect(middle, start, label="again")
    xml = _slides(deck.build(tmp_path, formats=("pptx",)).pptx)[0]  # type: ignore[arg-type]
    tree = etree.fromstring(xml.encode())
    p = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
    a = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    names = {int(nv.get("id")): nv.get("name") for nv in tree.iter(f"{p}cNvPr")}
    joined, both = [], 0
    for line in tree.iter(f"{p}cxnSp"):
        # Each line is the slide program's connector, ending in its own arrowhead, named for
        # what it joins, and joined to a block's shape at each end that meets the middle of
        # a side (its connection site): the two lines between Start and Middle meet theirs
        # above and below the middle, and are left free there.
        assert line.find(f".//{a}tailEnd") is not None and line.find(f".//{a}xfrm").get("rot") is None
        ends = [line.find(f".//{a}{tag}") for tag in ("stCxn", "endCxn")]
        assert {names[int(end.get("id"))] for end in ends if end is not None} <= {"Start", "Middle", "End"}
        both += all(end is not None for end in ends)
        joined.append(line.find(f".//{p}cNvPr").get("name"))
    assert sorted(joined) == ["Line from Middle to End", "Line from Middle to Start", "Line from Start to Middle"]
    assert both == 1
    assert "arrowhead" not in xml and "Line\"" not in xml

    # Groups only where they mean something: the figure, each block, a line with its words.
    def depth(element, level: int = 0) -> int:
        return max([level, *(depth(child, level + 1) for child in element.iter(f"{p}grpSp") if child is not element
                              and child.getparent() is element)])

    assert depth(tree.find(f"{p}cSld/{p}spTree")) <= 3
    assert not any(name.endswith((".root", ".components", ".connectors", ".content", ".steps")) or name == "slide1"
                   for name in names.values())
    # A line's words stand beside it, named for what they say: LibreOffice draws a joined
    # connector grouped within a group again, and wrongly.
    assert "Label “again”" in names.values()


def test_a_numbered_list_numbers_every_level_in_its_tier(tmp_path: Path) -> None:
    from flexo_talk.compose import render_slide
    from flexo_talk.deck import list_numbers

    assert list_numbers([0, 1, 1, 2, 2, 1, 0, 1]) == ["1.", "a.", "b.", "i.", "ii.", "c.", "2.", "a."]
    deck = Deck("tiers")
    with deck.slide("Steps") as slide:
        slide.bullets("Plan", ["Ask", "Listen"], "Build", numbered=True)
    svg = render_slide(deck, deck.slides[0]).svg
    assert [">1.<" in svg, ">a.<" in svg, ">b.<" in svg, ">2.<" in svg] == [True] * 4
    result = deck.build(tmp_path, formats=("pptx",))
    xml = _slides(result.pptx)[0]  # type: ignore[arg-type]
    assert xml.count('type="arabicPeriod"') == 2 and xml.count('type="alphaLcPeriod"') == 2


def test_a_quotations_mark_is_a_curled_quotation_mark_in_every_theme() -> None:
    import flexo
    from flexo.fonts import hb_font
    from flexo.outline import _outline
    from flexo.text import FontStack

    from flexo_talk import compose

    for theme in flexo.THEMES:
        deck = Deck("quoted", theme=theme)
        with deck.slide("Why") as slide:
            slide.quote("Nothing in biology makes sense", by="Dobzhansky")
        svg = compose.render_slide(deck, deck.slides[0]).svg
        mark = re.search(r'<text id="[^"]*\.mark"[^>]*font-family="([^"]+)"', svg)
        assert mark is not None, theme
        # Two wedges of straight strokes (Figtree's) read as "//" set so large: the mark is
        # drawn in a face whose mark is curled where the title face's is not.
        family = mark.group(1).split(",")[0].strip("'\" ")
        face = FontStack(deck.typography(40, title=True).with_family(family)).face(400, False)
        gid = hb_font(face, 400).get_nominal_glyph(ord("“"))
        assert any(segment.kind == "C" for segment in _outline(face, 400, gid)), theme
        # And carried with the slide, for a browser to draw it in that face too.
        assert f"font-family:'{family}'" in svg, theme


def test_a_caption_and_a_build_are_written_as_a_deck_document_writes_them() -> None:
    from flexo_talk.document import deck_document

    deck = Deck("written")
    with deck.slide("Cues") as slide:
        slide.table([["Receptor", "Cue"], ["Tar", "Aspartate"]], caption="Table 1.  Cues").build_in()
    (block,) = deck_document(deck)["slides"][0]["body"]
    assert block["caption"] == "Table 1. Cues" and block["build"] is True
    assert deck.slides[0].regions["body"].builds == {0}


def test_columns_moved_to_their_places_move_only_their_own_native_lists_tables_and_words() -> None:
    import xml.etree.ElementTree as ET

    from flexo_talk import compose

    deck = Deck("columns")
    with deck.slide("Tables", layout="two-columns") as slide:
        slide.regions["left"].table([["A", "B"], ["1", "2"]])
        slide.regions["right"].table([["C", "D"], ["3", "4"], ["5", "6"], ["7", "8"]])
    with deck.slide("Lists", layout="two-columns", align="middle") as slide:
        slide.regions["left"].bullets("One")
        slide.regions["left"].text(r"Rate $\frac{a}{b}$")
        slide.regions["right"].bullets("Two", "Three", "Four")
        slide.regions["right"].text(r"Rate $\frac{c}{d}$")
    for slide in deck.slides:
        rendered = compose.render_slide(deck, slide)
        root = ET.fromstring(rendered.svg)
        moved = {item.get("id"): float(re.search(r"translate\(0 ([-\d.]+)\)", item.get("transform"))[1])
                 for item in root.iter() if item.get("data-flexo-talk") == "region" and item.get("transform")}
        assert len(moved) == 2

        def down(identifier: str, moved: dict[str, float] = moved) -> float:
            return next((by for name, by in moved.items() if identifier.startswith(f"{name}.")), 0.0)

        def drawn(identifier: str, root: ET.Element = root) -> ET.Element:
            return next(item for item in root.iter() if item.get("id") == identifier)

        # Each native table, list and line of words (for the PowerPoint) where its column's
        # drawing is: moved by its own column alone, never by the one before it as well.
        for table in rendered.tables:
            top = float(re.search(r"M [-\d.]+ ([-\d.]+)", drawn(f"{table.id}.rule0").get("d"))[1])
            assert table.y == pytest.approx(top + down(table.id), abs=0.01)
        for layout in rendered.lists:
            assert layout.items[0][2] == pytest.approx(float(drawn(f"{layout.id}.0").get("y")) + down(layout.id))
        for words in rendered.worded:
            baseline = float(next(item.get("y") for item in drawn(words.id).iter() if item.get("y")))
            assert words.baseline == pytest.approx(baseline + down(words.id), abs=0.01)
        assert len(rendered.tables) + len(rendered.lists) + len(rendered.worded) in {2, 4}


def test_tables_side_by_side_share_one_top() -> None:
    from flexo_talk import compose

    deck = Deck("columns")
    with deck.slide("Tables", layout="two-columns") as slide:
        slide.regions["left"].table([["A", "B"], ["1", "2"]])
        slide.regions["right"].table([["C", "D"], ["3", "4"], ["5", "6"], ["7", "8"]])
    rendered = compose.render_slide(deck, deck.slides[0])
    # Columns of pictures or tables line up at the top, as Keynote's do, however tall each is.
    left, right = rendered.tables
    assert left.y == pytest.approx(right.y)
