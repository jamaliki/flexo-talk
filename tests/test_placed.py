"""A slide's object placed where it says (a block's ``at``, ``Region.place(at=)``), out of its
place's flow, as a shape is: read and written back as a document says it, what is wrong said
where it is, drawn there in every output -- over what the slide lays out, under its shapes, the
rest of its place laid out as if it were not there -- and drawn just as it was when it is placed
where it stood."""

from __future__ import annotations

import copy
import json
import re
import shutil
import subprocess
import xml.etree.ElementTree as ET
from pathlib import Path

import pytest
import yaml
from flexo.drawing import Drawing, Group, ink_bounds, read_drawing
from pptx import Presentation
from pptx.util import Emu

from flexo_talk import Deck
from flexo_talk.compose import PLACEHOLDERS, render_slide
from flexo_talk.deck import SettingError
from flexo_talk.document import DeckDocumentError, deck_document, deck_from_document, dump_document, save_document
from flexo_talk.studio import DeckKind

SVG = "{http://www.w3.org/2000/svg}"


def deck_of(*slides: dict, **deck: object) -> dict:
    return {"schema_version": 1, "deck": {"id": "talk", "theme": "paper", **deck}, "slides": list(slides)}


def drawn(document: dict, base: Path | str = ".", *, editing: bool = False, at: int = 0) -> tuple:
    deck = deck_from_document(document, base)
    token = PLACEHOLDERS.set(editing)
    try:
        rendered = render_slide(deck, deck.slides[at])
    finally:
        PLACEHOLDERS.reset(token)
    return rendered, ET.fromstring(rendered.svg)


def element(root: ET.Element, ident: str) -> ET.Element | None:
    return next((item for item in root.iter() if item.get("id") == ident), None)


def inked(svg: str) -> dict[str, tuple[float, ...]]:
    """Each object's ink on the slide, by its id: where it is drawn, its region's moves taken."""

    found: dict[str, tuple[float, ...]] = {}

    def walk(group: Group) -> None:
        for item in group.items:
            ident = getattr(item, "id", None) or ""
            if re.fullmatch(r"slide\d+\.(left|right|body|column\d+)\.\d+", ident):
                found[ident] = tuple(round(value, 1) for value in ink_bounds(Drawing(0.0, 0.0, Group(None, [item]))))
            elif isinstance(item, Group):
                walk(item)

    walk(read_drawing(svg).root)
    return found


def picture(tmp_path: Path) -> str:
    from PIL import Image

    Image.new("RGB", (400, 300), "#4a90d9").save(tmp_path / "cell.png")
    return "cell.png"


# -- the document --


def test_a_placed_object_is_read_and_written_back_as_written() -> None:
    body = [{"text": "In line"}, {"text": "Measured at 4 K", "at": [80, 440], "width": 300, "locked": True},
            {"bullets": ["One", "Two"], "at": [600.5, 120]}]
    document = deck_of({"title": "Placed", "body": body})
    deck = deck_from_document(document, ".")
    region = deck.slides[0].regions["body"]
    assert region.at == {1: (80.0, 440.0), 2: (600.5, 120.0)}
    assert region.locked == {1}
    written = deck_document(deck)
    assert written["slides"][0]["body"] == body
    assert yaml.safe_load(dump_document(written))["slides"][0]["body"] == body


def test_a_place_the_studio_writes_is_one_line_beside_the_objects_own_keys(tmp_path: Path) -> None:
    path = tmp_path / "talk.yaml"
    path.write_text(
        "schema_version: 1\ndeck:\n  id: talk\nslides:\n- title: T\n  body:\n  - text: In line\n  - text: Out\n"
    )
    document = deck_of({"title": "T", "body": [{"text": "In line"}, {"text": "Out", "at": [80, 440], "width": 300}]})
    save_document({**document, "deck": {"id": "talk"}}, path)
    assert "  - text: Out\n    at: [80, 440]\n    width: 300\n" in path.read_text()


def test_a_place_written_wrong_is_said_where_it_is() -> None:
    def read(block: dict) -> None:
        deck_from_document(deck_of({"title": "T", "body": [{"text": "A"}, block]}), ".")

    for wrong in ([80], [80, 120, 3], "80, 120", [80, True], [80, "120"], [float("nan"), 0], [20_000, 0]):
        with pytest.raises(DeckDocumentError, match="at must be") as raised:
            read({"text": "B", "at": wrong})
        assert raised.value.where == "slides[0].body[1] (text)"
    with pytest.raises(DeckDocumentError, match="locked must be true or false"):
        read({"text": "B", "locked": "yes"})
    with pytest.raises(DeckDocumentError, match="Unknown key"):
        read({"text": "B", "size_at": [1, 2]})


def test_an_object_is_placed_in_python_after_it_is_added() -> None:
    deck = Deck("placed")
    with deck.slide("T") as slide:
        slide.text("In line")
        slide.text("Measured at 4 K", width=300)
        slide.place(at=(80, 440))
    region = deck.slides[0].regions["body"]
    assert region.at == {1: (80.0, 440.0)}
    assert deck_document(deck)["slides"][0]["body"][1] == {"text": "Measured at 4 K", "width": 300, "at": [80, 440]}
    with pytest.raises(SettingError, match="at must be"):
        slide.place(at=[1])
    with pytest.raises(ValueError, match="horizontal, vertical, both, or at"):
        slide.place()


def test_an_agent_is_told_how_to_place_an_object() -> None:
    guide = DeckKind().guide()
    assert "Every block also takes: placeholder, build, horizontal, vertical, at, locked" in guide
    assert "at: [80, 440]" in guide


# -- the slide --


def test_a_placed_object_is_drawn_where_it_says_and_its_place_laid_out_without_it() -> None:
    body = [{"text": "First"}, {"text": "Moved away", "at": [600, 400], "width": 300}, {"text": "Third"}]
    rendered, root = drawn(deck_of({"title": "Placed", "body": body}))
    alone, _ = drawn(deck_of({"title": "Placed", "body": [body[0], body[2]]}))
    ink, without = inked(rendered.svg), inked(alone.svg)
    # Where it says: its words from its top left, wrapped in its width.
    left, top, right, _ = ink["slide1.body.1"]
    assert 600 <= left < 604 and 400 < top < 412 and right < 900
    # The rest as if it were not there: the third where the second would be.
    assert ink["slide1.body.0"] == without["slide1.body.0"]
    assert ink["slide1.body.2"] == without["slide1.body.1"]
    # Drawn over what the slide lays out: after its places, in a group of its own.
    content = [child.get("id") for child in element(root, "slide1.content")]
    assert content.index("slide1.placed") > content.index("slide1.body")
    assert element(element(root, "slide1.placed"), "slide1.body.1") is not None


def test_a_placed_object_is_under_the_shapes_and_over_the_slides_furniture() -> None:
    slide = {"title": "Placed", "body": [{"text": "Words", "at": [700, 480], "width": 200}],
             "shapes": [{"kind": "rectangle", "at": [690, 470], "size": [100, 40]}]}
    _, root = drawn(deck_of(slide, footer="A footer"))
    content = [child.get("id") for child in element(root, "slide1.content")]
    assert content.index("slide1.footer") < content.index("slide1.placed") < content.index("slide1.shapes")


def test_every_kind_placed_where_it_stood_is_drawn_just_as_it_was(tmp_path: Path) -> None:
    """The studio takes an object out of line by placing it at its frame (compose's
    ``data-flexo-frame``) with its frame's width: it does not move, nor change."""

    image = picture(tmp_path)
    figure = {"figure": {"id": "model"}, "nodes": [{"id": "a", "label": "Encoder"}, {"id": "b", "label": "Decoder"}],
              "edges": [{"from": "a", "to": "b"}]}
    slides = [
        {"title": "Words", "layout": "two-columns",
         "left": [{"text": "Some words, centred in their room, long enough to wrap onto a second line.",
                   "align": "middle"},
                  {"bullets": ["One", "Two", ["Under two"]]}, {"math": "E = mc^2"}, {"code": "x = 1\nprint(x)"}],
         "right": [{"image": image}, {"table": [["A", "B"], [1, 2]]},
                   {"stats": [{"value": "93%", "label": "right"}]}]},
        {"title": "More", "layout": "two-columns",
         "left": [{"quote": "Simplicity is prerequisite for reliability.", "by": "Dijkstra"},
                  {"callout": "A point.", "title": "Note"}],
         "right": [{"figure": figure}, {"gallery": [image, image]}]},
    ]
    document = deck_of(*slides)
    for number, slide in enumerate(slides):
        rendered, root = drawn(document, tmp_path, editing=True, at=number)
        was = inked(rendered.svg)
        for region in ("left", "right"):
            for index, block in enumerate(slide[region]):
                ident = f"slide{number + 1}.{region}.{index}"
                x, y, width, _ = (float(value) for value in element(root, ident).get("data-flexo-frame").split())
                moved = copy.deepcopy(document)
                placed = moved["slides"][number][region][index]
                placed |= {"at": [x, y], "width": width + 0.005}
                if "gallery" in placed:
                    # (Its pictures as tall as they were: the studio writes their height too.)
                    placed["height"] = float(element(root, ident).get("data-flexo-cells").split()[2])
                now = inked(drawn(moved, tmp_path, editing=True, at=number)[0].svg)
                assert now[ident] == pytest.approx(was[ident], abs=0.6), (ident, next(iter(block)))


def test_a_placed_figure_picture_or_table_keeps_its_proportions_as_wide_as_it_says(tmp_path: Path) -> None:
    image = picture(tmp_path)
    body = [{"image": image, "at": [100, 150], "width": 200},
            {"table": [["A", "B"], [1, 2]], "at": [500, 150], "width": 240}]
    rendered, _ = drawn(deck_of({"title": "T", "body": body}), tmp_path)
    ink = inked(rendered.svg)
    left, top, right, bottom = ink["slide1.body.0"]
    assert (left, top, right) == pytest.approx((100, 150, 300), abs=0.5)
    assert bottom - top == pytest.approx(150, abs=0.5)  # 4:3
    (table,) = rendered.tables
    assert (table.x, sum(table.widths)) == pytest.approx((500, 240), abs=0.5)
    # Without a width, as large as its place would set it alone.
    alone, _ = drawn(deck_of({"title": "T", "body": [{"image": image}]}), tmp_path)
    placed, _ = drawn(deck_of({"title": "T", "body": [{"image": image, "at": [48, 127]}]}), tmp_path)
    was, now = inked(alone.svg)["slide1.body.0"], inked(placed.svg)["slide1.body.0"]
    assert now[2] - now[0] == pytest.approx(was[2] - was[0], abs=0.5)


def test_a_placed_object_on_a_larger_slide_is_scaled_with_it_as_its_shapes_are() -> None:
    body = [{"text": "Words", "at": [480, 270], "width": 240}]
    _, root = drawn(deck_of({"title": "T", "body": body}, style={"width": 1920, "height": 1080}), editing=True)
    x, y, width, _ = (float(value) for value in element(root, "slide1.body.0").get("data-flexo-frame").split())
    assert (x, y, width) == pytest.approx((960, 540, 480))


def test_a_placed_object_on_a_slide_set_in_rows_leaves_its_cell_empty() -> None:
    slide = {"title": "Rows", "layout": "two-columns", "rows": ["One", "Two"],
             "left": [{"text": "A"}, {"text": "B"}],
             "right": [{"text": "C", "at": [600, 420], "width": 200}, {"text": "D"}]}
    rendered, _ = drawn(deck_of(slide))
    ink = inked(rendered.svg)
    assert ink["slide1.right.0"][0] == pytest.approx(600, abs=3)
    # D stays in the second row, beside B.
    assert ink["slide1.right.1"][1] == pytest.approx(ink["slide1.left.1"][1], abs=1)


def test_a_placed_object_builds_in_in_the_slides_order() -> None:
    body = [{"text": "First", "build": True}, {"text": "Second", "build": True, "at": [600, 400]},
            {"text": "Third", "build": True}]
    rendered, _ = drawn(deck_of({"title": "T", "body": body}))
    assert [ident for _, ident in rendered.builds] == ["slide1.body.0", "slide1.body.1", "slide1.body.2"]


def test_a_placed_placeholder_is_drawn_in_the_studio_alone() -> None:
    body = [{"text": "Here"}, {"table": [["", ""], ["", ""]], "placeholder": True, "at": [500, 300], "width": 200}]
    _, editing = drawn(deck_of({"title": "T", "body": body}), editing=True)
    _, shown = drawn(deck_of({"title": "T", "body": body}))
    assert element(editing, "slide1.body.1") is not None
    assert element(shown, "slide1.body.1") is None


# -- the outputs --


def test_a_placed_object_is_where_it_says_in_the_powerpoint_pdf_and_images(tmp_path: Path) -> None:
    body = [{"text": "In line"}, {"text": "Measured at 4 K", "at": [600, 400], "width": 300},
            {"table": [["A", "B"], [1, 2]], "at": [100, 300], "width": 200}]
    slide = {"title": "Placed", "body": body, "shapes": [{"kind": "rectangle", "at": [590, 390], "size": [60, 40]}]}
    written = DeckKind().export(deck_of(slide), tmp_path, "talk", ["pptx", "pdf", "svg", "png"], into=tmp_path / "out")
    assert {path.suffix for path in written} >= {".pptx", ".pdf", ".svg", ".png"}
    shapes = list(Presentation(str(next(path for path in written if path.suffix == ".pptx"))).slides[0].shapes)
    names = [shape.name for shape in shapes]
    words = shapes[names.index("Text “Measured at 4 K”")]
    table = shapes[names.index("Table")]
    assert (Emu(words.left).pt, Emu(words.top).pt) == pytest.approx((600, 400), abs=1)
    assert (Emu(table.left).pt, Emu(table.top).pt, Emu(table.width).pt) == pytest.approx((100, 300, 200), abs=1)
    # Stacked as drawn: over what is in line, under the shapes -- never grouped apart.
    order = [names.index(name) for name in ("Text “In line”", "Text “Measured at 4 K”", "Table", "Rectangle")]
    assert order == sorted(order)
    svg = next(path for path in written if path.suffix == ".svg").read_text()
    assert inked(svg)["slide1.body.1"][0] == pytest.approx(600, abs=4)


# -- the studio's sums (places.js) --


def _node(code: str) -> object:
    path = Path(__file__).parents[1] / "src/flexo_talk/studio/static/places.js"
    script = f"import * as P from {json.dumps(path.as_uri())};\n{code}"
    done = subprocess.run(["node", "--input-type=module", "-e", script], capture_output=True, text=True, check=True)
    return json.loads(done.stdout)


@pytest.mark.skipif(shutil.which("node") is None, reason="needs node")
def test_the_studio_takes_an_object_out_of_line_where_it_stands_and_puts_it_back() -> None:
    made = _node("""
      const text = P.stand({ text: "Words", placeholder: true }, [48, 127.004, 414.0049, 24], "text");
      const plot = P.stand({ plot: "plots.py:loss" }, [498, 127, 414, 276], "plot");
      const own = P.stand({ plot: "plots.py:loss", aspect: 2 }, [498, 127, 414, 276], "plot");
      const gallery = P.stand({ gallery: ["a.png"] }, [498, 127, 414, 120], "gallery", { pictures: 96.5 });
      const moved = P.movedBy(structuredClone(text), 12.004, -7);
      const back = P.inLine(structuredClone(moved), 1, 414);
      const narrow = P.inLine({ image: "a.png", at: [1, 2], width: 200 }, 1, 414);
      const larger = P.inLine({ image: "a.png", at: [1, 2], width: 200 }, 2, 864);
      const placed = [P.placed(text), P.placed(back), P.placed({ at: [1] }), P.placed({ at: ["a", 1] })];
      const reach = P.reach([[10, 20, 30, 40], null, [5, 25, 9, 30]]);
      console.log(JSON.stringify({ text, plot, own, gallery, moved, back, narrow, larger, placed, reach }));
    """)
    # Where it stands, as wide -- its width never a hair narrower than it was drawn -- and its
    # own, no placeholder.
    assert made["text"] == {"text": "Words", "at": [48, 127], "width": 414.01}
    # A plot keeps its proportions as its aspect, unless it has one; a gallery its pictures.
    assert made["plot"]["aspect"] == 1.5 and made["own"]["aspect"] == 2
    assert made["gallery"]["height"] == 96.5
    assert made["moved"]["at"] == [60, 120]
    # Back in line: filling its place, its width dropped; narrower, kept -- in the place's points.
    assert made["back"] == {"text": "Words"}
    assert made["narrow"] == {"image": "a.png", "width": 200}
    assert made["larger"] == {"image": "a.png", "width": 400}
    assert made["placed"] == [True, False, False, False]
    assert made["reach"] == [5, 20]
