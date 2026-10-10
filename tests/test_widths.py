"""Every object on a slide takes a width of its own (points): the most it takes across its
place, set there as its place sets it -- words wrapped in it, a plot, a gallery or a table
drawn to it -- in the slide, the PowerPoint and the file alike; and, while it is edited,
the slide says what an editor sizing it by its handles needs."""

from __future__ import annotations

import re
from pathlib import Path

import pytest

from flexo_talk import Deck
from flexo_talk.compose import PLACEHOLDERS, render_slide
from flexo_talk.document import BLOCKS, DeckDocumentError, deck_document, deck_from_document

PARAGRAPH = "A paragraph of words long enough to wrap across a few lines when it is given a width of its own."


def _picture(folder: Path, name: str = "face.png") -> str:
    from PIL import Image

    Image.new("RGB", (200, 200), (180, 120, 90)).save(folder / name)
    return name


def _blocks(folder: Path) -> list[dict]:
    """One of every kind, but a figure and a plot (made in Python, below)."""

    return [
        {"text": PARAGRAPH},
        {"bullets": ["A first item that runs on for a while so that it wraps", "A second item", ["nested"]]},
        {"quote": "The model's doubt was the most useful thing it told us.", "by": "a crystallographer"},
        {"callout": "A model that knows how sure it is tells us which structures to solve first.", "title": "Idea"},
        {"stats": [{"value": "200M", "label": "structures predicted"}, {"value": "0.1%", "label": "solved"}]},
        {"code": "def spread(samples):\n    return samples.std(axis=0).mean()"},
        {"math": "E = mc^2 + \\frac{a}{b}"},
        {"table": [["Method", "RMSD", "Calibrated"], ["Ensemble", "3.4", "no"], ["Ours", "2.2", "yes"]]},
        {"gallery": [_picture(folder), _picture(folder, "other.png")]},
        {"mechanism": [{"smiles": "[OH-:1].[CH3:2][Br:3]", "arrows": ["1 -> 2", "2-3 -> 3"]}]},
    ]


def _slide(folder: Path, block: dict, **extra: object) -> dict:
    return {"deck": {}, "slides": [{"title": "T", "body": [{**block, **extra}]}]}


def _editing(deck: Deck, index: int = 0) -> str:
    token = PLACEHOLDERS.set(True)
    try:
        return render_slide(deck, deck.slides[index]).svg
    finally:
        PLACEHOLDERS.reset(token)


def _span(svg: str, identifier: str) -> tuple[float, float]:
    found = re.search(rf'id="{re.escape(identifier)}"[^>]*?data-flexo-span="([-\d.]+) ([-\d.]+)"', svg)
    assert found, identifier
    return float(found.group(1)), float(found.group(2))


def test_every_kind_of_object_takes_a_width() -> None:
    assert all("width" in options for options in BLOCKS.values())


@pytest.mark.parametrize("index", range(10))
def test_an_object_given_a_width_is_set_in_it_where_its_place_sets_it(tmp_path: Path, index: int) -> None:
    """Narrower than its place, it takes that width: words wrapped in it, numbers, a listing's
    panel, an equation set smaller, a table's columns, a gallery's cells, a mechanism drawn
    smaller -- at the side its words are set against, or where a picture stands."""

    block = _blocks(tmp_path)[index]
    kind = next(iter(block))
    unsized = deck_from_document(_slide(tmp_path, block), tmp_path)
    _, wide = _span(_editing(unsized), "slide1.body.0") if kind != "mechanism" else (0.0, 0.0)
    asked = {"math": 60.0, "table": 260.0, "code": 230.0}.get(kind, 300.0)
    sized = deck_from_document(_slide(tmp_path, block, width=asked), tmp_path)
    svg = _editing(sized)
    if kind == "mechanism":
        width = float(re.search(r'id="slide1\.body\.0"[^>]*data-flexo-width="([\d.]+)"', svg).group(1))
        assert width == pytest.approx(asked, abs=1.0)
        return
    now_left, now = _span(svg, "slide1.body.0")
    assert now == pytest.approx(asked, abs=0.5), kind
    assert now < wide or kind == "table", kind
    margin = sized.style.margin
    if kind in {"text", "bullets", "quote", "callout", "code", "stats"}:
        # Set against the left of its place, as its words are.
        assert now_left == pytest.approx(margin, abs=0.5), kind
    else:
        # Centred, as an equation and the pictures and tables standing alone in their place are.
        assert now_left + now / 2 == pytest.approx(sized.style.width / 2, abs=0.5), kind
    if kind in {"text", "bullets", "quote", "callout"}:
        # Its words wrapped in it: on more lines than in its whole place.
        def lines(text: str) -> int:
            return len(re.findall(r'<tspan x=', text))

        assert lines(svg) > lines(_editing(unsized)), kind


def test_a_width_is_the_most_an_object_takes_and_its_place_the_most_it_is_given(tmp_path: Path) -> None:
    block = {"text": PARAGRAPH}
    deck = deck_from_document(_slide(tmp_path, block, width=3000), tmp_path)
    _, wide = _span(_editing(deck), "slide1.body.0")
    assert wide == pytest.approx(deck.style.width - 2 * deck.style.margin)


def test_an_object_as_wide_as_its_place_is_drawn_as_one_with_no_width(tmp_path: Path) -> None:
    """Given its place's width, words, a list, a quote, a callout, numbers or a gallery are
    drawn exactly as with no width of their own: no width is filling its place."""

    deck = Deck("place")
    place = deck.style.width - 2 * deck.style.margin
    for index in (0, 1, 2, 3, 4, 8):
        block = _blocks(tmp_path)[index]
        plain = deck_from_document(_slide(tmp_path, block), tmp_path)
        sized = deck_from_document(_slide(tmp_path, block, width=place), tmp_path)
        assert render_slide(sized, sized.slides[0]).svg == render_slide(plain, plain.slides[0]).svg, block


def test_a_centred_text_or_one_asked_to_stand_at_the_right_stands_there_given_a_width(tmp_path: Path) -> None:
    deck = deck_from_document(_slide(tmp_path, {"text": PARAGRAPH, "align": "middle"}, width=300), tmp_path)
    left, wide = _span(_editing(deck), "slide1.body.0")
    assert left + wide / 2 == pytest.approx(deck.style.width / 2, abs=0.5)
    deck = deck_from_document(_slide(tmp_path, {"text": PARAGRAPH}, width=300, horizontal="end"), tmp_path)
    left, wide = _span(_editing(deck), "slide1.body.0")
    assert left + wide == pytest.approx(deck.style.width - deck.style.margin, abs=0.5)
    # Words that read from the right are set against the right.
    deck = deck_from_document(_slide(tmp_path, {"text": "سلام دنیا، این یک جمله است"}, width=200), tmp_path)
    left, wide = _span(_editing(deck), "slide1.body.0")
    assert left + wide == pytest.approx(deck.style.width - deck.style.margin, abs=0.5)


def test_a_table_given_a_width_shares_it_among_its_columns(tmp_path: Path) -> None:
    rows = [["Method", "What it does"], ["Ensemble", "Several models trained apart"], ["Ours", "One model, sampled"]]
    natural = deck_from_document(_slide(tmp_path, {"table": rows}), tmp_path)
    (own,) = render_slide(natural, natural.slides[0]).tables
    for width in (500.0, sum(own.widths) - 60.0):
        deck = deck_from_document(_slide(tmp_path, {"table": rows}, width=width), tmp_path)
        (table,) = render_slide(deck, deck.slides[0]).tables
        assert sum(table.widths) == pytest.approx(width, abs=0.01)
        # Its words as large as they were; wider, each column in proportion to its words.
        assert table.size == own.size
        if width > sum(own.widths):
            assert [a / b for a, b in zip(table.widths, own.widths, strict=True)] == pytest.approx(
                [width / sum(own.widths)] * 2)
        else:
            # Narrower than its words, the long cell wraps; the short column keeps its width.
            assert table.widths[0] == pytest.approx(own.widths[0])
            assert max(table.heights) > max(own.heights)
        # In a slide program its columns are as wide as here: no room to grow into.
        assert table.room == (table.x, sum(table.widths))


def test_a_plot_given_a_width_is_drawn_that_wide_its_height_in_proportion() -> None:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    def drawn(width: float | None) -> tuple[float, float]:
        deck = Deck(f"plot{width}")
        figure, axes = plt.subplots()
        axes.plot([0, 1], [0, 1])
        with deck.slide("Plot") as slide:
            slide.plot(figure, width=width)
        render_slide(deck, deck.slides[0])
        plt.close(figure)
        # (As it was drawn: the SVG it was laid out as, at its size.)
        drawn = deck.slides[0].body.blocks[0].drawn
        found = re.search(r'<svg[^>]*?width="([\d.]+)pt" height="([\d.]+)pt"', drawn)
        assert found
        return float(found.group(1)), float(found.group(2))

    whole, half = drawn(None), drawn(400.0)
    assert half[0] == pytest.approx(400.0, abs=1.0)
    assert half[0] / half[1] == pytest.approx(whole[0] / whole[1], rel=0.01)


def test_a_width_is_written_and_read_back_and_said_when_it_is_wrong(tmp_path: Path) -> None:
    body = [{**block, "width": 300} for block in _blocks(tmp_path)]
    deck = deck_from_document({"deck": {}, "slides": [{"title": "T", "body": body}]}, tmp_path)
    written = deck_document(deck)["slides"][0]["body"]
    assert [block["width"] for block in written] == [300.0] * len(body)
    # A Python deck's objects take a width, and write it.
    made = Deck("python")
    with made.slide("T") as slide:
        slide.text(PARAGRAPH, width=320)
        slide.bullets("One", "Two", width=200)
        slide.code("x = 1", width=150)
        slide.table([["a", "b"]], width=200)
        slide.stats(("1", "one"), width=300)
    widths = [block.get("width") for block in deck_document(made)["slides"][0]["body"]]
    assert widths == [320.0, 200.0, 150.0, 200.0, 300.0]
    for wrong, said in ((0, "must be between 1 and 4032"), ("wide", "must be a number"), (True, "must be a number")):
        with pytest.raises(DeckDocumentError, match=rf"slides\[0\]\.body\[0\] \(text\).*width {said}"):
            deck_from_document(_slide(tmp_path, {"text": "Words"}, width=wrong), tmp_path)


def test_the_powerpoint_takes_the_widths(tmp_path: Path) -> None:
    """The slide program's list, table and text boxes are as wide as their objects were set,
    where they were set."""

    from pptx import Presentation
    from pptx.util import Emu

    body = [{"text": PARAGRAPH, "width": 300}, {"bullets": ["A first item that runs on and on"], "width": 250,
                                                 "horizontal": "end"},
            {"table": [["A", "B"], ["1", "2"]], "width": 400}]
    deck = deck_from_document({"deck": {}, "slides": [{"title": "T", "body": body}]}, tmp_path)
    result = deck.build(tmp_path / "out", formats=("pptx",))
    shapes = {shape.name.split(" ")[0]: shape for shape in Presentation(result.pptx).slides[0].shapes}
    margin, width = deck.style.margin, deck.style.width
    text, listed, table = shapes["Text"], shapes["List"], shapes["Table"]
    assert Emu(text.left).pt == pytest.approx(margin, abs=1.0)
    assert Emu(text.width).pt <= 300 + 2.0 + deck.style.body_size * 0.5
    assert Emu(listed.left).pt + Emu(listed.width).pt == pytest.approx(width - margin, abs=1.0)
    assert Emu(listed.width).pt == pytest.approx(250, abs=1.0)
    assert sum(Emu(column.width).pt for column in table.table.columns) == pytest.approx(400, abs=0.5)


def test_the_studio_s_drawing_says_what_sizing_an_object_needs(tmp_path: Path) -> None:
    """While edited, each object says where it stands across its place, the width it takes of
    itself and the least it can be; words, how far apart their lines are and where a line was
    begun by hand; a table, its columns; a gallery, its grid -- and nothing of it is exported."""

    body = [{"text": "One line\nand another begun by hand"}, *(_blocks(tmp_path)[index] for index in (1, 7, 8))]
    deck = deck_from_document({"deck": {}, "slides": [{"title": "T", "body": body}]}, tmp_path)
    svg = _editing(deck)
    text = re.search(r'<text id="slide1\.body\.0"[^>]*>', svg).group(0)
    for name in ("data-flexo-span", "data-flexo-share", "data-flexo-natural", "data-flexo-least", "data-flexo-line"):
        assert name in text
    assert 'data-flexo-breaks="1"' in text
    assert re.search(r'id="slide1\.body\.2"[^>]*data-flexo-columns="[\d.]+,[\d.]+ [\d.]+,[\d.]+ [\d.]+,[\d.]+"', svg)
    assert re.search(r'id="slide1\.body\.3"[^>]*data-flexo-cells="2 ', svg)
    exported = render_slide(deck, deck.slides[0]).svg
    assert "data-flexo-span" not in exported and "data-flexo-least" not in exported
