"""A drawing pasted onto a slide as shapes (flexo.svgimport): its gradients, its dashes of
their own, its fainter lines and the pictures that keep what no shape draws -- on the slide,
in the PDF and in the PowerPoint, as PowerPoint's own."""

from __future__ import annotations

import base64
import re
import xml.etree.ElementTree as ET
from pathlib import Path

import pytest
from flexo.svgimport import shapes_of
from pptx import Presentation

from flexo_talk.document import MissingFile, deck_from_document
from flexo_talk.studio import DeckKind

SVG = "{http://www.w3.org/2000/svg}"
PNG = base64.b64encode(bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
    "0000000d49444154789c6360f8cf00000301010018dd8db00000000049454e44ae426082"
)).decode()
DRAWING = (
    '<svg xmlns="http://www.w3.org/2000/svg" xmlns:xlink="http://www.w3.org/1999/xlink" width="600" height="300">'
    '<defs><linearGradient id="g" x1="0" x2="1"><stop offset="0" stop-color="#ff0000"/>'
    '<stop offset="1" stop-color="#ffff00"/></linearGradient>'
    '<clipPath id="c"><circle cx="450" cy="150" r="60"/></clipPath></defs>'
    '<rect x="20" y="20" width="200" height="100" rx="10" fill="url(#g)" stroke="#000000" stroke-width="2"/>'
    '<path d="M 20 200 H 300" stroke="#1a5d9b" stroke-width="3" stroke-dasharray="9 3 2 3" stroke-opacity="0.5"/>'
    '<rect x="20" y="230" width="100" height="50" fill="#dbebfd" stroke="#1a5d9b" stroke-opacity="0.4"/>'
    f'<image x="380" y="80" width="140" height="140" clip-path="url(#c)" xlink:href="data:image/png;base64,{PNG}"/>'
    "</svg>"
)


def imported(folder: Path) -> list[dict]:
    """The drawing as a slide's shapes, its pictures beside the deck."""

    count = [0]

    def keep(data: bytes, suffix: str) -> str:
        count[0] += 1
        (folder / "assets").mkdir(exist_ok=True)
        name = f"assets/drawing-part-{count[0]}{suffix}"
        (folder / name).write_bytes(data)
        return name

    return shapes_of(DRAWING, centre=[480, 270], fit=[880, 490], picture=keep).shapes


def deck(shapes: list[dict]) -> dict:
    return {"deck": {"id": "talk", "theme": "paper"}, "slides": [{"layout": "blank", "shapes": shapes}]}


def test_a_drawings_shapes_are_drawn_on_the_slide(tmp_path: Path) -> None:
    shapes = imported(tmp_path)
    group = shapes[0]["shapes"]
    assert [shape["kind"] for shape in group] == ["rounded-rectangle", "line", "rectangle", "picture"]
    assert group[0]["fill"]["gradient"] == "linear" and group[1]["dash"] == [pytest.approx(v, abs=0.01) for v in
                                                                              group[1]["dash"]]
    assert group[2]["line_opacity"] == 0.4 and group[2].get("fill", "accent-fill") == "accent-fill"
    built = deck_from_document(deck(shapes), tmp_path)
    # Its picture found from the deck's folder, as an image's is.
    assert Path(built.slides[0].shapes[0]["shapes"][3]["source"]).is_absolute()
    root = ET.fromstring(built.render()[0].svg)
    assert root.find(f".//{SVG}linearGradient") is not None
    image = next(item for item in root.iter(f"{SVG}image") if item.get("data-flexo-source"))
    assert image.get("href").startswith("data:image/svg+xml;base64,")
    dashed = [item for item in root.iter(f"{SVG}path") if item.get("stroke-dasharray")]
    assert dashed and dashed[0].get("stroke-opacity") is None  # (a line alone: its opacity the shape's)
    faint = [item for item in root.iter(f"{SVG}path") if item.get("stroke-opacity") == "0.4"]
    assert faint


def test_a_picture_not_beside_the_deck_is_said(tmp_path: Path) -> None:
    shapes = [{"kind": "picture", "at": [0, 0], "size": [10, 10], "source": "assets/gone.svg"}]
    with pytest.raises(MissingFile, match=r"gone\.svg"):
        deck_from_document(deck(shapes), tmp_path)
    missing: list = []
    deck_from_document(deck(shapes), tmp_path, errors=missing)
    assert missing and "gone.svg" in str(missing[0])


def test_a_drawings_shapes_are_powerpoints_own(tmp_path: Path) -> None:
    shapes = imported(tmp_path)
    written = DeckKind().export(deck(shapes), tmp_path, "talk", ["pptx", "pdf"], into=tmp_path / "out")
    pptx = next(path for path in written if path.suffix == ".pptx")
    slide = Presentation(str(pptx)).slides[0]
    from lxml import etree

    xml = etree.tostring(slide._element).decode()
    # The gradient as PowerPoint's own, left to right.
    assert re.search(r'<a:gradFill rotWithShape="1"><a:gsLst><a:gs pos="0"><a:srgbClr val="FF0000"/>', xml)
    assert '<a:lin ang="0" scaled="0"/>' in xml
    # Its dashes as its own lengths, a share of the line's width; its fainter line's alpha.
    assert "<a:custDash><a:ds d=\"300000\" sp=\"100000\"/>" in xml
    assert '<a:alpha val="40000"/>' in xml
    # The picture that keeps the clipped photograph, a picture.
    assert xml.count("<p:pic") == 1 and "svgBlip" in xml
    pdf = next(path for path in written if path.suffix == ".pdf").read_bytes()
    assert b"/ShadingType 2" in pdf
